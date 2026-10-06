"""Bounded XLSX I/O for interview registers. Never overwrites source workbooks."""
from __future__ import annotations

from copy import copy
import math
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile, BadZipFile

try:
    from flow1c_interview_policy import COLUMNS, REQUIRED, MAX_ROWS, InterviewError, apply_changes, text_value, validate_requirements, validate_rows
except ModuleNotFoundError:
    from scripts.flow1c_interview_policy import COLUMNS, REQUIRED, MAX_ROWS, InterviewError, apply_changes, text_value, validate_requirements, validate_rows


def excel_modules() -> tuple[Any, Any]:
    try:
        import openpyxl
        from openpyxl import styles
    except ImportError as exc:
        raise InterviewError("INTERVIEW_EXCEL_UNAVAILABLE", "Для Excel нужен openpyxl.",
                             "Установите зависимости requirements.txt через согласованный bootstrap; можно продолжить Markdown-черновик.") from exc
    return openpyxl, styles


def load_book(source: Path) -> Any:
    if source.suffix.lower() != ".xlsx" or source.stat().st_size > 25 * 1024 * 1024:
        raise InterviewError("INTERVIEW_UNSUPPORTED_BOOK", "Нужна XLSX-книга не более 25 MiB.")
    try:
        with ZipFile(source) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                raise InterviewError("INTERVIEW_UNSUPPORTED_BOOK", "Распакованная книга превышает 100 MiB.")
            # openpyxl cannot preserve all these parts. Reject instead of silently losing data.
            unsupported = ("xl/activeX/", "xl/embeddings/", "xl/slicer", "xl/pivot", "xl/threadedComments/", "xl/externalLinks/", "customXml/", "_xmlsignatures/")
            if any(name.startswith(unsupported) or name.endswith("vbaProject.bin")
                   or (name.startswith("xl/drawings/") and not re.fullmatch(r"xl/drawings/commentsDrawing\d+\.vml", name))
                   for name in archive.namelist()):
                raise InterviewError("INTERVIEW_UNSUPPORTED_BOOK", "В книге есть элементы, сохранность которых этот адаптер не гарантирует.",
                                     "Сохраните исходник; используйте специализированный инструмент с проверкой сохранности или создайте отдельный дополнительный реестр.")
            for entry in archive.infolist():
                if entry.filename.endswith(".xml") and re.search(br"<(?:[A-Za-z0-9_]+:)?extLst\b", archive.read(entry)):
                    raise InterviewError("INTERVIEW_UNSUPPORTED_BOOK", "Книга содержит неподдерживаемые расширения OOXML.")
        openpyxl, _ = excel_modules()
        return openpyxl.load_workbook(source, data_only=False, keep_links=True)
    except Exception as exc:  # XLSX/XML readers expose multiple corruption exception types.
        if isinstance(exc, InterviewError):
            raise
        raise InterviewError("INTERVIEW_BOOK_INVALID", f"Не удалось прочитать XLSX: {exc}") from exc


def read_register(book: Any, sheet: str | None = None) -> tuple[Any, dict[str, int], list[dict[str, str]], list[int]]:
    candidates = [book[sheet]] if sheet and sheet in book.sheetnames else list(book.worksheets) if not sheet else []
    matches = []
    for worksheet in candidates:
        if worksheet.max_column > 100 or worksheet.max_row > MAX_ROWS + 1:
            continue
        headers = [cell.value for cell in worksheet[1]]
        mapping = {key: headers.index(label) + 1 for key, label in COLUMNS.items() if label in headers}
        if all(key in mapping for key in REQUIRED):
            if any(headers.count(COLUMNS[key]) != 1 for key in mapping):
                raise InterviewError("INTERVIEW_SCHEMA_INVALID", "В шапке реестра повторены заголовки.")
            matches.append((worksheet, mapping))
    if len(matches) != 1:
        raise InterviewError("INTERVIEW_SCHEMA_NOT_DETECTED", "Не найден единственный лист с тремя парами код/название и вопросом.",
                             "Укажите sheet при нескольких реестрах; сопоставьте нестандартную шапку с register-schema.md. Исходник не изменён.")
    worksheet, mapping = matches[0]
    if worksheet.merged_cells.ranges:
        raise InterviewError("INTERVIEW_SCHEMA_INVALID", "Объединённые ячейки в реестре требуют отдельного преобразования.")
    rows, positions = [], []
    used = {str(worksheet.cell(index, mapping["question_id"]).value or "") for index in range(2, worksheet.max_row + 1)} if "question_id" in mapping else set()
    for index in range(2, worksheet.max_row + 1):
        if not any(worksheet.cell(index, mapping[key]).value is not None for key in REQUIRED):
            continue
        row = {}
        for key, column in mapping.items():
            cell = worksheet.cell(index, column)
            if cell.data_type == "f" and key in (*REQUIRED, "question_id"):
                raise InterviewError("INTERVIEW_SCHEMA_INVALID", "Коды, названия, вопросы и ID не могут быть формулами.")
            row[key] = str(cell.value) if cell.value is not None else ""
        if not row.get("question_id"):
            candidate = f"Q-{index - 1:04d}"
            while candidate in used:
                candidate += "-new"
            row["question_id"] = candidate
            used.add(candidate)
        rows.append(row)
        positions.append(index)
    return worksheet, mapping, validate_rows(rows), positions


def audit_book(book: Any, sheet: str | None = None) -> dict[str, Any]:
    worksheet, mapping, rows, positions = read_register(book, sheet)
    errors = []
    warnings = []
    if "question_id" not in mapping:
        warnings.append("question_ids_provisional: IDs из просмотра будут записаны только в новую копию.")
    if positions and positions != list(range(2, positions[-1] + 1)):
        errors.append("blank_rows_in_register")
    from openpyxl.utils.cell import range_boundaries
    ranges = []
    if worksheet.auto_filter.ref:
        ranges.append(("worksheet_filter", worksheet.auto_filter.ref))
    for table in worksheet.tables.values():
        ranges.append((f"table:{table.name}", table.ref))
        if table.autoFilter and table.autoFilter.ref:
            ranges.append((f"table_filter:{table.name}", table.autoFilter.ref))
    if not ranges:
        errors.append("filter_missing")
    expected_last = positions[-1] if positions else 1
    for name, reference in ranges:
        left, top, right, bottom = range_boundaries(reference)
        if top != 1 or left != 1 or bottom < expected_last or right < max(mapping.values()):
            errors.append(f"range_incomplete:{name}:{reference}")
    if any(worksheet.row_dimensions[index].hidden for index in positions):
        warnings.append("hidden_rows: проверьте, сохранён ли намеренный фильтр; автоматически не сбрасывать.")
    return {"schema_version": 1, "ready": not errors, "sheet": worksheet.title, "row_count": len(rows),
            "errors": errors, "warnings": warnings, "ranges": ranges,
            "limitations": ["Смысловая полнота интервью и визуальное оформление не проверены.", "Формулы не вычислялись."]}


def literal(cell: Any, value: str) -> None:
    cell.value = text_value(value)
    cell.data_type = "s"  # Even strings starting with '=' are user text, never formulas.


def style_sheet(worksheet: Any, widths: list[int]) -> None:
    _, styles = excel_modules()
    from openpyxl.utils import get_column_letter
    border = styles.Border(**{edge: styles.Side(style="thin", color="8FAADC") for edge in ("left", "right", "top", "bottom")})
    for row in worksheet:
        line_count = 1
        for cell in row:
            cell.font = styles.Font(name="Arial", size=10, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "222222")
            cell.fill = styles.PatternFill("solid", fgColor="4472C4" if cell.row == 1 else "FFFFFF")
            cell.border = border
            cell.alignment = styles.Alignment(wrap_text=True, vertical="top")
            width = widths[cell.column - 1] if cell.column <= len(widths) else 35
            line_count = max(line_count, sum(max(1, math.ceil(len(line) / max(1, width * 0.8))) for line in str(cell.value or "").split("\n")))
        worksheet.row_dimensions[row[0].row].height = min(409, max(32 if row[0].row == 1 else 45, line_count * 14 + 8))
    for index, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"


def add_sheet(book: Any, title: str, headers: list[str], values: list[list[str]]) -> Any:
    sheet = book.create_sheet(title)  # openpyxl allocates a distinct name on revisions.
    for column, value in enumerate(headers, 1):
        literal(sheet.cell(1, column), value)
    for row_index, row in enumerate(values, 2):
        for column, value in enumerate(row, 1):
            literal(sheet.cell(row_index, column), value)
    style_sheet(sheet, [35] * len(headers))
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def build_book(request: dict[str, Any], source: Path | None = None) -> tuple[Any, str, dict[str, Any]]:
    allowed = {"source", "sheet", "output_name", "rows", "patches", "requirements", "context", "sources", "open_questions"}
    if not isinstance(request, dict) or set(request) - allowed:
        raise InterviewError("INTERVIEW_INVALID_INPUT", "Неизвестные поля запроса реестра.")
    openpyxl, _ = excel_modules()
    book = load_book(source) if source else openpyxl.Workbook()
    if source:
        worksheet, mapping, existing, positions = read_register(book, request.get("sheet"))
        labels = [cell.value for cell in worksheet[1]]
        if any(not isinstance(label, str) or not label.strip() for label in labels) or len(labels) != len(set(labels)):
            raise InterviewError("INTERVIEW_SCHEMA_INVALID", "Каждая колонка, включаемая в фильтр, требует непустого уникального заголовка.")
        if positions and positions != list(range(2, positions[-1] + 1)):
            raise InterviewError("INTERVIEW_SCHEMA_INVALID", "Внутри реестра есть пустые строки; нужен отдельный план ремонта.")
        if len(worksheet.tables) > 1 or any(table.totalsRowCount or table.totalsRowShown for table in worksheet.tables.values()):
            raise InterviewError("INTERVIEW_UNSUPPORTED_BOOK", "Несколько таблиц или строка итогов требуют отдельного преобразования.")
        last = positions[-1] if positions else 1
        if any(cell.value is not None for data in worksheet.iter_rows(min_row=last + 1) for cell in data):
            raise InterviewError("INTERVIEW_UNSUPPORTED_BOOK", "Под реестром есть заметки или иные данные; нужен отдельный план, чтобы не связать их с новыми вопросами.")
    else:
        worksheet = book.active
        worksheet.title = "Реестр"
        mapping, existing, positions = {}, [], []
    rows = apply_changes(existing, request.get("rows", []), request.get("patches", []))
    if not rows:
        raise InterviewError("INTERVIEW_EMPTY_REGISTER", "Нужен хотя бы один процесс с вопросом.")
    requirements = validate_requirements(request.get("requirements", []), rows)
    for key, label in COLUMNS.items():
        if key not in mapping:
            mapping[key] = max(worksheet.max_column if source else 0, *mapping.values(), 0) + 1
            literal(worksheet.cell(1, mapping[key]), label)
    for index, row in enumerate(rows, 2):
        old = existing[index - 2] if index - 2 < len(existing) else {}
        if source and index > len(existing) + 1:
            worksheet.row_dimensions[index].height = worksheet.row_dimensions[2].height
        for key, column in mapping.items():
            if key not in row or (key in old and old[key] == row[key]):
                continue
            cell = worksheet.cell(index, column)
            if cell.data_type == "f":
                raise InterviewError("INTERVIEW_UNSAFE_PATCH", "Изменение затронуло формулу; исходник сохранён.")
            if source and index > len(existing) + 1:
                template = worksheet.cell(2, column)
                cell._style = copy(template._style)
                cell.alignment = copy(template.alignment)
            literal(cell, row[key])
            if key.endswith("_code") or key == "question_id":
                cell.number_format = "@"
    # IDs synthesized by the preview are absent in the input cells, so write them explicitly.
    for index, row in enumerate(rows, 2):
        literal(worksheet.cell(index, mapping["question_id"]), row["question_id"])
    from openpyxl.utils import get_column_letter
    reference = f"A1:{get_column_letter(worksheet.max_column)}{len(rows) + 1}"
    if not source:
        style_sheet(worksheet, [13, 30, 13, 34, 15, 38, 75, 32, 32, 65, 65, 18, 50, 45])
        from openpyxl.worksheet.table import Table, TableStyleInfo
        table = Table(displayName="InterviewRegister", ref=reference)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False, showColumnStripes=False)
        worksheet.add_table(table)
    else:
        for table in worksheet.tables.values():
            table.ref = reference
            if table.autoFilter:
                table.autoFilter.ref = reference
            # Preserve existing table column definitions and filters; add only new columns.
            from openpyxl.worksheet.table import TableColumn
            for column in range(len(table.tableColumns) + 1, worksheet.max_column + 1):
                table.tableColumns.append(TableColumn(id=column, name=str(worksheet.cell(1, column).value)))
    if worksheet.auto_filter.ref or not source:
        worksheet.auto_filter.ref = reference
    elif not worksheet.tables:
        worksheet.auto_filter.ref = reference
    context = request.get("context", {})
    if not isinstance(context, dict) or len(context) > 100:
        raise InterviewError("INTERVIEW_INVALID_INPUT", "context должен быть объектом до 100 полей.")
    metadata = [["Статус", "UNVERIFIED_DRAFT"], ["Generated by", "Flow1C"],
                ["Идентификаторы", "ID вопросов и draft_id — локальные ссылки интервью, не ID формального реестра."],
                ["Согласование", "Реестр не утверждает требования; вопрос сам по себе не является требованием."]]
    metadata.extend([[text_value(key), text_value(value)] for key, value in context.items()])
    add_sheet(book, "FLOW1C Контекст", ["Сведение", "Значение"], metadata)
    if requirements:
        fields = ("draft_id", "l3_code", "question_id", "answer", "text", "acceptance_criterion", "status", "owner", "priority", "implementation", "evidence")
        add_sheet(book, "FLOW1C Требования", ["Локальный ID черновика", "Код L3", "ID вопроса", "Исходный ответ", "Требование", "Критерий приёмки", "Статус", "Владелец", "Приоритет", "Способ реализации", "Основание проверки"],
                  [[item.get(field, "") for field in fields] for item in requirements])
    sources = request.get("sources", [])
    fields = ("id", "location", "section", "product_release", "checked_at", "availability", "conclusion", "related_questions")
    if not isinstance(sources, list) or len(sources) > 1000:
        raise InterviewError("INTERVIEW_INVALID_INPUT", "sources должен быть ограниченным списком.")
    if sources:
        if any(not isinstance(item, dict) or set(item) - set(fields) for item in sources):
            raise InterviewError("INTERVIEW_INVALID_INPUT", "Неизвестные поля источника.")
        add_sheet(book, "FLOW1C Источники", ["ID", "URL / файл", "Раздел", "Продукт / релиз", "Дата проверки", "Доступность", "Вывод", "Вопросы"],
                  [[text_value(item.get(field, "")) for field in fields] for item in sources])
    questions = request.get("open_questions", [])
    if not isinstance(questions, list) or len(questions) > MAX_ROWS:
        raise InterviewError("INTERVIEW_INVALID_INPUT", "open_questions должен быть ограниченным списком.")
    if questions:
        add_sheet(book, "FLOW1C Открытые вопросы", ["Что уточнить"], [[text_value(value)] for value in questions])
    book.properties.description = "UNVERIFIED_DRAFT — Flow1C interview register"
    return book, worksheet.title, {"rows": rows, "requirements": requirements}
