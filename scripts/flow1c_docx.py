"""Safe, bounded DOCX inspection and section writeback.

Only OOXML packages with a ``.docx`` suffix are accepted. The module never
executes document content, macros, external links or embedded objects. It
returns structured dictionaries so the CLI and agent adapters can expose the
same contract.
"""
from __future__ import annotations

import copy
import hashlib
import os
import re
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from typing import Any, Iterable

try:
    from docx import Document
    from docx.document import Document as DocumentObject
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
except ModuleNotFoundError:  # core CLI remains usable; doctor reports office capability separately
    Document = None  # type: ignore[assignment]
    DocumentObject = Any  # type: ignore[assignment,misc]
    Paragraph = Any  # type: ignore[assignment,misc]
    def qn(value: str) -> str:  # pragma: no cover - only reached without office dependencies
        return value

try:
    from flow1c_sections_policy import normalize_section_name
except ModuleNotFoundError:
    from scripts.flow1c_sections_policy import normalize_section_name

MAX_DOCX_BYTES = 50 * 1024 * 1024
MAX_UNPACKED_BYTES = 250 * 1024 * 1024
MAX_DOCX_PARTS = 4096
PLACEHOLDER_RE = re.compile(r"(?:\{\{[^}]+\}\}|\[\s*(?:placeholder|заполнить|вставить)[^]]*\]|placeholder|заполнить)", re.IGNORECASE)
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class DocxError(ValueError):
    def __init__(self, code: str, message: str, *, next_action: str = "", preserved_state: Any = None):
        self.code = code
        self.message = message
        self.next_action = next_action or "Correct the DOCX source or plan and retry."
        self.preserved_state = preserved_state
        super().__init__(message)

    def as_dict(self) -> dict[str, Any]:
        return {"code": self.code, "component": "docx-adapter", "message": self.message,
                "recoverable": True, "next_action": self.next_action, "preserved_state": self.preserved_state}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_docx_package(path: Path) -> None:
    if Document is None:
        raise DocxError("DOCX_RENDER_UNAVAILABLE", "Зависимость python-docx недоступна в текущем профиле.", next_action="Запустите bootstrap профиля office/documents и повторите операцию.")
    if path.suffix.casefold() != ".docx":
        raise DocxError("DOCX_INVALID", "Поддерживается только формат .docx; .doc и .docm запрещены.", next_action="Выберите обычный DOCX-файл.")
    if path.name.startswith("~$"):
        raise DocxError("DOCX_LOCK_FILE", "Файл является временным lock-файлом Word.", next_action="Выберите исходный документ, а не ~$-файл.")
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise DocxError("DOCX_SOURCE_NOT_FOUND", f"Не удалось прочитать DOCX: {path}", next_action="Проверьте путь и доступ к файлу.") from exc
    if size > MAX_DOCX_BYTES:
        raise DocxError("DOCX_INVALID", f"DOCX превышает ограничение размера {MAX_DOCX_BYTES} байт.", next_action="Уменьшите документ или используйте согласованный безопасный источник.")
    try:
        with zipfile.ZipFile(path) as package:
            names = package.namelist()
            if len(names) > MAX_DOCX_PARTS or len(set(names)) != len(names):
                raise DocxError("DOCX_INVALID", "OOXML has too many or duplicate parts.")
            if any(name.startswith(("/", "\\")) or "\\" in name or ".." in name.split("/") for name in names):
                raise DocxError("DOCX_INVALID", "Unsafe OOXML part path.")
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocxError("DOCX_INVALID", "В DOCX отсутствуют обязательные части OOXML.", next_action="Выберите валидный Word-документ.")
            total = sum(info.file_size for info in package.infolist())
            if total > MAX_UNPACKED_BYTES:
                raise DocxError("DOCX_INVALID", "Распакованный OOXML превышает безопасный лимит.", next_action="Используйте документ меньшего размера.")
            if package.testzip() is not None:
                raise DocxError("DOCX_INVALID", "OOXML-пакет поврежден.", next_action="Получите неповрежденную копию DOCX.")
            for info in package.infolist():
                if info.filename.endswith((".xml", ".rels")):
                    raw = package.read(info)
                    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                        raise DocxError("DOCX_INVALID", "DTD/entities are forbidden in OOXML.")
                    from lxml import etree
                    try:
                        tree = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True,
                                                                    load_dtd=False, huge_tree=False))
                        if tree.getroottree().docinfo.doctype:
                            raise DocxError("DOCX_INVALID", "DTD/entities are forbidden in all XML encodings.")
                    except etree.XMLSyntaxError as exc:
                        raise DocxError("DOCX_INVALID", "Malformed or over-budget OOXML XML.") from exc
            if any(name.casefold().endswith("vbaproject.bin") for name in names):
                raise DocxError("DOCX_INVALID", "Документы с макросами не поддерживаются.", next_action="Сохраните документ как .docx без макросов.")
    except zipfile.BadZipFile as exc:
        raise DocxError("DOCX_INVALID", "Файл не является валидным OOXML-пакетом.", next_action="Получите неповрежденный .docx.") from exc


def extract_template(path: Path) -> dict[str, Any]:
    """Generic template API; policy remains outside the OOXML adapter."""
    try:
        from flow1c_docx_templates import extract
    except ModuleNotFoundError:
        from scripts.flow1c_docx_templates import extract
    validate_docx_package(path)
    return extract(path)


def fill_template(path: Path, extraction: dict, operations: list[dict]) -> bytes:
    try:
        from flow1c_docx_templates import write_bytes
    except ModuleNotFoundError:
        from scripts.flow1c_docx_templates import write_bytes
    validate_docx_package(path)
    return write_bytes(path, extraction, operations)


def _body(document: DocumentObject):
    return document.element.body


def _direct_paragraphs(document: DocumentObject) -> list[Paragraph]:
    return [Paragraph(child, document) for child in _body(document) if child.tag == qn("w:p")]


def _outline_level(paragraph_properties: Any) -> int | None:
    if paragraph_properties is None:
        return None
    node = paragraph_properties.find(qn("w:outlineLvl"))
    if node is None:
        return None
    try:
        return int(node.get(qn("w:val"), "")) + 1
    except (TypeError, ValueError):
        return None


def _is_explicit_body_override(paragraph: Paragraph) -> bool:
    """Detect template paragraphs that deliberately neutralize a heading style.

    Some published templates reuse a numbered heading style for instructional
    body text and disable both numbering and bold directly on the paragraph.
    Treating the inherited outline level as a real boundary makes the target
    section appear empty and also hides newly inserted paragraphs from the
    write verifier.
    """
    properties = paragraph._p.pPr
    if properties is None:
        return False
    numbering = properties.find(qn("w:numPr"))
    num_id = numbering.find(qn("w:numId")) if numbering is not None else None
    run_properties = properties.find(qn("w:rPr"))
    bold = run_properties.find(qn("w:b")) if run_properties is not None else None
    bold_value = str(bold.get(qn("w:val"), "1") if bold is not None else "1").casefold()
    return num_id is not None and num_id.get(qn("w:val")) == "0" and bold_value in {"0", "false", "off"}


def _paragraph_level(paragraph: Paragraph) -> int | None:
    if _is_explicit_body_override(paragraph):
        return None
    style = getattr(paragraph, "style", None)
    style_name = str(getattr(style, "name", "") or "")
    match = re.search(r"heading\s*(\d+)", style_name, re.IGNORECASE)
    if match:
        return int(match.group(1))
    direct_level = _outline_level(paragraph._p.pPr)
    if direct_level is not None:
        return direct_level
    seen: set[str] = set()
    while style is not None:
        style_id = str(getattr(style, "style_id", "") or "")
        if style_id in seen:
            break
        seen.add(style_id)
        level = _outline_level(getattr(getattr(style, "element", None), "pPr", None))
        if level is not None:
            return level
        style = getattr(style, "base_style", None)
    return None


def _heading_match(text: str, section: dict[str, Any]) -> bool:
    candidates = [section["display_name"], section["section_id"], *section.get("aliases", []), *section.get("word", {}).get("heading_aliases", [])]
    normalized = normalize_section_name(text)
    return any(normalized == normalize_section_name(candidate) for candidate in candidates)


def _sdt_tag(element) -> str:
    tag = element.find(f".//{W_NS}sdtPr/{W_NS}tag")
    return str(tag.get(qn("w:val"), "")) if tag is not None else ""


def _bookmark_names(element) -> set[str]:
    return {str(node.get(qn("w:name"), "")) for node in element.iter(qn("w:bookmarkStart")) if node.get(qn("w:name"))}


def _bookmark_start(element, name: str):
    return next((node for node in element.iter(qn("w:bookmarkStart"))
                 if node.get(qn("w:name")) == name), None)


def _text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def _find_candidates(document: DocumentObject, section: dict[str, Any]) -> list[dict[str, Any]]:
    body = _body(document)
    children = list(body)
    paragraphs = _direct_paragraphs(document)
    para_by_el = {id(p._p): p for p in paragraphs}
    controls = []
    for index, child in enumerate(children):
        if child.tag != qn("w:sdt"):
            continue
        tag = _sdt_tag(child)
        if tag == f"flow1c-section:{section['section_id']}":
            text = _text_of(child)
            controls.append({"kind": "content_control", "candidate_id": tag, "title": section["display_name"], "element": child, "body_index": index, "anchor_text": text})
    if controls:
        return controls
    bookmark = f"flow1c_section_{section['section_id'].replace('-', '_')}"
    bookmarks = []
    for index, child in enumerate(children):
        start_node = _bookmark_start(child, bookmark)
        if start_node is not None:
            bookmark_id = start_node.get(qn("w:id"))
            end_node = None
            end_index = index
            for candidate_index in range(index, len(children)):
                end_node = next((node for node in children[candidate_index].iter(qn("w:bookmarkEnd"))
                                 if node.get(qn("w:id")) == bookmark_id), None)
                if end_node is not None:
                    end_index = candidate_index
                    break
            if end_node is None:
                continue
            paragraph = Paragraph(child, document) if child.tag == qn("w:p") else None
            level = _paragraph_level(paragraph) if paragraph is not None else None
            bookmarks.append({"kind": "bookmark", "candidate_id": bookmark, "title": _text_of(child),
                              "element": child, "body_index": index, "anchor_text": _text_of(child),
                              "bookmark_start": start_node, "bookmark_end": end_node,
                              "bookmark_end_index": end_index, "level": level,
                              "bookmark_is_heading": paragraph is not None and level is not None})
    if bookmarks:
        return bookmarks
    headings = []
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        paragraph = para_by_el.get(id(child)) or Paragraph(child, document)
        if _heading_match(paragraph.text, section):
            headings.append({"kind": "heading", "candidate_id": f"heading:{index}", "title": paragraph.text, "element": child, "body_index": index, "anchor_text": paragraph.text, "level": _paragraph_level(paragraph)})
    return headings


def _body_content_end(children: list[Any]) -> int:
    return next((index for index, child in enumerate(children) if child.tag == qn("w:sectPr")), len(children))


def _range_bounds(document: DocumentObject, candidate: dict[str, Any]) -> tuple[int, int]:
    children = list(_body(document))
    start = candidate["body_index"]
    end = _body_content_end(children)
    heading_anchor = candidate["kind"] == "heading" or candidate.get("bookmark_is_heading", False)
    if heading_anchor:
        level = candidate.get("level")
        for index in range(start + 1, len(children)):
            if children[index].tag != qn("w:p"):
                continue
            paragraph = Paragraph(children[index], document)
            paragraph_level = _paragraph_level(paragraph)
            if paragraph_level is not None and (level is None or paragraph_level <= level):
                end = index
                break
        return start + 1, end
    if candidate["kind"] == "bookmark":
        return start, min(int(candidate.get("bookmark_end_index", start)) + 1, end)
    return start + 1, end


def _range_text(document: DocumentObject, candidate: dict[str, Any]) -> str:
    if candidate["kind"] == "content_control":
        return candidate.get("anchor_text", "")
    children = list(_body(document))
    start, end = _range_bounds(document, candidate)
    return "\n".join(_text_of(child) for child in children[start:end] if _text_of(child)).strip()


def _status(existing: str) -> str:
    if not existing:
        return "FOUND_EMPTY"
    return "FOUND_PLACEHOLDER" if PLACEHOLDER_RE.search(existing) else "FOUND_CONTENT"


def _verification_text(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value or ""))).strip()


def _nodes_text(nodes: Iterable[Any]) -> str:
    return _verification_text("\n".join(_text_of(node) for node in nodes if _text_of(node)))


def inspect_docx(path: str | Path, sections: Iterable[dict[str, Any]]) -> dict[str, Any]:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise DocxError("DOCX_SOURCE_NOT_FOUND", f"DOCX не найден: {source}", next_action="Укажите существующий файл.")
    validate_docx_package(source)
    try:
        document = Document(str(source))
    except Exception as exc:  # python-docx exposes several parser-specific exceptions
        raise DocxError("DOCX_INVALID", f"Не удалось открыть DOCX: {source}", next_action="Получите неповрежденный документ.") from exc
    result_sections = []
    for section in sections:
        candidates = _find_candidates(document, section)
        if not candidates:
            result_sections.append({"section_id": section["section_id"], "status": "NOT_FOUND", "candidates": []})
            continue
        if len(candidates) > 1:
            result_sections.append({"section_id": section["section_id"], "status": "AMBIGUOUS", "candidates": [{"candidate_id": c["candidate_id"], "kind": c["kind"], "title": c["title"]} for c in candidates]})
            continue
        candidate = candidates[0]
        existing = _range_text(document, candidate)
        result_sections.append({"section_id": section["section_id"], "status": _status(existing), "candidate": {"candidate_id": candidate["candidate_id"], "kind": candidate["kind"], "title": candidate["title"]}, "existing_text": existing})
    return {"schema_version": 1, "source": str(source), "source_sha256": sha256_file(source), "sections": result_sections}


def _section_exemplars(document: DocumentObject, candidate: dict[str, Any]) -> tuple[Any | None, Any | None]:
    """Use formatting owned by the selected section, never by a fresh DOCX."""
    if candidate["kind"] == "content_control":
        content = candidate["element"].find(f".//{W_NS}sdtContent")
        children = list(content) if content is not None else []
    else:
        start, end = _range_bounds(document, candidate)
        children = list(_body(document))[start:end]
    paragraph = next((node for node in children if node.tag == qn("w:p") and _paragraph_level(Paragraph(node, document)) is None), None)
    table = next((node for node in children if node.tag == qn("w:tbl")), None)
    return paragraph, table


def _set_paragraph_text(node: Any, value: str, exemplar: Any | None) -> None:
    for child in list(node):
        if child.tag != qn("w:pPr"):
            node.remove(child)
    run = OxmlElement("w:r")
    if exemplar is not None:
        source_run = exemplar.find(qn("w:r"))
        properties = source_run.find(qn("w:rPr")) if source_run is not None else None
        if properties is not None:
            run.append(copy.deepcopy(properties))
    text_node = OxmlElement("w:t")
    text_node.text = value
    run.append(text_node)
    node.append(run)


def _set_formatted_paragraph_text(node: Any, value: str, exemplar: Any | None) -> None:
    """Turn supported inline Markdown into Word runs, without visible markup."""
    for child in list(node):
        if child.tag != qn("w:pPr"):
            node.remove(child)
    source_run = exemplar.find(qn("w:r")) if exemplar is not None else None
    source_properties = source_run.find(qn("w:rPr")) if source_run is not None else None
    value = re.sub(r"\\([#*`_])", r"\1", value)
    cursor = 0
    for match in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", value):
        if match.start() > cursor:
            _append_run(node, value[cursor:match.start()], source_properties)
        if match.group(1) is not None:
            _append_run(node, match.group(1), source_properties, bold=True)
        else:
            _append_run(node, f"«{match.group(2)}»", source_properties)
        cursor = match.end()
    if cursor < len(value):
        _append_run(node, value[cursor:], source_properties)


def _append_run(node: Any, value: str, source_properties: Any | None, *, bold: bool = False) -> None:
    if not value:
        return
    run = OxmlElement("w:r")
    properties = copy.deepcopy(source_properties) if source_properties is not None else None
    if bold:
        if properties is None:
            properties = OxmlElement("w:rPr")
        bold_node = properties.find(qn("w:b"))
        if bold_node is None:
            bold_node = OxmlElement("w:b")
            properties.append(bold_node)
        bold_node.set(qn("w:val"), "1")
    if properties is not None:
        run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = value
    run.append(text_node)
    node.append(run)


def _table_from_exemplar(document: DocumentObject, exemplar: Any, rows: list[list[str]]) -> Any:
    table = copy.deepcopy(exemplar)
    existing_rows = table.findall(qn("w:tr"))
    if not existing_rows or len(existing_rows[0].findall(qn("w:tc"))) != max(map(len, rows)):
        raise DocxError("DOCX_LAYOUT_UNSUPPORTED", "Число столбцов новой таблицы отличается от таблицы-образца раздела.",
                        next_action="Используйте число столбцов таблицы-образца или подготовьте подходящий шаблон.")
    while len(existing_rows) < len(rows):
        clone = copy.deepcopy(existing_rows[-1])
        table.append(clone)
        existing_rows.append(clone)
    for extra in existing_rows[len(rows):]:
        table.remove(extra)
    for row_node, values in zip(existing_rows, rows):
        for cell_node, value in zip(row_node.findall(qn("w:tc")), values):
            paragraphs = cell_node.findall(qn("w:p"))
            exemplar_paragraph = paragraphs[0] if paragraphs else None
            paragraph = copy.deepcopy(exemplar_paragraph) if exemplar_paragraph is not None else OxmlElement("w:p")
            _set_paragraph_text(paragraph, value, exemplar_paragraph)
            for child in list(cell_node):
                if child.tag != qn("w:tcPr"):
                    cell_node.remove(child)
            cell_node.append(paragraph)
    return table


def _new_table(document: DocumentObject, rows: list[list[str]]) -> Any:
    table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            table.cell(row_index, col_index).text = value
    node = copy.deepcopy(table._tbl)
    table._tbl.getparent().remove(table._tbl)
    return node


def _style_has_numbering(document: DocumentObject, style_name: str) -> bool:
    style = document.styles[style_name] if style_name in document.styles else None
    seen: set[str] = set()
    while style is not None:
        style_id = str(style.style_id)
        if style_id in seen:
            break
        seen.add(style_id)
        properties = style.element.pPr
        if properties is not None and properties.find(qn("w:numPr")) is not None:
            return True
        style = style.base_style
    return False


def _markdown_nodes(content: str, document: DocumentObject | None = None, candidate: dict[str, Any] | None = None,
                    *, section_id: str = ""):
    temp = Document()
    paragraph_exemplar, table_exemplar = _section_exemplars(document, candidate) if document is not None and candidate is not None else (None, None)
    lines = str(content).replace("\r\n", "\n").splitlines()
    index = 0
    heading_numbers: dict[int, int] = {}
    parent_number = ""
    if candidate is not None:
        parent_match = re.match(r"^\s*(\d+(?:\.\d+)*)[.)]?\s+", str(candidate.get("title", "")))
        parent_number = parent_match.group(1) if parent_match else ""
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1]):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                if document is not None:
                    temp.element.body.append(_table_from_exemplar(document, table_exemplar, rows) if table_exemplar is not None else _new_table(document, rows))
                else:
                    table = temp.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                    table.style = "Table Grid"
                    for row_index, row in enumerate(rows):
                        for col_index, value in enumerate(row):
                            table.cell(row_index, col_index).text = value
            continue
        style = None
        value = line
        heading = re.match(r"^\\?(#{1,6})\s+(.+)$", line)
        technical_group = (re.match(r"^(?:\d+|[a-z])[.)]\s+(.+)$", line, re.IGNORECASE)
                           if section_id == "technical-implementation" else None)
        if heading:
            anchor_level = int(candidate.get("level") or 1) if candidate is not None else 1
            heading_level = min(9, anchor_level + max(1, len(heading.group(1)) - 2))
            style, value = f"Heading {heading_level}", heading.group(2)
            heading_numbers[heading_level] = heading_numbers.get(heading_level, 0) + 1
            for deeper in [level for level in heading_numbers if level > heading_level]:
                del heading_numbers[deeper]
            if not re.match(r"^\d+(?:\.\d+)*[.)]\s+", value):
                if section_id == "technical-implementation" or document is None or not _style_has_numbering(document, style):
                    parts = [str(heading_numbers[level]) for level in sorted(heading_numbers)]
                    value = f"{parent_number + '.' if parent_number else ''}{'.'.join(parts)}. {value}"
        elif technical_group:
            anchor_level = int(candidate.get("level") or 1) if candidate is not None else 1
            heading_level = min(9, anchor_level + 1)
            heading_numbers[heading_level] = heading_numbers.get(heading_level, 0) + 1
            style = f"Heading {heading_level}"
            value = f"{heading_numbers[heading_level]}. {technical_group.group(1)}"
        elif re.match(r"^[-*+·•]\s+", line):
            style, value = "List Bullet", re.sub(r"^[-*+·•]\s+", "", line)
        elif re.match(r"^(?:\d+|[a-z])[.)]\s+", line, re.IGNORECASE):
            style, value = "List Number", re.sub(r"^(?:\d+|[a-z])[.)]\s+", "", line, flags=re.IGNORECASE)
        paragraph = temp.add_paragraph(style=style)
        if document is not None and paragraph_exemplar is not None and not (heading or technical_group):
            source_properties = paragraph_exemplar.find(qn("w:pPr"))
            if source_properties is not None:
                old = paragraph._p.pPr
                if old is not None:
                    paragraph._p.remove(old)
                paragraph._p.insert(0, copy.deepcopy(source_properties))
            if style is not None:
                if style not in document.styles:
                    raise DocxError("DOCX_LAYOUT_UNSUPPORTED", f"В исходном документе нет стиля списка {style}.",
                                    next_action="Добавьте стиль списка в шаблон или передайте текст без Markdown-списка.")
                properties = paragraph._p.get_or_add_pPr()
                numbering = properties.find(qn("w:numPr"))
                if numbering is not None:
                    properties.remove(numbering)
                properties.get_or_add_pStyle().val = document.styles[style].style_id
        if (technical_group or (heading and section_id == "technical-implementation")) and document is not None and _style_has_numbering(document, style):
            properties = paragraph._p.get_or_add_pPr()
            numbering = properties.find(qn("w:numPr"))
            if numbering is None:
                numbering = OxmlElement("w:numPr")
                properties.append(numbering)
            num_id = numbering.find(qn("w:numId"))
            if num_id is None:
                num_id = OxmlElement("w:numId")
                numbering.append(num_id)
            num_id.set(qn("w:val"), "0")
        _set_formatted_paragraph_text(paragraph._p, value, paragraph_exemplar if not (heading or technical_group) else None)
        index += 1
    return [copy.deepcopy(child) for child in temp.element.body if child.tag in {qn("w:p"), qn("w:tbl")}]


def _insert_after(anchor, nodes: list[Any]) -> None:
    current = anchor
    for node in nodes:
        current.addnext(node)
        current = node


def _replace_content_control(candidate: dict[str, Any], nodes: list[Any]) -> None:
    content = candidate["element"].find(f".//{W_NS}sdtContent")
    if content is None:
        raise DocxError("DOCX_WRITE_FAILED", "Content control не содержит sdtContent.")
    for child in list(content):
        content.remove(child)
    for node in nodes:
        content.append(node)


def _append_content_control(candidate: dict[str, Any], nodes: list[Any]) -> None:
    content = candidate["element"].find(f".//{W_NS}sdtContent")
    if content is None:
        raise DocxError("DOCX_WRITE_FAILED", "Content control не содержит sdtContent.")
    for node in nodes:
        content.append(node)


def _paragraph_in(node: Any, *, last: bool = False):
    paragraphs = [item for item in node.iter(qn("w:p"))]
    return paragraphs[-1 if last else 0] if paragraphs else None


def _attach_bookmark(nodes: list[Any], start_node: Any | None, end_node: Any | None) -> None:
    if not nodes:
        raise DocxError("DOCX_WRITE_FAILED", "Нельзя записать пустой диапазон bookmark.")
    first = _paragraph_in(nodes[0])
    last = _paragraph_in(nodes[-1], last=True)
    if first is None or last is None:
        raise DocxError("DOCX_WRITE_FAILED", "Bookmark должен содержать абзац для безопасной записи.")
    if start_node is not None:
        first.insert(0, copy.deepcopy(start_node))
    if end_node is not None:
        last.append(copy.deepcopy(end_node))


def write_docx(source: str | Path, output: str | Path, operations: list[dict[str, Any]], *, expected_source_sha256: str | None = None) -> dict[str, Any]:
    source_path = Path(source).expanduser().resolve()
    output_path = Path(output).expanduser().resolve()
    validate_docx_package(source_path)
    source_hash = sha256_file(source_path)
    if expected_source_sha256 and source_hash != expected_source_sha256:
        raise DocxError("DOCX_WRITE_FAILED", "Исходный DOCX изменился после инспекции.", next_action="Повторите docx-inspect и docx-write-plan.")
    if output_path.exists():
        raise DocxError("OUTPUT_ALREADY_EXISTS", f"Выходной файл уже существует: {output_path}", next_action="Выберите другой результат или используйте сохраненный проверенный результат.")
    try:
        document = Document(str(source_path))
        def operation_index(item: dict[str, Any]) -> int:
            raw = str(item.get("body_index", 0))
            try:
                return int(raw.rsplit(":", 1)[-1])
            except ValueError:
                return 0

        for operation in sorted(operations, key=operation_index, reverse=True):
            section = operation["section"]
            candidates = _find_candidates(document, section)
            if len(candidates) != 1:
                code = "TARGET_SECTION_NOT_FOUND" if not candidates else "TARGET_SECTION_AMBIGUOUS"
                raise DocxError(code, f"Целевой раздел {section['section_id']} больше не найден однозначно.", next_action="Повторите inspect/plan.")
            candidate = candidates[0]
            status = _status(_range_text(document, candidate))
            mode = operation.get("mode", "replace")
            if status == "FOUND_CONTENT" and mode not in {"replace", "append"}:
                raise DocxError("REPLACE_MODE_REQUIRED", f"Раздел {section['display_name']} уже заполнен.", next_action="Явно выберите replace или append.")
            nodes = _markdown_nodes(operation.get("content", ""), document, candidate,
                                    section_id=section["section_id"])
            if candidate["kind"] == "content_control":
                if mode == "append":
                    _append_content_control(candidate, nodes)
                else:
                    _replace_content_control(candidate, nodes)
                continue
            body = _body(document)
            children = list(body)
            start, end = _range_bounds(document, candidate)
            bookmark_range = candidate["kind"] == "bookmark" and not candidate.get("bookmark_is_heading", False)
            if bookmark_range and mode == "replace":
                _attach_bookmark(nodes, candidate.get("bookmark_start"), candidate.get("bookmark_end"))
            if mode == "replace":
                for child in children[start:end]:
                    body.remove(child)
                for offset, node in enumerate(nodes):
                    body.insert(start + offset, node)
            else:
                if bookmark_range:
                    end_marker = candidate.get("bookmark_end")
                    if end_marker is not None and end_marker.getparent() is not None:
                        end_marker.getparent().remove(end_marker)
                    _attach_bookmark(nodes, None, end_marker)
                for offset, node in enumerate(nodes):
                    body.insert(end + offset, node)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(prefix=f".{output_path.stem}-", suffix=".docx", dir=output_path.parent)
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            document.save(str(temporary_path))
            validate_docx_package(temporary_path)
            verified = Document(str(temporary_path))
            for operation in operations:
                candidates = _find_candidates(verified, operation["section"])
                if len(candidates) != 1:
                    raise DocxError("DOCX_VERIFY_FAILED", f"Целевой раздел {operation['section']['section_id']} не найден после записи.", next_action="Проверьте структуру целевого раздела.")
                actual_text = _verification_text(_range_text(verified, candidates[0]))
                expected_text = _nodes_text(_markdown_nodes(operation.get("content", ""), verified, candidates[0],
                                                            section_id=operation["section"]["section_id"]))
                mode = str(operation.get("mode", "replace"))
                verified_text = actual_text == expected_text if mode == "replace" else expected_text in actual_text
                if not expected_text or not verified_text:
                    raise DocxError(
                        "DOCX_VERIFY_FAILED",
                        f"Не удалось подтвердить текст раздела {operation['section']['section_id']} после {mode}.",
                        next_action="Повторите docx-inspect; проверьте границы раздела и создайте новый plan.",
                        preserved_state={
                            "section_id": operation["section"]["section_id"],
                            "mode": mode,
                            "expected_text_sha256": hashlib.sha256(expected_text.encode("utf-8")).hexdigest(),
                            "actual_text_sha256": hashlib.sha256(actual_text.encode("utf-8")).hexdigest(),
                        },
                    )
            os.replace(temporary_path, output_path)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()
    except DocxError:
        raise
    except Exception as exc:
        raise DocxError("DOCX_WRITE_FAILED", f"Ошибка записи DOCX: {exc}", next_action="Исходный файл сохранен; исправьте причину и повторите план.") from exc
    return {"state": "WRITTEN", "layout_state": "WRITTEN_LAYOUT_UNVERIFIED",
            "layout_warning": "Поддерживаемый renderer не настроен; выполнена только структурная OOXML-проверка.",
            "source": str(source_path), "output": str(output_path), "source_sha256": source_hash,
            "output_sha256": sha256_file(output_path),
            "section_ids": [str(item["section"]["section_id"]) for item in operations],
            "modes": {str(item["section"]["section_id"]): str(item.get("mode", "replace")) for item in operations}}
