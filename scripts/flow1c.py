#!/usr/bin/env python3
"""Flow1C command line utilities.

The CLI deliberately keeps project state in reviewable files. JSON is written to
files with .yaml suffix only where a human-facing manifest is expected; JSON is a
valid YAML subset and avoids a mandatory YAML runtime dependency.
"""

from __future__ import annotations

import argparse
import datetime as dt
import getpass
import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
import uuid
from types import SimpleNamespace
from pathlib import Path
from typing import Any, Iterable

try:
    import flow1c_interview_cli as interview_cli
    from flow1c_interview_policy import InterviewError
except ModuleNotFoundError:
    from scripts import flow1c_interview_cli as interview_cli
    from scripts.flow1c_interview_policy import InterviewError

try:
    import flow1c_templates_cli as templates_cli
    from flow1c_templates import TemplateService, readiness as template_readiness
    from flow1c_templates_policy import TemplateError
except ModuleNotFoundError:
    from scripts import flow1c_templates_cli as templates_cli
    from scripts.flow1c_templates import TemplateService, readiness as template_readiness
    from scripts.flow1c_templates_policy import TemplateError

try:
    from flow1c_policy import (MODES, TERMINAL, accept_deviation, apply_deviation, available_actions,
        allowed_tools_for_mode, assess_request, build_conditions,
        capability_readiness, load_capability_model, operation_profile, python_version_supported, readiness_state,
        extract_git_ref, reference_slug, resolve_capabilities, resolve_work_reference, select_request_mode,
        stage_errors, transition, validate_git_ref, validate_reference)
except ModuleNotFoundError:
    from scripts.flow1c_policy import (MODES, TERMINAL, accept_deviation, apply_deviation, available_actions,
        allowed_tools_for_mode, assess_request, build_conditions,
        capability_readiness, load_capability_model, operation_profile, python_version_supported, readiness_state,
        extract_git_ref, reference_slug, resolve_capabilities, resolve_work_reference, select_request_mode,
        stage_errors, transition, validate_git_ref, validate_reference)

try:
    from external_tools import bsl_smoke, load_manifest as load_external_tools_manifest, package_version, skills_connected, version_satisfies
except ModuleNotFoundError:
    from scripts.external_tools import bsl_smoke, load_manifest as load_external_tools_manifest, package_version, skills_connected, version_satisfies

try:
    from cc_query_inspector import CcInspectionError, CcInspectionUnavailable, resolve_input_path, run_inspection
except ModuleNotFoundError:
    from scripts.cc_query_inspector import CcInspectionError, CcInspectionUnavailable, resolve_input_path, run_inspection

try:
    from flow1c_query_policy import analyze_query, compare_optimization, infer_query_request_intent, select_query_intent
    from flow1c_query_schema import parse_metadata_xml
except ModuleNotFoundError:
    from scripts.flow1c_query_policy import analyze_query, compare_optimization, infer_query_request_intent, select_query_intent
    from scripts.flow1c_query_schema import parse_metadata_xml

try:
    from redmine_client import RedmineClient, RedmineError
    from redmine_policy import RedminePolicyError, dmsf_identifier, issue_number, normalize_base_url
    import redmine_credentials
except ModuleNotFoundError:
    from scripts.redmine_client import RedmineClient, RedmineError
    from scripts.redmine_policy import RedminePolicyError, dmsf_identifier, issue_number, normalize_base_url
    from scripts import redmine_credentials

try:
    import flow1c_git as git_analysis
    from gitea_client import GiteaClient, GiteaError
    from flow1c_git_policy import (canonical_action, migrate_record as migrate_git_record,
                                   semantic_fingerprint, validate_transition)
except ModuleNotFoundError:
    from scripts import flow1c_git as git_analysis
    from scripts.gitea_client import GiteaClient, GiteaError
    from scripts.flow1c_git_policy import (canonical_action, migrate_record as migrate_git_record,
                                   semantic_fingerprint, validate_transition)

try:
    from flow1c_sections_policy import SectionPolicyError, catalog_summary, resolve_section_names, ensure_write_authorized, validate_section_content
    from flow1c_sections import (approve_section, load_catalog, load_state, paths_for, save_section,
                                 append_audit, storage_dir)
    from flow1c_docx import DocxError, inspect_docx, sha256_file, write_docx
except ModuleNotFoundError:
    from scripts.flow1c_sections_policy import SectionPolicyError, catalog_summary, resolve_section_names, ensure_write_authorized, validate_section_content
    from scripts.flow1c_sections import (approve_section, load_catalog, load_state, paths_for, save_section,
                                          append_audit, storage_dir)
    from scripts.flow1c_docx import DocxError, inspect_docx, sha256_file, write_docx


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


ROOT = Path(__file__).resolve().parents[1]
CONFIG_FILE = ".flow1c.json"
LOCAL_CONFIG_FILE = ".flow1c.local.json"
ID_PATTERN = re.compile(r"^[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё0-9_.]*-\d+$")
VALID_ROLES = ("analyst", "functional-architect", "technical-architect", "tester")
PROFILE_NAMES = ("conversation", "project-basic", "documents", "analysis", "implementation", "full")
SETUP_VALIDATION_VERSION = 1
LOCAL_CONFIG_SCHEMA_VERSION = 2
AGENT_GATE_SCHEMA_VERSION = 1
POLICY_VERSION = 3
STAGES_FILE = ROOT / "config" / "stages.json"
CAPABILITIES_FILE = ROOT / "config" / "capabilities.json"
SETUP_DIR = ROOT / ".workspace" / "setup"
AGENT_GATES_DIR = ROOT / ".workspace" / "agent-gates"
ALLOWED_ARTIFACT_EXTENSIONS = {".docx", ".xlsx", ".pdf", ".txt", ".md", ".csv", ".xml", ".png", ".jpg", ".jpeg"}
DENIED_ARTIFACT_EXTENSIONS = {".exe", ".dll", ".ps1", ".bat", ".cmd", ".js", ".ts", ".py", ".zip", ".7z", ".rar"}
MAX_INTAKE_FILES = 200
MAX_INTAKE_BYTES = 500 * 1024 * 1024
MAX_STORED_ARTIFACT_PATH = 240
INDEPENDENT_DRAFT_DECISION_PATTERN = re.compile(
    r"(?:независим\w*\s+черновик|отдельн\w*\s+черновик|без\s+формальн\w*\s+привязк\w*|independent\s+draft)",
    re.IGNORECASE,
)


class WorkflowError(RuntimeError):
    pass


class RedmineOperationError(WorkflowError):
    """Safe, machine-readable failure from Redmine setup or disconnect."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        recoverable: bool,
        next_action: str,
        previous_connection_preserved: bool,
        additional_errors: list[dict[str, Any]] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        primary = {
            "code": code,
            "message": message,
            "recoverable": recoverable,
            "next_action": next_action,
            "previous_connection_preserved": previous_connection_preserved,
        }
        self.payload = {
            "schema_version": 1,
            "state": "FAILED",
            "code": code,
            "configured": None,
            "verified": False,
            "previous_connection_preserved": previous_connection_preserved,
            "errors": (
                list(additional_errors)
                if code == "REDMINE_ROLLBACK_FAILED" and additional_errors
                else [primary, *(additional_errors or [])]
            ),
            "next_action": next_action,
        }


def read_json(path: Path, default: Any = None) -> Any:
    if not path.exists():
        return default
    try:
        return sanitize_json_value(json.loads(path.read_text(encoding="utf-8-sig")))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise WorkflowError(f"Cannot read JSON {path}: {exc}") from exc


def sanitize_text(text: str) -> str:
    """Replace lone UTF-16 surrogates while preserving valid surrogate pairs."""
    result: list[str] = []
    index = 0
    while index < len(text):
        codepoint = ord(text[index])
        if 0xD800 <= codepoint <= 0xDBFF:
            if index + 1 < len(text):
                low = ord(text[index + 1])
                if 0xDC00 <= low <= 0xDFFF:
                    result.append(chr(0x10000 + ((codepoint - 0xD800) << 10) + (low - 0xDC00)))
                    index += 2
                    continue
            result.append("\uFFFD")
        elif 0xDC00 <= codepoint <= 0xDFFF:
            result.append("\uFFFD")
        else:
            result.append(text[index])
        index += 1
    return "".join(result)


def sanitize_json_value(value: Any) -> Any:
    """Return JSON-compatible data with no unpaired UTF-16 surrogates."""
    if isinstance(value, str):
        return sanitize_text(value)
    if isinstance(value, dict):
        return {
            sanitize_text(key) if isinstance(key, str) else key: sanitize_json_value(item)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [sanitize_json_value(item) for item in value]
    if isinstance(value, tuple):
        return tuple(sanitize_json_value(item) for item in value)
    return value


def configure_stdio_utf8() -> None:
    """Make the CLI JSON protocol UTF-8 even on Windows with an ANSI locale."""
    for stream_name in ("stdin", "stdout", "stderr"):
        stream = getattr(sys, stream_name, None)
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is not None:
            try:
                reconfigure(encoding="utf-8", errors="strict")
            except (io.UnsupportedOperation, ValueError):
                # Test harnesses and embedded callers may expose non-configurable streams.
                continue


def read_json_stdin() -> dict[str, Any]:
    try:
        request = json.load(sys.stdin)
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise WorkflowError(f"Invalid JSON stdin: {exc}") from exc
    if not isinstance(request, dict):
        raise WorkflowError("JSON stdin must contain an object.")
    return sanitize_json_value(request)


def atomic_write_bytes(path: Path, payload: bytes) -> None:
    """Write beside the target, flush completely, then atomically replace it."""
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                pass


def write_json(path: Path, value: Any) -> None:
    """Atomically replace a JSON file so a failed write cannot destroy valid state."""
    payload = json.dumps(sanitize_json_value(value), ensure_ascii=False, indent=2) + "\n"
    atomic_write_bytes(path, payload.encode("utf-8"))


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def load_config() -> tuple[dict[str, Any], dict[str, Any]]:
    config = read_json(ROOT / CONFIG_FILE)
    if config is None:
        raise WorkflowError(f"{CONFIG_FILE} not found. Run scripts/bootstrap.ps1 first.")
    local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
    return config, local


def project_root(local: dict[str, Any] | None = None) -> Path:
    """Return the external documentation repository, or ROOT for legacy projects."""
    if local is None:
        local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
    raw = str(local.get("documentation_path", "")).strip()
    return Path(raw).expanduser().resolve() if raw else ROOT


def command_path(name: str) -> str | None:
    return shutil.which(name)


def normalize_git_url(url: str) -> str:
    """Return a comparable host/path identity for common Git remote syntaxes."""
    value = str(url or "").strip().replace("\\", "/")
    if not value:
        return ""
    if "://" not in value and re.match(r"^[^/@]+@[^:]+:.+$", value):
        user_host, path = value.split(":", 1)
        value = f"ssh://{user_host}/{path}"
    parsed = urllib.parse.urlsplit(value)
    if parsed.scheme and parsed.hostname:
        host = parsed.hostname.casefold()
        path = parsed.path
    else:
        host = ""
        path = value
    path = re.sub(r"^/+|/+$", "", path)
    path = re.sub(r"\.git$", "", path, flags=re.IGNORECASE)
    return f"{host}/{path}".casefold().strip("/")


def git_remote_url(path: Path) -> str:
    if not command_path("git") or not (path / ".git").exists():
        return ""
    result = subprocess.run(
        ["git", "config", "--get", "remote.origin.url"],
        cwd=path,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )
    return result.stdout.strip()


def redact_url_credentials(url: str) -> str:
    """Return a display-safe remote URL without embedded HTTP credentials."""
    value = str(url or "").strip()
    if not value or "://" not in value:
        return value
    parsed = urllib.parse.urlsplit(value)
    if parsed.username is None and parsed.password is None:
        return value
    hostname = parsed.hostname or ""
    if ":" in hostname and not hostname.startswith("["):
        hostname = f"[{hostname}]"
    netloc = hostname
    if parsed.port is not None:
        netloc += f":{parsed.port}"
    return urllib.parse.urlunsplit((parsed.scheme, netloc, parsed.path, parsed.query, parsed.fragment))


def path_is_within(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except (OSError, ValueError):
        return False


def containing_workflow_root(path: Path) -> Path | None:
    candidate = path.resolve()
    if candidate.is_file():
        candidate = candidate.parent
    for current in (candidate, *candidate.parents):
        if (current / CONFIG_FILE).is_file() and (current / "scripts" / "flow1c.py").is_file():
            return current
    return None


def contains_1c_sources(path: Path) -> bool:
    if not path.is_dir():
        return False
    try:
        return any(candidate.suffix.casefold() in {".bsl", ".xml", ".mdo"} for candidate in path.rglob("*"))
    except OSError:
        return False


def resolve_1c_source_root(path: Path) -> Path:
    """Resolve a repository checkout to its single nested 1C XML source root."""
    root = path.resolve()
    if (root / "Configuration.xml").is_file():
        return root
    candidates: list[Path] = []
    excluded = {".git", ".tools", ".venv", ".workspace", "build", "node_modules"}
    try:
        for current, directories, files in os.walk(root):
            directories[:] = [name for name in directories if name.casefold() not in excluded]
            if "Configuration.xml" in files:
                candidates.append(Path(current).resolve())
                if len(candidates) > 1:
                    return root
    except OSError:
        return root
    return candidates[0] if candidates else root


def extension_source_state(local: dict[str, Any]) -> tuple[Path | None, str, str]:
    raw = str(local.get("extension_path", "")).strip()
    extension = Path(raw) if raw else None
    mode = str(local.get("extension_mode", "")).strip().casefold()
    has_sources = bool(extension and contains_1c_sources(extension))
    if mode == "git":
        expected_url = str(local.get("extension_repository_url", "")).strip()
        actual_url = git_remote_url(extension) if extension else ""
        valid = bool(has_sources and expected_url and actual_url) and (
            normalize_git_url(expected_url) == normalize_git_url(actual_url)
        )
        return extension, "OK" if valid else "ERROR", actual_url or "Git origin is not configured"
    if mode == "local-export":
        valid = has_sources and not str(local.get("extension_repository_url", "")).strip()
        detail = "local XML/BSL export" if valid else "local export must contain XML/BSL and have no repository URL"
        return extension, "OK" if valid else "ERROR", detail
    return extension, "ERROR", "set extension_mode to git or local-export"


def rlm_health_url(endpoint: str) -> str:
    parsed = urllib.parse.urlsplit(endpoint)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return ""
    return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, "/health", "", ""))


def endpoint_is_healthy(endpoint: str, timeout: float = 2.0) -> bool:
    health_url = rlm_health_url(endpoint)
    if not health_url:
        return False
    try:
        with urllib.request.urlopen(health_url, timeout=timeout) as response:
            return 200 <= int(response.status) < 300
    except (OSError, urllib.error.URLError, ValueError):
        return False


def rlm_mcp_call(
    endpoint: str,
    name: str,
    arguments: dict[str, Any],
    *,
    timeout: float = 5.0,
) -> tuple[dict[str, Any] | None, str]:
    """Call a stateless RLM MCP tool and decode its structured JSON result."""
    parsed = urllib.parse.urlsplit(endpoint)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc or parsed.path.rstrip("/") != "/mcp":
        return None, "invalid MCP endpoint"
    payload = json.dumps(
        {
            "jsonrpc": "2.0",
            "id": 1,
            "method": "tools/call",
            "params": {"name": name, "arguments": arguments},
        },
        ensure_ascii=False,
    ).encode("utf-8")
    request = urllib.request.Request(
        endpoint,
        data=payload,
        method="POST",
        headers={
            "Accept": "application/json, text/event-stream",
            "Content-Type": "application/json; charset=utf-8",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            body = response.read().decode("utf-8")
    except (OSError, UnicodeError, urllib.error.URLError, ValueError) as exc:
        return None, str(exc)

    data_lines = [line[6:] for line in body.splitlines() if line.startswith("data: ")]
    if not data_lines:
        return None, "MCP response contains no data event"
    try:
        envelope = json.loads(data_lines[-1])
        if envelope.get("error"):
            return None, str(envelope["error"].get("message") or envelope["error"])
        result = envelope["result"]
        raw = result.get("structuredContent", {}).get("result")
        if raw is None:
            content = result.get("content", [])
            raw = next((item.get("text") for item in content if item.get("type") == "text"), None)
        decoded = json.loads(raw) if isinstance(raw, str) else raw
        if not isinstance(decoded, dict):
            return None, "MCP tool returned a non-object result"
        return decoded, ""
    except (KeyError, TypeError, json.JSONDecodeError) as exc:
        return None, f"invalid MCP response: {exc}"


def rlm_endpoint_index_state(endpoint: str, source: Path) -> tuple[str, str]:
    result, error = rlm_mcp_call(endpoint, "rlm_index", {"action": "info", "path": str(source)})
    if error:
        return "ERROR", error
    assert result is not None
    if result.get("error"):
        return "ERROR", str(result["error"])
    modules = result.get("modules")
    if isinstance(modules, int):
        return "OK", f"available ({modules} modules)"
    return "ERROR", "endpoint did not return index information"


def rlm_session_execute(
    endpoint: str,
    source: Path,
    query: str,
    code: str,
    *,
    effort: str,
    max_output_chars: int,
    timeout: float,
    domains: list[str],
) -> tuple[dict[str, Any] | None, str]:
    """Open an RLM session, execute sandbox code, and always close the session."""
    start_args = {
        "query": query,
        "path": str(source),
        "effort": effort,
        "max_output_chars": max_output_chars,
    }
    start_result, start_error = rlm_mcp_call(
        endpoint,
        "rlm_start",
        start_args,
        timeout=timeout,
    )
    domain_error = start_error or str(start_result.get("error", "") if start_result else "")
    if "domains" in domain_error and ("явный выбор" in domain_error or "required" in domain_error):
        start_result, start_error = rlm_mcp_call(
            endpoint, "rlm_start", {**start_args, "domains": domains}, timeout=timeout,
        )
    if start_error:
        return None, start_error
    if start_result is None:
        return None, "RLM session start returned no result"
    if start_result.get("error"):
        return None, str(start_result["error"])

    session_id = str(start_result.get("session_id", "")).strip()
    if not session_id:
        return None, "endpoint did not create a sandbox session"

    execute_result: dict[str, Any] | None = None
    execute_error = ""
    try:
        execute_result, execute_error = rlm_mcp_call(
            endpoint,
            "rlm_execute",
            {"session_id": session_id, "code": code},
            timeout=timeout,
        )
    finally:
        end_result, end_error = rlm_mcp_call(
            endpoint,
            "rlm_end",
            {"session_id": session_id},
            timeout=10.0,
        )

    if execute_error:
        return None, execute_error
    if execute_result is None:
        return None, "RLM execution returned no result"
    if execute_result.get("error"):
        return None, str(execute_result["error"])
    if end_error:
        return None, f"query executed but session could not be closed: {end_error}"
    if end_result is None:
        return None, "query executed but session close returned no result"
    if end_result.get("error"):
        return None, f"query executed but session could not be closed: {end_result['error']}"
    return execute_result, ""


def rlm_endpoint_session_state(endpoint: str, source: Path) -> tuple[str, str]:
    marker = "FLOW1C_WORKFLOW_RLM_OK"
    result, error = rlm_session_execute(
        endpoint,
        source,
        "Flow1C health check",
        f"print({marker!r})",
        effort="low",
        max_output_chars=1000,
        timeout=10.0,
        domains=[],
    )
    if error:
        return "ERROR", error
    assert result is not None
    if marker not in str(result.get("stdout", "")):
        return "ERROR", "sandbox execution did not return the smoke-test marker"
    return "OK", "sandbox execution succeeded and session closed"


def rlm_index_state(index_command: str, source: Path) -> tuple[str, str]:
    if not index_command or not Path(index_command).is_file():
        return "ERROR", "rlm-bsl-index is not installed"
    # Receipts are issued only after a successful index job and strict check.
    # Reuse the index adapter's policy rather than treating calendar age as
    # content corruption or maintaining a separate doctor exception.
    try:
        from rlm_index_runtime import IndexManager
    except ModuleNotFoundError:
        from scripts.rlm_index_runtime import IndexManager
    try:
        manager = IndexManager(ROOT, source, command=[index_command])
        if manager.receipt_path.exists() or manager.job_path.exists():
            state = manager.status()
            return ("OK" if state["state"] == "FRESH" else "ERROR"), state["detail"]
    except Exception as exc:
        return "ERROR", f"Index validation failed: {exc}"
    try:
        result = subprocess.run(
            [index_command, "index", "info", str(source)],
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
            check=False,
            timeout=120,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return "ERROR", str(exc)
    output = "\n".join(part for part in (result.stdout.strip(), result.stderr.strip()) if part)
    match = re.search(r"^\s*Status:\s*([^\r\n]+)", output, flags=re.IGNORECASE | re.MULTILINE)
    status = match.group(1).strip() if match else ""
    if result.returncode == 0 and status.casefold() == "fresh":
        return "OK", status
    detail = status or (output.splitlines()[0] if output else f"exit code {result.returncode}")
    return "ERROR", detail


def run_git(args: list[str], *, check: bool = True, capture: bool = True) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        ["git", *args],
        cwd=project_root(),
        check=check,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=capture,
    )


def in_git_repository() -> bool:
    try:
        return run_git(["rev-parse", "--is-inside-work-tree"]).stdout.strip() == "true"
    except (FileNotFoundError, subprocess.CalledProcessError):
        return False


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def header_map_from_values(values: Iterable[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    for column, value in enumerate(values, start=1):
        key = normalized_header(value)
        if key and key not in result:
            result[key] = column
    return result


def header_map(sheet: Any, row: int) -> dict[str, int]:
    values = next(sheet.iter_rows(min_row=row, max_row=row, values_only=True), ())
    return header_map_from_values(values)


def find_column(headers: dict[str, int], *candidates: str) -> int | None:
    for candidate in candidates:
        found = headers.get(normalized_header(candidate))
        if found:
            return found
    return None


def normalized_cell_value(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return value


def row_cell_value(values: tuple[Any, ...], column: int | None) -> Any:
    if not column:
        return None
    index = column - 1
    return normalized_cell_value(values[index]) if index < len(values) else None


def cell_value(sheet: Any, row: int, column: int | None) -> Any:
    """Compatibility helper for callers that are not streaming a read-only sheet."""
    if not column:
        return None
    return normalized_cell_value(sheet.cell(row, column).value)


def split_ids(value: Any) -> list[str]:
    if value is None:
        return []
    return [item.strip().upper() for item in re.split(r"[,;\n]+", str(value)) if item.strip()]


def make_diff(previous: dict[str, Any], current: dict[str, Any]) -> dict[str, list[str]]:
    old_keys = set(previous)
    new_keys = set(current)
    return {
        "added": sorted(new_keys - old_keys),
        "changed": sorted(key for key in old_keys & new_keys if previous[key] != current[key]),
        "missing_from_source": sorted(old_keys - new_keys),
    }


def emit_doctor_report(rows: list[tuple[str, str, str]], *, json_output: bool = False) -> int:
    ready = not any(state == "ERROR" for _, state, _ in rows)
    if json_output:
        print(json.dumps({
            "schema_version": 1,
            "state": "READY" if ready else "BLOCKED",
            "ready": ready,
            "checks": [
                {"name": name, "state": state, "detail": detail}
                for name, state, detail in rows
            ],
        }, ensure_ascii=False, indent=2))
    else:
        width = max(len(row[0]) for row in rows)
        for name, state, detail in rows:
            print(f"{name:<{width}}  {state:<8}  {detail}")
    return 0 if ready else 1


def load_capabilities() -> dict[str, Any]:
    value = read_json(ROOT / "config" / "capabilities.json")
    try:
        return load_capability_model(value)
    except ValueError as exc:
        raise WorkflowError(f"Invalid capability configuration: {exc}") from exc


def _contains_secret(value: Any, path: str = "") -> bool:
    if isinstance(value, dict):
        for key, child in value.items():
            current = f"{path}.{key}" if path else str(key)
            if any(marker in str(key).casefold() for marker in ("token", "secret", "password", "credential", "api_key")):
                return True
            if _contains_secret(child, current):
                return True
    elif isinstance(value, list):
        return any(_contains_secret(child, path) for child in value)
    return False


def setup_checkpoint_path(setup_id: str) -> Path:
    if not re.fullmatch(r"[A-Za-z0-9-]{8,64}", str(setup_id or "")):
        raise WorkflowError("Invalid setup ID")
    return ROOT / ".workspace" / "setup" / f"{setup_id}.json"


def save_setup_checkpoint(checkpoint: dict[str, Any]) -> None:
    if _contains_secret(checkpoint):
        raise WorkflowError("Setup checkpoints must not contain tokens, passwords or secrets")
    if checkpoint.get("schema_version") != 1:
        raise WorkflowError("Unsupported setup checkpoint schema")
    checkpoint["updated_at"] = utc_now()
    write_json(setup_checkpoint_path(str(checkpoint.get("setup_id", ""))), checkpoint)


def create_setup_checkpoint(profile: str, confirmed_values: dict[str, Any], setup_id: str = "") -> dict[str, Any]:
    setup_id = setup_id or str(uuid.uuid4())
    path = setup_checkpoint_path(setup_id)
    existing = read_json(path, {})
    if isinstance(existing, dict) and existing.get("schema_version") == 1:
        existing["confirmed_values"] = {**existing.get("confirmed_values", {}), **confirmed_values}
        existing["requested_profile"] = profile
        save_setup_checkpoint(existing)
        return existing
    now = utc_now()
    checkpoint = {"schema_version": 1, "setup_id": setup_id, "requested_profile": profile,
                  "state": "IN_PROGRESS", "confirmed_values": confirmed_values,
                  "completed_steps": [], "pending_jobs": [], "errors": [],
                  "created_at": now, "updated_at": now}
    save_setup_checkpoint(checkpoint)
    return checkpoint


def _module_check(module: str, label: str) -> tuple[str, str, str]:
    try:
        __import__(module)
        return label, "OK", "installed"
    except ImportError:
        return label, "ERROR", "not installed; run bootstrap.ps1 with the required profile"


def _capability_checks(name: str, model: dict[str, Any]) -> list[tuple[str, str, str]]:
    """Probe only the requested capability. Dependencies are resolved by policy."""
    if name == "core":
        version = sys.version_info[:2]
        python_ok = python_version_supported(version, model)
        stages_ok = False
        try:
            stages = load_stages()
            stages_ok = bool(stages.get("operations"))
        except WorkflowError:
            pass
        skills = list((ROOT / ".agents" / "skills").glob("flow1c-*/SKILL.md"))
        policy_state, policy_detail = git_read_policy_check()
        return [
            ("Python", "OK" if python_ok else "ERROR", f"{sys.version.split()[0]}; supported >= {model['python']['minimum']}"),
            (".venv", "OK" if (ROOT / ".venv").is_dir() else "ERROR", str(ROOT / ".venv")),
            (CONFIG_FILE, "OK" if (ROOT / CONFIG_FILE).is_file() else "ERROR", str(ROOT / CONFIG_FILE)),
            ("Workflow stages", "OK" if stages_ok else "ERROR", str(ROOT / "config" / "stages.json")),
            ("Built-in flow1c skills", "OK" if skills else "ERROR", f"{len(skills)} discovered"),
            ("Git read policy", policy_state, policy_detail),
        ]
    if name == "git":
        git = command_path("git")
        return [("Git", "OK" if git else "ERROR", git or "not found")]
    if name == "project":
        rows: list[tuple[str, str, str]] = []
        local_path = ROOT / LOCAL_CONFIG_FILE
        if not local_path.is_file():
            return [(LOCAL_CONFIG_FILE, "ERROR", "not configured")]
        try:
            _, local = load_config()
        except WorkflowError as exc:
            return [(LOCAL_CONFIG_FILE, "ERROR", str(exc))]
        setup = local.get("setup", {}) if isinstance(local.get("setup"), dict) else {}
        configured_root = str(setup.get("workflow_root", "")).strip()
        validated = setup.get("validation_version") == SETUP_VALIDATION_VERSION and bool(configured_root) and Path(configured_root).resolve() == ROOT.resolve()
        rows.append(("Setup confirmation", "OK" if validated else "ERROR", "validated for this checkout" if validated else "run setup-resume or configure-project.ps1"))
        for key in ("documentation_path", "configuration_path", "extension_path"):
            raw = str(local.get(key, "")).strip()
            path = Path(raw) if raw else None
            valid = bool(path and path.exists() and not path_is_within(path, ROOT) and containing_workflow_root(path) is None)
            rows.append((key, "OK" if valid else "ERROR", raw or "not configured"))
        documentation = project_root(local)
        rows.append(("Docs Git repo", "OK" if (documentation / ".git").exists() else "ERROR", str(documentation)))
        _, state, detail = extension_source_state(local)
        rows.append(("Extension source", state, redact_url_credentials(detail)))
        return rows
    if name == "template-core":
        return [_module_check("jsonschema", "jsonschema")]
    if name == "template-markdown":
        return [_module_check("markdown_it", "markdown-it-py")]
    if name == "template-docx":
        return [_module_check("docx", "python-docx")]
    if name == "office":
        return [_module_check("openpyxl", "openpyxl"), _module_check("markitdown", "MarkItDown"),
                _module_check("docx", "python-docx"), _module_check("docxtpl", "docxtpl")]
    if name == "rlm":
        rows: list[tuple[str, str, str]] = []
        try:
            manifest = load_external_tools_manifest(ROOT / "config" / "external-tools.json")
            version = package_version("rlm-tools-bsl")
            requirement = manifest["rlm_tools_bsl"]["requirement"]
            compatible = bool(version and version_satisfies(version, requirement))
        except Exception as exc:
            return [("rlm-tools-bsl", "ERROR", str(exc))]
        rows.append(("rlm-tools-bsl", "OK" if compatible else "ERROR", f"{version or 'not installed'}; required {requirement}"))
        index = ROOT / ".venv" / "Scripts" / "rlm-bsl-index.exe"
        rows.append(("rlm-bsl-index", "OK" if index.is_file() else "ERROR", str(index)))
        local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
        config = read_json(ROOT / CONFIG_FILE, {})
        endpoint = str((local.get("rlm") or {}).get("endpoint") or (config.get("quality") or {}).get("rlm_endpoint") or "")
        healthy = endpoint_is_healthy(endpoint)
        rows.append(("RLM MCP endpoint", "OK" if healthy else "ERROR", endpoint or "not configured"))
        index_command = str((local.get("rlm") or {}).get("index_command", ""))
        for label, raw in (("configuration", local.get("configuration_path")), ("extension", local.get("extension_path"))):
            source = resolve_1c_source_root(Path(str(raw))) if raw else None
            state, detail = rlm_index_state(index_command, source) if source else ("ERROR", "source is not configured")
            rows.append((f"RLM {label} index", state, detail))
        return rows
    if name == "bsl-ls":
        local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
        bsl = local.get("bsl_language_server", {}) if isinstance(local.get("bsl_language_server"), dict) else {}
        executable = Path(str(bsl.get("command", "")))
        if not executable.is_file():
            return [("BSL LS", "ERROR", "not installed or configured")]
        smoke_ok, detail = bsl_smoke(executable)
        return [("BSL LS", "OK" if smoke_ok else "ERROR", detail)]
    if name == "cc-1c-skills":
        repository = ROOT / ".tools" / "cc-1c-skills"
        commit = ""
        if (repository / ".git").is_dir():
            result = subprocess.run(["git", "-C", str(repository), "rev-parse", "HEAD"], text=True, capture_output=True, check=False)
            commit = result.stdout.strip() if result.returncode == 0 else ""
        ok = bool(commit) and skills_connected(ROOT)
        return [("cc-1c-skills", "OK" if ok else "ERROR", commit or "not installed; explicit opt-in required")]
    return [(name, "ERROR", "unknown capability")]


def cmd_profile_doctor(profile: str, *, json_output: bool) -> int:
    model = load_capabilities()
    try:
        required = resolve_capabilities(model, profile)
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    capabilities: dict[str, Any] = {}
    all_rows: list[tuple[str, str, str]] = []
    states: dict[str, str] = {}
    optional_missing: list[str] = []
    for name, definition in model["capabilities"].items():
        if name not in required:
            optional_present = bool(
                name == "cc-1c-skills"
                and (ROOT / ".tools" / "cc-1c-skills" / ".git").is_dir()
                and skills_connected(ROOT)
            )
            state = ("OPTIONAL" if optional_present else "OPTIONAL_MISSING") if definition.get("optional") else "NOT_REQUESTED"
            checks: list[tuple[str, str, str]] = []
            if state == "OPTIONAL_MISSING":
                optional_missing.append(name)
        else:
            checks = _capability_checks(name, model)
            state = "READY" if all(row[1] != "ERROR" for row in checks) else "BLOCKED"
            all_rows.extend(checks)
        states[name] = state
        capabilities[name] = {"state": state, "checks": [
            {"name": label, "state": check_state, "detail": detail} for label, check_state, detail in checks
        ]}
    ready = capability_readiness(required, states)
    blocking = [item for item in all_rows if item[1] == "ERROR"]
    next_actions: list[str] = []
    blocked_names = {name for name in required if states.get(name) != "READY"}
    installable = blocked_names & {"core", "office", "rlm", "bsl-ls", "cc-1c-skills", "template-core", "template-markdown", "template-docx"}
    if installable:
        suffix = " -InstallCc1cSkills" if "cc-1c-skills" in installable else ""
        next_actions.append(f"bootstrap.ps1 -Profile {profile}{suffix}")
    if "project" in blocked_names:
        next_actions.append("resume setup or run configure-project.ps1 with confirmed project values")
    if "git" in blocked_names:
        next_actions.append("install Git or open a valid Git checkout")
    result = {
        "schema_version": 1,
        "state": "READY" if ready else "BLOCKED",
        "ready": ready,
        "checks": [{"name": n, "state": s, "detail": d} for n, s, d in all_rows],
        "requested_profile": profile,
        "required_capabilities": required,
        "capabilities": capabilities,
        "blocking_checks": [{"name": n, "state": s, "detail": d} for n, s, d in blocking],
        "optional_missing": optional_missing,
        "next_actions": next_actions,
    }
    if json_output:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        for item in result["checks"]:
            print(f"{item['name']}  {item['state']}  {item['detail']}")
    return 0 if ready else 1


def git_read_policy_check() -> tuple[str, str]:
    policy_path = ROOT / "config" / "git-read-policy.json"
    try:
        policy = read_json(policy_path)
        if not isinstance(policy, dict) or policy.get("schema_version") != 1:
            return "ERROR", "git read policy schema must be 1"
        commands = policy.get("commands", {})
        git_commands = set(commands.get("git", [])) if isinstance(commands, dict) else set()
        denied = set(policy.get("denied_git_subcommands", []))
        if not git_commands or git_commands & denied:
            return "ERROR", "Git read allowlist is empty or overlaps the mutation denylist"
        limits = policy.get("limits", {})
        if not all(isinstance(limits.get(key), int) and limits[key] > 0
                   for key in ("pipeline_segments", "timeout_seconds", "max_output_chars", "max_processes")):
            return "ERROR", "Git read policy resource limits are invalid"
        guard = ROOT / ".opencode" / "plugins" / "flow1c-guard.js"
        if not guard.is_file() or "validateSegment" not in guard.read_text(encoding="utf-8"):
            return "ERROR", "OpenCode Git guard validator is unavailable"
        return "OK", "schema 1; guarded read-only commands"
    except (OSError, WorkflowError) as exc:
        return "ERROR", str(exc)


def cmd_doctor(args: argparse.Namespace) -> int:
    requested_profile = getattr(args, "profile", None)
    requested_operation = getattr(args, "operation", None)
    template_gate = load_gate(args.gate_id) if getattr(args, "gate_id", None) else None
    if template_gate and not requested_operation:
        requested_operation = template_gate.get("operation")
    if requested_operation == "interview-preparation":
        selected_format = getattr(args, "format", None) or "xlsx"
        if selected_format not in {"xlsx", "markdown"}:
            raise InterviewError("INTERVIEW_INVALID_INPUT", "Сценарий интервью поддерживает XLSX или Markdown.")
        checks = [
            ("Python", "OK" if sys.version_info >= (3, 10) else "ERROR", sys.version.split()[0]),
            ("flow1c-interview-preparation", "OK" if (ROOT / ".agents/skills/flow1c-interview-preparation/SKILL.md").is_file() else "ERROR", "canonical skill"),
        ]
        if selected_format == "xlsx":
            checks.append(_module_check("openpyxl", "openpyxl"))
        ready = all(state != "ERROR" for _, state, _ in checks)
        result = {"schema_version": 1, "operation": requested_operation, "format": selected_format, "ready": ready,
                  "checks": [{"name": name, "state": state, "detail": detail} for name, state, detail in checks],
                  "next_actions": [] if ready else ["Restore the canonical skill or run the consented bootstrap.ps1 -Profile documents. No project sources or RLM are required."]}
        if getattr(args, "json", False):
            print(json.dumps(result, ensure_ascii=False, indent=2))
        else:
            for name, state, detail in checks:
                print(f"{name}  {state}  {detail}")
        return 0 if ready else 1
    if requested_operation in {"template-management", "template-document"}:
        pin_format = (template_gate or {}).get("template_pin", {}).get("source_format")
        selected_format = getattr(args, "format", None) or pin_format or "markdown"
        result = template_readiness(ROOT, read_json(ROOT / LOCAL_CONFIG_FILE, {}), selected_format, requested_operation)
        result["format_source"] = "explicit" if getattr(args, "format", None) else "gate-pin" if pin_format else "default"
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result["ready"] else 1
    if requested_operation:
        try:
            requested_profile = operation_profile(load_stages(), requested_operation)
        except ValueError as exc:
            raise WorkflowError(str(exc)) from exc
    if requested_profile:
        return cmd_profile_doctor(str(requested_profile), json_output=bool(getattr(args, "json", False)))
    rows: list[tuple[str, str, str]] = []
    policy_state, policy_detail = git_read_policy_check()
    rows.append(("Git read policy", policy_state, policy_detail))
    external_manifest: dict[str, Any] | None = None
    try:
        external_manifest = load_external_tools_manifest(ROOT / "config" / "external-tools.json")
        rows.append(("External tools manifest", "OK", "schema 1"))
    except Exception as exc:
        rows.append(("External tools manifest", "ERROR", str(exc)))
    rows.append(("Python", "OK" if sys.version_info[:2] >= (3, 10) else "ERROR", sys.version.split()[0] + "; supported >=3.10"))
    rows.append(("Git", "OK" if command_path("git") else "ERROR", command_path("git") or "not found"))
    try:
        import openpyxl  # noqa: F401
        rows.append(("openpyxl", "OK", "installed"))
    except ImportError:
        rows.append(("openpyxl", "ERROR", "run bootstrap.ps1"))

    try:
        import markitdown  # noqa: F401
        rows.append(("MarkItDown", "OK", "installed"))
    except ImportError:
        rows.append(("MarkItDown", "ERROR", "run bootstrap.ps1"))

    try:
        import docx  # noqa: F401
        rows.append(("python-docx", "OK", "installed"))
    except ImportError:
        rows.append(("python-docx", "ERROR", "run bootstrap.ps1"))

    try:
        import docxtpl  # noqa: F401
        rows.append(("docxtpl", "OK", "installed"))
    except ImportError:
        rows.append(("docxtpl", "ERROR", "run bootstrap.ps1"))

    try:
        import rlm_tools_bsl  # noqa: F401
        installed_rlm = package_version("rlm-tools-bsl")
        requirement = external_manifest["rlm_tools_bsl"]["requirement"] if external_manifest else ""
        compatible = bool(installed_rlm and external_manifest and version_satisfies(installed_rlm, requirement))
        rows.append(("rlm-tools-bsl", "OK" if compatible else "ERROR", f"{installed_rlm or 'unknown'}; required {requirement or 'valid manifest'}"))
    except ImportError:
        rows.append(("rlm-tools-bsl", "ERROR", "run bootstrap.ps1"))
    index_executable = ROOT / ".venv" / "Scripts" / "rlm-bsl-index.exe"
    rows.append(("rlm-bsl-index", "OK" if index_executable.is_file() else "ERROR", str(index_executable)))

    rows.append(("Java", "OK" if command_path("java") else "OPTIONAL", command_path("java") or "not found"))

    config_path = ROOT / CONFIG_FILE
    local_path = ROOT / LOCAL_CONFIG_FILE
    rows.append((CONFIG_FILE, "OK" if config_path.exists() else "ERROR", str(config_path)))
    rows.append((LOCAL_CONFIG_FILE, "OK" if local_path.exists() else "ERROR", str(local_path)))

    if config_path.exists():
        config, local = load_config()
        local_schema = local.get("schema_version")
        rows.append((
            "Local config schema",
            "OK" if local_schema == LOCAL_CONFIG_SCHEMA_VERSION else "ERROR",
            f"schema {local_schema}" if local_schema is not None else "run scripts/update.ps1 or /setup",
        ))
        setup = local.get("setup", {}) if isinstance(local.get("setup", {}), dict) else {}
        configured_root = str(setup.get("workflow_root", "")).strip()
        root_matches = bool(configured_root) and Path(configured_root).resolve() == ROOT.resolve()
        validation_ok = setup.get("validation_version") == SETUP_VALIDATION_VERSION and root_matches
        rows.append((
            "Setup confirmation",
            "OK" if validation_ok else "ERROR",
            "validated for this checkout" if validation_ok else "rerun configure-project.ps1 and confirm every value",
        ))

        configured_workflow_url = str(setup.get("workflow_repository_url", "")).strip()
        actual_workflow_url = git_remote_url(ROOT)
        workflow_url_ok = bool(configured_workflow_url and actual_workflow_url) and (
            normalize_git_url(configured_workflow_url) == normalize_git_url(actual_workflow_url)
        )
        rows.append((
            "Workflow checkout",
            "OK" if workflow_url_ok else "ERROR",
            redact_url_credentials(actual_workflow_url) or "origin is not configured",
        ))

        for key in ("documentation_path", "configuration_path", "extension_path"):
            raw = str(local.get(key, "")).strip()
            path = Path(raw) if raw else None
            exists = bool(path and path.exists())
            outside_workflow = bool(path and not path_is_within(path, ROOT))
            other_workflow = containing_workflow_root(path) if exists and path else None
            belongs_to_other_checkout = bool(other_workflow and other_workflow.resolve() != ROOT.resolve())
            state = "OK" if exists and outside_workflow and not belongs_to_other_checkout else "ERROR"
            detail = raw or "not configured"
            if exists and not outside_workflow:
                detail += " (must be outside the Flow1C checkout)"
            elif belongs_to_other_checkout:
                detail += f" (belongs to another Flow1C checkout: {other_workflow})"
            rows.append((key, state, detail))
        template_raw = str(local.get("functional_spec_template", "")).strip()
        if not template_raw:
            rows.append((
                "functional_spec_template",
                "OPTIONAL",
                "not configured; required only for formal functional-spec generation",
            ))
        else:
            template_path = Path(template_raw)
            exists = template_path.is_file()
            outside_workflow = not path_is_within(template_path, ROOT)
            other_workflow = containing_workflow_root(template_path) if exists else None
            belongs_to_other_checkout = bool(other_workflow and other_workflow.resolve() != ROOT.resolve())
            is_docx = template_path.suffix.lower() == ".docx"
            template_ok = exists and outside_workflow and not belongs_to_other_checkout and is_docx
            detail = template_raw
            if exists and not outside_workflow:
                detail += " (must be outside the Flow1C checkout)"
            elif belongs_to_other_checkout:
                detail += f" (belongs to another Flow1C checkout: {other_workflow})"
            elif exists and not is_docx:
                detail += " (must be a DOCX file)"
            if not template_ok:
                detail += "; ignored until formal functional-spec generation"
            rows.append(("functional_spec_template", "OK" if template_ok else "OPTIONAL", detail))
        documentation = project_root(local)
        extension, extension_state, extension_detail = extension_source_state(local)
        rows.append(("Docs Git repo", "OK" if (documentation / ".git").exists() else "ERROR", str(documentation)))
        documentation_url = str(local.get("documentation_repository_url", "")).strip()
        actual_documentation_url = git_remote_url(documentation)
        docs_initialized_locally = bool(setup.get("documentation_repository_initialized_locally"))
        docs_remote_matches = bool(documentation_url and actual_documentation_url) and (
            normalize_git_url(documentation_url) == normalize_git_url(actual_documentation_url)
        )
        docs_remote_ok = docs_remote_matches or (
            docs_initialized_locally and not documentation_url and not actual_documentation_url
        )
        rows.append((
            "Docs Git remote",
            "OK" if docs_remote_ok else "ERROR",
            redact_url_credentials(actual_documentation_url) or (
                "local repository initialization explicitly confirmed"
                if docs_initialized_locally and not documentation_url
                else "origin is not configured"
            ),
        ))
        shared_confirmed = bool(setup.get("allow_shared_workflow_documentation_repository"))
        repositories_differ = (
            docs_initialized_locally
            and documentation.resolve() != ROOT.resolve()
        ) or (
            bool(actual_workflow_url and actual_documentation_url)
            and normalize_git_url(actual_workflow_url) != normalize_git_url(actual_documentation_url)
        )
        rows.append((
            "Workflow/docs separation",
            "OK" if repositories_differ or shared_confirmed else "ERROR",
            "separate repositories" if repositories_differ else (
                "same repository explicitly confirmed" if shared_confirmed else "workflow and documentation repositories are the same"
            ),
        ))

        rows.append(("Extension source", extension_state, redact_url_credentials(extension_detail)))
        git_name = ""
        git_email = ""
        if command_path("git") and (documentation / ".git").exists():
            name_result = subprocess.run(
                ["git", "config", "user.name"], cwd=documentation, text=True, capture_output=True, check=False
            )
            email_result = subprocess.run(
                ["git", "config", "user.email"], cwd=documentation, text=True, capture_output=True, check=False
            )
            git_name = name_result.stdout.strip()
            git_email = email_result.stdout.strip()
        identity = f"{git_name} <{git_email}>" if git_name and git_email else "configure user.name and user.email"
        rows.append(("Git identity", "OK" if git_name and git_email else "ERROR", identity))
        bsl = local.get("bsl_language_server", {})
        command = str(bsl.get("command", "")).strip() if isinstance(bsl, dict) else ""
        legacy_jar = str(config.get("quality", {}).get("bsl_language_server_jar", "")).strip()
        executable = command or legacy_jar
        state = "OK" if executable and Path(executable).is_file() else "ERROR"
        rows.append(("BSL LS", state, executable or "not configured"))
        if state == "OK" and external_manifest:
            expected_version = external_manifest["bsl_language_server"]["version"]
            configured_version = str(bsl.get("version", "")) if isinstance(bsl, dict) else ""
            smoke_ok, smoke_detail = bsl_smoke(Path(executable))
            rows.append((
                "BSL LS manifest/smoke",
                "OK" if configured_version == expected_version and smoke_ok else "ERROR",
                f"configured {configured_version or 'unknown'}, expected {expected_version}; {smoke_detail}",
            ))

        skills_repository = ROOT / ".tools" / "cc-1c-skills"
        skills_commit = ""
        if (skills_repository / ".git").is_dir():
            commit_result = subprocess.run(
                ["git", "-C", str(skills_repository), "rev-parse", "HEAD"], text=True, capture_output=True, check=False
            )
            skills_commit = commit_result.stdout.strip() if commit_result.returncode == 0 else ""
        skills_ok = bool(skills_commit) and skills_connected(ROOT)
        rows.append((
            "cc-1c-skills",
            "OK" if skills_ok else "OPTIONAL",
            skills_commit if skills_ok else "repository/commit or connected .agents skills are missing",
        ))

        quality = config.get("quality", {}) if isinstance(config.get("quality", {}), dict) else {}
        rlm_endpoint = str(quality.get("rlm_endpoint", "")).strip()
        endpoint_healthy = endpoint_is_healthy(rlm_endpoint)
        rows.append((
            "RLM MCP endpoint",
            "OK" if endpoint_healthy else "ERROR",
            rlm_endpoint or "not configured",
        ))
        rlm = local.get("rlm", {}) if isinstance(local.get("rlm", {}), dict) else {}
        index_command = str(rlm.get("index_command", "")).strip()
        configuration_raw = str(local.get("configuration_path", "")).strip()
        configuration_source = resolve_1c_source_root(Path(configuration_raw)) if configuration_raw else None
        extension_source = resolve_1c_source_root(extension) if extension else None
        rlm_sources = (
            ("RLM configuration index", configuration_source),
            ("RLM extension index", extension_source),
        )
        endpoint_indexes_ok = endpoint_healthy
        for label, source in rlm_sources:
            index_state, index_detail = (
                rlm_index_state(index_command, source) if source else ("ERROR", "source is not configured")
            )
            rows.append((label, index_state, index_detail))
            endpoint_state, endpoint_detail = (
                rlm_endpoint_index_state(rlm_endpoint, source)
                if endpoint_healthy and source
                else ("ERROR", "endpoint or source is unavailable")
            )
            rows.append((f"{label} via MCP", endpoint_state, endpoint_detail))
            endpoint_indexes_ok = endpoint_indexes_ok and endpoint_state == "OK"
        session_state, session_detail = (
            rlm_endpoint_session_state(rlm_endpoint, configuration_source)
            if endpoint_indexes_ok and configuration_source
            else ("ERROR", "endpoint indexes are unavailable")
        )
        rows.append(("RLM MCP sandbox session", session_state, session_detail))
        gitea = {**config.get("gitea", {}), **local.get("gitea", {})}
        token_env = str(gitea.get("token_env", "FLOW1C_GITEA_TOKEN"))
        token_state = "OK" if os.environ.get(token_env) else "SETUP"
        rows.append(("Gitea token", token_state, token_env))

    return emit_doctor_report(rows, json_output=bool(getattr(args, "json", False)))


def cmd_registry_import(args: argparse.Namespace) -> int:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import column_index_from_string
    except ImportError as exc:
        raise WorkflowError("openpyxl is required. Run scripts/bootstrap.ps1.") from exc

    config, _ = load_config()
    registry = config.get("registry", {})
    source = Path(args.file).expanduser().resolve()
    if not source.is_file():
        raise WorkflowError(f"Registry file not found: {source}")

    report_path = project_root() / "registry" / "import-report.md"

    def fail_import(kind: str, message: str) -> int:
        write_text(report_path, "\n".join(["# Registry import report", "", f"Source: `{source.name}`", "",
                                                   "## Validation", "", f"- ERROR: {message}"]))
        write_json(report_path.parent / "status.json", {"status": "invalid", "source_sha256": sha256(source),
                                                        "report_path": str(report_path), "updated_at": utc_now()})
        print(json.dumps({"state": "BLOCKED", "report_path": str(report_path),
                          "errors": [{"kind": kind, "message": message}], "warnings": []},
                         ensure_ascii=False, indent=2))
        return 2

    try:
        workbook = load_workbook(source, data_only=False, read_only=True)
    except Exception as exc:  # openpyxl exposes several backend-specific corrupt-workbook exceptions
        return fail_import("invalid_workbook", f"Cannot open registry workbook: {exc}")
    req_sheet_name = registry.get("requirements_sheet", "Процессы требования")
    fs_sheet_name = registry.get("specifications_sheet", "Реестр ФС- не удалять")
    fatal_errors: list[str] = []
    errors: list[str] = []
    warnings: list[str] = []

    missing_sheets = [name for name in (req_sheet_name, fs_sheet_name) if name not in workbook.sheetnames]
    if missing_sheets:
        workbook.close()
        return fail_import("missing_sheet", "Registry sheets not found: " + ", ".join(missing_sheets))

    req_sheet = workbook[req_sheet_name]
    req_header_row = int(registry.get("requirements_header_row", 3))
    req_data_row = int(registry.get("requirements_data_row", req_header_row + 1))
    headers = header_map(req_sheet, req_header_row)
    id_column_letter = str(registry.get("requirement_id_column", "J")).upper()
    id_column = column_index_from_string(id_column_letter)
    req_header_values = next(
        req_sheet.iter_rows(min_row=req_header_row, max_row=req_header_row, values_only=True),
        (),
    )
    id_header_value = row_cell_value(req_header_values, id_column)
    id_header = normalized_header(id_header_value)
    if id_header not in {"idтребования", "id требования"}:
        fatal_errors.append(
            f"Column {id_column_letter} must contain the requirement ID header; found "
            f"'{id_header_value}'."
        )

    requirement_text_column = find_column(headers, "Требование")
    process_columns = {
        "code": find_column(headers, "Код БП"),
        "level_1": find_column(headers, "БП Уровень 1"),
        "level_2": find_column(headers, "БП Уровень 2"),
        "level_3": find_column(headers, "БП Уровень 3"),
        "block": find_column(headers, "Block", "Block for status"),
    }
    optional_columns = {
        "source": find_column(headers, "Источник"),
        "registered_at": find_column(headers, "Дата регистрации"),
        "owner": find_column(headers, "Владелец процесса"),
        "analyst": find_column(headers, "Аналитик/ФА (Ах)"),
        "criticality": find_column(headers, "Критичность"),
        "fit_gap": find_column(headers, "Покрывается стандартной функциональностью"),
        "fit_gap_comment": find_column(headers, "Комментарий к fit/gap анализу"),
    }

    requirements: dict[str, Any] = {}
    requirement_occurrences: dict[str, list[dict[str, Any]]] = {}
    duplicate_requirements: list[tuple[str, int]] = []
    req_columns = [
        id_column,
        requirement_text_column,
        *process_columns.values(),
        *optional_columns.values(),
    ]
    req_max_column = max(column for column in req_columns if column)
    req_rows = req_sheet.iter_rows(
        min_row=req_data_row,
        max_row=req_sheet.max_row,
        max_col=req_max_column,
        values_only=True,
    )
    for row, values in enumerate(req_rows, start=req_data_row):
        raw_id = row_cell_value(values, id_column)
        if raw_id in (None, ""):
            continue
        requirement_id = str(raw_id).strip().upper()
        if requirement_id.casefold() in {"строку не удалять", "строка не удалять"}:
            continue
        process = {key: row_cell_value(values, column) for key, column in process_columns.items()}
        process = {key: value for key, value in process.items() if value not in (None, "")}
        details = {key: row_cell_value(values, column) for key, column in optional_columns.items()}
        details = {key: value for key, value in details.items() if value not in (None, "")}
        requirement_record = {
            "id": requirement_id,
            "source_row": row,
            "text": row_cell_value(values, requirement_text_column),
            "process": process,
            "details": details,
        }
        requirement_occurrences.setdefault(requirement_id, []).append(requirement_record)
        if requirement_id in requirements:
            duplicate_requirements.append((requirement_id, row))
            continue
        if not ID_PATTERN.match(requirement_id):
            warnings.append(f"Unusual requirement ID '{requirement_id}' at row {row}.")
        requirements[requirement_id] = requirement_record

    if duplicate_requirements:
        errors.extend(f"Duplicate requirement ID {requirement_id} at row {row}." for requirement_id, row in duplicate_requirements)
    if not requirements:
        fatal_errors.append(f"No requirement IDs found in column {id_column_letter}, starting at row {req_data_row}.")

    fs_sheet = workbook[fs_sheet_name]
    fs_header_row = int(registry.get("specifications_header_row", 1))
    fs_data_row = int(registry.get("specifications_data_row", fs_header_row + 1))
    fs_headers = header_map(fs_sheet, fs_header_row)
    fs_code_column = find_column(fs_headers, "Код разработки (RICEF)", "Код разработки  (RICEF)", "Код ФС")
    fs_title_column = find_column(fs_headers, "Название разработки")
    fs_description_column = find_column(fs_headers, "Описание разработки")
    fs_status_column = find_column(fs_headers, "Статус ФС")
    fs_release_column = find_column(fs_headers, "Плановый релиз")
    fs_requirements_column = find_column(fs_headers, "Код требования", "Коды требований")
    if not fs_code_column:
        fatal_errors.append("The specifications register has no FS/development code column.")
    if not fs_requirements_column:
        fatal_errors.append("The specifications register has no requirement codes column.")

    specifications: dict[str, Any] = {}
    membership: dict[str, str] = {}
    membership_candidates: dict[str, list[str]] = {}
    specification_occurrences: dict[str, list[dict[str, Any]]] = {}
    if fs_code_column and fs_requirements_column:
        fs_columns = [
            fs_code_column,
            fs_title_column,
            fs_description_column,
            fs_status_column,
            fs_release_column,
            fs_requirements_column,
        ]
        fs_max_column = max(column for column in fs_columns if column)
        fs_rows = fs_sheet.iter_rows(
            min_row=fs_data_row,
            max_row=fs_sheet.max_row,
            max_col=fs_max_column,
            values_only=True,
        )
        for row, values in enumerate(fs_rows, start=fs_data_row):
            raw_code = row_cell_value(values, fs_code_column)
            if raw_code in (None, ""):
                continue
            code = str(raw_code).strip()
            requirement_ids = split_ids(row_cell_value(values, fs_requirements_column))
            specification_record = {
                "code": code,
                "source_row": row,
                "title": row_cell_value(values, fs_title_column),
                "description": row_cell_value(values, fs_description_column),
                "status_from_registry": row_cell_value(values, fs_status_column),
                "planned_release": row_cell_value(values, fs_release_column),
                "requirements": requirement_ids,
            }
            specification_occurrences.setdefault(code, []).append(specification_record)
            if code in specifications:
                errors.append(f"Duplicate FS code: {code}.")
            for requirement_id in requirement_ids:
                candidates = membership_candidates.setdefault(requirement_id, [])
                if code not in candidates:
                    candidates.append(code)
                if requirement_id not in requirements:
                    errors.append(f"{code} references missing requirement {requirement_id}.")
                if requirement_id in membership:
                    errors.append(
                        f"MVP relation violation: {requirement_id} belongs to both "
                        f"{membership[requirement_id]} and {code}."
                    )
                else:
                    membership[requirement_id] = code
            if code not in specifications:
                specifications[code] = specification_record

    data_root = project_root()
    normalized_dir = data_root / "registry" / "normalized"
    previous_requirements = read_json(normalized_dir / "requirements.json", {})
    previous_specifications = read_json(normalized_dir / "specifications.json", {})
    requirement_diff = make_diff(previous_requirements, requirements)
    specification_diff = make_diff(previous_specifications, specifications)

    impact: dict[str, list[str]] = {}
    changed_ids = set(requirement_diff["changed"]) | set(requirement_diff["missing_from_source"])
    for manifest_path in sorted((data_root / "work-items").glob("*/manifest.yaml")):
        manifest = read_json(manifest_path, {})
        affected = sorted(changed_ids & set(manifest.get("requirements", [])))
        if affected:
            impact[manifest.get("code", manifest_path.parent.name)] = affected

    report_lines = [
        "# Registry import report",
        "",
        f"Source: `{source.name}`",
        f"SHA-256: `{sha256(source)}`",
        f"Imported at: `{dt.datetime.now().astimezone().isoformat(timespec='seconds')}`",
        "",
        f"Requirements: {len(requirements)}",
        f"Specifications: {len(specifications)}",
        "",
        "## Changes",
        "",
        f"- Added requirements: {', '.join(requirement_diff['added']) or 'none'}",
        f"- Changed requirements: {', '.join(requirement_diff['changed']) or 'none'}",
        f"- Missing from source: {', '.join(requirement_diff['missing_from_source']) or 'none'}",
        f"- Added specifications: {', '.join(specification_diff['added']) or 'none'}",
        f"- Changed specifications: {', '.join(specification_diff['changed']) or 'none'}",
        "",
        "## Impact",
        "",
    ]
    if impact:
        report_lines.extend(f"- {code}: {', '.join(ids)}" for code, ids in sorted(impact.items()))
        report_lines.append("")
        report_lines.append("No PR was created. A user decision is required for each affected specification.")
    else:
        report_lines.append("No existing work item is affected.")
    report_lines.extend(["", "## Validation", ""])
    report_lines.extend(f"- ERROR: {message}" for message in [*fatal_errors, *errors])
    report_lines.extend(f"- WARNING: {message}" for message in warnings)
    if not fatal_errors and not errors and not warnings:
        report_lines.append("No validation issues found.")

    report_path = data_root / "registry" / "import-report.md"
    write_text(report_path, "\n".join(report_lines))
    workbook.close()

    def issue_record(message: str, *, warning: bool = False) -> dict[str, Any]:
        row_match = re.search(r"\brow (\d+)\b", message, flags=re.IGNORECASE)
        kind = "registry_warning" if warning else "registry_validation"
        lowered = message.casefold()
        if "duplicate requirement" in lowered:
            kind = "duplicate_requirement"
        elif "duplicate fs" in lowered:
            kind = "duplicate_specification"
        elif "missing requirement" in lowered:
            kind = "missing_requirement"
        record: dict[str, Any] = {"kind": kind, "message": message}
        if row_match:
            record["source_row"] = int(row_match.group(1))
        return record

    result = {
        "state": "BLOCKED" if fatal_errors else ("PARTIAL" if errors else "READY"),
        "report_path": str(report_path),
        "errors": [issue_record(message) for message in [*fatal_errors, *errors]],
        "warnings": [issue_record(message, warning=True) for message in warnings],
    }
    if fatal_errors:
        write_json(report_path.parent / "status.json", {"status": "invalid", "source_sha256": sha256(source),
                                                        "report_path": str(report_path), "updated_at": utc_now()})
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 2

    source_target = data_root / "registry" / "source" / "requirements.xlsx"
    source_target.parent.mkdir(parents=True, exist_ok=True)
    if source != source_target.resolve():
        shutil.copy2(source, source_target)
    target_dir = data_root / "registry" / ("partial" if errors else "normalized")
    write_json(target_dir / "requirements.json", requirements)
    write_json(target_dir / "specifications.json", specifications)
    write_json(target_dir / "links.json", {"requirement_to_specification": membership})
    if errors:
        write_json(target_dir / "scope-index.json", {
            "schema_version": 1,
            "source_sha256": sha256(source),
            "requirements": requirement_occurrences,
            "specifications": specification_occurrences,
            "requirement_to_specifications": membership_candidates,
            "errors": [issue_record(message) for message in errors],
        })
    write_json(data_root / "registry" / "status.json", {
        "status": "verified" if not errors else "partial", "source_sha256": sha256(source),
        "report_path": str(report_path), "errors_count": len(errors),
        "warnings_count": len(warnings), "updated_at": utc_now(),
    })
    result.update(requirements=len(requirements), specifications=len(specifications))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if errors else 0


def ensure_clean_tree() -> None:
    status = run_git(["status", "--porcelain"]).stdout.strip()
    if status:
        raise WorkflowError("Git working tree is not clean. Commit or stash existing changes first.")


def create_branch(branch: str) -> None:
    if not in_git_repository():
        raise WorkflowError("Current directory is not a Git repository.")
    ensure_clean_tree()
    default_branch = load_config()[0].get("project", {}).get("default_branch", "main")
    try:
        run_git(["fetch", "origin", default_branch])
        start_point = f"origin/{default_branch}"
    except subprocess.CalledProcessError:
        start_point = default_branch
    existing = run_git(["branch", "--list", branch]).stdout.strip()
    if existing:
        run_git(["switch", branch])
    else:
        run_git(["switch", "-c", branch, start_point])


def resolve_reference_args(args: argparse.Namespace, local: dict[str, Any] | None = None,
                           *, ignore_legacy_code: bool = False) -> dict[str, Any]:
    """Resolve new reference fields and all legacy CLI aliases."""
    if local is None:
        local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
    explicit = getattr(args, "task_reference", None)
    legacy_code = None if ignore_legacy_code else getattr(args, "code", None)
    legacy_g_number = getattr(args, "g_number", None)
    task = explicit if explicit not in (None, "") else legacy_code if legacy_code not in (None, "") else legacy_g_number
    project = getattr(args, "project_reference", None)
    if project in (None, ""):
        project = local.get("project_reference") if isinstance(local, dict) else None
    try:
        return resolve_work_reference(task, project)
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc


REFERENCE_KINDS = ("auto", "requirement", "specification")


def registry_scope_index(data_root: Path | None = None) -> dict[str, Any]:
    """Load the latest usable registry view without replacing the last verified indexes."""
    root = data_root or project_root()
    status = read_json(root / "registry" / "status.json", {})
    source_status = str(status.get("status") or "absent")
    if source_status == "invalid":
        return {"state": "invalid", "source_status": source_status, "status": status}
    if source_status == "partial":
        index = read_json(root / "registry" / "partial" / "scope-index.json", {})
        if not isinstance(index, dict) or index.get("schema_version") != 1:
            return {"state": "invalid", "source_status": "invalid", "status": status}
        return {**index, "state": "usable", "source_status": source_status, "status": status}
    requirements = read_json(root / "registry" / "normalized" / "requirements.json", {})
    specifications = read_json(root / "registry" / "normalized" / "specifications.json", {})
    links = read_json(root / "registry" / "normalized" / "links.json", {}).get(
        "requirement_to_specification", {}
    )
    legacy_verified = source_status == "absent" and bool(requirements) and bool(specifications)
    if (source_status != "verified" and not legacy_verified) or not isinstance(requirements, dict) or not isinstance(specifications, dict):
        return {"state": "absent", "source_status": source_status, "status": status}
    return {
        "schema_version": 1,
        "state": "usable",
        "source_status": "verified" if legacy_verified else source_status,
        "source_sha256": status.get("source_sha256"),
        "status": status,
        "requirements": {key: [value] for key, value in requirements.items()},
        "specifications": {key: [value] for key, value in specifications.items()},
        "requirement_to_specifications": {key: [value] for key, value in links.items()},
        "errors": [],
    }


def resolve_registry_reference(value: str, reference_kind: str = "auto",
                               data_root: Path | None = None) -> dict[str, Any]:
    """Resolve an FS number or requirement ID against the latest scoped registry view."""
    raw = str(value or "").strip()
    kind = str(reference_kind or "auto").strip().casefold()
    if kind not in REFERENCE_KINDS:
        raise WorkflowError(f"Unknown reference_kind: {reference_kind}")
    index = registry_scope_index(data_root)
    if index["state"] != "usable":
        return {"state": index["state"], "requested_reference": raw,
                "reference_kind": kind, "source_status": index.get("source_status")}
    requirements = index.get("requirements", {})
    specifications = index.get("specifications", {})
    requirement_key = raw.upper()
    requirement_match = requirement_key in requirements
    specification_match = raw in specifications
    if kind == "auto" and requirement_match and specification_match:
        return {
            "state": "ambiguous", "reason": "reference_kind_collision",
            "requested_reference": raw, "reference_kind": kind,
            "candidates": ["requirement", "specification"],
            "source_status": index["source_status"],
        }
    resolved_kind = kind
    if kind == "auto":
        resolved_kind = "requirement" if requirement_match else "specification" if specification_match else "auto"
    if resolved_kind == "requirement":
        if not requirement_match:
            return {"state": "not_found", "requested_reference": raw,
                    "reference_kind": resolved_kind, "source_status": index["source_status"]}
        candidates = list(dict.fromkeys(index.get("requirement_to_specifications", {}).get(requirement_key, [])))
        if not candidates:
            return {"state": "not_assigned", "requested_reference": raw,
                    "reference_kind": resolved_kind, "source_status": index["source_status"]}
        if len(candidates) > 1:
            return {"state": "ambiguous", "reason": "multiple_specifications",
                    "requested_reference": raw, "reference_kind": resolved_kind,
                    "candidates": candidates, "source_status": index["source_status"]}
        code = candidates[0]
    elif resolved_kind == "specification":
        if not specification_match:
            return {"state": "not_found", "requested_reference": raw,
                    "reference_kind": resolved_kind, "source_status": index["source_status"]}
        code = raw
    else:
        return {"state": "not_found", "requested_reference": raw,
                "reference_kind": kind, "source_status": index["source_status"]}
    occurrences = specifications.get(code, [])
    linked_ids = list(dict.fromkeys(
        requirement_id
        for occurrence in occurrences if isinstance(occurrence, dict)
        for requirement_id in occurrence.get("requirements", [])
    ))
    existing_ids = [requirement_id for requirement_id in linked_ids if requirement_id in requirements]
    if not existing_ids:
        return {"state": "no_existing_requirements", "requested_reference": raw,
                "reference_kind": resolved_kind, "code": code,
                "source_status": index["source_status"]}
    first = next((item for item in occurrences if isinstance(item, dict)), {})
    selected_requirements = {key: requirements[key][0] for key in existing_ids}
    return {
        "state": "resolved", "requested_reference": raw,
        "reference_kind": resolved_kind, "code": code,
        "requirements": existing_ids, "requirement_records": selected_requirements,
        "specification": {**first, "code": code, "requirements": existing_ids},
        "source_status": index["source_status"],
        "source_sha256": index.get("source_sha256") or index.get("status", {}).get("source_sha256"),
    }


def work_item_root(reference: str, *, for_create: bool = False) -> Path:
    """Map an opaque reference to a safe directory and discover legacy locations."""
    try:
        exact = validate_reference(reference, required=True)
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    assert exact is not None
    parent = project_root() / "work-items"
    slug = reference_slug(exact)
    direct = parent / slug
    if direct.exists():
        existing = read_json(direct / "manifest.yaml", {})
        existing_reference = existing.get("work_reference") or existing.get("task_reference") or existing.get("code")
        if existing_reference == exact:
            return direct
        if for_create:
            suffix = hashlib.sha256(exact.encode("utf-8")).hexdigest()[:8]
            return parent / reference_slug(exact, suffix=suffix)
    for manifest_path in parent.glob("*/manifest.yaml"):
        manifest = read_json(manifest_path, {})
        if exact in {manifest.get("work_reference"), manifest.get("task_reference"), manifest.get("code")}:
            return manifest_path.parent
    return direct


def cmd_fs_start(args: argparse.Namespace) -> int:
    reference = resolve_reference_args(args)
    requested_reference = reference["work_reference"]
    if not requested_reference:
        raise WorkflowError("A user-supplied project or task reference is required; it is never generated automatically.")

    data_root = project_root()
    resolution = resolve_registry_reference(
        str(requested_reference), str(getattr(args, "reference_kind", "auto") or "auto"), data_root,
    )
    if resolution["state"] == "resolved":
        code = str(resolution["code"])
        item = resolution["specification"]
        title = item.get("title") or args.title or code
        requirement_ids = resolution["requirements"]
        requirements = resolution["requirement_records"]
        reference = {**reference, "task_reference": code, "work_reference": code}
    elif resolution["state"] == "ambiguous":
        raise WorkflowError(
            f"Registry reference '{requested_reference}' is ambiguous ({resolution['reason']}): "
            + ", ".join(resolution["candidates"])
        )
    elif resolution["state"] == "invalid":
        raise WorkflowError("The latest registry import is structurally invalid; fix it or use an explicit registry_bypass provisional flow")
    else:
        code = str(requested_reference)
        title = args.title
        requirement_ids = split_ids(args.requirements)
        requirements = read_json(data_root / "registry" / "normalized" / "requirements.json", {})
        if not title or not requirement_ids:
            raise WorkflowError(
                f"{requested_reference} cannot be resolved to a usable registry scope. "
                "Provide --title and --requirements explicitly or use registry_bypass."
            )
    missing = [requirement_id for requirement_id in requirement_ids if requirement_id not in requirements]
    if missing:
        raise WorkflowError("Requirements absent from normalized registry: " + ", ".join(missing))

    target = work_item_root(code, for_create=True)
    branch = f"fs/{target.name}/specification"
    if args.create_branch:
        create_branch(branch)

    if target.exists():
        raise WorkflowError(f"Work item already exists: {target}")
    template = ROOT / "templates" / "work-item"
    if template.is_dir():
        shutil.copytree(template, target)
    else:
        for relative in ("input/attachments", "analysis", "specification", "reviews", "evidence"):
            (target / relative).mkdir(parents=True, exist_ok=True)

    manifest = {
        "schema_version": 1,
        "code": code,
        **reference,
        "work_item_slug": target.name,
        "title": title,
        "status": "clarification",
        "traceability_mode": "registry",
        "requirements_source": "registry",
        "requirements": requirement_ids,
        "requested_reference": str(requested_reference),
        "reference_resolution": resolution.get("reference_kind", "explicit"),
        "branches": {"documentation": branch, "extension": f"feature/{target.name}"},
        "approvals": {"functional_architect": "pending", "technical_architect": "pending"},
        "created_at": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "updated_at": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
    }
    write_json(target / "manifest.yaml", manifest)
    snapshot = {requirement_id: requirements[requirement_id] for requirement_id in requirement_ids}
    write_json(target / "input" / "requirements.snapshot.yaml", snapshot)
    snapshot_path = target / "input" / "requirements.snapshot.yaml"
    manifest["registry"] = {
        "status": "verified", "snapshot_sha256": sha256(snapshot_path),
        "source_status": resolution.get("source_status", "verified"),
        "source_sha256": resolution.get("source_sha256"),
        "scope": {"specification": code, "requirements": requirement_ids},
    }
    manifest["requirement_basis"] = {
        "type": "registry_snapshot", "path": "input/requirements.snapshot.yaml", "sha256": sha256(snapshot_path),
    }
    write_json(target / "manifest.yaml", manifest)
    replacements = {"{{WORK_REFERENCE}}": code, "{{FS_CODE}}": code, "{{FS_TITLE}}": str(title)}
    for path in target.rglob("*.md"):
        text = path.read_text(encoding="utf-8")
        for old, new in replacements.items():
            text = text.replace(old, new)
        path.write_text(text, encoding="utf-8")
    print(f"Created {target}")
    print(f"Requirements: {', '.join(requirement_ids)}")
    print(f"Documentation branch: {branch}")
    return 0


def create_provisional_work_item(reference: dict[str, Any], *, title: str, user_brief: str,
                                 deviation: dict[str, Any], requirements: list[str] | None = None) -> Path:
    """Create a full work item from user material without inventing any identifiers."""
    code = reference.get("work_reference")
    if not code:
        raise WorkflowError("A user-supplied task_reference or project_reference is required for provisional work")
    brief = str(user_brief or "").strip()
    if not brief:
        raise WorkflowError("A nonempty user brief is required for provisional work")
    target = work_item_root(str(code), for_create=True)
    if target.exists():
        raise WorkflowError(f"Work item already exists: {target}")
    template = ROOT / "templates" / "work-item"
    if template.is_dir():
        shutil.copytree(template, target)
    else:
        for relative in ("input/attachments", "analysis", "specification", "reviews", "evidence"):
            (target / relative).mkdir(parents=True, exist_ok=True)
    (target / "evidence").mkdir(parents=True, exist_ok=True)
    (target / "input" / "attachments").mkdir(parents=True, exist_ok=True)
    brief_path = target / "input" / "user-brief.md"
    write_text(brief_path, brief)
    artifacts_path = target / "input" / "artifacts.json"
    if not artifacts_path.exists():
        write_json(artifacts_path, {"artifacts": [], "confirmed_absent": []})
    confirmed_at = str(deviation.get("recorded_at") or utc_now())
    bypass = {
        "type": "registry_bypass",
        "reason": str(deviation.get("reason") or "Registry traceability was explicitly bypassed"),
        "actor": str(deviation.get("actor") or "user"),
        "user_statement": str(deviation.get("user_statement") or deviation.get("reason") or "").strip(),
        "confirmed_at": confirmed_at,
        "scope": str(deviation.get("scope") or "work-item"),
        "condition_ids": list(dict.fromkeys(str(item) for item in deviation.get("condition_ids", deviation.get("waived_conditions", [])))),
        "waived_conditions": list(dict.fromkeys(str(item) for item in deviation.get("waived_conditions", deviation.get("condition_ids", [])))),
        "unconfirmed_conditions": list(dict.fromkeys([
            "registry requirements", "registry assignment", "registry snapshot", *deviation.get("waived_conditions", []),
        ])),
    }
    branch = f"fs/{target.name}/specification"
    now = utc_now()
    manifest = {
        "schema_version": 1, "code": code, **reference, "work_item_slug": target.name,
        "title": str(title or code), "status": "clarification",
        "traceability_mode": "provisional", "requirements_source": "user_brief",
        "requirements": list(requirements or []),
        "registry": {"status": "bypassed", "snapshot_sha256": None, "bypass": bypass},
        "requirement_basis": {"type": "user_brief", "path": "input/user-brief.md", "sha256": sha256(brief_path)},
        "branches": {"documentation": branch, "extension": f"feature/{target.name}"},
        "approvals": {"functional_architect": "pending", "technical_architect": "pending"},
        "created_at": now, "updated_at": now,
    }
    write_json(target / "manifest.yaml", manifest)
    replacements = {"{{WORK_REFERENCE}}": str(code), "{{FS_CODE}}": str(code), "{{FS_TITLE}}": str(title or code)}
    for path in target.rglob("*.md"):
        if path == brief_path:
            continue
        text = path.read_text(encoding="utf-8")
        for old, new in replacements.items():
            text = text.replace(old, new)
        path.write_text(text, encoding="utf-8")
    return target


def load_manifest(code: str) -> tuple[Path, dict[str, Any]]:
    try:
        code = validate_reference(code, required=True) or ""
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    path = work_item_root(code) / "manifest.yaml"
    manifest = read_json(path)
    if manifest is None:
        raise WorkflowError(f"Work item not found: {code}")
    return path, manifest


def utc_now() -> str:
    return dt.datetime.now().astimezone().isoformat(timespec="seconds")


def load_stages() -> dict[str, Any]:
    path = ROOT / "config" / "stages.json"
    stages = read_json(path)
    if not isinstance(stages, dict) or not isinstance(stages.get("operations"), dict):
        raise WorkflowError(f"Invalid stage configuration: {path}")
    return stages


def meaningful_file(path: Path) -> bool:
    if not path.is_file() or path.stat().st_size == 0:
        return False
    if path.suffix.casefold() not in {".md", ".txt", ".yaml", ".json"}:
        return True
    if path.suffix.casefold() in {".yaml", ".json"}:
        try:
            return bool(json.loads(path.read_text(encoding="utf-8-sig")))
        except (OSError, json.JSONDecodeError):
            return False
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    table_separator_seen = False
    for line in text.splitlines():
        value = line.strip()
        if not value or value.startswith(("#", "> Generated", "<!--", "> **UNVERIFIED_DRAFT")):
            continue
        if value.startswith("|"):
            compact = value.replace("|", "").replace("-", "").replace(":", "").strip()
            if not compact:
                table_separator_seen = True
                continue
            if table_separator_seen:
                return True
            continue
        if "{{" in value and "}}" in value:
            continue
        return True
    return False


def validate_json_record(value: Any, schema_path: Path) -> list[str]:
    """Validate the deliberately small schema subset used by FLOW1C agent records."""
    schema = read_json(schema_path, {})
    if not isinstance(value, dict) or not isinstance(schema, dict):
        return ["record or schema is not a JSON object"]
    errors: list[str] = []
    for key in schema.get("required", []):
        if key not in value:
            errors.append(f"missing required property {key}")
    type_map = {"object": dict, "array": list, "string": str, "integer": int, "number": (int, float), "null": type(None)}
    for key, rules in schema.get("properties", {}).items():
        if key not in value or not isinstance(rules, dict):
            continue
        current = value[key]
        expected = rules.get("type")
        names = expected if isinstance(expected, list) else [expected] if expected else []
        expected_types = tuple(type_map[name] for name in names if name in type_map)
        if expected_types and not isinstance(current, expected_types):
            errors.append(f"property {key} has an invalid type")
            continue
        if "const" in rules and current != rules["const"]:
            errors.append(f"property {key} must equal {rules['const']}")
        if "enum" in rules and current not in rules["enum"]:
            errors.append(f"property {key} is outside the allowed enum")
        if isinstance(current, str) and rules.get("pattern") and not re.fullmatch(str(rules["pattern"]), current):
            errors.append(f"property {key} does not match {rules['pattern']}")
        if isinstance(current, str) and rules.get("minLength") and len(current) < int(rules["minLength"]):
            errors.append(f"property {key} is shorter than {rules['minLength']}")
        if isinstance(current, str) and rules.get("maxLength") and len(current) > int(rules["maxLength"]):
            errors.append(f"property {key} is longer than {rules['maxLength']}")
    return errors


def gate_path(gate_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9-]{8,64}", str(gate_id or ""), flags=re.IGNORECASE):
        raise WorkflowError("Invalid agent gate ID.")
    return ROOT / ".workspace" / "agent-gates" / f"{gate_id}.json"


def save_gate(gate: dict[str, Any]) -> None:
    sanitized = sanitize_json_value(gate)
    gate.clear()
    gate.update(sanitized)
    write_json(gate_path(str(gate["gate_id"])), gate)


def refresh_gate_actions(gate: dict[str, Any], *, persist: bool = False) -> dict[str, Any]:
    """Recompute the executable tool surface for the current gate state."""
    stage = load_stages().get("operations", {}).get(str(gate.get("operation")), {})
    actions = available_actions(mode=str(gate.get("mode", "formal")), operation=str(gate.get("operation", "")),
                                state=str(gate.get("state", "")), stage=stage if isinstance(stage, dict) else {})
    if gate.get("available_actions") != actions:
        gate["available_actions"] = actions
        if persist:
            save_gate(gate)
    return gate


def load_gate(gate_id: str, *, states: set[str] | None = None) -> dict[str, Any]:
    gate = read_json(gate_path(gate_id))
    if not isinstance(gate, dict) or gate.get("schema_version") != AGENT_GATE_SCHEMA_VERSION:
        raise WorkflowError("Agent gate does not exist or has an unsupported schema.")
    schema_errors = validate_json_record(gate, ROOT / "schemas" / "agent-gate.schema.json")
    if schema_errors:
        raise WorkflowError("Agent gate schema validation failed: " + "; ".join(schema_errors))
    if states is not None and gate.get("state") not in states:
        raise WorkflowError(f"Agent gate state {gate.get('state')} is not allowed for this action.")
    if gate.get("policy_version") != POLICY_VERSION:
        raise WorkflowError("Legacy gate requires re-assessment: call flow1c_begin with the saved operation, code and summary. Existing evidence is preserved.")
    refresh_gate_actions(gate, persist=True)
    return gate


def require_gate_tool(gate: dict[str, Any], tool_name: str) -> dict[str, Any]:
    if gate.get("mode", "formal") != "formal":
        allowed = allowed_tools_for_mode(str(gate.get("mode")), str(gate.get("operation", "")))
        if tool_name not in allowed:
            raise WorkflowError(f"Tool {tool_name} is unavailable in {gate.get('mode')} mode")
        return {"allowed_tools": allowed, "writable_targets": ["draft"] if gate.get("mode") == "draft" else []}
    stages = load_stages()
    stage = stages.get("operations", {}).get(gate.get("operation"))
    actions = available_actions(mode="formal", operation=str(gate.get("operation")), state=str(gate.get("state")),
                                stage=stage if isinstance(stage, dict) else {})
    if tool_name not in actions:
        raise WorkflowError(f"Tool {tool_name} is not allowed for operation {gate.get('operation')}.")
    return {**(stage if isinstance(stage, dict) else {}), "allowed_tools": actions}


def artifact_index_path(code: str | None) -> Path | None:
    if not code:
        return None
    return work_item_root(code) / "input" / "artifacts.json"


def load_artifact_index(code: str | None) -> dict[str, Any]:
    path = artifact_index_path(code)
    if path is None:
        artifacts: list[dict[str, Any]] = []
        confirmed_absent: set[str] = set()
        for manifest_path in sorted((project_root() / "inbox").glob("*/intake.json")):
            manifest = read_json(manifest_path, {})
            if isinstance(manifest, dict) and manifest.get("code") is None:
                artifacts.extend(item for item in manifest.get("artifacts", []) if isinstance(item, dict))
                confirmed_absent.update(str(value) for value in manifest.get("confirmed_absent", []))
        return {"schema_version": 1, "artifacts": artifacts, "confirmed_absent": sorted(confirmed_absent)}
    value = read_json(path, {})
    if not isinstance(value, dict):
        value = {}
    return {
        "schema_version": 1,
        "artifacts": list(value.get("artifacts", [])),
        "confirmed_absent": list(value.get("confirmed_absent", [])),
    }


def is_reparse_or_symlink(path: Path) -> bool:
    if path.is_symlink():
        return True
    try:
        attributes = getattr(path.stat(follow_symlinks=False), "st_file_attributes", 0)
        return bool(attributes & getattr(__import__("stat"), "FILE_ATTRIBUTE_REPARSE_POINT", 0))
    except OSError:
        return True


def enumerate_intake_files(
    sources: Iterable[str], *, excluded_roots: Iterable[Path] = ()
) -> tuple[list[tuple[Path, Path]], list[str]]:
    files: list[tuple[Path, Path]] = []
    skipped: list[str] = []
    excluded = tuple(path.resolve() for path in excluded_roots)

    def is_excluded(path: Path) -> bool:
        resolved = path.resolve()
        return any(resolved == root or path_is_within(resolved, root) for root in excluded)

    for raw in sources:
        source = Path(raw).expanduser().resolve()
        if not source.exists():
            raise WorkflowError(f"Artifact source does not exist: {source}")
        if is_excluded(source):
            skipped.append(str(source))
            continue
        if is_reparse_or_symlink(source):
            raise WorkflowError(f"Symlinks and reparse points are not accepted: {source}")
        if source.is_file():
            candidates = [(source, Path(source.name))]
        else:
            candidates = []
            for candidate in sorted(source.rglob("*")):
                if is_excluded(candidate):
                    if candidate.is_file():
                        skipped.append(str(candidate))
                    continue
                if is_reparse_or_symlink(candidate):
                    skipped.append(str(candidate))
                    continue
                if candidate.is_file():
                    candidates.append((candidate, candidate.relative_to(source)))
        for candidate, relative in candidates:
            if candidate.name.startswith("~$"):
                skipped.append(str(candidate))
                continue
            suffix = candidate.suffix.casefold()
            if suffix in DENIED_ARTIFACT_EXTENSIONS or suffix not in ALLOWED_ARTIFACT_EXTENSIONS:
                skipped.append(str(candidate))
                continue
            files.append((candidate, relative))
    return files, skipped


def paths_are_equal(left: Path, right: Path) -> bool:
    return os.path.normcase(str(left.resolve())) == os.path.normcase(str(right.resolve()))


def free_intake_target(root: Path, digest: str, source: Path) -> Path:
    parent = root / "sources" / digest
    target = parent / source.name
    if len(str(target)) <= MAX_STORED_ARTIFACT_PATH:
        return target
    target = parent / f"artifact-{digest[:12]}{source.suffix.casefold()}"
    if len(str(target)) > MAX_STORED_ARTIFACT_PATH:
        raise WorkflowError(
            "The intake destination path is too long. Configure a shorter documentation_path."
        )
    return target


def intake_destination(code: str | None, category: str, intake_id: str, stages: dict[str, Any]) -> Path:
    data_root = project_root()
    category_config = stages.get("artifact_categories", {}).get(category, {})
    destination = category_config.get("destination", "attachments")
    if not code or destination == "inbox":
        return data_root / "inbox" / intake_id / "originals"
    if destination == "meetings":
        return work_item_root(code) / "input" / "meetings" / intake_id
    return work_item_root(code) / "input" / "attachments" / intake_id


def write_derived_artifact(source: Path, relative: Path, code: str | None, intake_id: str) -> str | None:
    if not code or source.suffix.casefold() in {".xlsx", ".png", ".jpg", ".jpeg"}:
        return None
    derived_root = work_item_root(code) / "input" / "derived" / intake_id
    target = derived_root / relative.with_suffix(relative.suffix + ".md")
    try:
        if source.suffix.casefold() in {".txt", ".md", ".csv"}:
            converted = source.read_text(encoding="utf-8-sig", errors="replace")
        else:
            from markitdown import MarkItDown  # type: ignore[import-not-found]

            converted = MarkItDown().convert(str(source)).text_content
        write_text(
            target,
            "<!-- Generated by Flow1C artifact-intake. Rebuild instead of editing. -->\n\n" + converted,
        )
        return str(target.relative_to(project_root())).replace("\\", "/")
    except (ImportError, OSError, AttributeError, ValueError):
        return None


def persist_intake_files(
    files: list[tuple[Path, Path]],
    skipped: list[str],
    *,
    code: str | None,
    category: str,
    received_via: str,
    source_record: dict[str, Any] | None = None,
    index: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Store accepted files through the common inbox/work-item artifact path."""
    if len(files) > MAX_INTAKE_FILES:
        raise WorkflowError(f"Intake contains more than {MAX_INTAKE_FILES} files.")
    total_bytes = sum(path.stat().st_size for path, _ in files)
    if total_bytes > MAX_INTAKE_BYTES:
        raise WorkflowError(f"Intake contains {total_bytes} bytes; the limit is {MAX_INTAKE_BYTES}.")
    intake_id = dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:8]
    stages = load_stages()
    if category not in stages.get("artifact_categories", {}):
        raise WorkflowError(f"Unknown artifact category: {category}")
    destination = intake_destination(code, category, intake_id, stages)
    index = index or load_artifact_index(code)
    known_hashes = {item.get("sha256") for item in index["artifacts"]}
    copied: list[dict[str, Any]] = []
    for source, relative in files:
        digest = sha256(source)
        if digest in known_hashes:
            skipped.append(source.name)
            continue
        target = destination / relative
        if target.exists():
            target = target.with_name(f"{target.stem}-{digest[:8]}{target.suffix}")
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        derived = write_derived_artifact(target, relative, code, intake_id)
        relative_target = str(target.relative_to(project_root())).replace("\\", "/")
        record = {
            "category": category,
            "name": source.name,
            "relative_path": relative_target,
            "sha256": digest,
            "size": source.stat().st_size,
            "media_type": mimetypes.guess_type(source.name)[0] or "application/octet-stream",
            "received_at": utc_now(),
            "received_via": received_via,
            "derived_path": derived,
        }
        if source_record:
            record["source"] = source_record
        copied.append(record)
        index["artifacts"].append(record)
        known_hashes.add(digest)
    index["confirmed_absent"] = sorted(set(index["confirmed_absent"]))
    if code:
        write_json(artifact_index_path(code), index)
    else:
        manifest_path = destination.parent / "intake.json"
        manifest: dict[str, Any] = {
            "schema_version": 1,
            "intake_id": intake_id,
            "code": None,
            "created_at": utc_now(),
            "artifacts": copied,
            "confirmed_absent": index["confirmed_absent"],
        }
        if source_record:
            manifest["source"] = source_record
        write_json(manifest_path, manifest)
    return {
        "state": "ACCEPTED",
        "intake_id": intake_id,
        "code": code,
        "destination": str(destination),
        "copied": copied,
        "skipped": skipped,
        "confirmed_absent": index["confirmed_absent"],
    }


def extension_git_state(code: str, manifest: dict[str, Any]) -> tuple[dict[str, Any] | None, str]:
    _, local = load_config()
    raw = str(local.get("extension_path", "")).strip()
    extension = Path(raw).resolve() if raw else None
    if extension is None or not extension.is_dir():
        return None, "extension path is not configured"
    if not (extension / ".git").exists() or not command_path("git"):
        return None, "extension source is not a Git checkout"
    branch_result = subprocess.run(
        ["git", "branch", "--show-current"], cwd=extension, text=True, encoding="utf-8", errors="replace", capture_output=True
    )
    branch = branch_result.stdout.strip()
    expected = str(manifest.get("branches", {}).get("extension", f"feature/{code}"))
    default_branch = load_config()[0].get("project", {}).get("default_branch", "main")
    base_candidates = [f"origin/{default_branch}", str(default_branch)]
    base = next(
        (
            candidate
            for candidate in base_candidates
            if subprocess.run(
                ["git", "rev-parse", "--verify", candidate], cwd=extension, capture_output=True, check=False
            ).returncode
            == 0
        ),
        "",
    )
    if not base:
        return None, f"base branch {default_branch} is unavailable"
    head = subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=extension, text=True, encoding="utf-8", errors="replace", capture_output=True
    ).stdout.strip()
    names = subprocess.run(
        ["git", "diff", "--name-only", f"{base}...HEAD"],
        cwd=extension,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
    ).stdout.splitlines()
    return {
        "path": str(extension),
        "branch": branch,
        "expected_branch": expected,
        "base": base,
        "head": head,
        "files": [name for name in names if name.strip()],
    }, ""


def redmine_settings(local: dict[str, Any] | None = None) -> dict[str, Any]:
    local = local if local is not None else read_json(ROOT / LOCAL_CONFIG_FILE, {})
    value = local.get("redmine", {}) if isinstance(local, dict) else {}
    return value if isinstance(value, dict) else {}


def redmine_api_key(base_url: str) -> tuple[str, str]:
    from_environment = os.environ.get("FLOW1C_REDMINE_API_KEY", "").strip()
    if from_environment:
        return from_environment, "environment"
    try:
        stored = redmine_credentials.read_api_key(base_url)
    except (RedmineError, RedminePolicyError) as exc:
        raise WorkflowError(str(exc)) from exc
    if stored:
        return stored, "windows-credential-manager"
    raise WorkflowError(
        "Redmine API key is not configured. On Windows run `flow1c.py redmine configure --url <HTTPS URL>`; "
        "other environments must set FLOW1C_REDMINE_API_KEY for the agent process."
    )


def redmine_client_from_config(*, recover: bool = True) -> tuple[RedmineClient, str, str]:
    recovery_warnings = recover_redmine_transaction() if recover else []
    settings = redmine_settings()
    raw_url = str(settings.get("base_url", "")).strip()
    if not raw_url:
        raise WorkflowError("Redmine is not configured. Run `flow1c.py redmine configure --url <HTTPS URL>`.")
    try:
        base_url = normalize_base_url(raw_url)
        api_key, source = redmine_api_key(base_url)
        client = RedmineClient(base_url, api_key)
        client.recovery_warnings = recovery_warnings
        return client, base_url, source
    except (RedmineError, RedminePolicyError) as exc:
        raise WorkflowError(str(exc)) from exc


def redmine_cleanup_state_path() -> Path:
    return ROOT / ".workspace" / "redmine-pending-cleanup.json"


def redmine_transaction_state_path() -> Path:
    return ROOT / ".workspace" / "redmine-configure-transaction.json"


def redmine_pending_cleanup_urls() -> list[str]:
    state = read_json(redmine_cleanup_state_path(), {})
    values = state.get("urls", []) if isinstance(state, dict) else []
    result: list[str] = []
    for value in values if isinstance(values, list) else []:
        try:
            normalized = normalize_base_url(str(value))
        except (RedmineError, RedminePolicyError):
            continue
        if normalized not in result:
            result.append(normalized)
    return result


def save_redmine_pending_cleanup_urls(urls: Iterable[str]) -> None:
    path = redmine_cleanup_state_path()
    cleaned: list[str] = []
    for value in urls:
        normalized = normalize_base_url(value)
        if normalized not in cleaned:
            cleaned.append(normalized)
    if cleaned:
        write_json(path, {"schema_version": 1, "urls": cleaned})
    elif path.exists():
        path.unlink()


def clear_redmine_transaction_state() -> None:
    path = redmine_transaction_state_path()
    if path.exists():
        path.unlink()


def recover_redmine_transaction() -> list[dict[str, Any]]:
    """Finish cleanup after activation or roll back credentials before activation."""
    path = redmine_transaction_state_path()
    if not path.exists():
        return []
    try:
        transaction = read_json(path)
        if not isinstance(transaction, dict) or transaction.get("schema_version") != 1:
            raise ValueError("invalid state")
        transaction_id = str(transaction.get("transaction_id", ""))
        backup_target = redmine_credentials.transaction_backup_target_name(transaction_id)
        transaction_phase = str(transaction.get("phase", "prepared"))
        if transaction_phase not in {"prepared", "activated"}:
            raise ValueError("invalid transaction phase")
        new_url = normalize_base_url(str(transaction.get("new_url", "")))
        old_url_raw = str(transaction.get("old_url", "") or "")
        old_url = normalize_base_url(old_url_raw) if old_url_raw else None
        old_url_invalid = transaction.get("old_url_invalid") is True
        credential_managed = transaction.get("credential_managed") is True
        credential_backup = transaction.get("credential_backup") is True
        _ = backup_target
        config_path = ROOT / LOCAL_CONFIG_FILE
        local = read_json(config_path, {})
        if not isinstance(local, dict):
            raise ValueError("invalid local configuration")
        active_raw = str(redmine_settings(local).get("base_url", "")).strip()
        try:
            active_url = normalize_base_url(active_raw) if active_raw else None
        except (RedmineError, RedminePolicyError):
            active_url = None
    except (WorkflowError, ValueError, TypeError):
        raise RedmineOperationError(
            "REDMINE_ROLLBACK_FAILED",
            "An interrupted Redmine transaction could not be read safely; no credential was changed.",
            recoverable=True,
            next_action="Inspect `.workspace/redmine-configure-transaction.json` and the active local configuration, then retry redmine cleanup.",
            previous_connection_preserved=True,
        ) from None

    if credential_managed and sys.platform != "win32":
        raise RedmineOperationError(
            "REDMINE_ROLLBACK_FAILED",
            "An interrupted transaction requires Windows Credential Manager recovery.",
            recoverable=True,
            next_action="Run a Redmine command on the Windows profile that owns the Credential Manager entries.",
            previous_connection_preserved=active_url != new_url,
        ) from None

    warnings: list[dict[str, Any]] = []
    retain_transaction = False
    pending_readable = True
    try:
        pending_urls = redmine_pending_cleanup_urls()
    except WorkflowError:
        pending_urls = []
        pending_readable = False

    transaction_activated = active_url == new_url and (old_url != new_url or transaction_phase == "activated")
    if transaction_activated:
        if old_url_invalid:
            warnings.append({
                "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
                "message": "The new connection is active, but the previous invalid URL could not identify its saved credential for cleanup.",
                "recoverable": True,
                "next_action": "Review Windows Credential Manager after correcting the previous Redmine URL.",
            })
        if old_url and old_url != new_url and sys.platform == "win32":
            try:
                redmine_credentials.delete_api_key(old_url)
                pending_urls = [value for value in pending_urls if value != old_url]
            except Exception:
                if old_url not in pending_urls:
                    pending_urls.append(old_url)
                retain_transaction = True
                warnings.append({
                    "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
                    "message": "The new connection is active, but the previous saved credential could not be removed during recovery.",
                    "recoverable": True,
                    "next_action": "Run `redmine cleanup` to retry removal from the local recovery record.",
                })
        if pending_readable:
            try:
                save_redmine_pending_cleanup_urls(pending_urls)
                retain_transaction = False
            except Exception:
                retain_transaction = True
                warnings.append({
                    "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
                    "message": "The connection is active, but the cleanup recovery record could not be updated.",
                    "recoverable": True,
                    "next_action": "Repair `.workspace/redmine-pending-cleanup.json` and run `redmine cleanup`.",
                })
        if not pending_readable:
            retain_transaction = True
        if credential_backup:
            try:
                redmine_credentials.delete_transaction_backup(transaction_id)
            except Exception:
                retain_transaction = True
                warnings.append({
                    "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
                    "message": "The connection is active, but a temporary credential backup could not be removed.",
                    "recoverable": True,
                    "next_action": "Run `redmine cleanup` again on Windows to finish transaction recovery.",
                })
    elif credential_managed:
        try:
            backup = redmine_credentials.read_transaction_backup(transaction_id) if credential_backup else None
            if backup is not None:
                redmine_credentials.write_api_key(new_url, backup)
                if not redmine_credentials.verify_api_key(new_url, backup):
                    raise RedmineError("Credential restoration could not be verified.")
                redmine_credentials.delete_transaction_backup(transaction_id)
            elif not credential_backup:
                redmine_credentials.delete_api_key(new_url)
                if redmine_credentials.read_api_key(new_url) is not None:
                    raise RedmineError("Credential rollback could not be verified.")
            # If a prior credential was recorded but its backup was never written,
            # the canonical entry has not yet been touched and must be left alone.
        except Exception:
            raise RedmineOperationError(
                "REDMINE_ROLLBACK_FAILED",
                "The interrupted Redmine transaction could not restore its credential state.",
                recoverable=True,
                next_action="Check Windows Credential Manager access and rerun a Redmine command to retry recovery.",
                previous_connection_preserved=active_url != new_url,
            ) from None

    if not retain_transaction:
        try:
            clear_redmine_transaction_state()
        except OSError:
            warnings.append({
                "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
                "message": "Redmine transaction recovery completed, but its temporary state file could not be removed.",
                "recoverable": True,
                "next_action": "Run `redmine cleanup` again to remove the completed recovery record.",
            })
    if not pending_readable:
        warnings.append({
            "code": "REDMINE_CLEANUP_STATE_READ_FAILED",
            "message": "The interrupted transaction was recovered, but an existing credential cleanup record could not be read.",
            "recoverable": True,
            "next_action": "Repair `.workspace/redmine-pending-cleanup.json` and run `redmine cleanup`.",
        })
    return warnings


def redmine_config_snapshot() -> tuple[Path, bytes | None, dict[str, Any]]:
    path = ROOT / LOCAL_CONFIG_FILE
    try:
        original = path.read_bytes() if path.exists() else None
        local = json.loads(original.decode("utf-8-sig")) if original is not None else {}
    except (OSError, UnicodeError, json.JSONDecodeError):
        raise RedmineOperationError(
            "REDMINE_CONFIG_WRITE_FAILED",
            "The local configuration could not be read, so Redmine settings were left unchanged.",
            recoverable=True,
            next_action="Repair the local JSON configuration and retry redmine configure or disconnect.",
            previous_connection_preserved=True,
        ) from None
    if not isinstance(local, dict):
        raise RedmineOperationError(
            "REDMINE_CONFIG_WRITE_FAILED",
            "The local configuration must contain a JSON object; Redmine settings were left unchanged.",
            recoverable=True,
            next_action="Repair the local JSON configuration and retry the Redmine operation.",
            previous_connection_preserved=True,
        )
    return path, original, sanitize_json_value(local)


def restore_redmine_config_snapshot(path: Path, original: bytes | None) -> None:
    if original is None:
        if path.exists():
            path.unlink()
    else:
        atomic_write_bytes(path, original)


def redmine_error_record(
    code: str, message: str, next_action: str, *, previous_connection_preserved: bool = True
) -> dict[str, Any]:
    return {
        "code": code, "message": message, "recoverable": True, "next_action": next_action,
        "previous_connection_preserved": previous_connection_preserved,
    }


def raise_redmine_transaction_error(
    code: str,
    message: str,
    next_action: str,
    *,
    previous_connection_preserved: bool,
    rollback_failures: list[str] | None = None,
) -> None:
    if rollback_failures:
        original = redmine_error_record(
            code, message, next_action, previous_connection_preserved=previous_connection_preserved
        )
        recovery_action = "Inspect the reported state, restore the preserved connection if needed, and retry redmine cleanup."
        rollback_records = [
            redmine_error_record(
                "REDMINE_ROLLBACK_FAILED", detail, recovery_action,
                previous_connection_preserved=previous_connection_preserved,
            )
            for detail in rollback_failures
        ]
        raise RedmineOperationError(
            "REDMINE_ROLLBACK_FAILED",
            "Redmine setup failed and automatic rollback was incomplete.",
            recoverable=True,
            next_action=recovery_action,
            previous_connection_preserved=previous_connection_preserved,
            additional_errors=[original, *rollback_records],
        ) from None
    raise RedmineOperationError(
        code,
        message,
        recoverable=True,
        next_action=next_action,
        previous_connection_preserved=previous_connection_preserved,
    ) from None


def restore_redmine_credential(base_url: str, previous_value: str | None) -> None:
    if previous_value is None:
        redmine_credentials.delete_api_key(base_url)
        if redmine_credentials.read_api_key(base_url) is not None:
            raise RedmineError("Credential cleanup could not be verified.")
    else:
        redmine_credentials.write_api_key(base_url, previous_value)
        if not redmine_credentials.verify_api_key(base_url, previous_value):
            raise RedmineError("Credential restoration could not be verified.")


def previous_redmine_connection_preserved(previous_url: str | None, requested_url: str, rollback_failed: bool) -> bool:
    return not rollback_failed or previous_url != requested_url


def cmd_redmine_configure(args: argparse.Namespace) -> int:
    recovery_warnings = recover_redmine_transaction()
    if redmine_transaction_state_path().exists():
        raise RedmineOperationError(
            "REDMINE_ROLLBACK_FAILED",
            "A previous Redmine transaction still needs recovery; this configure request made no changes.",
            recoverable=True,
            next_action="Run `redmine cleanup` on Windows to finish recovery, then retry redmine configure.",
            previous_connection_preserved=True,
        ) from None
    config_path, original_config, local = redmine_config_snapshot()
    previous_settings = redmine_settings(local)
    raw_previous_url = str(previous_settings.get("base_url", "")).strip()
    previous_url: str | None = None
    previous_url_invalid = False
    if raw_previous_url:
        try:
            previous_url = normalize_base_url(raw_previous_url)
        except (RedmineError, RedminePolicyError):
            previous_url_invalid = True
    try:
        base_url = normalize_base_url(args.url)
    except (RedmineError, RedminePolicyError) as exc:
        raise_redmine_transaction_error(
            "REDMINE_VALIDATION_FAILED", str(exc),
            "Correct the HTTPS Redmine URL and retry; the existing configuration is unchanged.",
            previous_connection_preserved=True,
        )
    api_key = os.environ.get("FLOW1C_REDMINE_API_KEY", "").strip()
    save_to_credential_manager = not bool(api_key)
    if save_to_credential_manager:
        if sys.platform != "win32":
            raise_redmine_transaction_error(
                "REDMINE_VALIDATION_FAILED",
                "No Redmine API key source is available on this platform.",
                "Set FLOW1C_REDMINE_API_KEY in the agent environment and retry; the key is never accepted as a command-line argument or saved in project files.",
                previous_connection_preserved=True,
            )
        try:
            api_key = getpass.getpass("Redmine API key (input is hidden): ").strip()
        except (EOFError, OSError):
            raise_redmine_transaction_error(
                "REDMINE_VALIDATION_FAILED", "The hidden API-key prompt did not complete.",
                "Retry in an interactive terminal or set FLOW1C_REDMINE_API_KEY in the agent environment.",
                previous_connection_preserved=True,
            )
    if not api_key:
        raise_redmine_transaction_error(
            "REDMINE_VALIDATION_FAILED", "Redmine API key is empty.",
            "Provide a valid key through the hidden prompt or FLOW1C_REDMINE_API_KEY.",
            previous_connection_preserved=True,
        )
    try:
        client = RedmineClient(base_url, api_key)
        client.current_user()
    except Exception:
        raise_redmine_transaction_error(
            "REDMINE_VALIDATION_FAILED",
            "Redmine did not verify the current user with the supplied API key.",
            "Check the HTTPS URL, API key, REST API setting, and account access; the previous connection is unchanged.",
            previous_connection_preserved=True,
        )

    previous_new_url_credential: str | None = None
    credential_snapshot_available = False
    if save_to_credential_manager:
        try:
            previous_new_url_credential = redmine_credentials.read_api_key(base_url)
            credential_snapshot_available = True
        except Exception:
            raise_redmine_transaction_error(
                "REDMINE_CREDENTIAL_WRITE_FAILED",
                "The existing credential state could not be checked, so no credential or configuration was changed.",
                "Check Windows Credential Manager access and retry; the previous connection is unchanged.",
                previous_connection_preserved=True,
            )

    transaction_id = uuid.uuid4().hex
    transaction_state = {
        "schema_version": 1,
        "transaction_id": transaction_id,
        "phase": "prepared",
        "old_url": previous_url,
        "old_url_invalid": previous_url_invalid,
        "new_url": base_url,
        "credential_managed": save_to_credential_manager,
        "credential_backup": save_to_credential_manager and previous_new_url_credential is not None,
    }
    try:
        write_json(redmine_transaction_state_path(), transaction_state)
    except Exception:
        raise_redmine_transaction_error(
            "REDMINE_CONFIG_WRITE_FAILED",
            "Redmine transaction recovery state could not be written; no credential or configuration was changed.",
            "Check `.workspace` permissions and retry; the previous connection is unchanged.",
            previous_connection_preserved=True,
        )
    backup_staged = False
    if transaction_state["credential_backup"]:
        try:
            redmine_credentials.write_transaction_backup(transaction_id, previous_new_url_credential or "")
            backup_staged = True
        except Exception:
            rollback_failures: list[str] = []
            try:
                redmine_credentials.delete_transaction_backup(transaction_id)
            except Exception:
                rollback_failures.append("The temporary credential backup could not be removed.")
            try:
                clear_redmine_transaction_state()
            except Exception:
                rollback_failures.append("The interrupted transaction record could not be removed.")
            raise_redmine_transaction_error(
                "REDMINE_CREDENTIAL_WRITE_FAILED",
                "The existing credential could not be staged safely for recovery.",
                "Check Windows Credential Manager access and retry; the previous connection was preserved.",
                previous_connection_preserved=True,
                rollback_failures=rollback_failures,
            )

    credential_write_attempted = False
    if save_to_credential_manager:
        credential_write_attempted = True
        try:
            redmine_credentials.write_api_key(base_url, api_key)
        except Exception:
            failures: list[str] = []
            credential_restored = False
            try:
                restore_redmine_credential(base_url, previous_new_url_credential)
                credential_restored = True
            except Exception:
                failures.append("The credential for the requested URL could not be restored or removed.")
            if credential_restored and backup_staged:
                try:
                    redmine_credentials.delete_transaction_backup(transaction_id)
                except Exception:
                    failures.append("The temporary credential backup could not be removed.")
            if credential_restored and not failures:
                try:
                    clear_redmine_transaction_state()
                except Exception:
                    failures.append("The interrupted transaction record could not be removed.")
            raise_redmine_transaction_error(
                "REDMINE_CREDENTIAL_WRITE_FAILED",
                "Windows Credential Manager could not save the new API key.",
                "Check Credential Manager access and retry; the previous connection was preserved.",
                previous_connection_preserved=previous_redmine_connection_preserved(previous_url, base_url, not credential_restored),
                rollback_failures=failures,
            )
        try:
            verified = redmine_credentials.verify_api_key(base_url, api_key)
        except Exception:
            verified = False
        if not verified:
            failures = []
            credential_restored = False
            try:
                restore_redmine_credential(base_url, previous_new_url_credential)
                credential_restored = True
            except Exception:
                failures.append("The credential for the requested URL could not be restored or removed.")
            if credential_restored and backup_staged:
                try:
                    redmine_credentials.delete_transaction_backup(transaction_id)
                except Exception:
                    failures.append("The temporary credential backup could not be removed.")
            if credential_restored and not failures:
                try:
                    clear_redmine_transaction_state()
                except Exception:
                    failures.append("The interrupted transaction record could not be removed.")
            raise_redmine_transaction_error(
                "REDMINE_CREDENTIAL_VERIFY_FAILED",
                "The saved API key could not be read back and verified.",
                "Check Credential Manager access and retry; the previous connection was preserved.",
                previous_connection_preserved=previous_redmine_connection_preserved(previous_url, base_url, not credential_restored),
                rollback_failures=failures,
            )

    updated_local = dict(local)
    updated_settings = dict(previous_settings)
    updated_settings["base_url"] = base_url
    updated_local["redmine"] = updated_settings
    config_changed = updated_local != local
    try:
        if config_changed:
            write_json(config_path, updated_local)
        transaction_state["phase"] = "activated"
        write_json(redmine_transaction_state_path(), transaction_state)
    except Exception:
        rollback_failures = []
        config_restored = False
        try:
            restore_redmine_config_snapshot(config_path, original_config)
            config_restored = True
        except Exception:
            rollback_failures.append("The original local configuration could not be restored.")
        credential_restored = not credential_write_attempted
        if config_restored and credential_write_attempted and credential_snapshot_available:
            try:
                restore_redmine_credential(base_url, previous_new_url_credential)
                credential_restored = True
            except Exception:
                rollback_failures.append("The newly written credential could not be restored or removed.")
        if config_restored and credential_restored and backup_staged:
            try:
                redmine_credentials.delete_transaction_backup(transaction_id)
            except Exception:
                rollback_failures.append("The temporary credential backup could not be removed.")
        if config_restored and credential_restored and not rollback_failures:
            try:
                clear_redmine_transaction_state()
            except Exception:
                rollback_failures.append("The interrupted transaction record could not be removed.")
        raise_redmine_transaction_error(
            "REDMINE_CONFIG_WRITE_FAILED",
            "The Redmine configuration could not be activated; rollback was attempted.",
            "Retry after checking local file permissions; inspect the configuration and Credential Manager if rollback failed.",
            previous_connection_preserved=not rollback_failures or (
                "The original local configuration could not be restored." not in rollback_failures
                and previous_url != base_url
            ),
            rollback_failures=rollback_failures,
        )

    warnings: list[dict[str, Any]] = list(recovery_warnings)
    pending_state_readable = True
    try:
        pending_urls = redmine_pending_cleanup_urls()
    except WorkflowError:
        pending_urls = []
        pending_state_readable = False
        warnings.append({
            "code": "REDMINE_CLEANUP_STATE_READ_FAILED",
            "message": "The new connection is active, but the previous credential cleanup record could not be read.",
            "recoverable": True,
            "next_action": "Repair `.workspace/redmine-pending-cleanup.json` and run `redmine cleanup`.",
        })
    if previous_url_invalid and raw_previous_url:
        warnings.append({
            "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
            "message": "The new connection is active, but the previous URL was invalid and its saved credential could not be identified for cleanup.",
            "recoverable": True,
            "next_action": "Review the previous Redmine credential in Windows Credential Manager after correcting its URL.",
        })
    elif previous_url and previous_url != base_url and sys.platform == "win32":
        try:
            redmine_credentials.delete_api_key(previous_url)
            pending_urls = [value for value in pending_urls if value != previous_url]
        except Exception:
            if previous_url not in pending_urls:
                pending_urls.append(previous_url)
            warnings.append({
                "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
                "message": "The new connection is active, but the previous saved credential could not be removed.",
                "recoverable": True,
                "next_action": "Run `redmine cleanup` to retry removal from the local recovery record.",
            })
    pending_state_saved = False
    try:
        if pending_state_readable:
            save_redmine_pending_cleanup_urls(pending_urls)
            pending_state_saved = True
    except Exception:
        if not any(item["code"] == "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED" for item in warnings):
            warnings.append({
                "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
                "message": "The new connection is active, but the cleanup recovery record could not be updated.",
                "recoverable": True,
                "next_action": "Review the old credential in Windows Credential Manager; do not repeat configure solely to repair the record.",
            })

    if pending_urls and pending_state_readable and not any(
        item.get("code") == "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED" for item in warnings
    ):
        warnings.append({
            "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
            "message": "One or more previous saved Redmine credentials are still pending cleanup.",
            "recoverable": True,
            "next_action": "Run `redmine cleanup` to retry removal from the local recovery record.",
        })

    keep_transaction = not pending_state_saved
    if backup_staged:
        try:
            redmine_credentials.delete_transaction_backup(transaction_id)
        except Exception:
            keep_transaction = True
            warnings.append({
                "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
                "message": "The new connection is active, but its temporary credential backup could not be removed.",
                "recoverable": True,
                "next_action": "Run a Redmine command again to retry transaction cleanup.",
            })
    if not keep_transaction:
        try:
            clear_redmine_transaction_state()
        except OSError:
            warnings.append({
                "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
                "message": "The new connection is active, but its transaction recovery file could not be removed.",
                "recoverable": True,
                "next_action": "Run a Redmine command again to retry transaction cleanup.",
            })

    state = "CONNECTED_WITH_WARNING" if warnings else "CONNECTED"
    print(json.dumps({
        "schema_version": 1,
        "state": state,
        "configured": True,
        "verified": True,
        "previous_connection_preserved": True,
        "warnings": warnings,
        "next_action": warnings[0]["next_action"] if warnings else None,
        "base_url": base_url,
        "credential_source": "environment" if not save_to_credential_manager else "windows-credential-manager",
        "next": "The agent can now fetch issues and import their supported attachments.",
    }, ensure_ascii=False, indent=2))
    return 0


def cmd_redmine_status(args: argparse.Namespace) -> int:
    recovery_warnings = recover_redmine_transaction()
    settings = redmine_settings()
    raw_url = str(settings.get("base_url", "")).strip()
    if not raw_url:
        result = {"state": "NOT_CONFIGURED", "configured": False, "next_action": "Run redmine configure with the Redmine HTTPS URL."}
    else:
        try:
            base_url = normalize_base_url(raw_url)
        except (RedmineError, RedminePolicyError) as exc:
            result = {"state": "INVALID", "configured": True, "error": str(exc)}
        else:
            try:
                _, source = redmine_api_key(base_url)
                result = {"state": "READY", "configured": True, "base_url": base_url, "credential_source": source}
            except WorkflowError as exc:
                result = {"state": "NEEDS_SECRET", "configured": True, "base_url": base_url, "error": str(exc)}
    warnings = list(recovery_warnings)
    try:
        pending_urls = redmine_pending_cleanup_urls()
    except WorkflowError:
        pending_urls = []
        warnings.append({
            "code": "REDMINE_CLEANUP_STATE_READ_FAILED",
            "message": "The Redmine credential cleanup record could not be read.",
            "recoverable": True,
            "next_action": "Repair `.workspace/redmine-pending-cleanup.json` and run `redmine cleanup`.",
        })
    if pending_urls and not any(item.get("code") == "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED" for item in warnings):
        warnings.append({
            "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
            "message": "One or more previous saved Redmine credentials are still pending cleanup.",
            "recoverable": True,
            "next_action": "Run `redmine cleanup` to retry removal from the local recovery record.",
        })
    if warnings:
        result["warnings"] = warnings
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["state"] == "READY" else 1


def cmd_redmine_test(args: argparse.Namespace) -> int:
    client, base_url, source = redmine_client_from_config()
    try:
        user = client.current_user()
    except RedmineError as exc:
        raise WorkflowError(str(exc)) from exc
    result = {"state": "CONNECTED", "base_url": base_url, "credential_source": source, "verified": True}
    if getattr(client, "recovery_warnings", None):
        result["warnings"] = client.recovery_warnings
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def redmine_issue_summary(issue: dict[str, Any], base_url: str) -> dict[str, Any]:
    description = sanitize_text(str(issue.get("description", "")))
    subject = sanitize_text(str(issue.get("subject", "")))
    description_limit = 12000
    subject_limit = 500
    return {
        "id": issue["id"],
        "url": f"{base_url}/issues/{issue['id']}",
        "subject": subject[:subject_limit],
        "subject_truncated": len(subject) > subject_limit,
        "description": description[:description_limit],
        "description_truncated": len(description) > description_limit,
        "status": sanitize_text(str(issue.get("status", "")))[:200],
        "project": sanitize_text(str(issue.get("project", "")))[:200],
    }


def redmine_standard_attachment_summaries(issue: dict[str, Any]) -> list[dict[str, Any]]:
    result = []
    for attachment in issue.get("attachments", []):
        if not isinstance(attachment, dict):
            continue
        result.append({
            "id": attachment.get("id"),
            "name": sanitize_text(str(attachment.get("filename") or "attachment"))[:500],
            "size": attachment.get("filesize") if isinstance(attachment.get("filesize"), int) else None,
            "mime_type": sanitize_text(str(attachment.get("content_type") or "application/octet-stream"))[:128],
        })
    return result


def redmine_dms_error(exc: RedmineError) -> dict[str, Any]:
    return {
        "code": exc.code,
        "message": sanitize_text(str(exc))[:1000],
        "recoverable": exc.code not in {"REDMINE_DMSF_INVALID_ID", "REDMINE_DMSF_NOT_ATTACHED"},
        "next_action": "Check the Redmine issue, DMSF permissions and file metadata, then retry the same selection.",
    }


def redmine_upload_error(exc: RedmineError) -> dict[str, Any]:
    """Return a stable, secret-free error record for the confirmed DMSF write path."""
    next_actions = {
        "REDMINE_DMSF_UPLOAD_NOT_CONFIRMED": "Ask for explicit confirmation of this exact issue and local file, then retry with --confirmed.",
        "REDMINE_DMSF_UPLOAD_FILE_NOT_FOUND": "Check the absolute local file path and select the intended ordinary file.",
        "REDMINE_DMSF_UPLOAD_UNSUPPORTED_TYPE": "Choose a file with an extension allowed by the DMSF upload policy.",
        "REDMINE_DMSF_UPLOAD_SIZE_LIMIT": "Choose a file no larger than 500 MiB and check the Redmine server limit.",
        "REDMINE_DMSF_UPLOAD_ACCESS_DENIED": "Check issue visibility and the API user's DMSF upload and issue-link permissions.",
        "REDMINE_DMSF_UNAVAILABLE": "Check that DMSF REST API and issue attachment support are enabled for this project.",
        "REDMINE_DMSF_ISSUE_ATTACH_UNAVAILABLE": "Enable DMSF attachments for this exact issue/project or upgrade the plugin before retrying.",
        "REDMINE_DMSF_UPLOAD_NAME_CONFLICT": "Choose a different file name or obtain a separate explicit revision decision.",
        "REDMINE_DMSF_COMMIT_UNCERTAIN": "Inspect Redmine and DMSF manually; do not retry automatically.",
        "REDMINE_DMSF_UPLOAD_VERIFICATION_FAILED": "Inspect the issue and DMSF manually; do not retry automatically.",
    }
    return {
        "code": exc.code,
        "message": sanitize_text(str(exc))[:1000],
        "next_action": next_actions.get(exc.code, "Check Redmine availability and permissions, then retry only when the result is known."),
    }


def cmd_redmine_upload(args: argparse.Namespace) -> int:
    try:
        issue_id = issue_number(args.issue)
    except RedminePolicyError as exc:
        raise WorkflowError(str(exc)) from exc
    requested_path = Path(str(args.file)).expanduser()
    try:
        absolute_path = requested_path.resolve(strict=False)
    except OSError:
        absolute_path = requested_path.absolute()
    issue_summary: dict[str, Any] = {"id": issue_id}
    project_id: int | None = None
    local_file: dict[str, Any] = {"path": str(absolute_path), "name": absolute_path.name, "size": None}
    try:
        client, base_url, _ = redmine_client_from_config(recover=False)
        issue = client.issue(issue_id)
        issue_summary = redmine_issue_summary(issue, base_url)
        project_id = issue.get("project_id") if isinstance(issue.get("project_id"), int) else None
        upload = client.upload_dmsf_file(issue, absolute_path, confirmed=bool(args.confirmed))
        preflight = upload.get("preflight", {}) if isinstance(upload, dict) else {}
        local_file = {
            "path": str(preflight.get("path", absolute_path)),
            "name": str(preflight.get("name", absolute_path.name)),
            "size": preflight.get("size"),
        }
        project_id = preflight.get("project_id", project_id)
        result: dict[str, Any] = {
            "schema_version": 1,
            "state": upload.get("state", "ERROR"),
            "issue": issue_summary,
            "project_id": project_id,
            "local_file": local_file,
            "dmsf_file": upload.get("dms_file"),
            "verification": upload.get("verification", {"issue_reloaded": False, "dmsf_attached": False}),
            "warnings": upload.get("warnings", []),
            "errors": upload.get("errors", []),
        }
        if result["state"] == "NEEDS_CONFIRMATION":
            result["next_action"] = "After explicit consent for this exact issue and file, repeat the command with --confirmed. No Redmine change was made."
            exit_code = 1
        elif result["state"] == "UPLOADED":
            result["next_action"] = None
            exit_code = 0
        else:
            result["next_action"] = upload.get("next_action") or "Inspect Redmine and DMSF manually; do not retry automatically."
            exit_code = 2
    except RedmineError as exc:
        result = {
            "schema_version": 1,
            "state": "ERROR",
            "issue": issue_summary,
            "project_id": project_id,
            "local_file": local_file,
            "dmsf_file": None,
            "verification": {"issue_reloaded": False, "dmsf_attached": False},
            "warnings": [],
            "errors": [redmine_upload_error(exc)],
            "next_action": redmine_upload_error(exc)["next_action"],
        }
        exit_code = 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return exit_code


def cmd_redmine_revise(args: argparse.Namespace) -> int:
    """Update the revision of one explicitly selected issue DMSF document."""
    try:
        issue_id = issue_number(args.issue)
        file_id = dmsf_identifier(args.dms_file, label="DMSF file ID")
        previous_id = dmsf_identifier(args.expected_revision, label="DMSF revision ID")
    except RedminePolicyError as exc:
        raise WorkflowError(str(exc)) from exc
    result: dict[str, Any] = {"schema_version": 1, "state": "ERROR",
                              "issue": {"id": issue_id}, "dmsf_file": {"id": file_id},
                              "local_file": {"path": str(Path(args.file).resolve(strict=False))}}
    try:
        client, base_url, _ = redmine_client_from_config(recover=False)
        issue = client.issue(issue_id)
        result["issue"] = redmine_issue_summary(issue, base_url)
        revision = client.revise_dmsf_file(issue, file_id, args.file,
                                           expected_revision_id=previous_id, confirmed=bool(args.confirmed))
        result.update({"state": revision["state"], "preflight": revision.get("preflight"),
                       "dmsf_file": revision.get("dms_file", {"id": file_id}),
                       "verification": revision.get("verification"), "errors": revision.get("errors", [])})
        result["next_action"] = ("Confirm the exact file, issue and existing DMSF ID before using --confirmed."
                                 if revision["state"] == "NEEDS_CONFIRMATION" else
                                 "Inspect the issue and DMSF manually; do not retry automatically."
                                 if revision["state"] == "COMMITTED_UNVERIFIED" else None)
        exit_code = 0 if revision["state"] == "REVISED" else 1
    except RedmineError as exc:
        result["errors"] = [redmine_upload_error(exc)]
        result["next_action"] = result["errors"][0]["next_action"]
        exit_code = 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return exit_code


def redmine_dms_inventory(client: Any, issue: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if hasattr(client, "dmsf_file_inventory"):
        try:
            files, failures = client.dmsf_file_inventory(issue)
            return files, [
                {
                    "code": str(item.get("code", "REDMINE_DMSF_INVALID_METADATA")),
                    "message": sanitize_text(str(item.get("message", "DMSF metadata could not be read")))[:1000],
                    "recoverable": True,
                    "next_action": "Check DMSF permissions and file metadata, then retry the same issue.",
                    **({"dms_file_id": item["dms_file_id"]} if item.get("dms_file_id") is not None else {}),
                }
                for item in failures
            ]
        except RedmineError as exc:
            return [], [redmine_dms_error(exc)]
    try:
        return (client.current_dmsf_files(issue), []) if hasattr(client, "current_dmsf_files") else ([], [])
    except RedmineError as exc:
        return [], [redmine_dms_error(exc)]


def cmd_redmine_files(args: argparse.Namespace) -> int:
    try:
        issue_id = issue_number(args.issue)
    except RedminePolicyError as exc:
        raise WorkflowError(str(exc)) from exc
    client, base_url, _ = redmine_client_from_config(recover=False)
    try:
        issue = client.issue(issue_id)
    except RedmineError as exc:
        result = {"schema_version": 1, "state": "ERROR", "issue": {"id": issue_id}, "errors": [redmine_dms_error(exc)]}
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 2
    standard = redmine_standard_attachment_summaries(issue)
    dms_files: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    dms_files, warnings = redmine_dms_inventory(client, issue)
    if any(item["code"] == "REDMINE_DMSF_UNAVAILABLE" for item in warnings):
        state = "DMSF_NOT_AVAILABLE"
    elif not standard and not dms_files:
        state = "ERROR" if warnings else "NO_FILES"
    elif dms_files:
        state = "SELECTION_REQUIRED"
    else:
        state = "FILES_AVAILABLE"
    result = {
        "schema_version": 1,
        "state": state,
        "issue": {"id": issue_id, "url": f"{base_url}/issues/{issue_id}"},
        "standard_attachments": standard,
        "dms_files": dms_files,
        "warnings": warnings,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if warnings else 0


def cmd_redmine_relations(args: argparse.Namespace) -> int:
    try:
        issue_id = issue_number(args.issue)
    except RedminePolicyError as exc:
        raise WorkflowError(str(exc)) from exc
    client, _, _ = redmine_client_from_config(recover=False)
    try:
        inventory = client.related_issues(issue_id)
    except RedmineError as exc:
        result = {
            "schema_version": 1, "state": "ERROR", "issue": {"id": issue_id},
            "relation_count": 0, "relations": [],
            "errors": [{"code": exc.code, "message": sanitize_text(str(exc))[:1000]}],
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 2
    partial = any(record["error"] is not None for record in inventory["relations"])
    result = {
        "schema_version": 1,
        "state": "PARTIAL" if partial else "COMPLETE",
        "issue": inventory["issue"],
        "relation_count": len(inventory["relations"]),
        "relations": inventory["relations"],
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if partial else 0


def cmd_redmine_disconnect(args: argparse.Namespace) -> int:
    recover_redmine_transaction()
    config_path, original_config, local = redmine_config_snapshot()
    settings = redmine_settings(local)
    raw_url = str(settings.get("base_url", "")).strip()
    had_redmine_section = "redmine" in local
    active_url: str | None = None
    invalid_url = False
    if raw_url:
        try:
            active_url = normalize_base_url(raw_url)
        except (RedmineError, RedminePolicyError):
            invalid_url = True
    local.pop("redmine", None)
    try:
        if had_redmine_section:
            write_json(config_path, local)
    except Exception:
        rollback_failures = []
        try:
            restore_redmine_config_snapshot(config_path, original_config)
        except Exception:
            rollback_failures.append("The original local configuration could not be restored.")
        raise_redmine_transaction_error(
            "REDMINE_CONFIG_WRITE_FAILED",
            "The Redmine section could not be removed; the saved credential was left untouched.",
            "Check local file permissions and retry redmine disconnect.",
            previous_connection_preserved=not rollback_failures,
            rollback_failures=rollback_failures,
        )

    pending_state_readable = True
    try:
        candidates = redmine_pending_cleanup_urls()
    except WorkflowError:
        candidates = []
        pending_state_readable = False
    if active_url and active_url not in candidates:
        candidates.append(active_url)
    failed: list[str] = []
    removed_any = False
    warnings: list[dict[str, Any]] = []
    if invalid_url:
        warnings.append({
            "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
            "message": "The local Redmine configuration was removed, but its invalid previous URL could not identify a credential to remove.",
            "recoverable": True,
            "next_action": "Review Windows Credential Manager for the Flow1C Redmine entry associated with the previous URL.",
        })
    for candidate in candidates if sys.platform == "win32" and pending_state_readable else []:
        try:
            removed_any = redmine_credentials.delete_api_key(candidate) or removed_any
        except Exception:
            failed.append(candidate)
    if failed:
        warnings.append({
            "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
            "message": "Redmine is disconnected, but one or more saved credentials could not be removed.",
            "recoverable": True,
            "next_action": "Run `redmine cleanup` to retry removal from the local recovery record.",
        })
    try:
        if pending_state_readable and sys.platform == "win32":
            save_redmine_pending_cleanup_urls(failed)
        elif not pending_state_readable:
            warnings.append({
                "code": "REDMINE_CLEANUP_STATE_READ_FAILED",
                "message": "Redmine is disconnected, but the pending credential cleanup record could not be read.",
                "recoverable": True,
                "next_action": "Repair `.workspace/redmine-pending-cleanup.json` and run `redmine cleanup`.",
            })
    except Exception:
        warnings.append({
            "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
            "message": "Redmine is disconnected, but the credential cleanup recovery record could not be updated.",
            "recoverable": True,
            "next_action": "Review Windows Credential Manager and remove any remaining Flow1C Redmine credential manually.",
        })
    if redmine_transaction_state_path().exists():
        warnings.append({
            "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
            "message": "Redmine is disconnected, but an interrupted configure transaction still needs recovery.",
            "recoverable": True,
            "next_action": "Run `redmine cleanup` on Windows to finish transaction recovery.",
        })
    env_key_present = bool(os.environ.get("FLOW1C_REDMINE_API_KEY", "").strip())
    if env_key_present:
        warnings.append({
            "code": "REDMINE_ENVIRONMENT_KEY_PRESERVED",
            "message": "FLOW1C_REDMINE_API_KEY remains managed by the agent environment and was not changed.",
            "recoverable": True,
            "next_action": "Remove FLOW1C_REDMINE_API_KEY from the agent environment separately if it should no longer be available.",
        })
    state = "DISCONNECTED_WITH_WARNING" if warnings else "DISCONNECTED"
    print(json.dumps({
        "schema_version": 1,
        "state": state,
        "configured": False,
        "verified": False,
        "previous_connection_preserved": False,
        "credential_removed": removed_any,
        "environment_key_cleared": False,
        "warnings": warnings,
        "next_action": warnings[0]["next_action"] if warnings else None,
    }, ensure_ascii=False, indent=2))
    return 0


def cmd_redmine_cleanup(args: argparse.Namespace) -> int:
    recovery_warnings = recover_redmine_transaction()
    requested_url = str(getattr(args, "url", "") or "").strip()
    try:
        candidates = redmine_pending_cleanup_urls()
        if requested_url:
            normalized = normalize_base_url(requested_url)
            if normalized not in candidates:
                candidates.append(normalized)
    except (RedmineError, RedminePolicyError) as exc:
        raise RedmineOperationError(
            "REDMINE_VALIDATION_FAILED", str(exc), recoverable=True,
            next_action="Provide a valid HTTPS Redmine URL to retry credential cleanup.",
            previous_connection_preserved=False,
        ) from None
    except WorkflowError:
        raise RedmineOperationError(
            "REDMINE_CLEANUP_STATE_READ_FAILED",
            "The pending Redmine credential cleanup record could not be read.",
            recoverable=True,
            next_action="Repair `.workspace/redmine-pending-cleanup.json` or pass `--url` to retry a specific credential.",
            previous_connection_preserved=bool(redmine_settings().get("base_url")),
        ) from None
    if sys.platform != "win32":
        warnings = [{
            "code": "REDMINE_CREDENTIAL_WRITE_FAILED",
            "message": "Windows Credential Manager cleanup is available only on Windows.",
            "recoverable": True,
            "next_action": "Run redmine cleanup on the Windows profile that owns the credential.",
        }]
        print(json.dumps({
            "schema_version": 1, "state": "CLEANUP_WITH_WARNING", "configured": bool(redmine_settings().get("base_url")),
            "credential_removed": 0, "warnings": warnings, "next_action": warnings[0]["next_action"],
        }, ensure_ascii=False, indent=2))
        return 1
    failed: list[str] = []
    removed = 0
    warnings: list[dict[str, Any]] = list(recovery_warnings)
    for candidate in candidates:
        try:
            removed += int(bool(redmine_credentials.delete_api_key(candidate)))
        except Exception:
            failed.append(candidate)
    try:
        save_redmine_pending_cleanup_urls(failed)
    except Exception:
        warnings.append({
            "code": "REDMINE_CLEANUP_STATE_WRITE_FAILED",
            "message": "Credential cleanup ran, but the recovery record could not be updated.",
            "recoverable": True,
            "next_action": "Retry redmine cleanup after checking local file permissions.",
        })
    if failed:
        warnings.append({
            "code": "REDMINE_OLD_CREDENTIAL_CLEANUP_FAILED",
            "message": "One or more saved Redmine credentials remain and are recorded for retry.",
            "recoverable": True,
            "next_action": "Check Windows Credential Manager access and retry redmine cleanup.",
        })
    elif not redmine_transaction_state_path().exists():
        warnings = []
    print(json.dumps({
        "schema_version": 1,
        "state": "CLEANUP_WITH_WARNING" if warnings else "CLEANUP_COMPLETE",
        "configured": bool(redmine_settings().get("base_url")),
        "credential_removed": removed,
        "warnings": warnings,
        "next_action": warnings[0]["next_action"] if warnings else None,
    }, ensure_ascii=False, indent=2))
    return 1 if warnings else 0


def cmd_redmine_fetch(args: argparse.Namespace) -> int:
    try:
        issue_id = issue_number(args.issue)
    except RedminePolicyError as exc:
        raise WorkflowError(str(exc)) from exc
    supplied_code = str(getattr(args, "code", "") or "").strip() or None
    gate_id = str(getattr(args, "gate_id", "") or "").strip() or None
    code: str | None = None
    if gate_id:
        gate = load_gate(gate_id, states={
            "NEEDS_INPUT", "NEEDS_CONFIRMATION", "WAITING_USER", "READY",
            "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT", "BLOCKED",
        })
        require_gate_tool(gate, "flow1c_redmine_fetch")
        code = str(gate.get("work_reference") or gate.get("code") or "").strip() or None
        if supplied_code and supplied_code != code:
            raise WorkflowError("Redmine target does not match the active gate work item")
    elif supplied_code:
        raise WorkflowError("Importing Redmine attachments into a work item requires an active gate_id")
    if code:
        load_manifest(code)
    selected_values = list(getattr(args, "dms_file", []) or [])
    all_dms = bool(getattr(args, "all_dms", False))
    revision_id = getattr(args, "dms_revision", None)
    if all_dms and selected_values:
        raise WorkflowError("--all-dms cannot be combined with --dms-file")
    if revision_id is not None and len(selected_values) != 1:
        raise WorkflowError("--dms-revision requires exactly one --dms-file")
    try:
        selected_ids = [dmsf_identifier(value) for value in selected_values]
        revision_id = dmsf_identifier(revision_id, label="DMSF revision ID") if revision_id is not None else None
    except RedminePolicyError as exc:
        raise WorkflowError(str(exc)) from exc
    if len(set(selected_ids)) != len(selected_ids):
        raise WorkflowError("A DMSF file may be selected only once.")
    client, base_url, credential_source = redmine_client_from_config()
    try:
        issue = client.issue(issue_id)
    except RedmineError as exc:
        raise WorkflowError(str(exc)) from exc
    issue_summary = redmine_issue_summary(issue, base_url)
    dms_files: list[dict[str, Any]] = []
    dms_warnings: list[dict[str, Any]] = []
    dms_files, dms_warnings = redmine_dms_inventory(client, issue)
    active_by_id = {item.get("id"): item for item in dms_files}
    if selected_ids and any(file_id not in active_by_id for file_id in selected_ids):
        missing = next(file_id for file_id in selected_ids if file_id not in active_by_id)
        error = next((item for item in dms_warnings if item.get("dms_file_id") == missing), None)
        if error is None:
            error = next((item for item in dms_warnings if item.get("dms_file_id") is None), None)
        if error is None:
            error = redmine_dms_error(RedmineError(
                f"DMSF file {missing} is not currently attached to issue {issue_id}.", code="REDMINE_DMSF_NOT_ATTACHED"
            ))
        print(json.dumps({"schema_version": 1, "state": "ERROR", "issue": issue_summary, "errors": [error]}, ensure_ascii=False, indent=2))
        return 2
    if revision_id is not None:
        selected_metadata = active_by_id[selected_ids[0]]
        available_revisions = {item.get("id") for item in selected_metadata.get("revisions", []) if isinstance(item, dict)}
        if revision_id not in available_revisions:
            error = RedmineError("Requested DMSF revision is not available for this file.", code="REDMINE_DMSF_INVALID_METADATA")
            print(json.dumps({"schema_version": 1, "state": "ERROR", "issue": issue_summary, "errors": [redmine_dms_error(error)]}, ensure_ascii=False, indent=2))
            return 2
    requested_dms = list(active_by_id.values()) if all_dms else [active_by_id[file_id] for file_id in selected_ids]
    if not issue["attachments"] and not requested_dms:
        state = (
            "DMSF_NOT_AVAILABLE" if any(item.get("code") == "REDMINE_DMSF_UNAVAILABLE" for item in dms_warnings)
            else "ERROR" if dms_warnings else "SELECTION_REQUIRED" if dms_files else "NO_ATTACHMENTS"
        )
        result = {"state": state, "issue": issue_summary, "copied": [], "dms_files": dms_files, "warnings": dms_warnings}
        if getattr(client, "recovery_warnings", None):
            result["warnings"].extend(client.recovery_warnings)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 1 if dms_warnings else 0
    workspace = ROOT / ".workspace"
    try:
        workspace.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="redmine-", dir=workspace) as temporary:
            temp_root = Path(temporary)
            standard_root = temp_root / "standard"
            dms_root = temp_root / "dms"
            standard_root.mkdir()
            dms_root.mkdir()
            _, attachment_metadata = client.download_attachments(issue, standard_root) if issue["attachments"] else ([], [])
            dms_downloads: list[tuple[Path, dict[str, Any], dict[str, Any]]] = []
            remaining_bytes = max(0, MAX_INTAKE_BYTES - sum(item.get("size", 0) for item in attachment_metadata if isinstance(item.get("size"), int)))
            for dms_metadata in requested_dms:
                try:
                    fresh_issue = client.issue(issue_id)
                    if not client.is_dmsf_attached(fresh_issue, dms_metadata["id"]):
                        raise RedmineError(f"DMSF file {dms_metadata['id']} is no longer attached to issue {issue_id}.", code="REDMINE_DMSF_NOT_ATTACHED")
                    path, downloaded = client.download_dmsf_file(
                        dms_metadata, dms_root, revision_id=revision_id if len(requested_dms) == 1 else None,
                        max_bytes=remaining_bytes,
                    )
                    remaining_bytes -= downloaded["size"]
                    dms_downloads.append((path, downloaded, dms_metadata))
                except RedmineError as exc:
                    dms_warnings.append({**redmine_dms_error(exc), "dms_file_id": dms_metadata.get("id")})
            files, skipped_paths = enumerate_intake_files([standard_root])
            unsupported_names = sorted({Path(item).name for item in skipped_paths})
            if not files and not dms_downloads:
                if dms_warnings:
                    print(json.dumps({
                        "schema_version": 1, "state": "ERROR", "issue": issue_summary, "copied": [],
                        "dms_files": dms_files, "warnings": dms_warnings, "errors": dms_warnings,
                        "skipped_unsupported": unsupported_names,
                    }, ensure_ascii=False, indent=2))
                    return 2
                if dms_files:
                    print(json.dumps({"schema_version": 1, "state": "SELECTION_REQUIRED", "issue": issue_summary,
                                      "copied": [], "dms_files": dms_files, "warnings": [],
                                      "skipped_unsupported": unsupported_names}, ensure_ascii=False, indent=2))
                    return 0
                result = {
                    "state": "NO_SUPPORTED_ATTACHMENTS", "issue": issue_summary,
                    "downloaded": attachment_metadata, "skipped_unsupported": unsupported_names,
                    "supported_types": sorted(ALLOWED_ARTIFACT_EXTENSIONS),
                }
                if getattr(client, "recovery_warnings", None):
                    result["warnings"] = client.recovery_warnings
                print(json.dumps(result, ensure_ascii=False, indent=2))
                return 1
            source_record = {"system": "redmine", "base_url": base_url, "issue_id": issue_id, "issue_url": f"{base_url}/issues/{issue_id}"}
            result = {"state": "ACCEPTED", "intake_id": None, "code": code, "destination": None, "copied": [], "skipped": []}
            if files:
                result = persist_intake_files(
                    files, [], code=code, category="redmine_attachments", received_via="redmine", source_record=source_record,
                )
            dms_copied: list[dict[str, Any]] = []
            for path, downloaded, metadata in dms_downloads:
                provenance = {
                    **source_record,
                    "attachment_type": "dmsf",
                    "dmsf_file_id": metadata["id"],
                    "dmsf_revision_id": downloaded.get("revision", {}).get("id"),
                    "dmsf_project_id": metadata["project_id"],
                    "source_url": downloaded.get("download_url") or (
                        f"{base_url}/dmsf/files/{metadata['id']}/view?download={downloaded['revision']['id']}"
                        if revision_id is not None and len(requested_dms) == 1
                        else f"{base_url}/dmsf/files/{metadata['id']}/download"
                    ),
                }
                item_result = persist_intake_files(
                    [(path, Path(path.name))], [], code=code, category="redmine_attachments",
                    received_via="redmine", source_record=provenance,
                )
                dms_copied.extend(item_result["copied"])
                result["skipped"].extend(item_result["skipped"])
                if not result.get("destination"):
                    result.update({key: item_result[key] for key in ("intake_id", "destination")})
            result["copied"].extend(dms_copied)
            downloaded_all = attachment_metadata + [item for _, item, _ in dms_downloads]
            existing_names = result["skipped"]
            result["skipped"] = sorted(unsupported_names + existing_names)
            result.update(
                issue=issue_summary,
                downloaded=downloaded_all, dms_files=dms_files, credential_source=credential_source,
                skipped_unsupported=unsupported_names,
                skipped_existing=existing_names,
                warnings=dms_warnings,
            )
            if getattr(client, "recovery_warnings", None):
                result["warnings"].extend(client.recovery_warnings)
            if dms_warnings:
                result["state"] = "IMPORTED_WITH_WARNINGS" if result["copied"] else "ERROR"
            elif unsupported_names:
                result["state"] = "PARTIAL"
            elif dms_downloads:
                result["state"] = "IMPORTED"
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 0 if not unsupported_names and not dms_warnings else 1
    except RedmineError as exc:
        print(json.dumps({"schema_version": 1, "state": "ERROR", "issue": issue_summary, "errors": [redmine_dms_error(exc)]}, ensure_ascii=False, indent=2))
        return 2
    except OSError as exc:
        raise WorkflowError(f"Could not download or import Redmine attachments safely: {exc}") from exc


def rlm_readiness() -> tuple[bool, list[str]]:
    config, local = load_config()
    quality = config.get("quality", {})
    endpoint = str(quality.get("rlm_endpoint", "")).strip()
    errors: list[str] = []
    if not endpoint_is_healthy(endpoint):
        errors.append("RLM MCP endpoint is not healthy")
        return False, errors
    index_command = str(local.get("rlm", {}).get("index_command", "")) if isinstance(local.get("rlm"), dict) else ""
    for label, raw in (("configuration", local.get("configuration_path")), ("extension", local.get("extension_path"))):
        path = resolve_1c_source_root(Path(str(raw))) if raw else None
        if path is None:
            errors.append(f"RLM {label} source is not configured")
            continue
        state, detail = rlm_index_state(index_command, path)
        if state != "OK":
            errors.append(f"RLM {label} index: {detail}")
        endpoint_state, endpoint_detail = rlm_endpoint_index_state(endpoint, path)
        if endpoint_state != "OK":
            errors.append(f"RLM {label} MCP index: {endpoint_detail}")
    return not errors, errors


def build_input_message(found: list[dict[str, str]], required: list[dict[str, str]], conditional: list[dict[str, str]]) -> str:
    missing = required or conditional
    label = missing[0]["label"] if missing else "описание задачи"
    return (
        f"Для формального этапа нужно уточнить: {label}. "
        "Предоставим материал или продолжим независимым черновиком из описания в чате? "
        "Можно приложить файлы в чат или указать папку. Полный перечень сохранён в диагностике."
    )


def read_setup_state(profile: str = "analysis") -> dict[str, Any]:
    powershell = command_path("powershell") or command_path("pwsh")
    if not powershell:
        raise WorkflowError("PowerShell is not available; setup cannot be audited.")
    result = subprocess.run(
        [powershell, "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "setup-state.ps1"), "-Json", "-Profile", profile],
        cwd=ROOT,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
        timeout=120,
    )
    try:
        value = json.loads(result.stdout.lstrip("\ufeff"))
    except json.JSONDecodeError as exc:
        if result.returncode != 0:
            raise WorkflowError((result.stderr or result.stdout or "setup-state failed").strip()) from exc
        raise WorkflowError(f"setup-state returned invalid JSON: {exc}") from exc
    if not isinstance(value, dict):
        raise WorkflowError("setup-state returned an invalid payload.")
    return value


def setup_user_message(state: dict[str, Any]) -> str:
    if state.get("state") == "READY" and state.get("ready"):
        return f"Профиль {state.get('requested_profile', 'conversation')} готов; можно продолжать работу."
    if state.get("state") == "BLOCKED":
        errors = state.get("errors", []) if isinstance(state.get("errors"), list) else []
        details = "; ".join(str(item.get("message", "")) for item in errors if isinstance(item, dict))
        actions = state.get("next_actions", []) if isinstance(state.get("next_actions"), list) else []
        return "Аудит setup заблокирован: " + (details or "неизвестная ошибка") + ((" Следующий шаг: " + str(actions[0])) if actions else "")
    prerequisites = state.get("prerequisites", {}) if isinstance(state.get("prerequisites"), dict) else {}
    bootstrap = state.get("bootstrap", {}) if isinstance(state.get("bootstrap"), dict) else {}
    lines = ["Проверка окружения выполнена."]
    if prerequisites.get("missing"):
        lines.append("Отсутствуют обязательные компоненты: " + ", ".join(map(str, prerequisites["missing"])))
        lines.append("Подтвердите установку через setup-bootstrap либо укажите абсолютные пути к одобренным offline-установщикам.")
    elif bootstrap.get("missing"):
        lines.append("Требуется bootstrap: " + ", ".join(map(str, bootstrap["missing"])))
        lines.append("Подтвердите установку зависимостей выбранного профиля.")
    else:
        lines.append("Локальные зависимости готовы. Подтвердите значения внешних репозиториев, каталогов, Gitea и Git identity.")
    confirmations = state.get("required_confirmations", [])
    if confirmations:
        lines.append("Обязательные подтверждения: " + ", ".join(map(str, confirmations)))
    optional_inputs = state.get("optional_inputs", [])
    if optional_inputs:
        lines.append(
            "Необязательные значения: " + ", ".join(map(str, optional_inputs))
            + ". Шаблон DOCX можно подключить позже перед формальной подготовкой ТЗ."
        )
    lines.append("Пути передавайте абсолютными; существующие значения считаются только подсказками до явного подтверждения.")
    return "\n".join(lines)


def new_gate(operation: str, code: str | None, state: str, **values: Any) -> dict[str, Any]:
    gate = {
        "schema_version": AGENT_GATE_SCHEMA_VERSION,
        "policy_version": POLICY_VERSION,
        "mode": "formal",
        "gate_id": str(uuid.uuid4()),
        "operation": operation,
        "code": code,
        "compliance": "COMPLIANT",
        "deviations": [],
        "state": state,
        "created_at": utc_now(),
        **values,
    }
    stage = load_stages()["operations"].get(operation, {})
    mode = str(gate.get("mode", "formal"))
    gate.setdefault("conditions", [])
    gate.setdefault("remaining_blockers", [])
    gate.setdefault("available_actions", available_actions(mode=mode, operation=operation, state=state, stage=stage))
    gate.setdefault("clarification", None)
    if state in {"NEEDS_CODE", "NEEDS_INPUT", "NEEDS_CONFIRMATION", "BLOCKED"} and not gate["clarification"]:
        gate["clarification"] = {"reason": state.lower(), "question": gate.get("user_message", "Уточните следующий шаг."),
                                 "options": ["Предоставить данные", "Обсудить отдельно", "Независимый черновик"]}
    gate.setdefault("awaiting_user_input", bool(gate.get("clarification") and state in {
        "NEEDS_CODE", "NEEDS_INPUT", "NEEDS_CONFIRMATION", "WAITING_USER", "BLOCKED",
    }))
    refresh_gate_actions(gate)
    save_gate(gate)
    return gate


def request_root(gate: dict[str, Any]) -> Path:
    request_id = str(gate.get("request_id", gate["gate_id"]))
    if not re.fullmatch(r"[a-f0-9-]{8,64}", request_id):
        raise WorkflowError("Invalid request ID")
    local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
    # A request started before setup must remain readable after setup.
    if gate.get("storage_kind") == "workspace":
        base = ROOT / ".workspace" / "drafts"
    else:
        base = project_root(local) / "drafts" if local.get("documentation_path") else ROOT / ".workspace" / "drafts"
    return (base / request_id).resolve()


# Functional-specification sections are intentionally implemented as a small
# vertical slice outside the formal gate. They can be used from a natural
# language draft without a work-item, while the approval and DOCX commands
# remain explicit and independently auditable.
def _sections_catalog() -> dict[str, Any]:
    return load_catalog(ROOT / "standards" / "functional-specification-sections.md")


def _sections_root(work_reference: str | None) -> Path:
    local = read_json(ROOT / LOCAL_CONFIG_FILE, {}) or {}
    configured = project_root(local) if local.get("documentation_path") else None
    # A stale or temporarily unavailable documentation checkout must not make
    # an independent draft impossible; use the documented local fallback.
    if configured is not None and configured.exists():
        return configured
    return ROOT if work_reference else ROOT / ".workspace"


def _section_args(args: argparse.Namespace) -> tuple[str | None, str | None]:
    request_id = str(getattr(args, "request_id", "") or "").strip() or None
    work_reference = str(getattr(args, "work_reference", "") or "").strip() or None
    if not request_id and not work_reference:
        raise WorkflowError("Specify request_id for an independent draft or work_reference for a work-item section.")
    return request_id, work_reference


def _resolve_section_source(raw: str) -> Path:
    value = str(raw or "").strip()
    if not value:
        raise DocxError("DOCX_SOURCE_REQUIRED", "Источник DOCX не указан явно.", next_action="Укажите приложенный или явно выбранный путь к .docx.")
    local = read_json(ROOT / LOCAL_CONFIG_FILE, {}) or {}
    roots = [ROOT.resolve(), project_root(local).resolve()]
    roots.extend(Path(item).expanduser().resolve() for item in local.get("allowed_external_paths", []) if str(item).strip())
    candidate = Path(value).expanduser()
    candidate = (Path.cwd() / candidate).resolve() if not candidate.is_absolute() else candidate.resolve()
    if not any(path_is_within(candidate, root) for root in roots):
        raise DocxError("DOCX_PATH_FORBIDDEN", f"Путь DOCX находится вне разрешенных корней: {candidate}", next_action="Скопируйте файл в request/work-item или настройте разрешенный внешний корень.")
    if candidate.is_symlink():
        raise DocxError("DOCX_PATH_FORBIDDEN", "Симлинки для источника DOCX запрещены.", next_action="Укажите обычный файл внутри разрешенного корня.")
    return candidate


def _section_records(args: argparse.Namespace) -> list[dict[str, Any]]:
    requested = list(getattr(args, "section_id", []) or [])
    if not requested:
        requested = list(getattr(args, "sections", []) or [])
    if isinstance(requested, str):
        requested = [requested]
    if not requested:
        raise SectionPolicyError("SECTION_UNKNOWN", "Не указан section_id.", next_action="Вызовите section-catalog и укажите каноническое имя.")
    return resolve_section_names([str(item) for item in requested], _sections_catalog())


def _require_section_tool(args: argparse.Namespace, tool_name: str) -> None:
    gate_id = str(getattr(args, "gate_id", "") or "").strip()
    if not gate_id:
        return
    require_gate_tool(load_gate(gate_id), tool_name)


def cmd_section_catalog(args: argparse.Namespace) -> int:
    _require_section_tool(args, "flow1c_section")
    catalog = _sections_catalog()
    print(json.dumps({
        "schema_version": 1, "state": "READY", "sections": catalog_summary(catalog),
        "contracts": {section["section_id"]: section for section in catalog["sections"]},
    }, ensure_ascii=False, indent=2))
    return 0


def cmd_section_save(args: argparse.Namespace) -> int:
    _require_section_tool(args, "flow1c_section")
    catalog = _sections_catalog()
    records = _section_records(args)
    if len(records) != 1:
        raise SectionPolicyError("SECTION_AMBIGUOUS", "section-save сохраняет одну секцию за операцию.", next_action="Вызовите section-save отдельно для каждого section_id.")
    request_id, work_reference = _section_args(args)
    content = sys.stdin.read() if getattr(args, "content_stdin", False) else str(getattr(args, "content", "") or "")
    if not content.strip():
        raise WorkflowError("Section content must be non-empty")
    checklist = getattr(args, "checklist", []) or []
    if isinstance(checklist, str):
        checklist = json.loads(checklist)
    sources = getattr(args, "sources", []) or []
    result = save_section(_sections_root(work_reference), section=records[0], content=content, checklist=checklist,
                          sources=sources, request_id=request_id, work_reference=work_reference)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_section_approve(args: argparse.Namespace) -> int:
    _require_section_tool(args, "flow1c_section")
    records = _section_records(args)
    if len(records) != 1:
        raise SectionPolicyError("SECTION_AMBIGUOUS", "section-approve фиксирует одну секцию за операцию.")
    request_id, work_reference = _section_args(args)
    result = approve_section(_sections_root(work_reference), section=records[0],
                             approved_by=str(getattr(args, "approved_by", "user") or "user"),
                             approval_statement=str(getattr(args, "approval_statement", "") or ""),
                             request_id=request_id, work_reference=work_reference)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def _docx_inspection(args: argparse.Namespace) -> tuple[Path, list[dict[str, Any]], dict[str, Any]]:
    source = _resolve_section_source(getattr(args, "source", "") or getattr(args, "path", ""))
    sections = _section_records(args)
    inspection = inspect_docx(source, sections)
    return source, sections, inspection


def cmd_docx_inspect(args: argparse.Namespace) -> int:
    _require_section_tool(args, "flow1c_docx")
    _, _, inspection = _docx_inspection(args)
    print(json.dumps(inspection, ensure_ascii=False, indent=2))
    return 0


def _load_section_current(root: Path, section: dict[str, Any], request_id: str | None, work_reference: str | None) -> tuple[Path, str, dict[str, Any]]:
    content_path, state_path, content, state = load_state(root, request_id=request_id, work_reference=work_reference,
                                                          section_id=section["section_id"], catalog_section=section)
    return state_path, content, state


def _plan_directory(root: Path) -> Path:
    return (root / "docx-plans") if root.name.casefold() == ".workspace" else (root / ".workspace" / "docx-plans")


def _plan_path(raw: str, root: Path) -> Path:
    plan_dir = _plan_directory(root).resolve()
    if raw:
        path = Path(raw).expanduser().resolve()
        if not path_is_within(path, plan_dir):
            raise DocxError("DOCX_PATH_FORBIDDEN", f"План находится вне управляемого каталога: {path}")
        return path
    plan_dir.mkdir(parents=True, exist_ok=True)
    return plan_dir / f"{uuid.uuid4()}.json"


def _write_plan_json(path: Path, value: dict[str, Any]) -> None:
    write_json(path, value)


def cmd_docx_write_plan(args: argparse.Namespace) -> int:
    _require_section_tool(args, "flow1c_docx")
    source, sections, inspection = _docx_inspection(args)
    request_id, work_reference = _section_args(args)
    root = _sections_root(work_reference)
    modes: dict[str, str] = {}
    raw_modes = getattr(args, "modes", "") or ""
    if raw_modes:
        modes = json.loads(raw_modes) if isinstance(raw_modes, str) else dict(raw_modes)
    default_mode = str(getattr(args, "mode", "") or "").strip()
    by_id = {item["section_id"]: item for item in inspection["sections"]}
    operations = []
    for section in sections:
        item = by_id[section["section_id"]]
        status = item["status"]
        if status in {"NOT_FOUND", "AMBIGUOUS"}:
            code = "TARGET_SECTION_NOT_FOUND" if status == "NOT_FOUND" else "TARGET_SECTION_AMBIGUOUS"
            raise DocxError(code, f"Целевой раздел {section['display_name']} имеет статус {status}.", next_action="Уточните место раздела и повторите инспекцию.")
        mode = str(modes.get(section["section_id"], default_mode) or ("replace" if status in {"FOUND_EMPTY", "FOUND_PLACEHOLDER"} else "")).strip()
        if status == "FOUND_CONTENT" and mode not in {"replace", "append"}:
            raise DocxError("REPLACE_MODE_REQUIRED", f"Раздел {section['display_name']} уже содержит текст.", next_action="Покажите существующее содержимое и явно выберите replace или append.")
        state_path, content, state = _load_section_current(root, section, request_id, work_reference)
        validate_section_content(section["section_id"], content)
        if state.get("state") != "APPROVED":
            raise SectionPolicyError("DRAFT_NOT_APPROVED", f"Раздел {section['section_id']} не согласован.", state=state)
        current_hash = hashlib.sha256(content.replace("\r\n", "\n").encode("utf-8")).hexdigest()
        if current_hash != state.get("content_sha256") or current_hash != state.get("approved_content_sha256"):
            raise SectionPolicyError("APPROVAL_STALE", f"Согласование {section['section_id']} устарело.", next_action="Сохраните новую версию и согласуйте ее заново.", state=state)
        operations.append({"section_id": section["section_id"], "mode": mode,
                           "content_sha256": current_hash})
    output = str(getattr(args, "output", "") or "").strip()
    output_path = Path(output).expanduser().resolve() if output else source.with_name(f"{source.stem}-flow1c-filled.docx")
    if not (path_is_within(output_path, source.parent) or path_is_within(output_path, root.resolve())):
        raise DocxError("DOCX_PATH_FORBIDDEN", f"Выходной DOCX вне разрешенного каталога: {output_path}", next_action="Выберите результат рядом с источником или внутри request/work-item.")
    plan = {"schema_version": 1, "state": "READY", "source": str(source), "source_sha256": inspection["source_sha256"], "output": str(output_path), "request_id": request_id, "work_reference": work_reference, "sections": operations, "write_command_required": True}
    plan_path = _plan_path(str(getattr(args, "plan", "") or ""), root)
    _write_plan_json(plan_path, plan)
    append_audit(root, {"action": "docx-write-plan", "plan": str(plan_path), "source": str(source), "source_sha256": plan["source_sha256"], "section_ids": [item["section_id"] for item in sections]})
    print(json.dumps({"state": "READY", "plan": str(plan_path), "source": str(source), "source_sha256": plan["source_sha256"], "output": str(output_path), "section_ids": [item["section_id"] for item in sections], "write_command_required": True}, ensure_ascii=False, indent=2))
    return 0


def cmd_docx_write(args: argparse.Namespace) -> int:
    _require_section_tool(args, "flow1c_docx")
    raw_plan = str(getattr(args, "plan", "") or "").strip()
    if not raw_plan:
        raise DocxError("WRITE_COMMAND_REQUIRED", "Для docx-write требуется подтвержденный plan.", next_action="Сначала вызовите docx-write-plan.")
    plan_path = Path(raw_plan).expanduser().resolve()
    plan = read_json(plan_path)
    if not isinstance(plan, dict) or plan.get("state") not in {"READY", "WRITTEN"}:
        raise DocxError("DRAFT_NOT_READY", "План записи отсутствует или уже недействителен.", next_action="Создайте новый docx-write-plan.")
    root = _sections_root(plan.get("work_reference"))
    if not path_is_within(plan_path, _plan_directory(root).resolve()):
        raise DocxError("DOCX_PATH_FORBIDDEN", f"План находится вне управляемого каталога: {plan_path}",
                        next_action="Используйте plan, созданный командой docx-write-plan.")
    if plan.get("state") == "WRITTEN" and isinstance(plan.get("result"), dict):
        existing_result = plan["result"]
        existing_output = Path(str(existing_result.get("output") or plan.get("output") or "")).expanduser()
        expected_output_hash = str(existing_result.get("output_sha256") or "")
        if existing_output.is_file() and expected_output_hash and sha256_file(existing_output) == expected_output_hash:
            print(json.dumps({**existing_result, "state": "NO_CHANGE"}, ensure_ascii=False, indent=2))
            return 0
        raise DocxError("DOCX_WRITE_FAILED", "Проверенный результат из plan отсутствует.", next_action="Создайте новый plan; исходный DOCX не изменен.", preserved_state=plan)
    current_source = _resolve_section_source(str(plan.get("source") or ""))
    if not current_source.is_file() or sha256_file(current_source) != plan.get("source_sha256"):
        raise DocxError("DOCX_WRITE_FAILED", "Исходный DOCX изменился или исчез после plan.", next_action="Повторите docx-inspect и docx-write-plan.", preserved_state=plan)
    operations = []
    states: list[tuple[Path, dict[str, Any]]] = []
    plan_sections = plan.get("sections", [])
    if not isinstance(plan_sections, list) or not plan_sections:
        raise DocxError("DRAFT_NOT_READY", "План не содержит целевых разделов.", next_action="Создайте новый docx-write-plan.")
    section_ids = [str(item.get("section_id") or "") for item in plan_sections if isinstance(item, dict)]
    if len(section_ids) != len(plan_sections) or len(set(section_ids)) != len(section_ids):
        raise DocxError("DRAFT_NOT_READY", "План содержит некорректный список разделов.", next_action="Создайте новый docx-write-plan.")
    canonical_sections = {item["section_id"]: item for item in resolve_section_names(section_ids, _sections_catalog())}
    output_path = Path(str(plan.get("output") or "")).expanduser().resolve()
    if output_path == current_source or not (path_is_within(output_path, current_source.parent) or path_is_within(output_path, root.resolve())):
        raise DocxError("DOCX_PATH_FORBIDDEN", f"Выходной DOCX вне разрешенного каталога: {output_path}",
                        next_action="Создайте новый plan с результатом рядом с источником или внутри request/work-item.")
    for item in plan_sections:
        section = canonical_sections[item["section_id"]]
        if item.get("mode") not in {"replace", "append"}:
            raise DocxError("REPLACE_MODE_REQUIRED", f"План содержит недопустимый режим для {section['section_id']}.", next_action="Создайте новый docx-write-plan.")
        state_path, content, state = _load_section_current(root, section, plan.get("request_id"), plan.get("work_reference"))
        validate_section_content(section["section_id"], content)
        # This is the second, explicit authorization point. Merely saving an
        # approval never reaches this function.
        ensure_write_authorized(state, content, write_command=True)
        if content.replace("\r\n", "\n").encode("utf-8") and item["content_sha256"] != hashlib.sha256(content.replace("\r\n", "\n").encode("utf-8")).hexdigest():
            raise SectionPolicyError("APPROVAL_STALE", f"Текст {section['section_id']} изменился после plan.", next_action="Создайте новый plan.", state=state)
        states.append((state_path, state))
        operations.append({"section": section, "mode": item["mode"], "body_index": 0, "content": content})
    for state_path, state in states:
        state["state"] = "WRITE_PENDING"
        write_json(state_path, state)
    try:
        result = write_docx(current_source, output_path, operations, expected_source_sha256=plan["source_sha256"])
    except (DocxError, OSError) as exc:
        for state_path, state in states:
            state["state"] = "WRITE_FAILED"
            state["write_error"] = getattr(exc, "code", "DOCX_WRITE_FAILED")
            write_json(state_path, state)
        append_audit(root, {"action": "docx-write-failed", "plan": str(plan_path), "code": getattr(exc, "code", "DOCX_WRITE_FAILED")})
        raise
    for state_path, state in states:
        state["state"] = "WRITTEN"
        state["written_output"] = result["output"]
        state["output_sha256"] = result["output_sha256"]
        write_json(state_path, state)
    plan["state"] = "WRITTEN"
    plan["result"] = result
    write_json(plan_path, plan)
    append_audit(root, {"action": "docx-write", "plan": str(plan_path), **result})
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def reference_mismatch(code: str | None, requirements: list[str]) -> str:
    if not code:
        return ""
    try:
        validate_reference(code, required=True)
    except ValueError as exc:
        return str(exc)
    manifest = read_json(work_item_root(code) / "manifest.yaml", {})
    registered = read_json(project_root() / "registry/normalized/specifications.json", {}).get(code, {})
    item = manifest or registered
    if not item:
        return f"{code} отсутствует среди рабочих элементов и в импортированном реестре."
    unexpected = sorted(set(requirements) - set(item.get("requirements", [])))
    return f"Требования {', '.join(unexpected)} не привязаны к {code}." if unexpected else ""


def begin_free_request(args: argparse.Namespace) -> int:
    operation = str(args.operation).strip().casefold()
    if operation in {"template-management", "template-document"}:
        return templates_cli.begin(SimpleNamespace(**globals()), args)
    query_intent = None
    if operation == "query-analysis":
        try:
            query_intent = select_query_intent(str(args.summary or ""), getattr(args, "query_intent", None))
        except ValueError as exc:
            raise WorkflowError(str(exc)) from exc
    explicit_git_ref = getattr(args, "git_ref", None)
    contextual_git_ref = extract_git_ref(str(args.summary or "")) if operation == "code-review" else None
    legacy_code = getattr(args, "code", None)
    git_ref_value = explicit_git_ref or contextual_git_ref
    use_legacy_as_git_ref = bool(
        operation == "code-review"
        and not getattr(args, "task_reference", None)
        and not getattr(args, "project_reference", None)
        and legacy_code
        and (explicit_git_ref or str(legacy_code) == str(contextual_git_ref or ""))
    )
    if use_legacy_as_git_ref and not git_ref_value:
        git_ref_value = legacy_code
    try:
        git_ref = validate_git_ref(git_ref_value)
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    reference = resolve_reference_args(args, ignore_legacy_code=use_legacy_as_git_ref)
    code = reference["work_reference"]
    mismatch = str(getattr(args, "mismatch", "") or "") or reference_mismatch(code, getattr(args, "requirements", []) or [])
    assessment = assess_request(operation, args.mode, str(args.summary or ""), code=code, mismatch=mismatch)
    requested_git_refs = list(dict.fromkeys(str(item).strip() for item in
                              (getattr(args, "git_refs", None) or []) if str(item).strip()))
    if git_ref and git_ref not in requested_git_refs:
        requested_git_refs.insert(0, git_ref)
    target_ref = str(getattr(args, "target_ref", "") or "").strip() or "main"
    refresh_requested = bool(getattr(args, "refresh_git_refs", False))
    gate = new_gate(operation, None, assessment.pop("state"), **assessment, **reference,
                    reference_code=code, git_ref=git_ref, summary=str(args.summary or ""),
                    git_context={"repository": "extension", "source_refs": requested_git_refs,
                                 "target_ref": target_ref, "refresh_requested": refresh_requested,
                                 "freshness": "REFRESH_REQUESTED" if refresh_requested else "UNKNOWN"},
                    git_analysis_state={"schema_version": 1, "cache": {}, "terminal": {},
                                        "history_search_without_evidence": 0, "snapshots": {}},
                    mode_selection=getattr(args, "mode_selection", None),
                    mode_history=[getattr(args, "mode_selection", None)] if getattr(args, "mode_selection", None) else [],
                    requirements=list(getattr(args, "requirements", []) or []),
                    presented_paths=list(getattr(args, "path", []) or []), answers=[], notes=[], artifacts=[])
    gate["request_id"] = gate["gate_id"]
    gate["storage_kind"] = "documentation" if read_json(ROOT / LOCAL_CONFIG_FILE, {}).get("documentation_path") else "workspace"
    gate["available_actions"] = allowed_tools_for_mode(args.mode, operation)
    gate["output"] = "result.md" if args.mode == "draft" else None
    gate["skill"] = {"query-analysis": "flow1c-query-analysis", "interview-preparation": "flow1c-interview-preparation"}.get(operation, "flow1c-consultation")
    if query_intent:
        gate["query_intent"] = query_intent
        # Existing saved v0 gates remain compatible; every new query gate is checked.
        gate["query_contract_version"] = 1
    if getattr(args, "requested_operation", None):
        gate["requested_operation"] = args.requested_operation
        gate["routing_reason"] = "explicit_1c_query_request"
    # Free work must not inject an incompatible formal-only role skill.
    skill_path = ROOT / ".agents" / "skills" / str(gate["skill"]) / "SKILL.md"
    gate["skill_instructions"] = skill_path.read_text(encoding="utf-8") if skill_path.exists() else "Discuss the goal; use chat sources; draft without changing formal status."
    gate["evidence_path"] = str(request_root(gate) / "evidence.json")
    write_json(Path(gate["evidence_path"]), {"schema_version": 2, "gate_id": gate["gate_id"], "operation": operation,
               "code": None, "git_ref": git_ref, **reference, "artifacts": [], "rlm_queries": [], "cc_inspections": [],
               "source_reads": [], "git_reads": [], "git_analysis": [], "snapshots": [], "changed_files": [],
               "query_schemas": [], "query_checks": []})
    save_request(gate)
    print(json.dumps(gate, ensure_ascii=False, indent=2))
    return 0 if gate["state"] == "READY" else 1


def save_request(gate: dict[str, Any]) -> None:
    save_gate(gate)
    if gate.get("mode", "formal") != "formal":
        write_json(request_root(gate) / "request.json", {**gate, "generated_by": "Flow1C", "document_status": "UNVERIFIED_DRAFT"})


def cmd_agent_dialogue(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id)
    require_gate_tool(gate, "flow1c_dialogue")
    if gate["state"] in {"COMPLETE", "COMPLETE_WITH_DEVIATIONS", "CONSULTATION_COMPLETE", "NON_COMPLIANT"}:
        raise WorkflowError("Start a new request for a completed operation")
    action = args.action
    pending = (gate.get("clarification") or {}).get("reason")
    if action == "answer" and pending == "missing_work_item":
        answer = str(args.answer or "").strip()
        if (args.resolution != "independent-draft"
                or not INDEPENDENT_DRAFT_DECISION_PATTERN.search(answer)):
            raise WorkflowError("Select independent-draft and provide the user's explicit answer")
        reference = str(gate.get("work_reference") or gate.get("code") or "")
        selection = {
            "mode": "draft", "actor": "user", "reason": "explicit_independent_draft", "selected_at": utc_now(),
        }
        gate.setdefault("answers", []).append({
            "question": (gate.get("clarification") or {}).get("question", ""), "answer": answer,
        })
        gate.setdefault("deviations", []).append({
            "type": "process", "actor": "user", "reason": "Formal binding was declined",
            "user_statement": answer, "scope": "gate", "condition_ids": [],
            "waived_conditions": [], "recorded_at": utc_now(),
        })
        gate.update(
            mode="draft", code=None, reference_code=reference or None,
            request_id=gate.get("request_id") or gate["gate_id"],
            storage_kind="documentation" if read_json(ROOT / LOCAL_CONFIG_FILE, {}).get("documentation_path") else "workspace",
            output="result.md", state="READY_WITH_DEVIATIONS", compliance="DEVIATED",
            available_actions=allowed_tools_for_mode("draft", str(gate.get("operation", ""))), mode_selection=selection,
            mode_history=[*gate.get("mode_history", []), selection], awaiting_user_input=False,
            clarification=None, remaining_blockers=[],
            user_message="Формальная привязка отклонена. Работа продолжается как независимый UNVERIFIED_DRAFT.",
        )
        gate["evidence_path"] = str(request_root(gate) / "evidence.json")
        write_json(Path(gate["evidence_path"]), {
            "schema_version": 1, "gate_id": gate["gate_id"], "operation": gate["operation"],
            "code": None, "project_reference": gate.get("project_reference"),
            "task_reference": gate.get("task_reference"), "work_reference": gate.get("work_reference"),
            "artifacts": [], "rlm_queries": [], "source_reads": [], "changed_files": [],
            "deviations": gate["deviations"],
        })
        save_request(gate)
        print(json.dumps(gate, ensure_ascii=False, indent=2))
        return 0
    if action == "deviate":
        deviation_type = str(getattr(args, "deviation_type", "") or "")
        scope = str(getattr(args, "scope", "") or "")
        condition_ids = getattr(args, "condition_ids", None)
        user_statement = str(getattr(args, "user_statement", "") or "").strip()
        reason = str(args.answer or "").strip()
        if not deviation_type or not scope or not isinstance(condition_ids, list) or not condition_ids or not user_statement or not reason:
            raise WorkflowError("deviate requires nonempty answer/reason, user_statement, deviation_type, scope and condition_ids")
        stage = load_stages()["operations"].get(str(gate.get("operation")), {})
        try:
            gate = apply_deviation(gate, {
                "reason": reason, "user_statement": user_statement,
                "deviation_type": deviation_type, "scope": scope,
                "condition_ids": condition_ids, "recorded_at": utc_now(),
            }, stage)
        except ValueError as exc:
            raise WorkflowError(str(exc)) from exc
        reference = str(gate.get("work_reference") or gate.get("code") or "")
        has_work_item = bool(reference and (work_item_root(reference) / "manifest.yaml").is_file())
        if deviation_type == "registry_bypass" and not has_work_item:
            if not stage.get("provisional_allowed"):
                raise WorkflowError("This operation does not allow provisional registry bypass")
            gate.update(
                state="NEEDS_CONFIRMATION", work_item_exists=has_work_item, awaiting_user_input=False,
                user_message="registry_bypass зафиксирован. Запустите provisional-start с пользовательским идентификатором и непустым описанием.",
            )
        elif reference and gate.get("state") in {"BLOCKED", "NEEDS_INPUT", "READY_WITH_DEVIATIONS"}:
            if gate.get("evidence_path"):
                evidence_path, evidence = evidence_for_gate(gate)
            else:
                evidence_path = work_item_root(reference) / "evidence" / f"{gate['operation']}-{gate['gate_id']}.json"
                evidence = {
                    "schema_version": 1, "gate_id": gate["gate_id"], "operation": gate["operation"],
                    "code": reference, "project_reference": gate.get("project_reference"),
                    "task_reference": gate.get("task_reference"), "work_reference": reference,
                    "context_sha256": "", "artifacts": load_artifact_index(reference)["artifacts"],
                    "diff": None, "rlm_queries": [], "source_reads": [], "bsl_ls": None,
                    "changed_files": [], "created_at": utc_now(),
                }
                gate["evidence_path"] = str(evidence_path)
            evidence["deviations"] = gate["deviations"]
            evidence["conditions"] = gate.get("conditions", [])
            write_json(evidence_path, evidence)
        refresh_gate_actions(gate)
        save_request(gate)
        print(json.dumps(gate, ensure_ascii=False, indent=2))
        return 0
    if action == "record":
        if not str(args.answer or "").strip():
            raise WorkflowError("A nonempty user answer, assumption or open question is required")
        gate.setdefault("notes", []).append({"kind": args.kind, "text": args.answer, "created_at": utc_now()})
    else:
        if action == "answer" and pending == "profile":
            profile = str(getattr(args, "profile", "") or args.answer or "").strip().casefold()
            if profile not in load_capabilities()["profiles"]:
                raise WorkflowError("Choose a valid setup profile; analysis is recommended and full must be explicit")
            gate["state"] = "SUPERSEDED"
            gate.setdefault("answers", []).append({"question": gate["clarification"]["question"], "answer": args.answer})
            save_gate(gate)
            return cmd_agent_begin(argparse.Namespace(operation="setup", mode="formal", code=None,
                task_reference=None, project_reference=None, g_number=None, profile=profile,
                reference_kind="auto", git_ref=None,
                summary=gate.get("summary", ""), path=gate.get("presented_paths", []), requirements=[],
                mismatch="", allow_incomplete_draft=False))
        if action == "answer" and pending == "registry_reference_kind":
            resolution = gate.get("registry_resolution", {})
            selected_kind = str(getattr(args, "reference_kind", "") or "").strip().casefold()
            selected_reference = str(
                args.code or gate.get("requested_reference") or gate.get("work_reference") or ""
            ).strip()
            if resolution.get("reason") == "multiple_specifications":
                if selected_reference not in resolution.get("candidates", []):
                    raise WorkflowError("Select one of the specification numbers returned by the gate")
                selected_kind = "specification"
            elif selected_kind not in {"requirement", "specification"}:
                raise WorkflowError("Select reference_kind=requirement or reference_kind=specification")
            gate["state"] = "SUPERSEDED"
            gate.setdefault("answers", []).append({
                "question": gate["clarification"]["question"], "answer": args.answer,
            })
            save_gate(gate)
            return cmd_agent_begin(argparse.Namespace(
                operation=gate["operation"], mode="formal", code=None,
                task_reference=selected_reference, project_reference=gate.get("project_reference"),
                reference_kind=selected_kind, git_ref=gate.get("git_ref"), g_number=None,
                profile=None, summary=gate.get("summary", "") + "\nОтвет пользователя: " + str(args.answer or ""),
                path=gate.get("presented_paths", []), requirements=gate.get("requirements", []),
                mismatch="", allow_incomplete_draft=False,
            ))
        if action == "answer" and pending == "requirement_mismatch":
            if gate.get("mode", "formal") == "formal":
                if args.resolution not in {"independent-draft", "correct-code"}:
                    raise WorkflowError("Select independent-draft or correct-code")
                code = str(args.code or "").strip() if args.resolution == "correct-code" else None
                if args.resolution == "correct-code" and (not code or reference_mismatch(code, gate.get("requirements", []))):
                    raise WorkflowError("Supply a corrected code matching the requirements")
                gate["state"] = "SUPERSEDED"
                gate.setdefault("answers", []).append({"question": gate["clarification"]["question"], "answer": args.answer})
                save_gate(gate)
                return cmd_agent_begin(argparse.Namespace(operation=gate["operation"], mode="draft" if code is None else "formal",
                    code=code, summary=gate.get("summary", "") + "\nОтвет пользователя: " + args.answer,
                    path=gate.get("presented_paths", []), requirements=gate.get("requirements", []), mismatch="", allow_incomplete_draft=False))
            if args.resolution == "independent-draft":
                gate.update(mode="draft", code=None, reference_code=None, output="result.md")
            elif args.resolution == "correct-code":
                code = str(args.code or "").strip()
                if not code:
                    raise WorkflowError("Supply the corrected user-assigned code")
                mismatch = reference_mismatch(code, gate.get("requirements", []))
                if mismatch:
                    raise WorkflowError(mismatch)
                gate["reference_code"] = code
            else:
                raise WorkflowError("Select independent-draft or correct-code; the registry will not be changed")
        try:
            gate = transition(gate, action, question=str(args.question or ""), answer=str(args.answer or ""))
        except ValueError as exc:
            raise WorkflowError(str(exc)) from exc
    if getattr(args, "mode", None):
        if gate.get("mode", "formal") == "formal" or args.mode not in {"explore", "draft"}:
            raise WorkflowError("Mode changes here are only for free requests; start a new formal gate")
        if gate.get("mode") != args.mode:
            selection = {"mode": args.mode, "actor": "caller", "reason": "explicit_dialogue_change", "selected_at": utc_now()}
            gate["mode_selection"] = selection
            gate.setdefault("mode_history", []).append(selection)
        gate["mode"] = args.mode
        gate["output"] = "result.md" if args.mode == "draft" else None
    if gate.get("mode", "formal") != "formal":
        if gate["state"] == "DRAFT_COMPLETE":
            gate["state"] = "READY"
        gate["available_actions"] = allowed_tools_for_mode(str(gate["mode"]), str(gate.get("operation", "")))
    else:
        refresh_gate_actions(gate)
    save_request(gate)
    print(json.dumps(gate, ensure_ascii=False, indent=2))
    return 0


def safe_request_file(root: Path, relative: str, *, markdown: bool = False) -> Path:
    target = (root / relative).resolve()
    if not relative or not path_is_within(target, root) or target == root.resolve():
        raise WorkflowError("Path is outside the request directory")
    if markdown and (target.suffix.lower() != ".md" or target.relative_to(root.resolve()).parts[0] in {"sources", "evidence"}):
        raise WorkflowError("Draft outputs must be Markdown outside source/evidence directories")
    return target


def write_free_draft(gate: dict[str, Any], args: argparse.Namespace, content: str) -> int:
    if gate["mode"] != "draft" or args.target != "draft":
        raise WorkflowError("Free work can write only target=draft in draft mode")
    if not content.strip():
        raise WorkflowError("Draft content is empty")
    root = request_root(gate)
    target = safe_request_file(root, str(args.path or ""), markdown=True)
    if "UNVERIFIED_DRAFT" not in content:
        content = "> **UNVERIFIED_DRAFT** — рабочий черновик; открытые вопросы и допущения требуют уточнения.\n\n" + content
    write_text(target, "<!-- Generated by Flow1C. -->\n\n" + content)
    evidence_path, evidence = evidence_for_gate(gate)
    relative = str(target.relative_to(root)).replace("\\", "/")
    evidence.setdefault("changed_files", []).append({"target": "draft", "path": relative, "sha256": sha256(target)})
    write_json(evidence_path, evidence)
    print(json.dumps({"state": "WRITTEN", "absolute_path": str(target)}, ensure_ascii=False))
    return 0


def intake_free_request(gate: dict[str, Any], args: argparse.Namespace) -> int:
    if getattr(args, "code", None):
        raise WorkflowError("Use draft-promote to attach a free request to a user-selected work item")
    root = request_root(gate)
    promote_intake_id = str(getattr(args, "promote_intake_id", "") or "").strip()
    inbox_records: dict[str, dict[str, Any]] = {}
    if promote_intake_id:
        if getattr(args, "source", None):
            raise WorkflowError("Use either promote_intake_id or source paths in one intake call")
        if not re.fullmatch(r"[A-Za-z0-9-]{8,64}", promote_intake_id):
            raise WorkflowError("Invalid intake ID")
        inbox_root = project_root() / "inbox" / promote_intake_id
        manifest = read_json(inbox_root / "intake.json")
        if not isinstance(manifest, dict) or manifest.get("code") is not None:
            raise WorkflowError("Inbox intake is unavailable for this request")
        sources = []
        for item in manifest.get("artifacts", []):
            if not isinstance(item, dict):
                continue
            source = project_root() / str(item.get("relative_path", ""))
            if (not path_is_within(source, inbox_root / "originals") or is_reparse_or_symlink(source)
                    or not source.is_file() or sha256(source) != item.get("sha256")):
                raise WorkflowError("Inbox artifact changed; fetch it again before intake")
            sources.append(str(source))
            inbox_records[str(source.resolve())] = item
        if not sources:
            raise WorkflowError("Inbox intake has no verified artifacts")
        args.source = sources
    files, skipped = enumerate_intake_files(
        getattr(args, "source", []) or [], excluded_roots=(root,)
    )
    if not files:
        raise WorkflowError("No supported files found. The user can describe the task in chat instead.")
    if (len(files) > MAX_INTAKE_FILES or sum(p.stat().st_size for p, _ in files) > MAX_INTAKE_BYTES) and not args.confirm_large:
        raise WorkflowError("Large intake requires explicit confirmation")
    artifacts = gate.get("artifacts")
    if not isinstance(artifacts, list):
        artifacts = []
        gate["artifacts"] = artifacts
    copied = []
    for source, _ in files:
        digest = sha256(source)
        existing = next((item for item in artifacts if isinstance(item, dict) and item.get("sha256") == digest), None)
        if existing is not None:
            if source.suffix.lower() in {".docx", ".pdf"} and not existing.get("derived_path"):
                target = safe_request_file(root, str(existing["path"]))
                try:
                    from markitdown import MarkItDown
                    derived = target.with_suffix(target.suffix + ".md")
                    write_text(derived, "<!-- Generated from accepted source. -->\n" + MarkItDown().convert(str(target)).text_content)
                    existing["derived_path"] = str(derived.relative_to(root)).replace("\\", "/")
                    existing["derived_sha256"] = sha256(derived)
                    existing.pop("extraction_error", None)
                except (ImportError, OSError, ValueError, AttributeError) as exc:
                    existing["extraction_error"] = str(exc)
                existing["extraction_status"] = "ready" if existing.get("derived_path") else "failed"
                copied.append(existing)
            else:
                skipped.append(str(source))
            continue
        target = free_intake_target(root, digest, source)
        if paths_are_equal(source, target):
            skipped.append(str(source))
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        record = {"path": str(target.relative_to(root)).replace("\\", "/"), "source_path": str(source),
                  "name": source.name,
                  "sha256": digest, "category": args.category or "chat_material", "received_at": utc_now()}
        inbox_record = inbox_records.get(str(source.resolve()))
        if inbox_record:
            record["source"] = inbox_record.get("source")
            record["inbox_intake_id"] = promote_intake_id
            record["category"] = inbox_record.get("category", record["category"])
        if source.suffix.lower() in {".docx", ".pdf"}:
            try:
                from markitdown import MarkItDown
                derived = target.with_suffix(target.suffix + ".md")
                write_text(derived, "<!-- Generated from accepted source. -->\n" + MarkItDown().convert(str(target)).text_content)
                record["derived_path"] = str(derived.relative_to(root)).replace("\\", "/")
                record["derived_sha256"] = sha256(derived)
            except (ImportError, OSError, ValueError, AttributeError) as exc:
                record["extraction_error"] = str(exc)
            record["extraction_status"] = "ready" if record.get("derived_path") else "failed"
        artifacts.append(record)
        copied.append(record)
    evidence_path, evidence = evidence_for_gate(gate)
    evidence["artifacts"] = artifacts
    write_json(evidence_path, evidence)
    save_request(gate)
    print(json.dumps({"state": "ACCEPTED", "copied": copied, "skipped": skipped, "next": "continue_same_gate"}, ensure_ascii=False))
    return 0


def cmd_agent_inspect(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT", "WAITING_USER"})
    require_gate_tool(gate, "flow1c_inspect")
    root = ROOT.resolve() if args.scope == "workflow" else request_root(gate)
    limit = max(1, min(int(args.max_chars), 40000))
    excluded = {".git", ".venv", ".workspace", "node_modules", "__pycache__"}
    extensions = {".md", ".py", ".ps1", ".json", ".js", ".ts", ".txt", ".yaml", ".csv"}
    def allowed(p: Path) -> bool:
        relative = p.relative_to(root)
        return (p.suffix.lower() in extensions and not any(part in excluded for part in relative.parts)
                and not p.name.startswith((".flow1c.local", ".env"))
                and path_is_within(p, root) and not is_reparse_or_symlink(p))
    if args.path:
        target = safe_request_file(root, args.path)
        if not target.is_file() or not allowed(target):
            raise WorkflowError("Only workflow text and accepted request text can be read; configuration sources require RLM")
        candidates = [target]
    else:
        candidates = []
        for directory, dirs, names in os.walk(root, followlinks=False):
            dirs[:] = sorted(d for d in dirs if d not in excluded and not is_reparse_or_symlink(Path(directory) / d))
            for name in sorted(names):
                p = Path(directory) / name
                if allowed(p):
                    candidates.append(p)
            if len(candidates) >= 1000:
                break
    result = []
    remaining = limit
    for p in candidates[:1000]:
        if not args.path and not args.query:
            text = str(p.relative_to(root))
        else:
            if p.stat().st_size > 2 * 1024 * 1024:
                continue
            lines = p.read_text(encoding="utf-8-sig", errors="replace").splitlines()
            start = max(1, int(args.start_line))
            text = "\n".join(f"{i}: {line}" for i, line in enumerate(lines, 1) if i >= start and (not args.query or args.query.casefold() in line.casefold()))
        if not text:
            continue
        result.append({"path": str(p.relative_to(root)).replace("\\", "/"), "content": text[:remaining]})
        remaining -= len(text)
        if remaining <= 0:
            break
    print(json.dumps({"state": "READ", "items": result, "truncated": remaining <= 0 or len(candidates) >= 1000}, ensure_ascii=False))
    return 0


def git_inspection_repository(source: str) -> Path:
    if source == "workflow":
        repository = ROOT.resolve()
    else:
        local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
        raw = str(local.get("extension_path", "") or "").strip()
        if not raw:
            raise WorkflowError("extension_path is not configured")
        repository = Path(raw).resolve()
    if not repository.is_dir():
        raise WorkflowError(f"Git repository is unavailable: {repository}")
    probe = subprocess.run(
        ["git", "rev-parse", "--is-inside-work-tree"], cwd=repository, text=True,
        encoding="utf-8", errors="replace", capture_output=True, check=False, timeout=15,
    )
    if probe.returncode or probe.stdout.strip() != "true":
        raise WorkflowError(f"Path is not a Git repository: {repository}")
    return repository


def git_read(repository: Path, arguments: list[str], *, timeout: int = 30) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        ["git", "-c", "core.pager=cat", "--no-pager", *arguments], cwd=repository,
        text=True, encoding="utf-8", errors="replace", capture_output=True,
        check=False, timeout=timeout,
    )


def resolve_git_commit(repository: Path, raw_ref: str) -> str | None:
    try:
        ref = validate_git_ref(raw_ref)
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    if not ref:
        raise WorkflowError("git_ref is required")
    candidates = [f"refs/heads/{ref}", f"refs/remotes/origin/{ref}", ref]
    for candidate in dict.fromkeys(candidates):
        result = git_read(repository, ["rev-parse", "--verify", f"{candidate}^{{commit}}"])
        if result.returncode == 0:
            commit = result.stdout.strip().splitlines()[0] if result.stdout.strip() else ""
            if re.fullmatch(r"[0-9a-fA-F]{40,64}", commit):
                return commit.lower()
    return None


def git_commit_record(repository: Path, commit: str) -> dict[str, Any]:
    result = git_read(repository, ["show", "-s", "--format=%H%x00%P%x00%s", commit])
    if result.returncode:
        raise WorkflowError((result.stderr or "git show failed").strip())
    fields = result.stdout.rstrip("\n").split("\x00", 2)
    return {
        "commit": fields[0],
        "parents": fields[1].split() if len(fields) > 1 and fields[1] else [],
        "subject": fields[2] if len(fields) > 2 else "",
    }


def git_is_ancestor(repository: Path, ancestor: str, descendant: str) -> bool:
    result = git_read(repository, ["merge-base", "--is-ancestor", ancestor, descendant])
    if result.returncode not in {0, 1}:
        raise WorkflowError((result.stderr or "git merge-base --is-ancestor failed").strip())
    return result.returncode == 0


def find_git_integration(repository: Path, raw_ref: str, target_ref: str) -> dict[str, Any]:
    source = resolve_git_commit(repository, raw_ref)
    if not source:
        return {"git_ref": raw_ref, "status": "SOURCE_NOT_FOUND"}
    target = resolve_git_commit(repository, target_ref)
    if not target:
        return {"git_ref": raw_ref, "source_commit": source, "target_ref": target_ref,
                "status": "TARGET_NOT_FOUND"}
    common = {
        "git_ref": raw_ref, "source_commit": source,
        "target_ref": target_ref, "target_commit": target,
    }
    if not git_is_ancestor(repository, source, target):
        return {
            **common, "status": "NOT_MERGED",
            "reason": "The source commit is not an ancestor of the target; a squash or rebase cannot be proven from Git topology.",
        }

    first_parent_history = git_read(repository, ["rev-list", "--first-parent", target])
    if first_parent_history.returncode:
        raise WorkflowError((first_parent_history.stderr or "git rev-list --first-parent failed").strip())
    if source in set(first_parent_history.stdout.splitlines()):
        return {
            **common, "status": "NO_MERGE_COMMIT", "integration_method": "first-parent-or-fast-forward",
            "reason": "The source is present directly on the target first-parent history; no introducing merge commit exists.",
            "review_ref": source,
        }

    merges = git_read(repository, [
        "rev-list", "--first-parent", "--merges", "--ancestry-path", "--reverse", "--parents",
        f"{source}..{target}",
    ])
    if merges.returncode:
        raise WorkflowError((merges.stderr or "git rev-list failed").strip())
    for line in merges.stdout.splitlines():
        fields = line.split()
        if len(fields) < 3:
            continue
        merge_commit, first_parent, *merged_parents = fields
        if any(git_is_ancestor(repository, source, parent) for parent in merged_parents):
            record = git_commit_record(repository, merge_commit)
            return {
                **common, "status": "MERGE_FOUND", "integration_method": "merge",
                "merge_commit": merge_commit, "merge_parents": record["parents"],
                "merge_subject": record["subject"], "diff_base_commit": first_parent,
                "review_ref": merge_commit,
            }

    return {
        **common, "status": "NO_MERGE_COMMIT", "integration_method": "first-parent-or-fast-forward",
        "reason": "The source is present in the target, but no introducing merge commit exists on the target first-parent history.",
        "review_ref": source,
    }


def cmd_agent_git_inspect(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_git_inspect")
    if gate.get("mode") not in {"explore", "draft"}:
        raise WorkflowError("Git inspection is available only in explore or draft mode")
    repository = git_inspection_repository(str(args.repository))
    action = str(args.action)
    requested_refs = getattr(args, "git_refs", None) or []
    if action == "integration":
        raw_refs = list(dict.fromkeys(str(item).strip() for item in requested_refs if str(item).strip()))
        fallback_ref = str(getattr(args, "git_ref", "") or gate.get("git_ref") or "").strip()
        if not raw_refs and fallback_ref:
            raw_refs = [fallback_ref]
        if not raw_refs:
            raise WorkflowError("git_ref or git_refs is required")
        if len(raw_refs) > 20:
            raise WorkflowError("At most 20 git_refs may be inspected at once")
        target_ref = str(getattr(args, "target_ref", "") or
                         load_config()[0].get("project", {}).get("default_branch", "main")).strip()
        if not target_ref:
            raise WorkflowError("target_ref is required")
        integrations = [find_git_integration(repository, item, target_ref) for item in raw_refs]
        content = json.dumps(integrations, ensure_ascii=False, sort_keys=True)
        evidence_path, evidence = evidence_for_gate(gate)
        digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
        evidence.setdefault("git_reads", []).append({
            "action": action, "git_refs": raw_refs, "target_ref": target_ref,
            "sha256": digest, "created_at": utc_now(),
        })
        write_json(evidence_path, evidence)
        print(json.dumps({
            "state": "READ", "action": action, "target_ref": target_ref,
            "repository": str(repository), "integrations": integrations, "sha256": digest,
            "next_action": "Review MERGE_FOUND entries with action=diff and git_ref=review_ref; do not search history for NOT_MERGED or NO_MERGE_COMMIT entries.",
        }, ensure_ascii=False, indent=2))
        return 0
    if requested_refs:
        raise WorkflowError("git_refs is available only for action=integration")

    raw_ref = str(getattr(args, "git_ref", "") or gate.get("git_ref") or "").strip()
    commit = resolve_git_commit(repository, raw_ref)
    if not commit:
        print(json.dumps({"state": "NOT_FOUND", "git_ref": raw_ref, "repository": str(repository)}, ensure_ascii=False))
        return 0

    record = git_commit_record(repository, commit)
    parents = record["parents"]
    content = "\x00".join((record["commit"], " ".join(parents), record["subject"])) + "\n"
    payload: dict[str, Any] = {
        "state": "READ", "action": action, "git_ref": raw_ref, "commit": commit,
        "parents": parents, "is_merge": len(parents) > 1, "repository": str(repository),
    }
    if action == "log":
        count = max(1, min(int(args.max_count), 100))
        result = git_read(repository, ["log", f"--max-count={count}", "--format=%H%x00%P%x00%s", commit])
        if result.returncode:
            raise WorkflowError((result.stderr or "git log failed").strip())
        content = result.stdout
        payload["log"] = content[:int(args.max_chars)]
    elif action == "latest-merge":
        result = git_read(repository, ["log", "--merges", "--max-count=1", "--format=%H%x00%P%x00%s", commit])
        if result.returncode:
            raise WorkflowError((result.stderr or "git log --merges failed").strip())
        content = result.stdout
        payload["merge"] = content[:int(args.max_chars)] or None
    elif action == "diff":
        base_ref = str(getattr(args, "base_ref", "") or "").strip()
        base = resolve_git_commit(repository, base_ref) if base_ref else (parents[0] if len(parents) > 1 else None)
        if not base:
            default_branch = str(load_config()[0].get("project", {}).get("default_branch", "main"))
            base = resolve_git_commit(repository, default_branch)
            if base:
                merge_base = git_read(repository, ["merge-base", base, commit])
                base = merge_base.stdout.strip() if merge_base.returncode == 0 else base
        if not base:
            raise WorkflowError("A base commit is unavailable; provide base_ref")
        result = git_read(repository, ["diff", "--no-ext-diff", "--find-renames", "--unified=3", base, commit], timeout=60)
        names = git_read(repository, ["diff", "--no-ext-diff", "--name-status", base, commit])
        if result.returncode or names.returncode:
            raise WorkflowError((result.stderr or names.stderr or "git diff failed").strip())
        content = result.stdout
        limit = max(1, min(int(args.max_chars), 100000))
        payload.update(base_commit=base, files=names.stdout.splitlines(), diff=content[:limit], truncated=len(content) > limit)
    else:
        payload["subject"] = record["subject"]

    evidence_path, evidence = evidence_for_gate(gate)
    digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
    evidence.setdefault("git_reads", []).append({
        "action": action, "git_ref": raw_ref, "commit": commit,
        "base_commit": payload.get("base_commit"), "sha256": digest, "created_at": utc_now(),
    })
    write_json(evidence_path, evidence)
    payload["sha256"] = digest
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


def _git_payload_digest(payload: Any) -> str:
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _git_analysis_state(gate: dict[str, Any]) -> dict[str, Any]:
    state = gate.get("git_analysis_state")
    if not isinstance(state, dict):
        state = {"schema_version": 1, "cache": {}, "terminal": {},
                 "history_search_without_evidence": 0, "snapshots": {}}
        gate["git_analysis_state"] = state
    state.setdefault("cache", {})
    state.setdefault("terminal", {})
    state.setdefault("snapshots", {})
    return state


def _record_git_analysis(gate: dict[str, Any], *, action: str, fingerprint: str,
                         parameters: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    evidence_path, evidence = evidence_for_gate(gate)
    digest = _git_payload_digest(payload)
    evidence_id = f"GIT-{len(evidence.get('git_analysis', [])) + 1:03d}"
    record = {"schema_version": 2, "evidence_id": evidence_id, "action": action,
              "fingerprint": fingerprint, "parameters": parameters, "output_sha256": digest,
              "result": payload, "created_at": utc_now()}
    evidence.setdefault("git_analysis", []).append(record)
    legacy_ref = parameters.get("git_ref") or ((parameters.get("git_refs") or [None])[0])
    evidence.setdefault("git_reads", []).append({"action": parameters.get("requested_action", action), "git_ref": legacy_ref,
        "git_refs": parameters.get("git_refs", []), "target_ref": parameters.get("target_ref"),
        "commit": payload.get("commit"), "base_commit": payload.get("base_commit"),
        "sha256": digest, "evidence_id": evidence_id, "created_at": utc_now()})
    write_json(evidence_path, evidence)
    state = _git_analysis_state(gate)
    state["cache"][fingerprint] = {"evidence_id": evidence_id, "result": payload}
    if action == "merge-search":
        state["terminal"]["merge-search"] = {"fingerprint": fingerprint, "evidence_id": evidence_id,
                                                "refs": parameters.get("git_refs", [])}
    if action == "history-search":
        state["history_search_without_evidence"] = int(state.get("history_search_without_evidence", 0)) + 1
    save_request(gate)
    return {**payload, "evidence_id": evidence_id, "sha256": digest, "cache_hit": False}


def cmd_agent_git_refresh(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_git_refresh")
    if gate.get("mode") not in {"explore", "draft"}:
        raise WorkflowError("Git refresh is available only in explore or draft mode")
    context = gate.get("git_context", {}) if isinstance(gate.get("git_context"), dict) else {}
    if not context.get("refresh_requested"):
        raise WorkflowError("Git refresh requires the user's saved refresh_git_refs intent")
    repository = git_inspection_repository(str(args.repository))
    refs = list(getattr(args, "refs", None) or context.get("source_refs", []))
    target = str(context.get("target_ref", "") or "").strip()
    prefix = (str(read_json(ROOT / LOCAL_CONFIG_FILE, {}).get("extension_branch_prefix", "") or "").strip()
              if args.repository == "extension" else "")
    try:
        discovery = git_analysis.discover_issue_refs(repository, refs, prefix=prefix)
        if discovery["state"] != "READ":
            payload = {**discovery, "recoverable": True, "repository_changed": False,
                       "next_action": "check-origin-or-select-full-branch"}
        else:
            resolved = discovery["resolved"]
            refresh_material = list(dict.fromkeys([*resolved.values(), *([target] if target else [])]))
            payload = git_analysis.refresh_refs(repository, remote=str(args.remote), refs=refresh_material)
            payload["branch_discovery"] = discovery
            if payload.get("state") in {"UPDATED", "UNCHANGED"}:
                context["resolved_refs"] = resolved
    except git_analysis.GitAnalysisError as exc:
        print(json.dumps({"schema_version": 2, "state": "BLOCKED", "errors": [exc.payload]}, ensure_ascii=False, indent=2))
        return 2
    context["freshness"] = "FRESH" if payload.get("state") in {"UPDATED", "UNCHANGED"} else payload.get("state")
    context["refresh_result_sha256"] = _git_payload_digest(payload)
    gate["git_context"] = context
    fingerprint = semantic_fingerprint({"repository": args.repository, "action": "git-refresh",
                                        "remote": args.remote, "refs": refs})
    result = _record_git_analysis(gate, action="git-refresh", fingerprint=fingerprint,
                                  parameters={"repository": args.repository, "remote": args.remote, "refs": refs},
                                  payload=payload)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if payload.get("state") in {"UPDATED", "UNCHANGED"} else 1


def _gitea_pr_evidence(repository: Path, source_refs: list[str], target_ref: str,
                        *, allow_targeted_fetch: bool) -> tuple[dict[str, list[dict[str, Any]]], list[dict[str, Any]]]:
    config, _local = load_config()
    settings = config.get("gitea", {}) if isinstance(config.get("gitea"), dict) else {}
    token_env = str(settings.get("token_env", "") or "")
    if not token_env or not os.environ.get(token_env):
        return {}, []
    required = [str(settings.get(key, "") or "").strip() for key in ("base_url", "owner", "repository")]
    if not all(required):
        return {}, [{"code": "GITEA_CONFIG_INCOMPLETE", "recoverable": True,
                     "next_action": "continue-with-git-evidence"}]
    client = GiteaClient(required[0], required[1], required[2], token_env=token_env)
    records: dict[str, list[dict[str, Any]]] = {}
    limitations = []
    target_short = target_ref.rsplit("/", 1)[-1]
    for source in source_refs:
        try:
            response = client.merged_pulls(source, target_short)
        except GiteaError as exc:
            limitations.append(exc.as_dict())
            continue
        records[source] = response["records"]
        for item in records[source]:
            merge_sha = str(item.get("merge_commit_sha", ""))
            if allow_targeted_fetch and merge_sha and not git_analysis.resolve_ref(repository, merge_sha).get("selected"):
                fetched = git_analysis.refresh_refs(repository, remote="origin", refs=[merge_sha])
                if fetched.get("state") not in {"UPDATED", "UNCHANGED"}:
                    limitations.append({"code": fetched.get("code", "GIT_FETCH_FAILED"),
                                        "recoverable": True, "next_action": fetched.get("next_action")})
    return records, limitations


def cmd_agent_git_inspect_v2(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_git_inspect")
    if gate.get("mode") not in {"explore", "draft"}:
        raise WorkflowError("Git inspection is available only in explore or draft mode")
    repository = git_inspection_repository(str(args.repository))
    requested_action = str(args.action)
    action = canonical_action(requested_action)
    raw_refs = list(dict.fromkeys(str(item).strip() for item in (getattr(args, "git_refs", None) or []) if str(item).strip()))
    raw_ref = str(getattr(args, "git_ref", "") or gate.get("git_ref") or "").strip()
    if not raw_refs and raw_ref:
        raw_refs = [raw_ref]
    context = gate.get("git_context", {}) if isinstance(gate.get("git_context"), dict) else {}
    resolved_refs = context.get("resolved_refs", {}) if isinstance(context.get("resolved_refs"), dict) else {}
    raw_refs = [str(resolved_refs.get(item, item)) for item in raw_refs]
    raw_ref = str(resolved_refs.get(raw_ref, raw_ref))
    target_ref = str(getattr(args, "target_ref", "") or context.get("target_ref") or
                     load_config()[0].get("project", {}).get("default_branch", "main")).strip()
    parameters = {
        "repository": str(args.repository), "action": action, "git_ref": raw_ref,
        "git_refs": raw_refs, "target_ref": target_ref,
        "base_ref": str(getattr(args, "base_ref", "") or ""),
        "detail": str(getattr(args, "detail", "patch") or "patch"),
        "paths": list(getattr(args, "paths", None) or []),
        "path": str(getattr(args, "path", "") or ""),
        "start_line": int(getattr(args, "start_line", 1)),
        "max_files": int(getattr(args, "max_files", 100)),
        "output_max_chars": int(args.max_chars),
        "subject_query": str(getattr(args, "subject_query", "") or ""),
        "regex": bool(getattr(args, "regex", False)),
        "merges_only": bool(getattr(args, "merges_only", False)),
        "first_parent": bool(getattr(args, "first_parent", False)),
        "min_parents": getattr(args, "min_parents", None),
        "max_parents": getattr(args, "max_parents", None),
        "since": str(getattr(args, "since", "") or ""),
        "until": str(getattr(args, "until", "") or ""),
        "max_count": int(args.max_count),
        "include_pr_evidence": bool(getattr(args, "include_pr_evidence", True)),
        "include_patch_evidence": bool(getattr(args, "include_patch_evidence", True)),
        "cursor": str(getattr(args, "cursor", "") or ""),
    }
    fingerprint = semantic_fingerprint(parameters)
    parameters["requested_action"] = requested_action
    state = _git_analysis_state(gate)
    try:
        transition = validate_transition(state, action, fingerprint)
    except ValueError as exc:
        code, _, message = str(exc).partition(":")
        print(json.dumps({"schema_version": 2, "state": code, "code": code, "message": message.strip(),
                          "next_actions": ["use-saved-evidence", "complete"]}, ensure_ascii=False, indent=2))
        return 0 if code == "ALREADY_RESOLVED" else 2
    if transition == "CACHE_HIT":
        cached = state["cache"][fingerprint]
        print(json.dumps({**cached["result"], "evidence_id": cached["evidence_id"], "cache_hit": True},
                         ensure_ascii=False, indent=2))
        return 0
    freshness_known = context.get("freshness") == "FRESH"
    try:
        if action == "resolve":
            if not raw_ref:
                raise WorkflowError("git_ref is required")
            resolution = git_analysis.resolve_ref(repository, raw_ref,
                                                  role="target" if raw_ref == target_ref else "source")
            selected = resolution.get("selected")
            payload = {"state": "READ" if selected else "NOT_FOUND", "action": action, **resolution}
            if selected:
                record = git_analysis.commit_record(repository, selected["commit"])
                payload.update(commit=selected["commit"], parents=record["parents"],
                               subject=record["subject"], is_merge=len(record["parents"]) > 1,
                               git_ref=raw_ref, repository=str(repository))
        elif action == "merge-search":
            if not raw_refs:
                raise WorkflowError("git_ref or git_refs is required")
            if len(raw_refs) > 20:
                raise WorkflowError("At most 20 git_refs may be inspected at once")
            known = freshness_known or git_analysis.run_git(repository, ["remote", "get-url", "origin"]).returncode != 0
            integrations = git_analysis.merge_search(repository, raw_refs, target_ref,
                                                       freshness_known=known, include_patch_evidence=False)
            unresolved = [item["git_ref"] for item in integrations
                          if item["integration_status"] != "FAST_FORWARD" and
                          not any(candidate["evidence_level"] == "EXACT_TOPOLOGY"
                                  for candidate in item["merge_candidates"])]
            pr_evidence: dict[str, list[dict[str, Any]]] = {}
            pr_limitations: list[dict[str, Any]] = []
            if unresolved and bool(getattr(args, "include_pr_evidence", True)):
                pr_evidence, pr_limitations = _gitea_pr_evidence(
                    repository, unresolved, target_ref,
                    allow_targeted_fetch=bool(context.get("refresh_requested")),
                )
            if unresolved:
                fallback = git_analysis.merge_search(repository, unresolved, target_ref,
                                                      freshness_known=known,
                                                      include_patch_evidence=bool(getattr(args, "include_patch_evidence", True)),
                                                      pr_evidence=pr_evidence)
                replacements = {item["git_ref"]: item for item in fallback}
                integrations = [replacements.get(item["git_ref"], item) for item in integrations]
            if requested_action == "integration":
                legacy_statuses = {"FAST_FORWARD": "NO_MERGE_COMMIT", "NO_INTEGRATION_EVIDENCE": "NOT_MERGED"}
                for integration in integrations:
                    integration["legacy_status"] = legacy_statuses.get(integration["integration_status"], integration["integration_status"])
                    integration["status"] = integration["legacy_status"]
            if pr_limitations:
                for integration in integrations:
                    integration.setdefault("limitations", []).extend(pr_limitations)
            payload = {"state": "READ", "action": requested_action, "canonical_action": action,
                       "target_ref": target_ref, "repository": str(repository), "integrations": integrations,
                       "next_actions": list(dict.fromkeys(item for result in integrations for item in result["next_actions"]))}
        elif action == "branch-changes":
            if not raw_ref:
                raise WorkflowError("git_ref is required")
            payload = git_analysis.branch_changes(repository, raw_ref, target_ref,
                                                   freshness_known=freshness_known or
                                                   git_analysis.run_git(repository, ["remote", "get-url", "origin"]).returncode != 0)
            payload["action"] = action
        elif action == "history-search":
            payload = git_analysis.history_search(repository, target_ref,
                subject_query=str(getattr(args, "subject_query", "") or ""), regex=bool(getattr(args, "regex", False)),
                merges_only=bool(getattr(args, "merges_only", False)), first_parent=bool(getattr(args, "first_parent", False)),
                min_parents=getattr(args, "min_parents", None), max_parents=getattr(args, "max_parents", None),
                path=str(getattr(args, "path", "") or ""), since=str(getattr(args, "since", "") or ""),
                until=str(getattr(args, "until", "") or ""), cursor=str(getattr(args, "cursor", "") or ""),
                max_count=int(args.max_count))
            payload["action"] = action
        elif action == "read-at-ref":
            if not raw_ref or not getattr(args, "path", None):
                raise WorkflowError("git_ref and path are required")
            payload = git_analysis.read_at_ref(repository, raw_ref, str(args.path),
                                               start_line=int(getattr(args, "start_line", 1)), max_chars=int(args.max_chars))
            payload["action"] = action
        elif action == "diff":
            if not raw_ref:
                raise WorkflowError("git_ref is required")
            payload = git_analysis.diff(repository, raw_ref, base_ref=parameters["base_ref"] or None,
                                        detail=parameters["detail"], paths=parameters["paths"],
                                        cursor=parameters["cursor"], max_files=int(getattr(args, "max_files", 100)),
                                        max_chars=int(args.max_chars))
            payload["action"] = action
        elif action in {"log", "latest-merge"}:
            # Preserve the published legacy shape while using the hardened runner/ref resolver.
            resolved = git_analysis.resolve_ref(repository, raw_ref)
            selected = resolved.get("selected")
            if not selected:
                payload = {"state": "NOT_FOUND", "git_ref": raw_ref, "action": action}
            else:
                git_args = ["log", f"--max-count={1 if action == 'latest-merge' else min(int(args.max_count), 100)}",
                            "--format=%H%x00%P%x00%s"]
                if action == "latest-merge":
                    git_args.append("--merges")
                git_args.append(selected["commit"])
                result = git_analysis.run_git(repository, git_args)
                content = result.stdout.decode("utf-8", errors="replace")[:int(args.max_chars)]
                payload = {"state": "READ", "action": action, "git_ref": raw_ref,
                           "commit": selected["commit"], "merge" if action == "latest-merge" else "log": content or None}
        else:
            raise WorkflowError(f"Unsupported Git action: {action}")
    except git_analysis.GitAnalysisError as exc:
        print(json.dumps({"schema_version": 2, "state": "BLOCKED", "errors": [exc.payload]}, ensure_ascii=False, indent=2))
        return 2
    result = _record_git_analysis(gate, action=action, fingerprint=fingerprint,
                                  parameters=parameters, payload=payload)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_agent_git_snapshot(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_git_snapshot")
    if gate.get("mode") not in {"explore", "draft"}:
        raise WorkflowError("Git snapshots are available only in explore or draft mode")
    snapshot_action = str(getattr(args, "action", "create") or "create")
    if snapshot_action == "cleanup":
        payload = git_analysis.cleanup_snapshots(ROOT / ".workspace", ttl_hours=int(getattr(args, "ttl_hours", 168)))
        fingerprint = semantic_fingerprint({"action": "snapshot-cleanup", "ttl_hours": int(getattr(args, "ttl_hours", 168))})
        result = _record_git_analysis(gate, action="snapshot-cleanup", fingerprint=fingerprint,
                                      parameters={"action": "cleanup", "ttl_hours": int(getattr(args, "ttl_hours", 168))},
                                      payload=payload)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    repository = git_inspection_repository(str(args.repository))
    raw_ref = str(args.git_ref or "").strip()
    paths = list(getattr(args, "paths", None) or [])
    fingerprint = semantic_fingerprint({"repository": args.repository, "action": "snapshot",
                                        "git_ref": raw_ref, "paths": paths})
    state = _git_analysis_state(gate)
    if fingerprint in state["cache"]:
        cached = state["cache"][fingerprint]
        print(json.dumps({**cached["result"], "evidence_id": cached["evidence_id"], "cache_hit": True}, ensure_ascii=False, indent=2))
        return 0
    try:
        payload = git_analysis.create_snapshot(repository, ROOT / ".workspace", str(gate["request_id"]), raw_ref, paths)
    except git_analysis.GitAnalysisError as exc:
        print(json.dumps({"schema_version": 2, "state": "BLOCKED", "errors": [exc.payload]}, ensure_ascii=False, indent=2))
        return 2
    evidence_path, evidence = evidence_for_gate(gate)
    evidence.setdefault("snapshots", []).append({**payload, "created_at": utc_now()})
    write_json(evidence_path, evidence)
    result = _record_git_analysis(gate, action="snapshot", fingerprint=fingerprint,
                                  parameters={"repository": args.repository, "git_ref": raw_ref, "paths": paths},
                                  payload=payload)
    state = _git_analysis_state(gate)
    state["snapshots"][payload["snapshot_id"]] = {"commit": payload["commit"], "path": payload["source_path"]}
    save_request(gate)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def complete_free_request(gate: dict[str, Any], args: argparse.Namespace) -> int:
    if gate.get("operation") in {"template-management", "template-document"}:
        return templates_cli.complete(SimpleNamespace(**globals()), gate, args)
    output = None
    if gate["mode"] == "draft" and gate.get("interview_output") and (not args.output or str(args.output).lower().endswith(".xlsx")):
        record = gate["interview_output"]
        if args.output and args.output != record["path"]:
            raise InterviewError("INTERVIEW_OUTPUT_CHANGED", "Завершение требует последней проверенной книги этого запроса.")
        _, evidence = evidence_for_gate(gate)
        output = interview_cli.completion(request_root(gate), record, evidence.get("changed_files", []))
        gate.update(state="DRAFT_COMPLETE", document_status="UNVERIFIED_DRAFT", output=record["path"])
    elif gate["mode"] == "draft":
        output = safe_request_file(request_root(gate), args.output or "result.md", markdown=True)
        if not meaningful_file(output) or "UNVERIFIED_DRAFT" not in output.read_text(encoding="utf-8"):
            raise WorkflowError("A nonempty UNVERIFIED_DRAFT document is required")
        evidence_path, evidence = evidence_for_gate(gate)
        relative = str(output.relative_to(request_root(gate))).replace("\\", "/")
        if not any(item.get("path") == relative and item.get("sha256") == sha256(output) for item in evidence["changed_files"]):
            raise WorkflowError("The draft must match its recorded output")
        gate.update(state="DRAFT_COMPLETE", document_status="UNVERIFIED_DRAFT", output=relative)
    else:
        summary = str(getattr(args, "summary", "") or "").strip()
        if not summary:
            raise WorkflowError("Supply a consultation summary to record its result")
        if gate.get("operation") == "query-analysis" and gate.get("query_contract_version") == 1:
            _, evidence = evidence_for_gate(gate)
            checks = evidence.get("query_checks", [])
            if not checks:
                raise WorkflowError("Run query-check on the final query text before completing query-analysis")
            latest = checks[-1]
            candidate_path = request_root(gate) / "query-candidate.txt"
            if not candidate_path.is_file() or sha256(candidate_path) != latest.get("text_sha256"):
                raise WorkflowError("The final query changed after query-check; check the exact text again")
            if gate.get("query_intent") == "optimize":
                baseline_path = request_root(gate) / "query-baseline.txt"
                if not baseline_path.is_file() or sha256(baseline_path) != latest.get("baseline_sha256"):
                    raise WorkflowError("The original query changed after query-check; check both versions again")
            if gate.get("query_intent") in {"create", "optimize"} and latest.get("check_result") == "FAIL":
                raise WorkflowError("The candidate has a static error; correct it and run query-check again")
            for item in latest.get("schema_sources", []):
                source_path = Path(str(item.get("path", "")))
                if not source_path.is_file() or sha256(source_path) != item.get("sha256"):
                    raise WorkflowError("Metadata changed since query-check; refresh schema and check the candidate again")
            gate["query_verification"] = {
                "schema_version": 1, "verification_level": latest["verification_level"],
                "check_result": latest["check_result"], "text_sha256": latest["text_sha256"],
                "platform_executed": False, "performance_measured": False,
                "limitations": latest["limitations"],
            }
        gate.update(state="CONSULTATION_COMPLETE", result_summary=summary)
    gate["completed_at"] = utc_now()
    save_request(gate)
    result = {"state": gate["state"], "document_status": gate.get("document_status"),
                      "compliance": gate.get("compliance", "COMPLIANT"),
                      "deviations": gate.get("deviations", []),
                      "output": str(output) if output else None}
    if gate.get("query_verification"):
        result["query_verification"] = gate["query_verification"]
        result["query_text"] = (request_root(gate) / "query-candidate.txt").read_bytes().decode("utf-8")
    print(json.dumps(result, ensure_ascii=False))
    return 0


def cmd_interview_register(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_interview")
    if gate.get("operation") != "interview-preparation" or gate.get("mode") not in {"draft", "explore"}:
        raise InterviewError("INTERVIEW_GATE_REQUIRED", "Откройте interview-preparation gate в draft или explore.")
    if args.action == "write" and gate["mode"] != "draft":
        raise InterviewError("INTERVIEW_DRAFT_REQUIRED", "Запись книги доступна только в draft.")
    evidence_path, evidence = evidence_for_gate(gate)
    result = interview_cli.execute(args.action, args.request, request_root(gate), gate.get("artifacts", []), evidence.get("changed_files", []))
    if result["state"] == "WRITTEN":
        record = result["record"]
        if not any(item.get("path") == record["path"] and item.get("sha256") == record["sha256"] for item in evidence["changed_files"]):
            evidence["changed_files"].append({"target": "draft", **record})
        write_json(evidence_path, evidence)
        gate["interview_output"] = record
        gate["output"] = record["path"]
        save_request(gate)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_draft_promote(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"DRAFT_COMPLETE"})
    require_gate_tool(gate, "flow1c_promote")
    if not args.confirmed:
        raise WorkflowError("Explicit user agreement to attach this draft is required")
    reference = resolve_reference_args(args)
    code = reference["work_reference"]
    if not code:
        raise WorkflowError("A user-assigned project or task reference is required")
    _, manifest = load_manifest(code)
    requirements = list(args.requirements or gate.get("requirements", []))
    if not requirements:
        raise WorkflowError("Specify the requirement IDs this draft addresses")
    mismatch = reference_mismatch(code, requirements)
    if mismatch:
        raise WorkflowError(mismatch)
    links = read_json(project_root() / "registry/normalized/links.json", {}).get("requirement_to_specification", {})
    known = read_json(project_root() / "registry/normalized/requirements.json", {})
    if any(req not in known or (links.get(req) and links[req] != code) for req in requirements):
        raise WorkflowError("Requirement registry is missing or assigns a requirement to another work item")
    source = request_root(gate).resolve()
    target_root = (work_item_root(code) / "input" / "drafts").resolve()
    target = target_root / gate["request_id"]
    if not path_is_within(target, target_root):
        raise WorkflowError("Invalid promotion destination")
    if target.exists():
        raise WorkflowError("Draft has already been attached; existing materials are preserved")
    if any(is_reparse_or_symlink(p) for p in source.rglob("*")):
        raise WorkflowError("Draft contains a symlink or reparse point")
    shutil.copytree(source, target)
    write_json(target / "provenance.json", {"generated_by": "Flow1C", "request_id": gate["request_id"],
               "source": str(source), "code": code, "requirements": requirements, "document_status": "UNVERIFIED_DRAFT", "attached_at": utc_now()})
    print(json.dumps({"state": "ATTACHED", "destination": str(target), "document_status": "UNVERIFIED_DRAFT", "work_item_status": manifest["status"]}, ensure_ascii=False))
    return 0


def publication_snapshot(code: str, *, include_extension: bool = False) -> dict[str, Any]:
    root = work_item_root(code)
    files = {str(p.relative_to(root)).replace("\\", "/"): sha256(p) for p in sorted(root.rglob("*"))
             if p.is_file() and p.relative_to(root).parts[0] not in {"evidence", "context", ".git"}}
    snapshot: dict[str, Any] = {"files": files}
    if include_extension:
        _, manifest = load_manifest(code)
        diff, error = extension_git_state(code, manifest)
        if error or diff is None:
            raise WorkflowError(error or "Extension Git state is unavailable")
        path = Path(diff["path"])
        def git_output(arguments: list[str]) -> str:
            result = subprocess.run(["git", *arguments], cwd=path, capture_output=True, text=True, encoding="utf-8", errors="replace")
            if result.returncode:
                raise WorkflowError("Cannot capture extension Git state")
            return result.stdout
        snapshot["extension"] = {"head": diff["head"], "branch": diff["branch"],
            "base_head": git_output(["rev-parse", diff["base"]]).strip(),
            "dirty": git_output(["status", "--porcelain"]),
            "working_diff_sha256": hashlib.sha256(git_output(["diff", "--no-ext-diff", "--binary", "HEAD"]).encode("utf-8")).hexdigest()}
    return snapshot


def validate_publication(code: str, phase: str) -> None:
    _, manifest = load_manifest(code)
    if manifest.get("traceability_mode", "registry") != "registry" or manifest.get("registry", {}).get("status", "verified") != "verified":
        raise WorkflowError("Provisional or unverified-registry work items cannot be published")
    stage = load_stages()["operations"]["publish"]
    if manifest.get("status") not in stage["allowed_statuses"]:
        raise WorkflowError("Current work-item status does not allow publication")
    phases = {"specification": {"functional-spec", "functional-review"},
              "technical": {"technical-design", "technical-implementation", "code-review", "development"},
              "acceptance": {"testing"}}
    if phase not in phases:
        raise WorkflowError("Unknown publication phase")
    config, local = load_config()
    gitea = {**config.get("gitea", {}), **local.get("gitea", {})}
    reviewer_role = "technical" if phase == "technical" else "functional"
    if not gitea.get("reviewers", {}).get(reviewer_role):
        raise WorkflowError(f"No {reviewer_role} reviewers configured")
    index = load_artifact_index(code)
    present = {a.get("category") for a in index["artifacts"]} | set(index["confirmed_absent"])
    if any(category not in present for category in stage.get("confirm_absence", [])):
        raise WorkflowError("Organizational approvals must be supplied or their absence recorded before publication")
    for path in (work_item_root(code) / "evidence").glob("*.json"):
        evidence = read_json(path, {})
        if evidence.get("completion_state") != "COMPLETE" or evidence.get("operation") not in phases[phase]:
            continue
        snapshot = evidence.get("validation_snapshot")
        if not isinstance(snapshot, dict) or not evidence.get("gate_id"):
            continue
        completed_stage = load_stages()["operations"][evidence["operation"]]
        if any(manifest.get("approvals", {}).get(role) != status for role, status in completed_stage.get("required_approvals", {}).items()):
            continue
        if snapshot == publication_snapshot(code, include_extension="extension" in snapshot):
            if snapshot.get("extension", {}).get("dirty"):
                continue
            return
    raise WorkflowError("No current validated evidence matches this phase, documents and Git state. Run the relevant review again.")


def cmd_agent_begin(args: argparse.Namespace) -> int:
    stages = load_stages()
    operation = str(args.operation).strip().casefold()
    if operation == "consultation" and getattr(args, "mode", None) != "draft":
        inferred_intent = infer_query_request_intent(str(getattr(args, "summary", "") or ""))
        if inferred_intent:
            args.requested_operation = operation
            operation = "query-analysis"
            args.operation = operation
            if not getattr(args, "query_intent", None):
                args.query_intent = inferred_intent
    try:
        mode_selection = select_request_mode(operation, str(getattr(args, "summary", "") or ""), getattr(args, "mode", None))
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    mode_selection["selected_at"] = utc_now()
    mode = str(mode_selection["mode"])
    args.mode = mode
    args.mode_selection = mode_selection
    if mode != "formal":
        return begin_free_request(args)
    stage = stages["operations"].get(operation)
    if operation in {"consultation", "query-analysis", "workflow-review", "interview-preparation"}:
        raise WorkflowError("Use a free mode for consultation, query analysis, and workflow review")
    if not isinstance(stage, dict):
        payload = new_gate(
            operation or "unknown",
            None,
            "NEEDS_CONFIRMATION",
            summary=str(getattr(args, "summary", "") or ""),
            presented_paths=list(getattr(args, "path", []) or []),
            user_message="Намерение неоднозначно. Задайте пользователю один вопрос, различающий возможные операции.",
        )
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 1
    try:
        git_ref = validate_git_ref(getattr(args, "git_ref", None))
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    reference = resolve_reference_args(args)
    code = reference["work_reference"]
    registry_resolution: dict[str, Any] | None = None
    if code and stage.get("code_required"):
        registry_resolution = resolve_registry_reference(
            str(code), str(getattr(args, "reference_kind", "auto") or "auto")
        )
        if registry_resolution["state"] == "ambiguous":
            payload = new_gate(
                operation, str(code), "WAITING_USER", summary=args.summary, **reference,
                reference_kind=str(getattr(args, "reference_kind", "auto") or "auto"),
                registry_resolution=registry_resolution,
                mode_selection=mode_selection, mode_history=[mode_selection],
                clarification={
                    "reason": "registry_reference_kind",
                    "question": "Ссылка неоднозначна. Уточните, это ID требования или номер ФС; "
                                "если требование назначено нескольким ФС, укажите выбранный номер ФС.",
                    "options": registry_resolution["candidates"],
                },
                user_message="Требуется выбрать однозначную реестровую ссылку.",
            )
            print(json.dumps(payload, ensure_ascii=False, indent=2))
            return 1
        if registry_resolution["state"] == "resolved":
            requested_reference = str(code)
            code = str(registry_resolution["code"])
            reference = {
                **reference, "task_reference": code, "work_reference": code,
                "requested_reference": requested_reference,
            }
    assessment = assess_request(operation, mode, str(args.summary or ""), code=code,
                                code_required=bool(stage.get("code_required")), mismatch=str(getattr(args, "mismatch", "") or ""))
    if assessment["state"] == "WAITING_USER":
        payload = new_gate(operation, code, "WAITING_USER", summary=args.summary, **reference, **{k: v for k, v in assessment.items() if k != "state"})
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 1
    if stage.get("code_required") and not code:
        payload = new_gate(
            operation,
            None,
            "BLOCKED",
            summary=str(getattr(args, "summary", "") or ""),
            presented_paths=list(getattr(args, "path", []) or []),
            **reference,
            git_ref=git_ref,
            work_item_exists=False,
            mode_selection=mode_selection,
            mode_history=[mode_selection],
            clarification={"reason": "missing_work_item", "question": "Выберите способ продолжения без рабочего элемента.",
                           "options": ["provide-reference", "create-provisional", "independent-draft"]},
            user_message="Укажите пользовательский идентификатор, создайте provisional work-item после registry_bypass или продолжите независимым черновиком. Flow1C никогда не создаёт идентификатор автоматически.",
            conditions=[{
                "id": "missing:work_reference", "category": "work_reference",
                "message": "Пользовательский project/task reference отсутствует.", "source": "stage",
                "waivable": False, "blocking": True,
            }],
            remaining_blockers=[{"id": "missing:work_reference", "category": "work_reference",
                                 "message": "Пользовательский project/task reference отсутствует.", "source": "stage",
                                 "waivable": False, "blocking": True}],
        )
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 2

    if operation == "setup":
        profile = str(getattr(args, "profile", "") or "")
        if not profile:
            payload = new_gate(
                operation, None, "NEEDS_CONFIRMATION", **reference,
                summary=str(getattr(args, "summary", "") or ""),
                presented_paths=list(getattr(args, "path", []) or []),
                clarification={"reason": "profile", "question": "Какой уровень работы нужен? Рекомендуемый профиль для аналитика/архитектора — analysis; full выбирается только явно.",
                               "options": ["analysis", "conversation", "project-basic", "documents", "implementation", "full"]},
                user_message="Какой уровень работы нужен? Рекомендуемый профиль — analysis; full автоматически не выбирается.",
            )
            print(json.dumps(payload, ensure_ascii=False, indent=2))
            return 1
        setup_skill_path = ROOT / ".agents" / "skills" / str(stage["skill"]) / "SKILL.md"
        try:
            setup_state = read_setup_state(profile)
        except WorkflowError as exc:
            payload = new_gate(
                operation,
                None,
                "BLOCKED",
                summary=str(getattr(args, "summary", "") or ""),
                presented_paths=list(getattr(args, "path", []) or []),
                errors=[str(exc)],
                user_message=f"Настройка заблокирована: {exc}",
            )
            print(json.dumps(payload, ensure_ascii=False, indent=2))
            return 2
        setup_ready = bool(setup_state.get("ready")) and setup_state.get("state") == "READY"
        setup_gate_state = "READY" if setup_ready else "BLOCKED" if setup_state.get("state") == "BLOCKED" else "NEEDS_CONFIRMATION"
        payload = new_gate(
            operation,
            None,
            setup_gate_state,
            summary=str(getattr(args, "summary", "") or ""),
            presented_paths=list(getattr(args, "path", []) or []),
            skill=str(stage["skill"]),
            skill_instructions=setup_skill_path.read_text(encoding="utf-8") if setup_skill_path.is_file() else "",
            setup_state=setup_state,
            requested_profile=profile,
            action_completed="setup-audit" if setup_ready else None,
            user_message=setup_user_message(setup_state),
        )
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 1

    manifest: dict[str, Any] = {}
    item_root: Path | None = None
    if code:
        try:
            _, manifest = load_manifest(code)
            item_root = work_item_root(code)
        except WorkflowError as exc:
            payload = new_gate(
                operation,
                code,
                "BLOCKED",
                errors=[str(exc)],
                **reference,
                git_ref=git_ref,
                work_item_exists=False,
                mode_selection=mode_selection,
                mode_history=[mode_selection],
                registry_resolution=registry_resolution,
                reference_kind=(registry_resolution or {}).get("reference_kind", getattr(args, "reference_kind", "auto")),
                clarification={"reason": "missing_work_item", "question": "Рабочий элемент не найден. Как продолжить?",
                               "options": ["provide-reference", "create-provisional", "independent-draft"]},
                user_message=(
                    f"Рабочий элемент {code} не создан. "
                    + ("Выбранная связка найдена в реестре; запустите обычный fs-start."
                       if (registry_resolution or {}).get("state") == "resolved"
                       else "Исправьте/предоставьте реестр, подтвердите registry_bypass и создайте provisional, либо выберите независимый черновик.")
                ),
                conditions=[{
                    "id": "missing:registry_traceability", "category": "registry_traceability",
                    "message": (f"Связка для {code} пригодна, но work-item ещё не создан."
                                if (registry_resolution or {}).get("state") == "resolved"
                                else f"Рабочий элемент {code} отсутствует в пригодной области реестра."),
                    "source": "registry", "waivable": bool(stage.get("provisional_allowed"))
                    and (registry_resolution or {}).get("state") != "resolved", "blocking": True,
                }],
                remaining_blockers=[{"id": "missing:registry_traceability", "category": "registry_traceability",
                                     "message": (f"Связка для {code} пригодна, но work-item ещё не создан."
                                                 if (registry_resolution or {}).get("state") == "resolved"
                                                 else f"Рабочий элемент {code} отсутствует в пригодной области реестра."),
                                     "source": "registry", "waivable": bool(stage.get("provisional_allowed"))
                                     and (registry_resolution or {}).get("state") != "resolved", "blocking": True}],
            )
            print(json.dumps(payload, ensure_ascii=False, indent=2))
            return 2

    errors: list[str] = stage_errors(stage, manifest)
    traceability_mode = str(manifest.get("traceability_mode") or "registry") if manifest else "registry"
    if traceability_mode == "provisional" and not stage.get("provisional_allowed"):
        errors.append("provisional traceability is not allowed for this operation")
    if traceability_mode == "provisional" and stage.get("requires_registry_traceability"):
        errors.append("validated registry traceability is required for this operation")

    found_inputs: list[dict[str, str]] = []
    required_missing: list[dict[str, str]] = []
    if item_root:
        for relative in stage.get("required_files", []):
            path = item_root / relative
            if not meaningful_file(path):
                required_missing.append({"id": relative, "label": f"файл {relative}", "path": str(path)})
            else:
                found_inputs.append({"id": relative, "label": f"файл {relative}", "path": str(path)})
        for input_set in stage.get("input_sets", []):
            alternatives = list(input_set.get("alternatives", []))
            expected = "input/user-brief.md" if traceability_mode == "provisional" else "input/requirements.snapshot.yaml"
            acceptable = [expected] if expected in alternatives else alternatives
            selected = next((relative for relative in acceptable if meaningful_file(item_root / relative)), None)
            if selected:
                found_inputs.append({"id": str(input_set.get("id")), "label": f"источник {selected}", "path": str(item_root / selected)})
            else:
                required_missing.append({"id": str(input_set.get("id")), "label": "альтернативный источник требований",
                                         "path": " | ".join(str(item_root / relative) for relative in acceptable)})
    _, local = load_config()
    for key in stage.get("required_config", []):
        raw = str(local.get(key, "")).strip()
        configured_file = Path(raw) if raw else None
        valid = bool(configured_file and configured_file.is_file())
        if key == "functional_spec_template" and valid:
            valid = configured_file.suffix.lower() == ".docx"
        if key == "functional_spec_template" and local.get("documentation_path"):
            try:
                library = TemplateService(ROOT, local)
                usable = [v for v in library.list()["templates"] if v["document_type"] == "functional-spec" and v["ready"]]
                if usable:
                    valid = True
                    raw = str(library.store.root)
            except (TemplateError, OSError, ValueError):
                pass  # legacy input remains compatible; library CLI reports specific recovery diagnostics
        if not valid:
            required_missing.append({"id": key, "label": f"настройка {key}", "path": raw or str(ROOT / LOCAL_CONFIG_FILE)})
        else:
            found_inputs.append({"id": key, "label": f"настройка {key}", "path": raw})

    artifact_index = load_artifact_index(code)
    present_categories = {item.get("category") for item in artifact_index["artifacts"]}
    for item in artifact_index["artifacts"]:
        if isinstance(item, dict):
            found_inputs.append(
                {
                    "id": str(item.get("category", "artifact")),
                    "label": str(stages.get("artifact_categories", {}).get(item.get("category"), {}).get("label", item.get("name", "материал"))),
                    "path": str(project_root() / str(item.get("relative_path", ""))),
                }
            )
    confirmed_absent = set(artifact_index["confirmed_absent"])
    for category in stage.get("required_artifacts", []):
        if category in present_categories:
            continue
        category_config = stages.get("artifact_categories", {}).get(category, {})
        destination = intake_destination(code, category, "<intake-id>", stages)
        required_missing.append(
            {"id": category, "label": category_config.get("label", category), "path": str(destination)}
        )
    conditional_missing: list[dict[str, str]] = []
    for category in stage.get("confirm_absence", []):
        if category in present_categories or category in confirmed_absent:
            continue
        category_config = stages.get("artifact_categories", {}).get(category, {})
        destination = intake_destination(code, category, "<intake-id>", stages)
        conditional_missing.append(
            {"id": category, "label": category_config.get("label", category), "path": str(destination)}
        )

    diff: dict[str, Any] | None = None
    if code and (stage.get("requires_diff") or stage.get("requires_extension_branch")):
        diff, diff_error = extension_git_state(code, manifest)
        if diff_error:
            errors.append(diff_error)
        elif stage.get("requires_extension_branch") and diff and diff["branch"] != diff["expected_branch"]:
            errors.append(f"extension branch is {diff['branch']}; expected {diff['expected_branch']}")
        if stage.get("requires_diff") and diff is not None and not diff["files"]:
            required_missing.append(
                {"id": "extension_diff", "label": "точный Git diff расширения", "path": str(diff["path"])}
            )

    # Infrastructure is checked by the source/analysis action, not before a useful question.
    if stage.get("requires_validated_output") and code:
        evidence_root = work_item_root(code) / "evidence"
        validated = []
        for candidate in evidence_root.glob("*.json"):
            value = read_json(candidate, {})
            if isinstance(value, dict) and value.get("completion_state") == "COMPLETE":
                validated.append(candidate.name)
        if not validated:
            errors.append("no previously validated COMPLETE evidence is available for publication")
        config, local_config = load_config()
        gitea = {**config.get("gitea", {}), **local_config.get("gitea", {})}
        reviewers = gitea.get("reviewers", {}) if isinstance(gitea.get("reviewers"), dict) else {}
        if not any(isinstance(value, list) and value for value in reviewers.values()):
            errors.append("Gitea reviewer list is not configured")

    allow_incomplete = bool(getattr(args, "allow_incomplete_draft", False))
    if allow_incomplete and "extension" in stage.get("writable_targets", []):
        errors.append("an incomplete draft may not be used for extension development")
    if allow_incomplete and stage.get("requires_validated_output"):
        errors.append("an incomplete draft may not be published")
    state = readiness_state(errors, required_missing, conditional_missing, allow_incomplete)
    conditions = build_conditions(stage, errors=errors, required_missing=required_missing,
                                  conditional_missing=conditional_missing)
    registry_choice = operation == "registry" and any(item.get("id") == "requirements_workbook" for item in required_missing)
    if registry_choice:
        state = "NEEDS_CONFIRMATION"
    skill_path = ROOT / ".agents" / "skills" / str(stage["skill"]) / "SKILL.md"
    payload = new_gate(
        operation,
        code,
        state,
        **reference,
        git_ref=git_ref,
        summary=str(getattr(args, "summary", "") or ""),
        presented_paths=list(getattr(args, "path", []) or []),
        mode_selection=mode_selection,
        mode_history=[mode_selection],
        registry_resolution=registry_resolution,
        reference_kind=(registry_resolution or {}).get("reference_kind", getattr(args, "reference_kind", "auto")),
        skill=str(stage["skill"]),
        skill_instructions=skill_path.read_text(encoding="utf-8") if skill_path.is_file() else "",
        traceability_mode=traceability_mode,
        work_item_exists=bool(item_root),
        manifest_status=manifest.get("status") if manifest else None,
        manifest_approvals=manifest.get("approvals", {}) if manifest else {},
        context_role=stage.get("context_role"),
        output=stage.get("output"),
        errors=errors,
        found_inputs=found_inputs,
        missing_required=required_missing,
        missing_conditional=conditional_missing,
        clarification=({"reason": "registry_unavailable", "question": "Реестр не предоставлен. Как продолжить?",
                        "options": ["provide-or-fix-registry", "create-provisional", "independent-draft"]}
                       if registry_choice else None),
        diff=diff,
        conditions=conditions,
        remaining_blockers=[item for item in conditions if item.get("blocking")],
        user_message=(
            "Реестр не предоставлен. Можно предоставить/исправить реестр, продолжить provisional после явного registry_bypass или создать независимый черновик."
            if registry_choice
            else build_input_message(found_inputs, required_missing, conditional_missing)
            if state == "NEEDS_INPUT"
            else (
                "Техническая блокировка: "
                + "; ".join(errors)
                + (("\n\n" + build_input_message(found_inputs, required_missing, conditional_missing)) if required_missing or conditional_missing else "")
                if errors
                else ("Создан непроверенный черновой gate." if state == "UNVERIFIED_DRAFT" else "Входные проверки пройдены.")
            )
        ),
    )
    if traceability_mode == "provisional":
        bypass = manifest.get("registry", {}).get("bypass", {})
        if bypass:
            payload.update(compliance="DEVIATED", deviations=[bypass])
        if state == "READY":
            state = "READY_WITH_DEVIATIONS"
            payload.update(state=state, remaining_blockers=[])
    payload["state"] = state
    payload["available_actions"] = available_actions(mode="formal", operation=operation, state=state, stage=stage)
    if state in {"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT", "NEEDS_INPUT", "BLOCKED"} and code:
        basis = manifest.get("requirement_basis", {})
        evidence = {
            "schema_version": 1,
            "operation": operation,
            "code": code,
            "project_reference": reference.get("project_reference"),
            "task_reference": reference.get("task_reference"),
            "work_reference": code,
            "requested_reference": reference.get("requested_reference"),
            "reference_resolution": (registry_resolution or {}).get("reference_kind"),
            "context_sha256": "",
            "artifacts": artifact_index["artifacts"],
            "diff": None,
            "diff_inventory": diff,
            "rlm_queries": [],
            "source_reads": [],
            "bsl_ls": None,
            "changed_files": [],
            "created_at": utc_now(),
        }
        if basis.get("sha256"):
            evidence.update(
                traceability_mode=traceability_mode,
                requirement_basis={"type": basis.get("type", "registry_snapshot"), "sha256": basis["sha256"]},
                registry_snapshot_sha256=manifest.get("registry", {}).get("snapshot_sha256"),
                revalidation_required=traceability_mode == "provisional",
            )
        if payload.get("deviations"):
            evidence.update(deviations=payload.get("deviations", []), revalidation_required=traceability_mode == "provisional")
        evidence["gate_id"] = payload["gate_id"]
        evidence_path = work_item_root(code) / "evidence" / f"{operation}-{payload['gate_id']}.json"
        write_json(evidence_path, evidence)
        payload["evidence_path"] = str(evidence_path)
        save_gate(payload)
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if state == "READY" else (2 if state == "BLOCKED" else 1)


def cmd_artifact_intake(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"NEEDS_CODE", "NEEDS_INPUT", "NEEDS_CONFIRMATION", "READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT", "WAITING_USER", "BLOCKED"})
    require_gate_tool(gate, "flow1c_intake")
    if gate.get("mode", "formal") != "formal":
        return intake_free_request(gate, args)
    stages = load_stages()
    category = str(args.category)
    if category not in stages.get("artifact_categories", {}):
        raise WorkflowError(f"Unknown artifact category: {category}")
    code = str(getattr(args, "task_reference", "") or getattr(args, "code", "") or gate.get("work_reference") or gate.get("code") or "").strip() or None
    if code:
        load_manifest(code)
    index = load_artifact_index(code)
    promote_intake_id = str(getattr(args, "promote_intake_id", "") or "").strip()
    if promote_intake_id:
        if not code:
            raise WorkflowError("Inbox promotion requires a user-assigned task reference.")
        if not re.fullmatch(r"[A-Za-z0-9-]{8,64}", promote_intake_id):
            raise WorkflowError("Invalid intake ID for promotion.")
        inbox_root = project_root() / "inbox" / promote_intake_id
        inbox_manifest_path = inbox_root / "intake.json"
        inbox_manifest = read_json(inbox_manifest_path)
        if not isinstance(inbox_manifest, dict) or inbox_manifest.get("code") is not None:
            raise WorkflowError(f"Inbox intake is unavailable for promotion: {promote_intake_id}")
        records = [item for item in inbox_manifest.get("artifacts", []) if isinstance(item, dict)]
        if not records:
            raise WorkflowError("Inbox intake contains no artifacts to promote.")
        categories = {str(item.get("category", "")) for item in records}
        if len(categories) != 1:
            raise WorkflowError("An inbox intake with mixed categories cannot be promoted as one package.")
        original_root = inbox_root / "originals"
        destination = intake_destination(code, categories.pop(), promote_intake_id, stages)
        if destination.exists():
            raise WorkflowError(f"Promotion destination already exists: {destination}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(original_root), str(destination))
        promoted: list[dict[str, Any]] = []
        for record in records:
            old_path = project_root() / str(record.get("relative_path", ""))
            relative = old_path.relative_to(original_root)
            target = destination / relative
            promoted_record = {
                **record,
                "relative_path": str(target.relative_to(project_root())).replace("\\", "/"),
                "received_via": "inbox-promotion",
                "promoted_at": utc_now(),
                "derived_path": write_derived_artifact(target, relative, code, promote_intake_id),
            }
            if promoted_record.get("sha256") not in {item.get("sha256") for item in index["artifacts"]}:
                index["artifacts"].append(promoted_record)
            promoted.append(promoted_record)
        write_json(artifact_index_path(code), index)
        inbox_manifest.update({"state": "promoted", "code": code, "promoted_at": utc_now(), "artifacts": promoted})
        write_json(inbox_manifest_path, inbox_manifest)
        print(json.dumps({"state": "ACCEPTED", "intake_id": promote_intake_id, "code": code, "destination": str(destination), "copied": promoted, "skipped": []}, ensure_ascii=False, indent=2))
        return 0
    for value in getattr(args, "confirm_absence", []) or []:
        if value not in stages.get("artifact_categories", {}):
            raise WorkflowError(f"Unknown artifact category: {value}")
        if value not in index["confirmed_absent"]:
            index["confirmed_absent"].append(value)
    files, skipped = enumerate_intake_files(getattr(args, "source", []) or [])
    if (getattr(args, "source", []) or []) and not files and not (getattr(args, "confirm_absence", []) or []):
        raise WorkflowError(
            "No acceptable artifacts were found. Executables, scripts, archives, Office temporary files, symlinks and unsupported formats are rejected."
        )
    total_bytes = sum(path.stat().st_size for path, _ in files)
    if (len(files) > MAX_INTAKE_FILES or total_bytes > MAX_INTAKE_BYTES) and not args.confirm_large:
        raise WorkflowError(
            f"Intake contains {len(files)} files and {total_bytes} bytes; explicit large-batch confirmation is required."
        )
    received_via = str(getattr(args, "received_via", "") or ("folder" if any(Path(raw).is_dir() for raw in (getattr(args, "source", []) or [])) else "file"))
    result = persist_intake_files(files, skipped, code=code, category=category, received_via=received_via, index=index)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_context_build(args: argparse.Namespace) -> int:
    code = args.code
    _, manifest = load_manifest(code)
    data_root = project_root()
    item_root = work_item_root(code)
    traceability_mode = str(manifest.get("traceability_mode") or "registry")
    snapshot = read_json(item_root / "input" / "requirements.snapshot.yaml", {})
    brief_path = item_root / "input" / "user-brief.md"
    role_files = {
        "analyst": ["analysis/questions.md", "analysis/answers.md", "analysis/decisions.md"],
        "functional-architect": ["analysis/traceability.md", "specification/functional-spec.md"],
        "technical-architect": ["analysis/traceability.md", "specification/functional-spec.md", "specification/technical-design.md"],
        "tester": ["analysis/traceability.md", "specification/functional-spec.md", "testing/test-plan.md"],
    }
    lines = [
        f"# Context: {code} — {args.role}",
        "",
        "> Generated file. Rebuild it instead of editing it manually.",
        "",
        f"Status: `{manifest.get('status')}`",
        f"Title: {manifest.get('title')}",
        f"Requirements basis: {'user brief' if traceability_mode == 'provisional' else 'registry snapshot'}",
        f"Traceability: {'PROVISIONAL / UNVERIFIED_DRAFT' if traceability_mode == 'provisional' else 'REGISTRY'}",
        "",
        "## Requirements basis",
        "",
    ]
    if traceability_mode == "provisional":
        lines.extend([brief_path.read_text(encoding="utf-8-sig", errors="replace") if brief_path.is_file() else "User brief is missing.", ""])
        if manifest.get("requirements"):
            lines.extend(["User-supplied identifiers (not registry-confirmed): " + ", ".join(manifest["requirements"]), ""])
    else:
        for requirement_id in manifest.get("requirements", []):
            requirement = snapshot.get(requirement_id, {})
            lines.extend(
                [f"### {requirement_id}", "", str(requirement.get("text") or "No requirement text in the registry."), "",
                 "Process: `" + json.dumps(requirement.get("process", {}), ensure_ascii=False) + "`", ""]
            )
    lines.extend(["## Work item excerpts", ""])
    for relative in role_files[args.role]:
        path = item_root / relative
        if path.exists():
            value = path.read_text(encoding="utf-8").strip()
            if value:
                lines.extend([f"### {relative}", "", value, ""])
    artifact_index = load_artifact_index(code)
    lines.extend(["## Accepted artifacts", ""])
    if artifact_index["artifacts"]:
        for artifact in artifact_index["artifacts"]:
            lines.append(
                f"- `{artifact.get('category')}`: `{artifact.get('relative_path')}` "
                f"(SHA-256 `{artifact.get('sha256')}`)"
            )
    else:
        lines.append("No user artifacts have been accepted for this work item.")
    if artifact_index["confirmed_absent"]:
        lines.extend(["", "Confirmed absent: " + ", ".join(artifact_index["confirmed_absent"]), ""])
    target = item_root / "context" / f"{args.role}.md"
    write_text(target, "\n".join(lines))
    print(target)
    return 0


def evidence_for_gate(gate: dict[str, Any]) -> tuple[Path, dict[str, Any]]:
    raw = str(gate.get("evidence_path", "")).strip()
    if not raw:
        raise WorkflowError("This gate has no evidence record.")
    path = Path(raw).resolve()
    expected_root = request_root(gate) if gate.get("mode", "formal") != "formal" else (work_item_root(str(gate.get("work_reference") or gate.get("code"))) / "evidence").resolve()
    if not path_is_within(path, expected_root):
        raise WorkflowError("Evidence path is outside the current work item.")
    evidence = read_json(path)
    if not isinstance(evidence, dict):
        raise WorkflowError("Evidence record is missing or invalid.")
    if evidence.get("gate_id") and evidence["gate_id"] != gate["gate_id"]:
        raise WorkflowError("Evidence belongs to a different gate")
    if int(evidence.get("schema_version", 1) or 1) < 2:
        evidence["schema_version"] = 2
        migrated = []
        for item in evidence.get("git_analysis", []):
            if isinstance(item, dict) and isinstance(item.get("result"), dict):
                item = {**item, "result": migrate_git_record(item["result"])}
            migrated.append(item)
        evidence["git_analysis"] = migrated
        evidence.setdefault("snapshots", [])
    return path, evidence


def cmd_agent_context(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_context")
    code = str(gate.get("code") or "")
    role = str(gate.get("context_role") or "")
    if not code or role not in VALID_ROLES:
        raise WorkflowError("This operation does not define a role context.")
    command_args = argparse.Namespace(code=code, role=role)
    output = io.StringIO()
    from contextlib import redirect_stdout

    with redirect_stdout(output):
        cmd_context_build(command_args)
    target = Path(output.getvalue().strip()).resolve()
    evidence_path, evidence = evidence_for_gate(gate)
    isolated_target = evidence_path.parent / f"context-{gate['gate_id']}.md"
    shutil.copy2(target, isolated_target)
    target = isolated_target
    digest = sha256(target)
    evidence_id = f"CTX-{digest[:12]}"
    evidence["context_sha256"] = digest
    evidence["context_path"] = str(target)
    evidence["context_evidence_id"] = evidence_id
    write_json(evidence_path, evidence)
    gate["context_sha256"] = digest
    save_gate(gate)
    print(json.dumps({"state": "READY", "evidence_id": evidence_id, "path": str(target), "sha256": digest, "content": target.read_text(encoding="utf-8")}, ensure_ascii=False, indent=2))
    return 0


def cmd_agent_diff(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_diff")
    code = str(gate.get("work_reference") or gate.get("task_reference") or gate.get("code") or "")
    if not code:
        raise WorkflowError("A user-assigned task reference is required for an extension diff.")
    _, manifest = load_manifest(code)
    diff_state, error = extension_git_state(code, manifest)
    if error or diff_state is None:
        raise WorkflowError(error or "Extension diff is unavailable.")
    extension = Path(diff_state["path"])
    result = subprocess.run(
        ["git", "diff", "--no-ext-diff", "--unified=80", f"{diff_state['base']}...HEAD"],
        cwd=extension,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        raise WorkflowError(result.stderr.strip() or "Cannot build extension diff.")
    if len(result.stdout) > int(args.max_chars):
        raise WorkflowError(f"Bounded diff exceeds {args.max_chars} characters; narrow it before analysis.")
    digest = hashlib.sha256(result.stdout.encode("utf-8")).hexdigest()
    evidence_id = f"DIFF-{digest[:12]}"
    evidence_path, evidence = evidence_for_gate(gate)
    evidence["diff"] = {**diff_state, "sha256": digest, "evidence_id": evidence_id}
    write_json(evidence_path, evidence)
    print(json.dumps({"state": "READY", "evidence_id": evidence_id, "diff": diff_state, "sha256": digest, "content": result.stdout}, ensure_ascii=False, indent=2))
    return 0


def cmd_agent_source_read(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_source_read")
    if args.source != "extension":
        raise WorkflowError("Direct configuration reads are forbidden; use source-query.")
    try:
        local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
    except WorkflowError as exc:
        print(json.dumps({
            "state": "NEEDS_INPUT",
            "available_actions": ["flow1c_dialogue"],
            "user_message": "Локальная настройка источника расширения недоступна. Можно продолжить общую часть черновика или уточнить подключение.",
            "errors": [{"code": "EXTENSION_CONFIG_INVALID", "message": str(exc)}],
        }, ensure_ascii=False))
        return 1
    root_raw = str(local.get("extension_path", "")).strip()
    if not root_raw:
        print(json.dumps({
            "state": "NEEDS_INPUT",
            "available_actions": ["flow1c_dialogue"],
            "user_message": "Путь к checkout расширения не настроен. Можно продолжить общую часть черновика или уточнить подключение.",
            "errors": [{"code": "EXTENSION_PATH_NOT_CONFIGURED", "message": "extension_path is not configured"}],
        }, ensure_ascii=False))
        return 1
    root = Path(root_raw).resolve()
    if not root.is_dir():
        print(json.dumps({
            "state": "NEEDS_INPUT",
            "available_actions": ["flow1c_dialogue"],
            "user_message": "Настроенный checkout расширения недоступен. Можно продолжить общую часть черновика или уточнить подключение.",
            "errors": [{"code": "EXTENSION_PATH_INVALID", "message": "extension_path is not an existing directory"}],
        }, ensure_ascii=False))
        return 1
    target = (root / args.path).resolve()
    if not path_is_within(target, root) or not target.is_file():
        raise WorkflowError("Requested extension file is outside the configured checkout or does not exist.")
    if target.suffix.casefold() not in {".bsl", ".xml", ".json", ".md"}:
        raise WorkflowError("Only BSL, XML, JSON and Markdown extension sources may be read.")
    content = target.read_text(encoding="utf-8-sig", errors="replace")
    if len(content) > int(args.max_chars):
        raise WorkflowError(f"Source exceeds {args.max_chars} characters; request a narrower file or use RLM.")
    digest = sha256(target)
    evidence_id = f"SRC-{digest[:12]}"
    evidence_path, evidence = evidence_for_gate(gate)
    relative = str(target.relative_to(root)).replace("\\", "/")
    source_reads = evidence.setdefault("source_reads", [])
    existing = next((item for item in source_reads if isinstance(item, dict)
                     and item.get("source") == "extension" and item.get("path") == relative
                     and item.get("sha256") == digest), None)
    if existing:
        evidence_id = str(existing.get("evidence_id") or evidence_id)
    else:
        source_reads.append(
            {"evidence_id": evidence_id, "source": "extension", "path": relative, "sha256": digest, "created_at": utc_now()}
        )
        write_json(evidence_path, evidence)
    print(json.dumps({"state": "READY", "evidence_id": evidence_id, "path": str(target), "sha256": digest, "content": content}, ensure_ascii=False, indent=2))
    return 0


def cmd_source_query(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_source_query")
    config, local = load_config()
    endpoint = str(config.get("quality", {}).get("rlm_endpoint", "")).strip()
    source_selector = args.source
    source_kind = str(source_selector.get("kind", "")) if isinstance(source_selector, dict) else str(source_selector or "")
    source_provenance: dict[str, Any] = {"kind": source_kind}
    if source_kind != "git_snapshot":
        ready, errors = rlm_readiness()
        if not ready:
            query_work = gate.get("operation") == "query-analysis"
            print(json.dumps({"state": "NEEDS_INPUT", "code": "RLM_SOURCE_UNAVAILABLE",
                              "available_actions": gate["available_actions"] if query_work else ["flow1c_dialogue", "flow1c_complete"],
                              "next_actions": (["continue-static-query-check"] if query_work else
                                               ["bootstrap-rlm-if-missing", "start-rlm", "ensure-source-indexes", "retry-source-query"]),
                              "user_message": ("RLM недоступен. Продолжите независимый анализ текста через flow1c_query_check; "
                                               "метаданные останутся неподтверждёнными." if query_work else
                                               "RLM недоступен. Восстановите установленный сервис и индексы, затем повторите запрос к источнику."),
                              "errors": errors}, ensure_ascii=False))
            return 1
    if source_kind == "git_snapshot":
        snapshot_id = str(source_selector.get("snapshot_id", ""))
        snapshot = _git_analysis_state(gate).get("snapshots", {}).get(snapshot_id)
        if not isinstance(snapshot, dict):
            raise WorkflowError("The requested Git snapshot does not belong to this gate")
        requested_commit = str(source_selector.get("commit", "") or snapshot.get("commit", ""))
        if requested_commit != snapshot.get("commit"):
            raise WorkflowError("Snapshot commit does not match the saved manifest")
        source_root = Path(str(snapshot.get("path", ""))).resolve()
        snapshot_root = (ROOT / ".workspace" / "git-analysis").resolve()
        if not path_is_within(source_root, snapshot_root) or not source_root.is_dir():
            raise WorkflowError("Snapshot source is unavailable or outside the managed cache")
        source_provenance.update(snapshot_id=snapshot_id, commit=requested_commit,
                                 repository=source_selector.get("repository", "extension"), path=str(source_root))
    else:
        if source_kind not in {"configuration", "extension"}:
            raise WorkflowError("source must be configuration, extension, or a git_snapshot selector")
        raw_path = str(local.get("configuration_path" if source_kind == "configuration" else "extension_path", "")).strip()
        if not raw_path:
            raise WorkflowError(f"{source_kind} source path is not configured.")
        source_root = resolve_1c_source_root(Path(raw_path))
        source_provenance["path"] = str(source_root)
    if source_kind == "git_snapshot":
        errors = [] if endpoint_is_healthy(endpoint) else ["RLM MCP endpoint is not healthy"]
    else:
        errors = []
    if errors:
        query_work = gate.get("operation") == "query-analysis"
        print(json.dumps({"state": "RECOVERABLE_ERROR", "code": "RLM_SNAPSHOT_INDEX_FAILED" if source_kind == "git_snapshot" else "RLM_SOURCE_UNAVAILABLE",
                          "source": source_provenance, "git_evidence_preserved": True,
                          "available_actions": gate["available_actions"] if query_work else ["flow1c_dialogue", "flow1c_complete"],
                          "next_actions": (["continue-static-query-check"] if query_work else
                                           ["bootstrap-rlm-if-missing", "start-rlm", "retry-source-query"]),
                          "user_message": ("RLM недоступен. Продолжите независимый анализ текста через flow1c_query_check; "
                                           "метаданные останутся неподтверждёнными." if query_work else
                                           "RLM недоступен. Восстановите сервис и повторите запрос; статический Git-анализ сохранён."),
                          "errors": errors}, ensure_ascii=False))
        return 1
    code = str(args.code or "").strip()
    if not code:
        raise WorkflowError(
            "RLM execution code is required. The query describes the evidence goal; "
            "code must be Python for rlm_execute and must print the result."
        )
    result, error = rlm_session_execute(
        endpoint,
        source_root,
        args.query,
        code,
        effort=args.effort,
        max_output_chars=args.max_chars,
        timeout=120.0,
        domains=["весь каталог"],
    )
    if error or result is None or result.get("error"):
        message = error or str(result.get("error") if result else "RLM query failed")
        evidence_path, evidence = evidence_for_gate(gate)
        evidence.setdefault("rlm_queries", []).append({"source": source_provenance, "query": args.query,
            "status": "recoverable_error", "code": "RLM_SNAPSHOT_INDEX_FAILED" if source_kind == "git_snapshot" else "RLM_QUERY_FAILED",
            "message": message, "created_at": utc_now()})
        write_json(evidence_path, evidence)
        query_work = gate.get("operation") == "query-analysis"
        print(json.dumps({"state": "RECOVERABLE_ERROR", "code": "RLM_SNAPSHOT_INDEX_FAILED" if source_kind == "git_snapshot" else "RLM_QUERY_FAILED",
                          "source": source_provenance, "message": message, "git_evidence_preserved": True,
                          **({"available_actions": gate["available_actions"]} if query_work else {}),
                          "next_actions": (["continue-static-query-check"] if query_work else
                                           ["check-rlm-health-and-index", "retry-source-query", "complete-static-analysis-if-recovery-fails"])}, ensure_ascii=False, indent=2))
        return 1
    if not str(result.get("stdout", "")).strip():
        raise WorkflowError("RLM query execution produced no output; code must print a result.")
    evidence_path, evidence = evidence_for_gate(gate)
    query_id = f"RLM-{len(evidence.get('rlm_queries', [])) + 1:03d}"
    record = {
        "id": query_id,
        "source": source_provenance,
        "query": args.query,
        "code": code,
        "reason": args.reason,
        "status": "retrieved_unvalidated" if gate.get("operation") == "query-analysis" else "confirmed",
        "result_sha256": hashlib.sha256(json.dumps(result, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest(),
        "result": result,
        "created_at": utc_now(),
    }
    evidence.setdefault("rlm_queries", []).append(record)
    write_json(evidence_path, evidence)
    print(json.dumps({"state": "RETRIEVED_UNVALIDATED" if gate.get("operation") == "query-analysis" else "CONFIRMED", "evidence_id": query_id, "source": source_provenance,
                      "result": result}, ensure_ascii=False, indent=2))
    return 0


def cmd_cc_inspect(args: argparse.Namespace) -> int:
    """Run one read-only CC inspector against accepted or configured XML."""
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_cc_inspect")
    if gate.get("operation") != "query-analysis":
        raise WorkflowError("CC query inspection is available only in the query-analysis scenario")
    local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
    configuration_raw = str(local.get("configuration_path", "") or "").strip()
    extension_raw = str(local.get("extension_path", "") or "").strip()
    try:
        target, source = resolve_input_path(
            source=args.source,
            relative_path=args.path,
            request_root=request_root(gate),
            request_artifacts=[item for item in gate.get("artifacts", []) if isinstance(item, dict)],
            configuration_root=resolve_1c_source_root(Path(configuration_raw)) if configuration_raw else None,
            extension_root=resolve_1c_source_root(Path(extension_raw)) if extension_raw else None,
        )
        result = run_inspection(
            checkout=ROOT / ".tools" / "cc-1c-skills", operation=args.operation,
            target=target, name=args.name, max_chars=args.max_chars,
        )
    except CcInspectionUnavailable as exc:
        print(json.dumps({
            "state": "NEEDS_INPUT", "limitation": "CC_SKILLS_UNAVAILABLE",
            "user_message": f"{exc}. Продолжите независимый анализ текста запроса; эти метаданные не подтверждены.",
            "available_actions": gate["available_actions"],
        }, ensure_ascii=False, indent=2))
        return 1
    except CcInspectionError as exc:
        raise WorkflowError(str(exc)) from exc
    evidence_path, evidence = evidence_for_gate(gate)
    inspection_id = f"CC-{len(evidence.get('cc_inspections', [])) + 1:03d}"
    record = {
        "id": inspection_id, "source": source,
        "source_kind": "local_xml_export" if source != "request" else "accepted_xml_artifact",
        "path": str(target), "skill": result["skill"], "operation": result["operation"],
        "result_sha256": hashlib.sha256(result["output"].encode("utf-8")).hexdigest(),
        "created_at": utc_now(),
    }
    evidence.setdefault("cc_inspections", []).append(record)
    write_json(evidence_path, evidence)
    print(json.dumps({
        "state": "CONFIRMED_LOCAL_XML", "evidence_id": inspection_id,
        "verification_level": "metadata", "source_kind": record["source_kind"],
        "database_verified": False, "execution_verified": False, "result": result,
    }, ensure_ascii=False, indent=2))
    return 0


def cmd_query_schema(args: argparse.Namespace) -> int:
    """Record exact fields from one XML export, with source provenance."""
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_query_schema")
    if gate.get("operation") != "query-analysis":
        raise WorkflowError("query-schema is available only for query-analysis")
    source = str(args.source or "")
    local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
    configuration_raw = str(local.get("configuration_path", "") or "").strip()
    extension_raw = str(local.get("extension_path", "") or "").strip()
    configuration_root = resolve_1c_source_root(Path(configuration_raw)) if source == "configuration" and configuration_raw else None
    extension_root = resolve_1c_source_root(Path(extension_raw)) if source == "extension" and extension_raw else None
    try:
        target, _ = resolve_input_path(
            source=source, relative_path=str(args.path or ""), request_root=request_root(gate),
            request_artifacts=[item for item in gate.get("artifacts", []) if isinstance(item, dict)],
            configuration_root=configuration_root, extension_root=extension_root,
        )
    except CcInspectionError as exc:
        raise WorkflowError(str(exc)) from exc
    evidence_path, evidence = evidence_for_gate(gate)
    rlm_id = str(getattr(args, "rlm_evidence_id", "") or "").strip()
    if source != "request":
        selected = next((item for item in evidence.get("rlm_queries", []) if item.get("id") == rlm_id), None)
        if not selected or selected.get("source", {}).get("kind") != source or selected.get("status") not in {"retrieved_unvalidated", "confirmed"}:
            raise WorkflowError("First discover this XML path with flow1c_source_query for the same source")
        selected_root = Path(str(selected.get("source", {}).get("path", ""))).resolve()
        expected_root = configuration_root if source == "configuration" else extension_root
        if expected_root is None or selected_root != expected_root.resolve():
            raise WorkflowError("The RLM result belongs to a different source checkout")
        stdout = str(selected.get("result", {}).get("stdout", "")).replace("\\", "/").casefold()
        if str(args.path or "").replace("\\", "/").casefold() not in stdout:
            raise WorkflowError("The selected RLM evidence does not identify this exact XML path")
    raw = target.read_bytes()
    try:
        schema = parse_metadata_xml(raw)
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    digest = hashlib.sha256(raw).hexdigest()
    for saved in evidence.get("query_schemas", []):
        if saved.get("path") == str(target) and saved.get("sha256") == digest and saved.get("source") == source:
            print(json.dumps({"state": "CONFIRMED_IN_XML", "schema": saved}, ensure_ascii=False, indent=2))
            return 0
    schema = {**schema, "id": f"QM-{len(evidence.get('query_schemas', [])) + 1:03d}",
              "source": source, "path": str(target), "relative_path": str(args.path or ""), "sha256": digest,
              "rlm_evidence_id": rlm_id or None, "created_at": utc_now()}
    evidence.setdefault("query_schemas", []).append(schema)
    write_json(evidence_path, evidence)
    print(json.dumps({"state": "CONFIRMED_IN_XML", "schema": schema}, ensure_ascii=False, indent=2))
    return 0


def cmd_query_check(args: argparse.Namespace) -> int:
    """Save and inspect the exact candidate that the agent will present."""
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    require_gate_tool(gate, "flow1c_query_check")
    if gate.get("operation") != "query-analysis":
        raise WorkflowError("query-check is available only for query-analysis")
    candidate = str(args.text or "")
    baseline = str(getattr(args, "baseline_text", "") or "")
    if gate.get("query_intent") == "optimize" and gate.get("query_contract_version") == 1 and not baseline.strip():
        raise WorkflowError("Optimization requires baseline_text so the original query remains reviewable")
    schema_ids = getattr(args, "schema_ids", []) or []
    if not isinstance(schema_ids, list) or any(not isinstance(item, str) for item in schema_ids):
        raise WorkflowError("schema_ids must be an array of evidence IDs")
    assumptions = getattr(args, "assumptions", []) or []
    if not isinstance(assumptions, list) or any(not isinstance(item, str) for item in assumptions):
        raise WorkflowError("assumptions must be an array of strings")
    changes = getattr(args, "changes", []) or []
    if not isinstance(changes, list) or any(not isinstance(item, str) for item in changes):
        raise WorkflowError("changes must be an array of strings")
    expected_result = str(getattr(args, "expected_result", "") or "").strip()
    if gate.get("query_contract_version") == 1 and gate.get("query_intent") in {"create", "optimize"} and not expected_result:
        raise WorkflowError("Describe the expected row grain and result in expected_result")
    if gate.get("query_contract_version") == 1 and gate.get("query_intent") == "optimize" and not changes:
        raise WorkflowError("Explain proposed structural changes in changes")
    evidence_path, evidence = evidence_for_gate(gate)
    available = {str(item.get("id")): item for item in evidence.get("query_schemas", [])}
    if len(set(schema_ids)) != len(schema_ids) or any(item not in available for item in schema_ids):
        raise WorkflowError("schema_ids must uniquely identify schemas recorded in this gate")
    schemas = [available[item] for item in schema_ids]
    schema_sources = []
    for schema in schemas:
        source_path = Path(str(schema.get("path", "")))
        if not source_path.is_file() or sha256(source_path) != schema.get("sha256"):
            raise WorkflowError(f"Metadata evidence {schema.get('id')} is stale; refresh it before query-check")
        schema_sources.append({"id": schema["id"], "path": str(source_path), "sha256": schema["sha256"]})
    try:
        report = analyze_query(candidate, schemas)
        baseline_report = analyze_query(baseline, schemas) if baseline else None
    except ValueError as exc:
        raise WorkflowError(str(exc)) from exc
    if baseline_report:
        report["diagnostics"].extend(compare_optimization(baseline_report, report))
    if not expected_result:
        report["limitations"].append("Ожидаемый состав строк и колонок не описан; соответствие бизнес-требованию не проверено.")
    report.update({"id": f"QC-{len(evidence.get('query_checks', [])) + 1:03d}",
                   "intent": gate.get("query_intent", "review"), "schema_ids": schema_ids,
                   "schema_sources": schema_sources,
                   "baseline_sha256": hashlib.sha256(baseline.encode("utf-8")).hexdigest() if baseline else None,
                   "expected_result": expected_result, "assumptions": assumptions, "changes": changes,
                   "created_at": utc_now()})
    candidate_path = request_root(gate) / "query-candidate.txt"
    previous = (evidence.get("query_checks") or [None])[-1]
    baseline_path = request_root(gate) / "query-baseline.txt"
    if (isinstance(previous, dict) and candidate_path.is_file()
            and sha256(candidate_path) == report["text_sha256"]
            and (not baseline or (baseline_path.is_file() and sha256(baseline_path) == report["baseline_sha256"]))
            and all(previous.get(key) == report.get(key) for key in (
                "text_sha256", "intent", "schema_ids", "schema_sources", "baseline_sha256",
                "expected_result", "assumptions", "changes"))):
        print(json.dumps({"state": "CHECKED", "candidate_path": str(candidate_path),
                          "query_text": candidate, "report": previous}, ensure_ascii=False, indent=2))
        return 0
    atomic_write_bytes(candidate_path, candidate.encode("utf-8"))
    if baseline:
        atomic_write_bytes(baseline_path, baseline.encode("utf-8"))
    evidence.setdefault("query_checks", []).append(report)
    write_json(evidence_path, evidence)
    print(json.dumps({"state": "CHECKED", "candidate_path": str(candidate_path),
                      "query_text": candidate, "report": report}, ensure_ascii=False, indent=2))
    return 0


def cmd_agent_write(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"})
    stage = require_gate_tool(gate, "flow1c_write")
    code = str(gate.get("code") or "")
    content = sys.stdin.read() if args.content_stdin else str(args.content or "")
    if gate.get("mode", "formal") != "formal":
        return write_free_draft(gate, args, content)
    if args.target == "extension":
        if "extension" not in stage.get("writable_targets", []):
            raise WorkflowError("This operation may not modify the extension checkout.")
        _, local = load_config()
        root = Path(str(local.get("extension_path", ""))).resolve()
        _, manifest = load_manifest(code)
        errors = stage_errors(stage, manifest)
        item_root = work_item_root(code)
        errors.extend(f"required input is missing: {relative}" for relative in stage.get("required_files", []) if not meaningful_file(item_root / relative))
        if errors:
            raise WorkflowError("Extension write prerequisites changed: " + "; ".join(errors))
        diff_state, error = extension_git_state(code, manifest)
        if error or diff_state is None or diff_state["branch"] != diff_state["expected_branch"]:
            raise WorkflowError(error or "Extension branch does not match the work-item manifest.")
    else:
        if not code or not any(value.startswith("work-item") for value in stage.get("writable_targets", [])):
            raise WorkflowError("This operation may not modify work-item documentation.")
        root = work_item_root(code)
    target = (root / args.path).resolve()
    if not path_is_within(target, root):
        raise WorkflowError("Write target is outside the permitted root.")
    if args.target == "work-item":
        relative_target = str(target.relative_to(root)).replace("\\", "/")
        allowed_prefixes = {
            "functional-spec": ("analysis/", "specification/functional-spec.md"),
            "functional-review": ("reviews/functional-review.md",),
            "technical-design": ("specification/technical-design.md", "reviews/technical-review.md"),
            "technical-implementation": ("specification/technical-implementation.md",),
            "code-review": ("implementation/code-review.md",),
            "testing": ("testing/",),
        }.get(str(gate.get("operation")), ())
        if not any(relative_target == value.rstrip("/") or relative_target.startswith(value) for value in allowed_prefixes):
            raise WorkflowError(f"Documentation path {relative_target} is not allowed for {gate.get('operation')}.")
    if gate["state"] in {"UNVERIFIED_DRAFT", "READY_WITH_DEVIATIONS"} and args.target != "work-item":
        raise WorkflowError("An unverified draft may not modify extension sources.")
    if gate["state"] in {"UNVERIFIED_DRAFT", "READY_WITH_DEVIATIONS"} and "UNVERIFIED_DRAFT" not in content:
        content = "> **UNVERIFIED_DRAFT:** обязательные источники отсутствуют; документ нельзя согласовывать или публиковать.\n\n" + content
    if args.target == "work-item" and str(gate.get("operation")) == "technical-implementation":
        validate_section_content("technical-implementation", content)
    write_text(target, content)
    evidence_path, evidence = evidence_for_gate(gate)
    relative = str(target.relative_to(root)).replace("\\", "/")
    record = {"target": args.target, "path": relative, "sha256": sha256(target)}
    evidence.setdefault("changed_files", []).append(record)
    write_json(evidence_path, evidence)
    print(json.dumps({"state": "WRITTEN", **record, "absolute_path": str(target)}, ensure_ascii=False, indent=2))
    return 0


def cmd_agent_analyze_bsl(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT",
                                           "BLOCKED", "NEEDS_INPUT", "NEEDS_CONFIRMATION", "NEEDS_CODE"})
    require_gate_tool(gate, "flow1c_analyze_bsl")
    if gate.get("operation") not in {"development", "technical-implementation", "code-review"}:
        raise WorkflowError("BSL analysis is not enabled for this operation.")
    selector = getattr(args, "source", None) or {"kind": "extension"}
    if not isinstance(selector, dict):
        raise WorkflowError("BSL source must be a typed source selector.")
    kind = str(selector.get("kind", ""))
    if kind == "git_snapshot":
        snapshot_id = str(selector.get("snapshot_id", ""))
        snapshot = _git_analysis_state(gate).get("snapshots", {}).get(snapshot_id)
        if not isinstance(snapshot, dict):
            raise WorkflowError("The requested Git snapshot does not belong to this gate")
        commit = str(snapshot.get("commit", ""))
        if selector.get("commit") and selector["commit"] != commit:
            raise WorkflowError("Snapshot commit does not match the saved manifest")
        source_path = Path(str(snapshot.get("path", ""))).resolve()
        if not path_is_within(source_path, (ROOT / ".workspace" / "git-analysis").resolve()) or not source_path.is_dir():
            raise WorkflowError("Snapshot source is unavailable or outside the managed cache")
        provenance = {"kind": kind, "snapshot_id": snapshot_id, "commit": commit, "path": str(source_path)}
    elif kind in {"extension", "configuration"}:
        local = read_json(ROOT / LOCAL_CONFIG_FILE, {})
        raw = str(local.get(f"{kind}_path", "")).strip()
        if not raw:
            raise WorkflowError(f"{kind} source path is not configured.")
        source_path = resolve_1c_source_root(Path(raw))
        if not source_path.is_dir():
            raise WorkflowError(f"{kind} source directory is unavailable: {source_path}")
        provenance = {"kind": kind, "path": str(source_path)}
    else:
        raise WorkflowError("BSL source must be extension, configuration, or a git_snapshot selector")
    powershell = command_path("powershell") or command_path("pwsh")
    if not powershell:
        raise WorkflowError("PowerShell is not available.")
    evidence_id = f"BSL-{uuid.uuid4().hex[:12]}"
    report_path = (ROOT / ".workspace" / "diagnostics" / "bsl-ls" / evidence_id).resolve()
    result = subprocess.run(
        [powershell, "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "analyze-bsl.ps1"),
         "-SourcePath", str(source_path), "-OutputPath", str(report_path)],
        cwd=ROOT,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
        timeout=600,
    )
    output = "\n".join(value.strip() for value in (result.stdout, result.stderr) if value.strip())
    if gate.get("evidence_path"):
        evidence_path, evidence = evidence_for_gate(gate)
    else:
        evidence_path = (ROOT / ".workspace" / "diagnostics" / "bsl-ls" / str(gate["gate_id"]) / "evidence.json").resolve()
        evidence = {"schema_version": 2, "gate_id": gate["gate_id"], "operation": gate["operation"],
                    "code": None, "artifacts": [], "rlm_queries": [], "source_reads": [], "changed_files": []}
    record = {
        "evidence_id": evidence_id,
        "state": "passed" if result.returncode == 0 else "failed",
        "exit_code": result.returncode,
        "output": output[-8000:],
        "source": provenance,
        "report_path": str(report_path),
        "created_at": utc_now(),
    }
    evidence["bsl_ls"] = record
    write_json(evidence_path, evidence)
    print(json.dumps({"state": "PASSED" if result.returncode == 0 else "FAILED", **record}, ensure_ascii=False, indent=2))
    return 0 if result.returncode == 0 else 2


def output_has_unverified_claims(text: str, evidence: dict[str, Any]) -> bool:
    claim = re.search(r"\b(проверен\w*|подтвержден\w*|verified|confirmed)\b", text, flags=re.IGNORECASE)
    if not claim:
        return False
    evidence_ids = [str(evidence.get("context_evidence_id", ""))]
    evidence_ids.extend(str(item.get("id", "")) for item in evidence.get("rlm_queries", []) if isinstance(item, dict))
    if isinstance(evidence.get("diff"), dict):
        evidence_ids.append(str(evidence["diff"].get("evidence_id", "")))
    if isinstance(evidence.get("bsl_ls"), dict):
        evidence_ids.append(str(evidence["bsl_ls"].get("evidence_id", "")))
    evidence_ids.extend(str(item.get("evidence_id", "")) for item in evidence.get("source_reads", []) if isinstance(item, dict))
    return not any(evidence_id and evidence_id in text for evidence_id in evidence_ids)


def recoverable_update_action_gate(gate: dict[str, Any]) -> bool:
    """Accept one retry for an update closed by an unrecorded successful action."""
    return (
        gate.get("state") == "NON_COMPLIANT"
        and gate.get("mode", "formal") == "formal"
        and gate.get("operation") == "update"
        and not gate.get("action_completed")
        and gate.get("completion_errors") == ["the requested workflow action has not completed"]
    )


def parse_updater_result(stdout: str) -> dict[str, Any] | None:
    """Read the updater JSON, including output from the older RLM launcher."""
    lines = stdout.lstrip("\ufeff").strip().splitlines()
    while lines and re.match(r"^(?:Compatible shared )?RLM MCP (?:is already running|started): ", lines[0]):
        lines.pop(0)
    try:
        result = json.loads("\n".join(lines))
    except json.JSONDecodeError:
        return None
    return result if isinstance(result, dict) else None


def run_update_command(command: list[str]) -> subprocess.CompletedProcess[str]:
    """Capture updater output without waiting for background services to close pipes."""
    runtime = ROOT / ".workspace"
    runtime.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryFile(mode="w+t", encoding="utf-8", errors="replace", dir=runtime) as stdout_file, \
            tempfile.TemporaryFile(mode="w+t", encoding="utf-8", errors="replace", dir=runtime) as stderr_file:
        try:
            result = subprocess.run(
                command, cwd=ROOT, stdout=stdout_file, stderr=stderr_file,
                check=False, timeout=600,
            )
        except subprocess.TimeoutExpired as exc:
            raise WorkflowError(
                "scripts/update.ps1 exceeded 600 seconds. Inspect .workspace/backups "
                "and the RLM index jobs before retrying this update gate."
            ) from exc
        stdout_file.seek(0)
        stderr_file.seek(0)
        return subprocess.CompletedProcess(
            command, result.returncode, stdout_file.read(), stderr_file.read()
        )


def cmd_agent_complete(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "UNVERIFIED_DRAFT", "NON_COMPLIANT"})
    legacy_update_retry = (
        gate["state"] == "NON_COMPLIANT"
        and gate.get("mode", "formal") == "formal"
        and gate.get("operation") == "update"
        and gate.get("action_completed") == "update"
        and len(gate.get("completion_errors", [])) == 1
        and str(gate["completion_errors"][0]).startswith("output override is not allowed;")
    )
    if gate["state"] == "NON_COMPLIANT" and not legacy_update_retry:
        raise WorkflowError("Agent gate state NON_COMPLIANT is not allowed for this action.")
    if not legacy_update_retry:
        require_gate_tool(gate, "flow1c_complete")
    if gate.get("mode", "formal") != "formal":
        return complete_free_request(gate, args)
    stages = load_stages()
    stage = stages["operations"][gate["operation"]]
    configured_output = str(stage.get("output") or "").strip()
    requested_output = str(args.output or "").strip()
    if requested_output and requested_output != configured_output:
        raise WorkflowError(
            f"flow1c_complete output override is not allowed; expected {configured_output or 'no output file'}. "
            "Omit output and retry flow1c_complete on the same gate."
        )
    if gate["operation"] == "update" and not gate.get("action_completed"):
        raise WorkflowError("The update result was not recorded in this gate. Retry flow1c_action action=update on the same gate, then retry flow1c_complete without output.")
    errors: list[str] = []
    if gate["operation"] in {"setup", "update", "registry", "status", "publish"} and not gate.get("action_completed"):
        errors.append("the requested workflow action has not completed")
    evidence: dict[str, Any] = {}
    evidence_path: Path | None = None
    traceability_mode = str(gate.get("traceability_mode") or "registry")
    if gate.get("code"):
        evidence_path, evidence = evidence_for_gate(gate)
        _, completion_manifest = load_manifest(str(gate.get("work_reference") or gate.get("code")))
        traceability_mode = str(completion_manifest.get("traceability_mode") or traceability_mode)
        errors.extend(f"evidence schema: {value}" for value in validate_json_record(evidence, ROOT / "schemas" / "evidence.schema.json"))
        if (gate.get("manifest_status") is None
                or completion_manifest.get("status") != gate.get("manifest_status")
                or completion_manifest.get("approvals", {}) != gate.get("manifest_approvals", {})):
            errors.append("current work-item status or approvals changed after gate opening")
    output_raw = configured_output
    output_path: Path | None = None
    output_text = ""
    if output_raw:
        root = work_item_root(str(gate.get("work_reference") or gate["code"])) if gate.get("work_reference") or gate.get("code") else project_root()
        output_path = (root / output_raw).resolve() if not Path(output_raw).is_absolute() else Path(output_raw).resolve()
        if not path_is_within(output_path, root) or not meaningful_file(output_path):
            errors.append(f"required output is missing or empty: {output_path}")
        elif output_path.suffix.casefold() in {".md", ".txt"}:
            output_text = output_path.read_text(encoding="utf-8-sig", errors="replace")
            if gate.get("operation") == "technical-implementation":
                try:
                    validate_section_content("technical-implementation", output_text)
                except SectionPolicyError as exc:
                    errors.append(exc.message)
    if stage.get("context_role") and not evidence.get("context_sha256"):
        errors.append("role context was not built")
    if gate.get("template_result"):
        try:
            template_service = TemplateService(ROOT, read_json(ROOT / LOCAL_CONFIG_FILE, {}))
            template_root = templates_cli.document_root(SimpleNamespace(**globals()), gate, template_service)
            template_service.document_validate({"plan_id": gate["template_result"]["plan_id"]}, template_root)
            templates_cli.check_formal_content(SimpleNamespace(**globals()), gate, template_service,
                                               {"plan_id": gate["template_result"]["plan_id"]}, template_root)
        except (TemplateError, OSError, ValueError) as exc:
            errors.append("template output integrity/policy: " + str(exc))
    if evidence.get("context_path"):
        context_path = Path(evidence["context_path"])
        if not context_path.is_file() or sha256(context_path) != evidence.get("context_sha256"):
            errors.append("recorded role context has changed")
    if evidence.get("diff"):
        _, manifest = load_manifest(str(gate["code"]))
        current_diff, error = extension_git_state(str(gate["code"]), manifest)
        if error or not current_diff or any(current_diff.get(k) != evidence["diff"].get(k) for k in ("head", "base", "branch")):
            errors.append("extension Git state changed after the recorded diff")
    if stage.get("requires_rlm") and not evidence.get("rlm_queries"):
        errors.append("no gated RLM evidence was recorded")
    if stage.get("requires_diff") and not evidence.get("diff"):
        errors.append("bounded extension diff was not recorded")
    if stage.get("requires_bsl_ls") and (evidence.get("bsl_ls") or {}).get("state") != "passed":
        errors.append("BSL Language Server did not pass")
    if output_text and output_has_unverified_claims(output_text, evidence):
        errors.append("output contains verified/confirmed claims without evidence")
    if gate["state"] == "READY_WITH_DEVIATIONS" or traceability_mode == "provisional":
        if output_text and "UNVERIFIED_DRAFT" not in output_text:
            errors.append("deviated output watermark is missing")
        integrity_markers = (
            "required output is missing or empty", "deviated output watermark is missing",
            "recorded role context has changed",
            "extension Git state changed", "output contains verified/confirmed claims",
            "evidence schema:", "current work-item status or approvals changed",
        )
        integrity_errors = [error for error in errors if any(marker in error for marker in integrity_markers)]
        waived_condition_ids = {
            str(condition_id)
            for deviation in gate.get("deviations", []) if isinstance(deviation, dict)
            for condition_id in deviation.get("condition_ids", deviation.get("waived_conditions", []))
        }
        unwaived_errors = []
        for error in errors:
            error_conditions = build_conditions(stage, errors=[error])
            if not error_conditions or any(item["id"] not in waived_condition_ids for item in error_conditions):
                unwaived_errors.append(error)
        state = "NON_COMPLIANT" if integrity_errors or unwaived_errors else "COMPLETE_WITH_DEVIATIONS"
    elif gate["state"] == "UNVERIFIED_DRAFT":
        if output_text and "UNVERIFIED_DRAFT" not in output_text:
            errors.append("unverified draft watermark is missing")
        state = "UNVERIFIED_DRAFT" if not errors else "NON_COMPLIANT"
    else:
        state = "COMPLETE" if not errors else "NON_COMPLIANT"
    snapshot = None
    if state == "COMPLETE" and gate.get("code"):
        snapshot = publication_snapshot(str(gate["code"]), include_extension=bool(stage.get("requires_diff") or stage.get("requires_extension_branch")))
    gate["state"] = state
    gate["completed_at"] = utc_now()
    gate["completion_errors"] = errors
    save_gate(gate)
    if evidence_path is not None:
        evidence["completed_at"] = gate["completed_at"]
        evidence["completion_state"] = state
        evidence["completion_errors"] = errors
        if state == "COMPLETE":
            evidence["validation_snapshot"] = snapshot
            evidence["validated_output"] = str(output_path) if output_path else None
        write_json(evidence_path, evidence)
    result = {"state": state, "ready": state == "COMPLETE", "compliance": "DEVIATED" if state == "COMPLETE_WITH_DEVIATIONS" else gate.get("compliance", "COMPLIANT"), "errors": errors, "deviations": gate.get("deviations", []), "output": str(output_path) if output_path else None}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if state in {"COMPLETE", "COMPLETE_WITH_DEVIATIONS", "UNVERIFIED_DRAFT"} else 2


def run_captured(handler: Any, namespace: argparse.Namespace) -> tuple[int, str]:
    output = io.StringIO()
    from contextlib import redirect_stdout

    with redirect_stdout(output):
        exit_code = int(handler(namespace))
    return exit_code, output.getvalue().strip()


def reconcile_registry(code: str, source: str, mappings: dict[str, str] | None = None) -> tuple[int, dict[str, Any]]:
    """Validate a replacement registry and reconcile only explicit or exact identifiers."""
    manifest_path, manifest = load_manifest(code)
    if manifest.get("traceability_mode") != "provisional":
        raise WorkflowError("registry-reconcile requires a provisional work item")
    exit_code, output = run_captured(
        cmd_registry_import, argparse.Namespace(file=source, allow_errors=False),
    )
    try:
        import_result = json.loads(output)
    except json.JSONDecodeError:
        import_result = {"state": "BLOCKED", "errors": [{"kind": "registry_import", "message": output}]}
    if exit_code:
        return exit_code, import_result

    known = read_json(project_root() / "registry" / "normalized" / "requirements.json", {})
    proposed: dict[str, str] = {}
    conflicts: list[dict[str, Any]] = []
    supplied = mappings or {}
    candidate_ids = [str(value) for value in manifest.get("requirements", [])]
    candidate_ids.extend(str(value) for value in supplied)
    brief_path = work_item_root(code) / "input" / "user-brief.md"
    brief_text = brief_path.read_text(encoding="utf-8-sig", errors="replace") if brief_path.is_file() else ""
    candidate_ids.extend(re.findall(r"[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё0-9_.]*-\d+", brief_text))
    for candidate in dict.fromkeys(candidate_ids):
        target = supplied.get(candidate, candidate)
        if target in known:
            proposed[candidate] = target
        else:
            conflicts.append({"source": candidate, "proposed": supplied.get(candidate), "reason": "no confirmed registry match"})
    if not proposed:
        conflicts.append({"source": "user_brief", "proposed": None, "reason": "explicit requirement mapping is required"})

    item_root = work_item_root(code)
    report_path = item_root / "analysis" / "registry-reconcile-report.md"
    report_lines = ["# Registry reconciliation report", "", f"Work item: `{code}`", "", "## Proposed mappings", ""]
    report_lines.extend(f"- `{source_id}` → `{target_id}`" for source_id, target_id in proposed.items())
    if not proposed:
        report_lines.append("No safe automatic matches were found.")
    report_lines.extend(["", "## Conflicts", ""])
    report_lines.extend(f"- `{item['source']}`: {item['reason']}" for item in conflicts)
    if not conflicts:
        report_lines.append("None.")
    write_text(report_path, "\n".join(report_lines))
    if conflicts:
        return 1, {"state": "NEEDS_CONFIRMATION", "report_path": str(report_path),
                   "proposed_mappings": proposed, "conflicts": conflicts}

    selected_ids = list(dict.fromkeys(proposed.values()))
    snapshot = {requirement_id: known[requirement_id] for requirement_id in selected_ids}
    snapshot_path = item_root / "input" / "requirements.snapshot.yaml"
    write_json(snapshot_path, snapshot)
    mapping_path = item_root / "input" / "registry-mapping.json"
    write_json(mapping_path, {"confirmed": True, "mapped_at": utc_now(), "mappings": proposed})
    manifest.update(
        traceability_mode="registry", requirements_source="registry", requirements=selected_ids,
        registry={"status": "verified", "snapshot_sha256": sha256(snapshot_path)},
        requirement_basis={"type": "registry_snapshot", "path": "input/requirements.snapshot.yaml", "sha256": sha256(snapshot_path)},
        updated_at=utc_now(),
    )
    write_json(manifest_path, manifest)
    for evidence_path in (item_root / "evidence").glob("*.json"):
        evidence = read_json(evidence_path, {})
        if isinstance(evidence, dict):
            evidence["revalidation_required"] = True
            evidence["reconciled_at"] = utc_now()
            write_json(evidence_path, evidence)
    return 0, {"state": "RECONCILED", "report_path": str(report_path), "mapping_path": str(mapping_path),
               "requirements": selected_ids, "revalidation_required": True}


def cmd_registry_reconcile(args: argparse.Namespace) -> int:
    try:
        mappings = json.loads(str(getattr(args, "mappings_json", "") or "{}"))
    except json.JSONDecodeError as exc:
        raise WorkflowError(f"Invalid mappings JSON: {exc}") from exc
    if not isinstance(mappings, dict):
        raise WorkflowError("mappings must be a JSON object")
    exit_code, result = reconcile_registry(str(args.code), str(args.file), {str(k): str(v) for k, v in mappings.items()})
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return exit_code


def cmd_agent_action(args: argparse.Namespace) -> int:
    gate = load_gate(args.gate_id, states={"READY", "READY_WITH_DEVIATIONS", "NEEDS_INPUT", "NEEDS_CONFIRMATION", "WAITING_BACKGROUND", "BLOCKED", "NON_COMPLIANT"})
    if recoverable_update_action_gate(gate):
        if args.action != "update":
            raise WorkflowError("Only an update retry is allowed for this recoverable gate.")
    else:
        if gate["state"] == "NON_COMPLIANT":
            raise WorkflowError("Agent gate state NON_COMPLIANT is not allowed for this action.")
        require_gate_tool(gate, "flow1c_action")
    operation = str(gate.get("operation"))
    action = str(args.action)
    try:
        parameters = json.loads(args.parameters_json or "{}")
    except json.JSONDecodeError as exc:
        raise WorkflowError(f"Invalid action parameters JSON: {exc}") from exc
    if not isinstance(parameters, dict):
        raise WorkflowError("Action parameters must be a JSON object.")
    if action == "update" and gate.get("update_id") and not parameters.get("new_run"):
        for flag, alias in (("allow_external_updates", "allowExternalUpdates"), ("skip_external_tool_updates", "skipExternalToolUpdates"), ("repair_prerequisites", "repairPrerequisites")):
            if flag not in parameters and alias in parameters:
                parameters[flag] = parameters[alias]
            if flag not in parameters and (gate.get("update_result") or {}).get("options", {}).get(flag):
                parameters[flag] = True

    allowed = {
        "setup": {"setup-audit", "setup-plan", "setup-bootstrap", "setup-configure", "setup-status", "setup-resume", "doctor"},
        "update": {"update", "doctor"},
        "registry": {"registry-import"},
        "functional-spec": {"fs-start", "provisional-start", "registry-reconcile"},
        "functional-review": {"provisional-start", "registry-reconcile"},
        "technical-design": {"provisional-start", "registry-reconcile"},
        "technical-implementation": {"provisional-start", "registry-reconcile"},
        "code-review": {"provisional-start", "registry-reconcile"},
        "testing": {"provisional-start", "registry-reconcile"},
        "status": {"status"},
        "publish": {"publish"},
    }
    if action not in allowed.get(operation, set()):
        raise WorkflowError(f"Action {action} is not allowed for operation {operation}.")
    if action == "provisional-start":
        bypasses = [item for item in gate.get("deviations", []) if item.get("type") == "registry_bypass"]
        if not bypasses:
            raise WorkflowError("provisional-start requires an explicit registry_bypass deviation")
        try:
            reference = resolve_work_reference(
                parameters.get("task_reference") or gate.get("task_reference") or gate.get("work_reference"),
                parameters.get("project_reference") or gate.get("project_reference"),
            )
        except ValueError as exc:
            raise WorkflowError(str(exc)) from exc
        brief = str(parameters.get("user_brief") or gate.get("summary") or "")
        requirements = parameters.get("requirements") or []
        if not isinstance(requirements, list):
            raise WorkflowError("requirements must be a list of user-supplied identifiers")
        target = create_provisional_work_item(
            reference, title=str(parameters.get("title") or gate.get("summary") or reference.get("work_reference") or ""),
            user_brief=brief, deviation=bypasses[-1], requirements=[str(item) for item in requirements],
        )
        code = str(reference["work_reference"])
        manifest = read_json(target / "manifest.yaml", {})
        basis = manifest["requirement_basis"]
        gate.update(**reference, code=code, work_item_exists=True, traceability_mode="provisional",
                    state="SUPERSEDED", compliance="DEVIATED")
        evidence_path = target / "evidence" / f"{operation}-{gate['gate_id']}.json"
        evidence = {
            "schema_version": 1, "gate_id": gate["gate_id"], "operation": operation, "code": code,
            **reference, "traceability_mode": "provisional",
            "requirement_basis": {"type": "user_brief", "sha256": basis["sha256"]},
            "registry_snapshot_sha256": None, "deviations": gate["deviations"], "revalidation_required": True,
            "context_sha256": "", "artifacts": [], "diff": None, "rlm_queries": [], "source_reads": [],
            "bsl_ls": None, "changed_files": [], "created_at": utc_now(),
        }
        write_json(evidence_path, evidence)
        gate["evidence_path"] = str(evidence_path)
        refresh_gate_actions(gate)
        save_gate(gate)
        print(json.dumps({"state": "ACTION_COMPLETE", "next": "flow1c_begin", "work_item": str(target),
                          "traceability_mode": "provisional"}, ensure_ascii=False, indent=2))
        return 0
    if action == "registry-reconcile":
        source = str(parameters.get("source") or "").strip()
        code = str(parameters.get("task_reference") or parameters.get("code") or gate.get("work_reference") or gate.get("code") or "")
        if not source or not code:
            raise WorkflowError("registry-reconcile requires source and a user-supplied work reference")
        mappings = parameters.get("mappings") or {}
        if not isinstance(mappings, dict):
            raise WorkflowError("mappings must be an object")
        exit_code, result = reconcile_registry(code, source, {str(k): str(v) for k, v in mappings.items()})
        if exit_code == 0:
            gate.update(action_completed=action, state="NEEDS_INPUT", traceability_mode="registry")
            refresh_gate_actions(gate)
            save_gate(gate)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return exit_code
    if action in {"setup-status", "setup-resume"}:
        setup_id = str(parameters.get("setup_id") or gate.get("setup_id") or "")
        checkpoint = read_json(setup_checkpoint_path(setup_id))
        if not isinstance(checkpoint, dict):
            raise WorkflowError(f"Setup checkpoint not found: {setup_id}")
        if checkpoint.get("state") == "COMPLETE":
            print(json.dumps(checkpoint, ensure_ascii=False, indent=2))
            return 0
        if action == "setup-status":
            jobs = checkpoint.get("pending_jobs", [])
            powershell_status = command_path("powershell") or command_path("pwsh")
            if jobs and not powershell_status:
                raise WorkflowError("PowerShell is not available to inspect background setup jobs.")
            refreshed: list[dict[str, Any]] = []
            for job in jobs:
                source = str(job.get("source") or job.get("source_path") or "") if isinstance(job, dict) else ""
                result = subprocess.run(
                    [powershell_status, "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "rlm-index.ps1"),
                     "-Action", "Status", "-SourcePath", source, "-Json"],
                    cwd=ROOT, text=True, encoding="utf-8", errors="replace", capture_output=True, check=False,
                )
                try:
                    status = json.loads(result.stdout.lstrip("\ufeff"))
                except json.JSONDecodeError:
                    status = {"kind": "rlm-index", "source": source, "state": "FAILED",
                              "detail": (result.stderr or result.stdout or "invalid status JSON").strip()}
                status["kind"] = "rlm-index"
                status["source"] = status.get("source_path") or source
                refreshed.append(status)
            states = {str(job.get("state")) for job in refreshed}
            if not refreshed or states <= {"FRESH"}:
                checkpoint["state"] = "IN_PROGRESS"
                checkpoint["pending_jobs"] = []
                checkpoint["next_action"] = "setup-resume"
            elif states & {"FAILED"}:
                checkpoint["state"] = "BLOCKED"
                checkpoint["pending_jobs"] = refreshed
                checkpoint["next_action"] = "repair-index-then-setup-resume"
            else:
                checkpoint["state"] = "WAITING_BACKGROUND"
                checkpoint["pending_jobs"] = refreshed
                checkpoint["next_action"] = "setup-status"
            save_setup_checkpoint(checkpoint)
            print(json.dumps(checkpoint, ensure_ascii=False, indent=2))
            return 2 if checkpoint["state"] == "BLOCKED" else 0
        parameters = {**checkpoint.get("confirmed_values", {}), **parameters, "setup_id": setup_id, "confirmed": True}
        action = "setup-configure"
    powershell = command_path("powershell") or command_path("pwsh")
    if action in {"setup-bootstrap", "setup-configure", "update", "publish"} and not parameters.get("confirmed"):
        result = {"state": "NEEDS_CONFIRMATION", "user_message": "Для этого изменения требуется явное подтверждение пользователя."}
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 1
    if action == "status":
        gate["action_completed"] = "status"
        save_gate(gate)
        print(json.dumps({"state": "ACTION_COMPLETE", "next": "flow1c_complete", "content": render_status()}, ensure_ascii=False, indent=2))
        return 0
    if action == "doctor":
        exit_code, output = run_captured(cmd_doctor, argparse.Namespace(
            json=True, profile=parameters.get("profile"), operation=parameters.get("operation")))
        print(output)
        return exit_code
    if action == "registry-import":
        source = str(parameters.get("source", "")).strip()
        if not source:
            raise WorkflowError("registry-import requires source")
        exit_code, output = run_captured(
            cmd_registry_import,
            argparse.Namespace(file=source, allow_errors=bool(parameters.get("allow_errors", False))),
        )
        try:
            registry_result = json.loads(output)
        except json.JSONDecodeError:
            registry_result = {"state": "BLOCKED", "errors": [{"kind": "registry_import", "message": output}], "warnings": []}
        if exit_code in {0, 1} and registry_result.get("state") in {"READY", "PARTIAL"}:
            gate["state"] = "READY"
            gate["action_completed"] = action
            gate["registry_import"] = registry_result
            refresh_gate_actions(gate)
            save_gate(gate)
        else:
            gate["state"] = "NEEDS_CONFIRMATION"
            gate["registry_import"] = registry_result
            refresh_gate_actions(gate)
            save_gate(gate)
        accepted = exit_code in {0, 1} and registry_result.get("state") in {"READY", "PARTIAL"}
        print(json.dumps({**registry_result, "state": "ACTION_COMPLETE" if accepted else "NEEDS_CONFIRMATION",
                          "next": "flow1c_complete" if accepted else "registry-bypass-or-fix"}, ensure_ascii=False, indent=2))
        return exit_code
    if action == "fs-start":
        code = str(parameters.get("task_reference") or parameters.get("code") or gate.get("work_reference") or gate.get("code") or "")
        exit_code, output = run_captured(
            cmd_fs_start,
            argparse.Namespace(
                code=code,
                task_reference=parameters.get("task_reference"),
                project_reference=parameters.get("project_reference"),
                g_number=None,
                reference_kind=parameters.get("reference_kind", gate.get("reference_kind", "auto")),
                title=parameters.get("title"),
                requirements=parameters.get("requirements"),
                create_branch=bool(parameters.get("create_branch", False)),
            ),
        )
        print(json.dumps({"state": "ACTION_COMPLETE" if exit_code == 0 else "BLOCKED", "next": "flow1c_begin" if exit_code == 0 else None, "output": output}, ensure_ascii=False, indent=2))
        return exit_code
    if action == "publish":
        if gate["state"] != "READY":
            raise WorkflowError("Publication requires a READY gate; confirmation does not override missing checks")
        work_reference = str(gate.get("work_reference") or gate.get("code") or parameters.get("task_reference") or parameters.get("code") or "")
        validate_publication(work_reference, str(parameters.get("phase", "technical")))
        exit_code, output = run_captured(
            cmd_pr_create,
            argparse.Namespace(
                code=work_reference,
                phase=str(parameters.get("phase", "technical")),
                title=parameters.get("title"),
            ),
        )
        if exit_code == 0:
            gate["state"] = "READY"
            gate["action_completed"] = action
            save_gate(gate)
        print(json.dumps({"state": "ACTION_COMPLETE" if exit_code == 0 else "BLOCKED", "next": "flow1c_complete" if exit_code == 0 else None, "output": output}, ensure_ascii=False, indent=2))
        return exit_code
    if not powershell:
        raise WorkflowError("PowerShell is not available.")
    if action == "setup-audit":
        command = [powershell, "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "setup-state.ps1"),
                   "-Json", "-Profile", str(parameters.get("profile") or "analysis")]
    elif action == "update":
        command = [powershell, "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "update.ps1"), "-Json"]
        if parameters.get("allow_external_updates") or parameters.get("allowExternalUpdates"):
            command.append("-AllowExternalUpdates")
        if parameters.get("skip_external_tool_updates") or parameters.get("skipExternalToolUpdates"):
            command.append("-SkipExternalToolUpdates")
        if parameters.get("repair_prerequisites") or parameters.get("repairPrerequisites"):
            command.append("-RepairPrerequisites")
        if parameters.get("retry_failed_indexes"):
            command.append("-RetryFailedIndexes")
        update_id = parameters.get("update_id") or (None if parameters.get("new_run") else gate.get("update_id"))
        if update_id:
            command.extend(["-UpdateId", str(update_id)])
            wait_seconds = parameters.get("wait_seconds", 30)
            if isinstance(wait_seconds, bool) or not isinstance(wait_seconds, int) or not 0 <= wait_seconds <= 30:
                raise WorkflowError("update wait_seconds must be an integer from 0 to 30")
            command.extend(["-WaitSeconds", str(wait_seconds)])
    elif action in {"setup-plan", "setup-bootstrap"}:
        profile = str(parameters.get("profile") or "analysis")
        command = [
            powershell,
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(ROOT / "scripts" / "bootstrap.ps1"),
            "-Profile", profile, "-Json",
        ]
        if action == "setup-plan":
            command.append("-Plan")
        if parameters.get("resume"):
            command.append("-Resume")
        if parameters.get("install_prerequisites"):
            command.append("-InstallPrerequisites")
        if parameters.get("install_cc_1c_skills"):
            command.append("-InstallCc1cSkills")
        if parameters.get("wheelhouse"):
            command.extend(["-Wheelhouse", str(parameters["wheelhouse"])])
        if parameters.get("no_index"):
            command.append("-NoIndex")
    else:
        field_map = {
            "workflow_repository_url": "-WorkflowRepositoryUrl",
            "documentation_repository_url": "-DocumentationRepositoryUrl",
            "documentation_path": "-DocumentationPath",
            "extension_mode": "-ExtensionMode",
            "extension_repository_url": "-ExtensionRepositoryUrl",
            "extension_branch_prefix": "-ExtensionBranchPrefix",
            "extension_path": "-ExtensionPath",
            "configuration_path": "-ConfigurationPath",
            "functional_spec_template": "-FunctionalSpecTemplate",
            "gitea_base_url": "-GiteaBaseUrl",
            "gitea_owner": "-GiteaOwner",
            "gitea_repository": "-GiteaRepository",
            "git_user_name": "-GitUserName",
            "git_user_email": "-GitUserEmail",
            "project_reference": "-ProjectReference",
        }
        profile = str(parameters.get("profile") or "analysis")
        safe_values = {key: value for key, value in parameters.items() if key not in {"confirmed", "setup_id"}}
        checkpoint = create_setup_checkpoint(profile, safe_values, str(parameters.get("setup_id") or ""))
        gate["setup_id"] = checkpoint["setup_id"]
        save_gate(gate)
        command = [powershell, "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "configure-project.ps1")]
        for key, flag in field_map.items():
            value = parameters.get(key)
            if value not in (None, ""):
                command.extend([flag, str(value)])
        if parameters.get("allow_shared_repository"):
            command.append("-AllowSharedWorkflowDocumentationRepository")
        if parameters.get("initialize_documentation_repository"):
            command.append("-InitializeDocumentationRepository")
        command.extend(["-Profile", profile, "-SetupId", checkpoint["setup_id"], "-Confirmed", "-Json"])
    result = run_update_command(command) if action == "update" else subprocess.run(
        command, cwd=ROOT, text=True, encoding="utf-8", errors="replace",
        capture_output=True, check=False, timeout=600,
    )
    output = "\n".join(value.strip() for value in (result.stdout, result.stderr) if value.strip())
    structured: dict[str, Any] | None = None
    if action == "update":
        structured = parse_updater_result(result.stdout)
    else:
        try:
            structured = json.loads(result.stdout.lstrip("\ufeff")) if result.stdout.strip() else None
        except json.JSONDecodeError:
            structured = None
    next_action = None
    if action == "update" and structured and structured.get("update_id"):
        gate["update_id"] = structured["update_id"]
        gate["update_result"] = structured
        gate.pop("action_completed", None)
        if structured.get("state") == "WAITING_BACKGROUND":
            gate["state"] = "WAITING_BACKGROUND"
        elif gate.get("state") == "WAITING_BACKGROUND":
            gate["state"] = "READY"
        save_gate(gate)
    if action == "setup-configure" and structured and structured.get("state") == "WAITING_BACKGROUND":
        checkpoint = read_json(setup_checkpoint_path(str(gate.get("setup_id"))), {})
        checkpoint["state"] = "WAITING_BACKGROUND"
        checkpoint["pending_jobs"] = structured.get("jobs", [])
        save_setup_checkpoint(checkpoint)
        gate["state"] = "WAITING_BACKGROUND"
        save_gate(gate)
        print(json.dumps(structured, ensure_ascii=False, indent=2))
        return 0
    if result.returncode == 0:
        if action == "setup-configure":
            gate["state"] = "READY"
            gate["action_completed"] = action
            save_gate(gate)
            next_action = "flow1c_complete"
        elif action == "update" and structured and structured.get("state") == "READY" and structured.get("ready") is True:
            gate["state"] = "READY"
            gate["action_completed"] = action
            save_gate(gate)
            next_action = "flow1c_complete"
        elif action == "setup-bootstrap":
            next_action = "flow1c_begin"
    if structured is not None:
        print(json.dumps(structured, ensure_ascii=False, indent=2))
        return result.returncode
    if action == "update":
        raise WorkflowError("Updater returned no valid JSON result; update completion was not recorded. Inspect the command output and retry this gate safely. " + output[-1000:])
    print(json.dumps({"state": "ACTION_COMPLETE" if result.returncode == 0 else "BLOCKED", "next": next_action, "exit_code": result.returncode, "output": output}, ensure_ascii=False, indent=2))
    return result.returncode


def render_status() -> str:
    rows = []
    for path in sorted((project_root() / "work-items").glob("*/manifest.yaml")):
        manifest = read_json(path, {})
        rows.append(
            (
                manifest.get("code", path.parent.name),
                manifest.get("title", ""),
                manifest.get("status", "unknown"),
                ", ".join(manifest.get("requirements", [])),
                manifest.get("approvals", {}).get("functional_architect", "pending"),
                manifest.get("approvals", {}).get("technical_architect", "pending"),
            )
        )
    lines = [
        "# Project status",
        "",
        "| FS | Title | Status | Requirements | Functional approval | Technical approval |",
        "|---|---|---|---|---|---|",
    ]
    lines.extend("| " + " | ".join(str(value).replace("|", "\\|") for value in row) + " |" for row in rows)
    if not rows:
        lines.append("| — | No work items | — | — | — | — |")
    return "\n".join(lines) + "\n"


def cmd_status(args: argparse.Namespace) -> int:
    value = render_status()
    if args.write:
        target = project_root() / "wiki" / "status.md"
        write_text(target, value)
        print(target)
    else:
        print(value, end="")
    return 0


def cmd_git_commit(args: argparse.Namespace) -> int:
    code = args.code
    load_manifest(code)
    if not in_git_repository():
        raise WorkflowError("Current directory is not a Git repository.")
    item_relative = work_item_root(code).relative_to(project_root()).as_posix()
    allowed = [item_relative, "wiki/status.md"]
    if args.include_registry:
        allowed.append("registry")
    run_git(["add", "--", *allowed])
    staged = run_git(["diff", "--cached", "--name-only"]).stdout.strip()
    if not staged:
        raise WorkflowError("No allowed changes are staged for commit.")
    unexpected = [line for line in staged.splitlines() if not any(line == root or line.startswith(root + "/") for root in allowed)]
    if unexpected:
        raise WorkflowError("Unexpected staged paths: " + ", ".join(unexpected))
    result = run_git(["commit", "-m", args.message])
    print(result.stdout.strip())
    return 0


def api_request(url: str, token: str, *, method: str = "GET", payload: Any = None) -> Any:
    body = None if payload is None else json.dumps(payload).encode("utf-8")
    request = urllib.request.Request(url, data=body, method=method)
    request.add_header("Accept", "application/json")
    request.add_header("Authorization", f"token {token}")
    if body is not None:
        request.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise WorkflowError(f"Gitea API returned HTTP {exc.code}: {details}") from exc
    except urllib.error.URLError as exc:
        raise WorkflowError(f"Cannot reach Gitea API: {exc}") from exc


def cmd_pr_create(args: argparse.Namespace) -> int:
    code = args.code
    validate_publication(code, args.phase)
    _, manifest = load_manifest(code)
    config, local = load_config()
    gitea = {**config.get("gitea", {}), **local.get("gitea", {})}
    required = ["base_url", "owner", "repository"]
    missing = [key for key in required if not str(gitea.get(key, "")).strip()]
    if missing:
        raise WorkflowError("Missing Gitea settings: " + ", ".join(missing))
    token_env = gitea.get("token_env", "FLOW1C_GITEA_TOKEN")
    token = os.environ.get(token_env, "")
    if not token:
        raise WorkflowError(f"Environment variable {token_env} is not set.")
    if not in_git_repository():
        raise WorkflowError("Current directory is not a Git repository.")
    branch = run_git(["branch", "--show-current"]).stdout.strip()
    if not branch or branch in {"main", "master"}:
        raise WorkflowError("Refusing to create a PR from the default branch.")
    if run_git(["status", "--porcelain"]).stdout.strip():
        raise WorkflowError("Commit all changes before creating a PR.")
    run_git(["push", "--set-upstream", "origin", branch], capture=True)

    phase_reviewers = {
        "specification": gitea.get("reviewers", {}).get("functional", []),
        "technical": gitea.get("reviewers", {}).get("technical", []),
        "acceptance": gitea.get("reviewers", {}).get("functional", []),
    }
    reviewers = phase_reviewers.get(args.phase, [])
    title = args.title or f"{code}: {manifest.get('title')} [{args.phase}]"
    body = (
        f"FS: {code}\n\n"
        f"Requirements: {', '.join(manifest.get('requirements', []))}\n\n"
        f"Status: {manifest.get('status')}\n\n"
        "Prepared by Flow1C. Human review and approval are required."
    )
    base_url = str(gitea["base_url"]).rstrip("/")
    owner = urllib.parse.quote(str(gitea["owner"]), safe="")
    repository = urllib.parse.quote(str(gitea["repository"]), safe="")
    url = f"{base_url}/api/v1/repos/{owner}/{repository}/pulls"
    payload = {
        "base": config.get("project", {}).get("default_branch", "main"),
        "head": branch,
        "title": title,
        "body": body,
        "reviewers": reviewers,
    }
    response = api_request(url, token, method="POST", payload=payload)
    print(response.get("html_url") or response.get("url") or json.dumps(response, ensure_ascii=False))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Flow1C CLI")
    subparsers = parser.add_subparsers(dest="command", required=True)

    doctor = subparsers.add_parser("doctor", help="Check local prerequisites and paths")
    doctor.add_argument("--json", action="store_true", help="Emit a machine-readable readiness gate")
    doctor.add_argument("--format", choices=("docx", "markdown"), help="Selected template format for targeted readiness")
    doctor.add_argument("--gate-id", help="Infer the selected template format from a saved request pin")
    doctor_scope = doctor.add_mutually_exclusive_group()
    doctor_scope.add_argument("--profile", choices=PROFILE_NAMES + ("template-markdown", "template-docx"))
    doctor_scope.add_argument("--operation", choices=tuple(load_stages()["operations"]))
    doctor.set_defaults(handler=cmd_doctor)

    registry_import = subparsers.add_parser("registry-import", help="Validate and normalize an Excel registry")
    registry_import.add_argument("--file", required=True)
    registry_import.add_argument("--allow-errors", action="store_true", help="Deprecated compatibility flag; partial indexes are written by default")
    registry_import.set_defaults(handler=cmd_registry_import)

    registry_reconcile = subparsers.add_parser("registry-reconcile", help="Reconcile a provisional work item with a corrected registry")
    registry_reconcile.add_argument("--file", required=True)
    registry_reconcile.add_argument("--code", required=True, help="User-supplied work reference")
    registry_reconcile.add_argument("--mappings-json", default="{}", help="Confirmed user-brief to registry ID mapping")
    registry_reconcile.set_defaults(handler=cmd_registry_reconcile)

    fs_start = subparsers.add_parser("fs-start", help="Create a work item for a user-supplied reference")
    fs_start.add_argument("--task-reference")
    fs_start.add_argument("--project-reference")
    fs_start.add_argument("--code", help="Deprecated alias for --task-reference")
    fs_start.add_argument("--g-number", help=argparse.SUPPRESS)
    fs_start.add_argument("--reference-kind", choices=REFERENCE_KINDS, default="auto")
    fs_start.add_argument("--title")
    fs_start.add_argument("--requirements", help="Comma-separated IDs when the FS is absent from the registry")
    fs_start.add_argument("--create-branch", action="store_true")
    fs_start.set_defaults(handler=cmd_fs_start)

    context = subparsers.add_parser("context-build", help="Build a compact role-specific context")
    context.add_argument("--code", required=True)
    context.add_argument("--role", required=True, choices=VALID_ROLES)
    context.set_defaults(handler=cmd_context_build)

    agent_begin = subparsers.add_parser("agent-begin", help="Create a machine-readable gate for a natural-language request")
    agent_begin.add_argument("--json-stdin", action="store_true")
    agent_begin.add_argument("--operation", choices=("ambiguous", *tuple(load_stages()["operations"])))
    agent_begin.add_argument("--code")
    agent_begin.add_argument("--task-reference")
    agent_begin.add_argument("--project-reference")
    agent_begin.add_argument("--reference-kind", choices=REFERENCE_KINDS, default="auto")
    agent_begin.add_argument("--git-ref")
    agent_begin.add_argument("--git-refs", nargs="*", default=[])
    agent_begin.add_argument("--target-ref", default="")
    agent_begin.add_argument("--refresh-git-refs", action="store_true")
    agent_begin.add_argument("--g-number", help=argparse.SUPPRESS)
    agent_begin.add_argument("--summary", default="")
    agent_begin.add_argument("--query-intent", choices=("create", "review", "optimize"))
    agent_begin.add_argument("--path", action="append", default=[])
    agent_begin.add_argument("--allow-incomplete-draft", action="store_true")
    agent_begin.add_argument("--mode", choices=MODES, default=None)
    agent_begin.add_argument("--requirements", nargs="*", default=[])
    agent_begin.add_argument("--mismatch", default="", help="Observed semantic mismatch; ask instead of silently remapping")
    agent_begin.add_argument("--profile", choices=PROFILE_NAMES)
    agent_begin.set_defaults(handler=cmd_agent_begin)

    dialogue = subparsers.add_parser("agent-dialogue", help="Persist a question, user answer or decision on the same request")
    dialogue.add_argument("--json-stdin", action="store_true")
    dialogue.add_argument("--gate-id")
    dialogue.add_argument("--action", choices=("ask", "answer", "record", "deviate"))
    dialogue.add_argument("--question", default="")
    dialogue.add_argument("--answer", default="")
    dialogue.add_argument("--kind", choices=("user_answer", "assumption", "open_question", "decision"), default="user_answer")
    dialogue.add_argument("--deviation-type", choices=("process", "registry_bypass"))
    dialogue.add_argument("--actor", default="user", help=argparse.SUPPRESS)
    dialogue.add_argument("--scope", choices=("gate", "work-item"))
    dialogue.add_argument("--condition-ids", nargs="*", default=None)
    dialogue.add_argument("--user-statement", default="")
    dialogue.add_argument("--resolution", choices=("independent-draft", "correct-code"))
    dialogue.add_argument("--reference-kind", choices=REFERENCE_KINDS)
    dialogue.add_argument("--code")
    dialogue.add_argument("--mode", choices=("explore", "draft"))
    dialogue.add_argument("--profile", choices=PROFILE_NAMES)
    dialogue.set_defaults(handler=cmd_agent_dialogue)

    inspect = subparsers.add_parser("agent-inspect", help="Bounded read/search of workflow or accepted request text")
    inspect.add_argument("--json-stdin", action="store_true")
    inspect.add_argument("--gate-id")
    inspect.add_argument("--scope", choices=("workflow", "request"), default="request")
    inspect.add_argument("--path", default="")
    inspect.add_argument("--query", default="")
    inspect.add_argument("--start-line", type=int, default=1)
    inspect.add_argument("--max-chars", type=int, default=12000)
    inspect.set_defaults(handler=cmd_agent_inspect)

    git_refresh = subparsers.add_parser("agent-git-refresh", help="Refresh explicitly allowed remote-tracking refs")
    git_refresh.add_argument("--json-stdin", action="store_true")
    git_refresh.add_argument("--gate-id")
    git_refresh.add_argument("--repository", choices=("workflow", "extension"), default="extension")
    git_refresh.add_argument("--remote", default="origin")
    git_refresh.add_argument("--refs", nargs="*", default=[])
    git_refresh.set_defaults(handler=cmd_agent_git_refresh)

    git_inspect = subparsers.add_parser("agent-git-inspect", help="Read Git refs, history and diffs for a free review")
    git_inspect.add_argument("--json-stdin", action="store_true")
    git_inspect.add_argument("--gate-id")
    git_inspect.add_argument("--repository", choices=("workflow", "extension"), default="extension")
    git_inspect.add_argument("--action", choices=("resolve", "log", "latest-merge", "integration", "merge-search",
                                                   "branch-changes", "history-search", "diff", "read-at-ref"), default="resolve")
    git_inspect.add_argument("--git-ref", default="")
    git_inspect.add_argument("--git-refs", nargs="*", default=[])
    git_inspect.add_argument("--base-ref", default="")
    git_inspect.add_argument("--target-ref", default="")
    git_inspect.add_argument("--detail", choices=("names", "stat", "patch", "first-parent-patch", "remerge-diff", "combined"), default="patch")
    git_inspect.add_argument("--paths", nargs="*", default=[])
    git_inspect.add_argument("--path", default="")
    git_inspect.add_argument("--cursor", default="")
    git_inspect.add_argument("--max-files", type=int, default=100)
    git_inspect.add_argument("--start-line", type=int, default=1)
    git_inspect.add_argument("--subject-query", default="")
    git_inspect.add_argument("--regex", action="store_true")
    git_inspect.add_argument("--merges-only", action="store_true")
    git_inspect.add_argument("--first-parent", action="store_true")
    git_inspect.add_argument("--min-parents", type=int)
    git_inspect.add_argument("--max-parents", type=int)
    git_inspect.add_argument("--since", default="")
    git_inspect.add_argument("--until", default="")
    git_inspect.add_argument("--include-pr-evidence", action=argparse.BooleanOptionalAction, default=True)
    git_inspect.add_argument("--include-patch-evidence", action=argparse.BooleanOptionalAction, default=True)
    git_inspect.add_argument("--max-count", type=int, default=20)
    git_inspect.add_argument("--max-chars", type=int, default=40000)
    git_inspect.set_defaults(handler=cmd_agent_git_inspect_v2)

    git_snapshot = subparsers.add_parser("agent-git-snapshot", help="Create an isolated historical Git snapshot")
    git_snapshot.add_argument("--json-stdin", action="store_true")
    git_snapshot.add_argument("--gate-id")
    git_snapshot.add_argument("--repository", choices=("workflow", "extension"), default="extension")
    git_snapshot.add_argument("--action", choices=("create", "cleanup"), default="create")
    git_snapshot.add_argument("--git-ref", required=False, default="")
    git_snapshot.add_argument("--paths", nargs="*", default=[])
    git_snapshot.add_argument("--ttl-hours", type=int, default=168)
    git_snapshot.set_defaults(handler=cmd_agent_git_snapshot)

    promote = subparsers.add_parser("draft-promote", help="Attach a completed free draft without changing status or approvals")
    promote.add_argument("--json-stdin", action="store_true")
    promote.add_argument("--gate-id")
    promote.add_argument("--code")
    promote.add_argument("--task-reference")
    promote.add_argument("--project-reference")
    promote.add_argument("--g-number", help=argparse.SUPPRESS)
    promote.add_argument("--requirements", nargs="*", default=[])
    promote.add_argument("--confirmed", action="store_true")
    promote.set_defaults(handler=cmd_draft_promote)

    intake = subparsers.add_parser("artifact-intake", help="Validate and copy user artifacts into a controlled intake")
    intake.add_argument("--json-stdin", action="store_true")
    intake.add_argument("--gate-id")
    intake.add_argument("--source", action="append", default=[])
    intake.add_argument("--code")
    intake.add_argument("--category", choices=tuple(load_stages().get("artifact_categories", {})))
    intake.add_argument("--confirm-absence", action="append", default=[])
    intake.add_argument("--confirm-large", action="store_true")
    intake.add_argument("--received-via", choices=("chat-attachment", "file", "folder"), default="")
    intake.add_argument("--promote-intake-id")
    intake.set_defaults(handler=cmd_artifact_intake)

    agent_context = subparsers.add_parser("agent-context", help="Build and return a gated role context")
    agent_context.add_argument("--json-stdin", action="store_true")
    agent_context.add_argument("--gate-id")
    agent_context.set_defaults(handler=cmd_agent_context)

    agent_diff = subparsers.add_parser("agent-diff", help="Return and record a bounded extension diff")
    agent_diff.add_argument("--json-stdin", action="store_true")
    agent_diff.add_argument("--gate-id")
    agent_diff.add_argument("--max-chars", type=int, default=120000)
    agent_diff.set_defaults(handler=cmd_agent_diff)

    source_read = subparsers.add_parser("agent-source-read", help="Read a gated extension source file")
    source_read.add_argument("--json-stdin", action="store_true")
    source_read.add_argument("--gate-id")
    source_read.add_argument("--source", choices=("extension", "configuration"), default="extension")
    source_read.add_argument("--path")
    source_read.add_argument("--max-chars", type=int, default=80000)
    source_read.set_defaults(handler=cmd_agent_source_read)

    source_query = subparsers.add_parser("source-query", help="Query indexed BSL sources and record evidence")
    source_query.add_argument("--json-stdin", action="store_true")
    source_query.add_argument("--gate-id")
    source_query.add_argument("--source")
    source_query.add_argument("--query")
    source_query.add_argument("--code")
    source_query.add_argument("--reason")
    source_query.add_argument("--effort", choices=("low", "medium", "high"), default="medium")
    source_query.add_argument("--max-chars", type=int, default=12000)
    source_query.set_defaults(handler=cmd_source_query)

    cc_inspect = subparsers.add_parser("cc-inspect", help="Run a bounded read-only cc-1c-skills metadata inspector")
    cc_inspect.add_argument("--json-stdin", action="store_true")
    cc_inspect.add_argument("--gate-id")
    cc_inspect.add_argument("--source", choices=("request", "configuration", "extension"))
    cc_inspect.add_argument("--operation", choices=("meta-overview", "meta-full", "meta-item", "skd-overview", "skd-query", "skd-fields", "skd-params", "skd-links", "skd-full"))
    cc_inspect.add_argument("--path")
    cc_inspect.add_argument("--name", default="")
    cc_inspect.add_argument("--max-chars", type=int, default=12000)
    cc_inspect.set_defaults(handler=cmd_cc_inspect)

    query_schema = subparsers.add_parser("query-schema", help="Extract exact fields from one gated 1C XML export")
    query_schema.add_argument("--json-stdin", action="store_true")
    query_schema.add_argument("--gate-id")
    query_schema.add_argument("--source", choices=("request", "configuration", "extension"))
    query_schema.add_argument("--path")
    query_schema.add_argument("--rlm-evidence-id", default="")
    query_schema.set_defaults(handler=cmd_query_schema)

    query_check = subparsers.add_parser("query-check", help="Save and statically inspect the final 1C query text")
    query_check.add_argument("--json-stdin", action="store_true")
    query_check.add_argument("--gate-id")
    query_check.add_argument("--text", default="")
    query_check.add_argument("--baseline-text", default="")
    query_check.add_argument("--schema-ids", nargs="*", default=[])
    query_check.add_argument("--expected-result", default="")
    query_check.add_argument("--assumptions", nargs="*", default=[])
    query_check.add_argument("--changes", nargs="*", default=[])
    query_check.set_defaults(handler=cmd_query_check)

    agent_write = subparsers.add_parser("agent-write", help="Write only within the target permitted by a gate")
    agent_write.add_argument("--json-stdin", action="store_true")
    agent_write.add_argument("--gate-id")
    agent_write.add_argument("--target", choices=("work-item", "extension", "draft"))
    agent_write.add_argument("--path")
    agent_write.add_argument("--content", default="")
    agent_write.add_argument("--content-stdin", action="store_true")
    agent_write.set_defaults(handler=cmd_agent_write)

    analyze = subparsers.add_parser("agent-analyze-bsl", help="Run BSL Language Server and record evidence")
    analyze.add_argument("--json-stdin", action="store_true")
    analyze.add_argument("--gate-id")
    analyze.add_argument("--source", type=json.loads, help="Typed BSL source selector")
    analyze.set_defaults(handler=cmd_agent_analyze_bsl)

    complete = subparsers.add_parser("agent-complete", help="Validate evidence and close an agent gate")
    complete.add_argument("--json-stdin", action="store_true")
    complete.add_argument("--gate-id")
    complete.add_argument("--output")
    complete.add_argument("--summary", default="")
    complete.set_defaults(handler=cmd_agent_complete)

    action = subparsers.add_parser("agent-action", help="Run a gated workflow action")
    action.add_argument("--json-stdin", action="store_true")
    action.add_argument("--gate-id")
    action.add_argument(
        "--action",
        choices=("setup-audit", "setup-plan", "setup-bootstrap", "setup-configure", "setup-status", "setup-resume", "doctor", "update", "registry-import", "fs-start", "provisional-start", "registry-reconcile", "status", "publish"),
    )
    action.add_argument("--parameters-json", default="{}")
    action.set_defaults(handler=cmd_agent_action)

    redmine = subparsers.add_parser("redmine", help="Configure Redmine, inspect issues, or explicitly upload a DMSF document")
    redmine_commands = redmine.add_subparsers(dest="redmine_command", required=True)
    redmine_configure = redmine_commands.add_parser("configure", help="Connect this workflow checkout to Redmine")
    redmine_configure.add_argument("--url", required=True, help="Redmine HTTPS base URL, optionally with a path prefix")
    redmine_configure.set_defaults(handler=cmd_redmine_configure)
    redmine_status = redmine_commands.add_parser("status", help="Show local Redmine connection readiness")
    redmine_status.set_defaults(handler=cmd_redmine_status)
    redmine_test = redmine_commands.add_parser("test", help="Test the configured API key with Redmine")
    redmine_test.set_defaults(handler=cmd_redmine_test)
    redmine_disconnect = redmine_commands.add_parser("disconnect", help="Remove this checkout's Redmine configuration")
    redmine_disconnect.set_defaults(handler=cmd_redmine_disconnect)
    redmine_cleanup = redmine_commands.add_parser("cleanup", help="Retry removal of saved Redmine credentials")
    redmine_cleanup.add_argument("--url", help="Optional HTTPS URL whose saved credential should be removed")
    redmine_cleanup.set_defaults(handler=cmd_redmine_cleanup)
    redmine_files = redmine_commands.add_parser("files", help="List standard and currently attached DMSF issue files without downloading")
    redmine_files.add_argument("issue", help="Numeric Redmine issue number")
    redmine_files.add_argument("--json", action="store_true", help="Emit a machine-readable file inventory")
    redmine_files.set_defaults(handler=cmd_redmine_files)
    redmine_relations = redmine_commands.add_parser("relations", help="Show direct issue relations with related issue tracker and status")
    redmine_relations.add_argument("issue", help="Numeric Redmine issue number")
    redmine_relations.add_argument("--json", action="store_true", help="Emit a machine-readable relation inventory")
    redmine_relations.set_defaults(handler=cmd_redmine_relations)
    redmine_upload = redmine_commands.add_parser("upload", help="Upload one local file to DMSF and attach it to an issue")
    redmine_upload.add_argument("issue", help="Numeric Redmine issue number")
    redmine_upload.add_argument("--file", required=True, help="Explicit local file path to upload")
    redmine_upload.add_argument("--confirmed", action="store_true", help="Confirm this exact issue and local file for the external write")
    redmine_upload.add_argument("--json", action="store_true", help="Emit a machine-readable upload result")
    redmine_upload.set_defaults(handler=cmd_redmine_upload)
    redmine_revise = redmine_commands.add_parser("revise", help="Add a revision to one DMSF document attached to an issue")
    redmine_revise.add_argument("issue", help="Numeric Redmine issue number")
    redmine_revise.add_argument("--dms-file", required=True, help="Exact existing DMSF file ID")
    redmine_revise.add_argument("--expected-revision", required=True, help="Current revision ID for race protection")
    redmine_revise.add_argument("--file", required=True, help="Local replacement file with the same name")
    redmine_revise.add_argument("--confirmed", action="store_true", help="Confirm revision of this exact document")
    redmine_revise.add_argument("--json", action="store_true", help="Emit a machine-readable result")
    redmine_revise.set_defaults(handler=cmd_redmine_revise)
    redmine_fetch = redmine_commands.add_parser("fetch", help="Find an issue by number and import its supported attachments")
    redmine_fetch.add_argument("issue", help="Numeric Redmine issue number")
    redmine_fetch.add_argument("--code", help="Existing Flow1C task reference; omit to place files in inbox")
    redmine_fetch.add_argument("--gate-id", help="Active Flow1C gate; derives and validates the target work item")
    redmine_fetch.add_argument("--dms-file", action="append", default=[], help="Explicit DMSF file ID to import; repeat for multiple files")
    redmine_fetch.add_argument("--dms-revision", help="DMSF revision ID; requires exactly one --dms-file")
    redmine_fetch.add_argument("--all-dms", action="store_true", help="Import every DMSF file currently attached to the issue")
    redmine_fetch.set_defaults(handler=cmd_redmine_fetch)

    status = subparsers.add_parser("status", help="Show or write the project status")
    status.add_argument("--write", action="store_true")
    status.set_defaults(handler=cmd_status)

    commit = subparsers.add_parser("git-commit", help="Commit only files allowed for one work item")
    commit.add_argument("--code", required=True)
    commit.add_argument("--message", required=True)
    commit.add_argument("--include-registry", action="store_true")
    commit.set_defaults(handler=cmd_git_commit)

    pr = subparsers.add_parser("pr-create", help="Push the current branch and create a Gitea PR")
    pr.add_argument("--code", required=True)
    pr.add_argument("--phase", required=True, choices=("specification", "technical", "acceptance"))
    pr.add_argument("--title")
    pr.set_defaults(handler=cmd_pr_create)

    section_catalog = subparsers.add_parser("section-catalog", help="List canonical functional-specification sections")
    section_catalog.add_argument("--json-stdin", action="store_true")
    section_catalog.add_argument("--gate-id")
    section_catalog.set_defaults(handler=cmd_section_catalog)

    section_save = subparsers.add_parser("section-save", help="Save a versioned functional-specification section draft")
    section_save.add_argument("--json-stdin", action="store_true")
    section_save.add_argument("--gate-id")
    section_save.add_argument("--request-id")
    section_save.add_argument("--work-reference")
    section_save.add_argument("--section-id", action="append", default=[])
    section_save.add_argument("--sections", nargs="*", default=[])
    section_save.add_argument("--content", default="")
    section_save.add_argument("--content-stdin", action="store_true")
    section_save.add_argument("--checklist", default="[]")
    section_save.add_argument("--sources", nargs="*", default=[])
    section_save.set_defaults(handler=cmd_section_save)

    section_approve = subparsers.add_parser("section-approve", help="Record explicit approval of the current section hash")
    section_approve.add_argument("--json-stdin", action="store_true")
    section_approve.add_argument("--gate-id")
    section_approve.add_argument("--request-id")
    section_approve.add_argument("--work-reference")
    section_approve.add_argument("--section-id", action="append", default=[])
    section_approve.add_argument("--sections", nargs="*", default=[])
    section_approve.add_argument("--approved-by", default="user")
    section_approve.add_argument("--approval-statement", required=False, default="")
    section_approve.set_defaults(handler=cmd_section_approve)

    docx_inspect = subparsers.add_parser("docx-inspect", help="Inspect a selected DOCX and locate section anchors")
    docx_inspect.add_argument("--json-stdin", action="store_true")
    docx_inspect.add_argument("--gate-id")
    docx_inspect.add_argument("--source", default="")
    docx_inspect.add_argument("--path", default="")
    docx_inspect.add_argument("--section-id", action="append", default=[])
    docx_inspect.add_argument("--sections", nargs="*", default=[])
    docx_inspect.set_defaults(handler=cmd_docx_inspect)

    docx_plan = subparsers.add_parser("docx-write-plan", help="Create a checked, explicit DOCX write plan")
    docx_plan.add_argument("--json-stdin", action="store_true")
    docx_plan.add_argument("--gate-id")
    docx_plan.add_argument("--source", default="")
    docx_plan.add_argument("--path", default="")
    docx_plan.add_argument("--section-id", action="append", default=[])
    docx_plan.add_argument("--sections", nargs="*", default=[])
    docx_plan.add_argument("--request-id")
    docx_plan.add_argument("--work-reference")
    docx_plan.add_argument("--mode", choices=("replace", "append"), default="")
    docx_plan.add_argument("--modes", default="")
    docx_plan.add_argument("--output", default="")
    docx_plan.add_argument("--plan", default="")
    docx_plan.set_defaults(handler=cmd_docx_write_plan)

    docx_write = subparsers.add_parser("docx-write", help="Write a previously checked DOCX plan to a new file")
    docx_write.add_argument("--json-stdin", action="store_true")
    docx_write.add_argument("--gate-id")
    docx_write.add_argument("--plan", default="")
    docx_write.set_defaults(handler=cmd_docx_write)
    interview = subparsers.add_parser("interview-register", help="Inspect, audit or write an interview XLSX draft without overwriting sources")
    interview.add_argument("--json-stdin", action="store_true")
    interview.add_argument("--gate-id")
    interview.add_argument("--action", choices=("inspect", "audit", "write"), default="inspect")
    interview.set_defaults(request={}, handler=cmd_interview_register)
    templates_cli.register_parser(subparsers, SimpleNamespace(**globals()))
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    configure_stdio_utf8()
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        if getattr(args, "json_stdin", False):
            request = read_json_stdin()
            args._json_input_fields = set(request)
            for key, value in request.items():
                attribute = str(key).replace("-", "_")
                if not hasattr(args, attribute):
                    raise WorkflowError(f"Unknown JSON input field: {key}")
                setattr(args, attribute, value)
        if getattr(args, "g_number", None):
            print("WARNING: --g-number is deprecated; use --task-reference. The value is treated as an opaque user reference.", file=sys.stderr)
        return int(args.handler(args))
    except RedmineOperationError as exc:
        print(json.dumps(exc.payload, ensure_ascii=False, indent=2))
        return 2
    except (SectionPolicyError, DocxError, TemplateError, InterviewError) as exc:
        payload = {"schema_version": 1, "state": "BLOCKED", "ready": False,
                   "errors": [exc.as_dict()]}
        if getattr(args, "json_stdin", False) or getattr(args, "json", False):
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        else:
            print(f"ERROR [{exc.code}]: {exc}", file=sys.stderr)
        return 2
    except WorkflowError as exc:
        if getattr(args, "json_stdin", False) or getattr(args, "json", False):
            print(json.dumps({"schema_version": 1, "state": "BLOCKED", "ready": False, "errors": [{
                "code": "WORKFLOW_ERROR", "message": str(exc), "recoverable": True,
                "next_action": "Correct the reported input or prerequisite and retry the same operation."
            }]}, ensure_ascii=False, indent=2))
        else:
            print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    except subprocess.CalledProcessError as exc:
        detail = (exc.stderr or exc.stdout or str(exc)).strip()
        if getattr(args, "json_stdin", False) or getattr(args, "json", False):
            print(json.dumps({"schema_version": 1, "state": "BLOCKED", "ready": False, "errors": [{
                "code": "COMMAND_FAILED", "message": detail, "recoverable": True,
                "next_action": "Inspect the failed command output, correct the environment, and retry."
            }]}, ensure_ascii=False, indent=2))
        else:
            print(f"ERROR: command failed: {detail}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
