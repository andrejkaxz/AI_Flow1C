from __future__ import annotations
from flow1c import context as svc_context
from flow1c import readiness as svc_readiness
from flow1c import sources as svc_sources
from flow1c import work_items as svc_work_items
import sys
import io
import json
import shutil
import sys
import unittest
from contextlib import redirect_stdout, redirect_stderr
from pathlib import Path
from unittest import mock
import test_flow1c as fixtures
from scripts.flow1c_interview_policy import (
    COLUMNS,
    InterviewError,
    apply_changes,
    validate_requirements,
    validate_rows,
)
from scripts.flow1c_interview_workbook import audit_book, load_book, read_register

flow1c = fixtures.flow1c


def row(question_id="Q-0001", question="Кто создаёт поставщика и проверяет дубли?"):
    return {
        "l1_code": "01",
        "l1_name": "Закупки",
        "l2_code": "01.01",
        "l2_name": "Управление поставщиками",
        "l3_code": "01.01.01",
        "l3_name": "Создание и учёт поставщиков",
        "question_id": question_id,
        "question": question,
        "answer": "",
        "owner": "Уточнить",
        "interviewer": "Уточнить",
    }


class InterviewPolicyTests(unittest.TestCase):

    def test_same_process_can_have_distinct_questions(self):
        self.assertEqual(len(validate_rows([row(), row("Q-0002", "Как изменяют данные?")])), 2)

    def test_conflicting_hierarchy_duplicate_ids_and_questions_rejected(self):
        for other in (
            {**row("Q-0002", "Новый вопрос"), "l3_name": "Другое имя"},
            row(),
            row("Q-0002"),
        ):
            with self.subTest(other=other), self.assertRaises(InterviewError):
                validate_rows([row(), other])

    def test_answer_patch_keeps_codes_question_and_unmentioned_fields(self):
        original = {**row(), "answer": "Исходный ответ", "source": "Письмо заказчика"}
        result = apply_changes(
            [original], [], [{"question_id": "Q-0001", "fields": {"decision": "Уточнить"}}]
        )
        self.assertEqual(result[0]["answer"], original["answer"])
        self.assertEqual(result[0]["source"], original["source"])
        self.assertEqual(original.get("decision"), None)
        with self.assertRaises(InterviewError):
            apply_changes(
                [original], [], [{"question_id": "Q-0001", "fields": {"question": "Новый смысл"}}]
            )

    def test_requirements_need_answer_trace_and_cannot_be_approved(self):
        requirement = {
            "draft_id": "DRAFT-1",
            "question_id": "Q-0001",
            "text": "Проверять дубли",
            "acceptance_criterion": "Дубль обнаружен",
        }
        with self.assertRaises(InterviewError):
            validate_requirements([requirement], [row()])
        answered = {**row(), "answer": "Нужно проверять ИНН"}
        result = validate_requirements([requirement], [answered])[0]
        self.assertEqual(result["answer"], answered["answer"])
        self.assertEqual(result["l3_code"], answered["l3_code"])
        with self.assertRaises(InterviewError):
            validate_requirements([{**requirement, "status": "Утверждено"}], [answered])
        with self.assertRaises(InterviewError):
            apply_changes([], [{**row(), "requirement": "Предположение без ответа"}], [])
        with self.assertRaises(InterviewError):
            validate_requirements(
                [{**requirement, "implementation": "Типовой механизм подтверждён"}], [answered]
            )


class InterviewGateTests(unittest.TestCase):
    setUp = fixtures.NaturalLanguageGateTests.setUp
    tearDown = fixtures.NaturalLanguageGateTests.tearDown

    def call(self, command, *, expected=0, **data):
        output, errors = (io.StringIO(), io.StringIO())
        with (
            mock.patch.object(sys, "stdin", io.StringIO(json.dumps(data))),
            redirect_stdout(output),
            redirect_stderr(errors),
        ):
            code = flow1c.main([command, "--json-stdin"])
        self.assertEqual(code, expected, errors.getvalue() or output.getvalue())
        return json.loads(output.getvalue())

    def begin(self, mode=None):
        skill = self.root / ".agents/skills/flow1c-interview-preparation"
        skill.mkdir(parents=True, exist_ok=True)
        shutil.copy2(
            self.old_root / ".agents/skills/flow1c-interview-preparation/SKILL.md",
            skill / "SKILL.md",
        )
        return self.call(
            "agent-begin",
            operation="interview-preparation",
            summary="Подготовь реестр для интервью по закупкам",
            **{"mode": mode} if mode else {},
        )

    def write(self, gate, **request):
        return self.call(
            "interview-register",
            gate_id=gate["gate_id"],
            action="write",
            request={"rows": [row()], **request},
        )

    def test_chat_only_gate_writes_xlsx_and_completes_without_rlm_or_registry(self):
        (self.root / svc_context.LOCAL_CONFIG_FILE).unlink()
        (self.root / svc_context.CONFIG_FILE).unlink()
        with (
            mock.patch.object(svc_sources, "rlm_readiness", side_effect=AssertionError("No RLM")),
            mock.patch.object(
                svc_work_items, "load_manifest", side_effect=AssertionError("No work item")
            ),
        ):
            gate = self.begin()
            self.assertEqual(gate["mode"], "draft")
            self.assertEqual(gate["skill"], "flow1c-interview-preparation")
            self.assertIn("Что уточняем на интервью", gate["skill_instructions"])
            self.assertIn("flow1c_interview", gate["available_actions"])
            result = self.write(gate)
            done = self.call("agent-complete", gate_id=gate["gate_id"])
        self.assertEqual(done["state"], "DRAFT_COMPLETE")
        self.assertEqual(done["document_status"], "UNVERIFIED_DRAFT")
        book = load_book(Path(result["absolute_path"]))
        self.assertTrue(audit_book(book, "Реестр")["ready"])
        self.assertEqual(book["Реестр"]["A2"].value, "01")
        self.assertEqual(book["Реестр"]["A2"].data_type, "s")
        self.assertIn(".workspace", done["output"])
        book.close()

    def test_write_is_idempotent_and_never_replaces_output(self):
        gate = self.begin()
        first = self.write(gate, output_name="first.xlsx")
        self.assertTrue(self.write(gate, output_name="first.xlsx")["reused"])
        before = Path(first["absolute_path"]).read_bytes()
        error = self.call(
            "interview-register",
            expected=2,
            gate_id=gate["gate_id"],
            action="write",
            request={"rows": [row("Q-0002", "Другая тема")], "output_name": "first.xlsx"},
        )
        self.assertEqual(error["errors"][0]["code"], "INTERVIEW_OUTPUT_EXISTS")
        self.assertEqual(Path(first["absolute_path"]).read_bytes(), before)

    def test_update_keeps_custom_cells_comments_formulas_validation_and_filters(self):
        import openpyxl
        from openpyxl.comments import Comment
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.filters import FilterColumn, Filters

        gate = self.begin()
        source_result = self.write(
            gate, output_name="first.xlsx", rows=[{**row(), "answer": "Ответ заказчика"}]
        )
        source = Path(source_result["absolute_path"])
        book = openpyxl.load_workbook(source)
        sheet = book["Реестр"]
        sheet.cell(1, 15).value = "Пользовательская формула"
        sheet.cell(2, 15).value = "=1+2"
        sheet["J2"].comment = Comment("Сохранить", "Заказчик")
        sheet["J2"].fill = openpyxl.styles.PatternFill("solid", fgColor="FFFF00")
        validation = DataValidation(type="list", formula1='"Да,Нет"')
        validation.add("H2:H10")
        sheet.add_data_validation(validation)
        sheet.auto_filter.filterColumn.append(
            FilterColumn(colId=1, filters=Filters(filter=["Закупки"]))
        )
        source = self.root / "customer.xlsx"
        book.save(source)
        book.close()
        intake = self.call("artifact-intake", gate_id=gate["gate_id"], source=[str(source)])
        accepted = intake["copied"][0]["path"]
        before = source.read_bytes()
        updated = self.call(
            "interview-register",
            gate_id=gate["gate_id"],
            action="write",
            request={
                "source": accepted,
                "output_name": "second.xlsx",
                "rows": [row("Q-0002", "Как исправляют реквизиты?")],
                "patches": [
                    {"question_id": "Q-0001", "fields": {"decision": "Проверить типовой механизм"}}
                ],
            },
        )
        self.assertEqual(source.read_bytes(), before)
        copied = openpyxl.load_workbook(updated["absolute_path"])
        sheet = copied["Реестр"]
        self.assertEqual(sheet["J2"].value, "Ответ заказчика")
        self.assertEqual(sheet["J2"].comment.text, "Сохранить")
        self.assertEqual(sheet["J2"].fill.fgColor.rgb, "00FFFF00")
        self.assertEqual(sheet.cell(2, 15).value, "=1+2")
        self.assertEqual(sheet.auto_filter.ref, "A1:O3")
        self.assertEqual(sheet.tables["InterviewRegister"].ref, "A1:O3")
        self.assertEqual(sheet.tables["InterviewRegister"].autoFilter.ref, "A1:O3")
        self.assertEqual(sheet.auto_filter.filterColumn[0].filters.filter, ["Закупки"])
        self.assertEqual(str(sheet.data_validations.dataValidation[0].sqref), "H2:H10")
        self.assertEqual(sheet["L2"].value, "Q-0001")
        copied.close()

    def test_completion_rejects_changed_book(self):
        gate = self.begin()
        result = self.write(gate)
        Path(result["absolute_path"]).write_bytes(b"changed")
        blocked = self.call("agent-complete", gate_id=gate["gate_id"], expected=2)
        self.assertEqual(blocked["state"], "BLOCKED")

    def test_source_must_be_accepted_and_paths_contained(self):
        gate = self.begin()
        for request in (
            {"source": "../outside.xlsx"},
            {"source": "not-accepted.xlsx"},
            {"output_name": "../outside.xlsx", "rows": [row()]},
        ):
            with self.subTest(request=request):
                result = self.call(
                    "interview-register",
                    gate_id=gate["gate_id"],
                    action="write",
                    request=request,
                    expected=2,
                )
                self.assertEqual(result["state"], "BLOCKED")

    def test_explore_cannot_write_and_formal_is_rejected(self):
        gate = self.begin("explore")
        self.call(
            "interview-register",
            gate_id=gate["gate_id"],
            action="write",
            request={"rows": [row()]},
            expected=2,
        )
        self.call(
            "agent-begin",
            operation="interview-preparation",
            mode="formal",
            summary="Подготовь интервью",
            expected=2,
        )

    def test_new_text_is_literal_not_formula(self):
        gate = self.begin()
        result = self.write(gate, rows=[{**row(), "answer": '=HYPERLINK("https://example.test")'}])
        book = load_book(Path(result["absolute_path"]))
        self.assertEqual(book["Реестр"]["J2"].data_type, "s")
        book.close()

    def test_intake_legacy_register_allocates_ids_only_in_new_copy(self):
        import openpyxl

        source = self.root / "legacy.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        keys = list(COLUMNS)[:11]
        sheet.append([COLUMNS[key] for key in keys])
        sheet.append([row().get(key, "") for key in keys])
        sheet.auto_filter.ref = "A1:K1"
        book.save(source)
        book.close()
        before = source.read_bytes()
        gate = self.begin()
        intake = self.call("artifact-intake", gate_id=gate["gate_id"], source=[str(source)])
        accepted = intake["copied"][0]["path"]
        preview = self.call(
            "interview-register",
            gate_id=gate["gate_id"],
            action="inspect",
            request={"source": accepted},
        )
        self.assertFalse(preview["audit"]["ready"])
        self.assertEqual(preview["rows"][0]["question_id"], "Q-0001")
        result = self.call(
            "interview-register",
            gate_id=gate["gate_id"],
            action="write",
            request={"source": accepted},
        )
        self.assertEqual(source.read_bytes(), before)
        self.assertTrue(result["record"]["audit"]["ready"])

    def test_requirements_and_unavailable_sources_are_saved(self):
        gate = self.begin()
        result = self.write(
            gate,
            rows=[{**row(), "answer": "Нужен контроль дублей"}],
            requirements=[
                {
                    "draft_id": "DRAFT-1",
                    "question_id": "Q-0001",
                    "text": "Проверять дубли",
                    "acceptance_criterion": "Дубль обнаружен",
                }
            ],
            sources=[
                {
                    "id": "S1",
                    "location": "ИТС",
                    "availability": "Недоступен",
                    "conclusion": "Требует проверки",
                }
            ],
        )
        book = load_book(Path(result["absolute_path"]))
        self.assertEqual(book["FLOW1C Требования"]["B2"].value, "01.01.01")
        self.assertEqual(book["FLOW1C Требования"]["D2"].value, "Нужен контроль дублей")
        self.assertEqual(book["FLOW1C Требования"]["G2"].value, "Требует согласования")
        self.assertEqual(book["FLOW1C Источники"]["F2"].value, "Недоступен")
        book.close()

    def test_doctor_reports_missing_excel_without_project_readiness(self):
        self.begin()
        (self.root / svc_context.LOCAL_CONFIG_FILE).unlink()
        output = io.StringIO()
        with (
            mock.patch.object(
                svc_readiness, "_module_check", return_value=("openpyxl", "ERROR", "missing")
            ),
            redirect_stdout(output),
        ):
            code = flow1c.main(["doctor", "--json", "--operation", "interview-preparation"])
        self.assertEqual(code, 1)
        self.assertFalse(json.loads(output.getvalue())["ready"])

    def test_markdown_readiness_and_draft_do_not_need_excel(self):
        gate = self.begin()
        output = io.StringIO()
        with (
            mock.patch.object(
                svc_readiness, "_module_check", side_effect=AssertionError("No Excel check")
            ),
            redirect_stdout(output),
        ):
            code = flow1c.main(
                ["doctor", "--json", "--operation", "interview-preparation", "--format", "markdown"]
            )
        self.assertEqual(code, 0)
        self.assertTrue(json.loads(output.getvalue())["ready"])
        self.call(
            "agent-write",
            gate_id=gate["gate_id"],
            target="draft",
            path="result.md",
            content="# Реестр интервью\n\nЗакупки: уточнить владельца процесса.",
        )
        done = self.call("agent-complete", gate_id=gate["gate_id"], output="result.md")
        self.assertEqual(done["state"], "DRAFT_COMPLETE")

    def test_malformed_and_unsupported_books_fail_without_losing_source(self):
        from zipfile import ZipFile

        gate = self.begin()
        created = self.write(gate)
        original = Path(created["absolute_path"])
        customer = self.root / "embedded.xlsx"
        shutil.copy2(original, customer)
        with ZipFile(customer, "a") as archive:
            archive.writestr("xl/embeddings/document.bin", b"untrusted")
        before = customer.read_bytes()
        intake = self.call("artifact-intake", gate_id=gate["gate_id"], source=[str(customer)])
        blocked = self.call(
            "interview-register",
            expected=2,
            gate_id=gate["gate_id"],
            action="write",
            request={"source": intake["copied"][0]["path"]},
        )
        self.assertEqual(blocked["errors"][0]["code"], "INTERVIEW_UNSUPPORTED_BOOK")
        self.assertEqual(customer.read_bytes(), before)
        corrupt = self.root / "corrupt.xlsx"
        corrupt.write_bytes(b"not a workbook")
        intake = self.call("artifact-intake", gate_id=gate["gate_id"], source=[str(corrupt)])
        blocked = self.call(
            "interview-register",
            expected=2,
            gate_id=gate["gate_id"],
            action="inspect",
            request={"source": intake["copied"][0]["path"]},
        )
        self.assertEqual(blocked["errors"][0]["code"], "INTERVIEW_BOOK_INVALID")

    def test_blank_headers_and_trailing_notes_are_not_silently_reassigned(self):
        import openpyxl

        gate = self.begin()
        created = self.write(gate)
        for kind in ("blank-header", "note"):
            book = openpyxl.load_workbook(created["absolute_path"])
            sheet = book["Реестр"]
            if kind == "blank-header":
                sheet.cell(2, 15).value = "Custom value without header"
            else:
                sheet["J4"] = "Заметка за реестром"
            customer = self.root / (kind + ".xlsx")
            book.save(customer)
            book.close()
            intake = self.call("artifact-intake", gate_id=gate["gate_id"], source=[str(customer)])
            with self.subTest(kind=kind):
                self.call(
                    "interview-register",
                    expected=2,
                    gate_id=gate["gate_id"],
                    action="write",
                    request={
                        "source": intake["copied"][0]["path"],
                        "rows": [row("Q-0002", "Новый вопрос")],
                    },
                )


class InterviewAdapterTests(unittest.TestCase):

    def test_schema_matches_rows_and_answer_patch_contract(self):
        from jsonschema import Draft202012Validator

        root = Path(__file__).resolve().parents[1]
        schema = json.loads(
            (root / "schemas/interview-register.schema.json").read_text(encoding="utf-8")
        )
        Draft202012Validator.check_schema(schema)
        validator = Draft202012Validator(schema)
        request = {
            "gate_id": "12345678",
            "action": "write",
            "request": {"rows": [row()], "output_name": "draft.xlsx"},
        }
        self.assertEqual(list(validator.iter_errors(request)), [])
        request["request"]["patches"] = [{"question_id": "Q-0001", "fields": {"l3_code": "99"}}]
        self.assertTrue(list(validator.iter_errors(request)))

    def test_all_clients_route_to_single_canonical_skill(self):
        root = Path(__file__).resolve().parents[1]
        adapter = (root / ".opencode/tools/flow1c.ts").read_text(encoding="utf-8")
        self.assertIn("export const interview = tool(", adapter)
        self.assertIn('"interview-preparation"', adapter)
        self.assertIn('"interview-register", "--json-stdin"', adapter)
        wrapper = (root / ".claude/skills/flow1c-interview-preparation/SKILL.md").read_text(
            encoding="utf-8"
        )
        self.assertIn("../../../.agents/skills/flow1c-interview-preparation/SKILL.md", wrapper)
        for file in (
            "AGENTS.md",
            "docs/intent-routing.md",
            ".opencode/agents/flow1c-controller.md",
        ):
            self.assertIn("interview-preparation", (root / file).read_text(encoding="utf-8"))
        self.assertTrue((root / "CLAUDE.md").read_text(encoding="utf-8").startswith("@AGENTS.md\n"))


if __name__ == "__main__":
    unittest.main()
