from __future__ import annotations
from flow1c import context as svc_context
from flow1c import documents as svc_documents
from scripts import external_tools as ext_scripts_external_tools
from scripts import flow1c_docx as ext_scripts_flow1c_docx
from scripts import flow1c_policy as ext_scripts_flow1c_policy
from flow1c import intake as svc_intake
from flow1c import readiness as svc_readiness
from flow1c import registry as svc_registry
from flow1c import sources as svc_sources
from flow1c import system as svc_system
from flow1c import work_items as svc_work_items
from flow1c.workflow import git_actions as svc_workflow_git_actions
from flow1c.workflow import state as svc_workflow_state
from flow1c.errors import WorkflowError
from scripts import flow1c_git as git_analysis
from flow1c.storage import read_json
from flow1c.storage import sha256
import subprocess
import sys
from flow1c.storage import write_json
import argparse
import importlib.util
import io
import json
import os
import shutil
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock
from openpyxl import Workbook
from flow1c import cli as flow1c

MODULE_PATH = Path(__file__).resolve().parents[1] / "scripts/flow1c.py"


class Args:

    def __init__(self, file: Path, allow_errors: bool = False):
        self.file = str(file)
        self.allow_errors = allow_errors


def make_registry(path: Path, duplicate_link: bool = False, missing_link: bool = False) -> None:
    workbook = Workbook()
    requirements = workbook.active
    requirements.title = "Процессы требования"
    requirements.append([])
    requirements.append([])
    requirements.append(
        [
            "Код БП",
            "БП Уровень 1",
            "BP lvl1 for status",
            "Block",
            "БП Уровень 2",
            "БП Уровень 3",
            "Ключ",
            "Контроль БП",
            "ID CR",
            "IDтребования",
            "Требование",
        ]
    )
    requirements.append(["строку не удалять"])
    requirements.append(
        [
            "01.01",
            "Sales",
            "Orders",
            "Order processing",
            "",
            "",
            "",
            "",
            "",
            "SLS-001",
            "Create an order",
        ]
    )
    requirements.append(
        [
            "01.01",
            "Sales",
            "Orders",
            "Order processing",
            "",
            "",
            "",
            "",
            "",
            "SLS-002",
            "Approve an order",
        ]
    )
    specifications = workbook.create_sheet("Реестр ФС- не удалять")
    specifications.append(
        [
            "Система",
            "Релиз",
            "Код разработки  (RICEF)",
            "Название разработки",
            "Описание разработки",
            "Приоритет",
            "Плановый релиз",
            "Критичность разработки",
            "Статус ФС",
            "Код требования",
        ]
    )
    ids = "SLS-001, SLS-999" if missing_link else "SLS-001, SLS-002"
    specifications.append(
        ["1C", "R1", "G-001", "Order changes", "Description", 1, "R1", "High", "Draft", ids]
    )
    if duplicate_link:
        specifications.append(
            ["1C", "R1", "G-002", "Second", "Description", 2, "R1", "Low", "Draft", "SLS-001"]
        )
    workbook.save(path)


class RegistryImportTests(unittest.TestCase):

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.old_root = flow1c.ROOT
        flow1c.ROOT = self.root
        (self.root / ".flow1c.json").write_text(
            json.dumps(
                {
                    "version": 1,
                    "project": {"name": "test", "default_branch": "main"},
                    "registry": {
                        "requirements_sheet": "Процессы требования",
                        "requirements_header_row": 3,
                        "requirements_data_row": 5,
                        "requirement_id_column": "J",
                        "specifications_sheet": "Реестр ФС- не удалять",
                        "specifications_header_row": 1,
                        "specifications_data_row": 2,
                    },
                    "gitea": {
                        "base_url": "https://example.invalid",
                        "owner": "x",
                        "repository": "y",
                        "token_env": "TOKEN",
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

    def tearDown(self) -> None:
        flow1c.ROOT = self.old_root
        self.temp.cleanup()

    def test_valid_registry_is_normalized(self) -> None:
        source = self.root / "input.xlsx"
        make_registry(source)
        self.assertEqual(flow1c.cmd_registry_import(Args(source)), 0)
        requirements = json.loads(
            (self.root / "registry/normalized/requirements.json").read_text(encoding="utf-8")
        )
        specifications = json.loads(
            (self.root / "registry/normalized/specifications.json").read_text(encoding="utf-8")
        )
        self.assertEqual(sorted(requirements), ["SLS-001", "SLS-002"])
        self.assertEqual(specifications["G-001"]["requirements"], ["SLS-001", "SLS-002"])

    def test_read_only_import_streams_rows_without_random_cell_access(self) -> None:
        source = self.root / "input.xlsx"
        make_registry(source)
        with mock.patch(
            "openpyxl.worksheet._read_only.ReadOnlyWorksheet.cell",
            side_effect=AssertionError("random cell access is quadratic in read-only mode"),
        ):
            self.assertEqual(flow1c.cmd_registry_import(Args(source)), 0)

    def test_missing_requirement_creates_partial_import(self) -> None:
        source = self.root / "input.xlsx"
        make_registry(source, missing_link=True)
        output = io.StringIO()
        with redirect_stdout(output):
            self.assertEqual(flow1c.cmd_registry_import(Args(source)), 1)
        result = json.loads(output.getvalue())
        self.assertEqual(result["state"], "PARTIAL")
        self.assertTrue(result["report_path"].endswith("import-report.md"))
        self.assertTrue(result["errors"])
        self.assertEqual(read_json(self.root / "registry/status.json")["status"], "partial")
        self.assertFalse((self.root / "registry/normalized/requirements.json").exists())
        self.assertTrue((self.root / "registry/partial/scope-index.json").exists())

    def test_requirement_cannot_belong_to_two_specs(self) -> None:
        source = self.root / "input.xlsx"
        make_registry(source, duplicate_link=True)
        self.assertEqual(flow1c.cmd_registry_import(Args(source)), 1)
        report = (self.root / "registry/import-report.md").read_text(encoding="utf-8")
        self.assertIn("belongs to both G-001 and G-002", report)

    def test_partial_registry_resolves_good_scope_by_fs_and_requirement(self) -> None:
        source = self.root / "partial.xlsx"
        make_registry(source, missing_link=True)
        self.assertEqual(flow1c.cmd_registry_import(Args(source)), 1)
        by_fs = svc_registry.resolve_registry_reference("G-001", product_root=flow1c.ROOT)
        by_requirement = svc_registry.resolve_registry_reference(
            "SLS-001", product_root=flow1c.ROOT
        )
        self.assertEqual(by_fs["state"], "resolved")
        self.assertEqual(by_fs["requirements"], ["SLS-001"])
        self.assertEqual(by_requirement["code"], "G-001")
        self.assertEqual(by_requirement["source_status"], "partial")

    def test_partial_registry_starts_scoped_work_item_and_preserves_provenance(self) -> None:
        source = self.root / "partial.xlsx"
        make_registry(source, missing_link=True)
        self.assertEqual(flow1c.cmd_registry_import(Args(source)), 1)
        args = argparse.Namespace(
            task_reference="SLS-001",
            project_reference=None,
            code=None,
            g_number=None,
            reference_kind="auto",
            title=None,
            requirements=None,
            create_branch=False,
        )
        self.assertEqual(flow1c.cmd_fs_start(args), 0)
        manifest = read_json(self.root / "work-items/G-001/manifest.yaml")
        self.assertEqual(manifest["work_reference"], "G-001")
        self.assertEqual(manifest["requested_reference"], "SLS-001")
        self.assertEqual(manifest["requirements"], ["SLS-001"])
        self.assertEqual(manifest["registry"]["status"], "verified")
        self.assertEqual(manifest["registry"]["source_status"], "partial")

    def test_formal_gate_resolves_requirement_to_fs_before_work_item_creation(self) -> None:
        from routing_fixture import install_route_assets
        install_route_assets(self.root)
        source = self.root / "partial.xlsx"
        make_registry(source, missing_link=True)
        self.assertEqual(flow1c.cmd_registry_import(Args(source)), 1)
        args = argparse.Namespace(
            operation="functional-spec",
            mode="formal",
            summary="Подготовить ФС",
            code=None,
            task_reference="SLS-001",
            project_reference=None,
            reference_kind="auto",
            git_ref=None,
            g_number=None,
            path=[],
            allow_incomplete_draft=False,
            requirements=[],
            mismatch="",
            profile=None,
        )
        output = io.StringIO()
        stages = json.loads(
            (MODULE_PATH.parents[1] / "config/stages.json").read_text(encoding="utf-8")
        )
        with (
            mock.patch.object(svc_context, "load_stages", return_value=stages),
            redirect_stdout(output),
        ):
            self.assertEqual(flow1c.cmd_agent_begin(args), 2)
        gate = json.loads(output.getvalue())
        self.assertEqual(gate["work_reference"], "G-001")
        self.assertEqual(gate["requested_reference"], "SLS-001")
        self.assertEqual(gate["registry_resolution"]["state"], "resolved")
        self.assertIn("fs-start", gate["user_message"])

    def test_multiple_fs_assignments_and_kind_collision_are_ambiguous(self) -> None:
        source = self.root / "multiple.xlsx"
        make_registry(source, duplicate_link=True)
        self.assertEqual(flow1c.cmd_registry_import(Args(source)), 1)
        multiple = svc_registry.resolve_registry_reference("SLS-001", product_root=flow1c.ROOT)
        self.assertEqual(multiple["state"], "ambiguous")
        self.assertEqual(multiple["candidates"], ["G-001", "G-002"])
        index = read_json(self.root / "registry/partial/scope-index.json")
        index["requirements"]["G-001"] = [{"id": "G-001", "source_row": 99}]
        write_json(self.root / "registry/partial/scope-index.json", index)
        collision = svc_registry.resolve_registry_reference("G-001", product_root=flow1c.ROOT)
        self.assertEqual(collision["reason"], "reference_kind_collision")

    def test_partial_import_preserves_last_verified_indexes(self) -> None:
        clean = self.root / "clean.xlsx"
        make_registry(clean)
        self.assertEqual(flow1c.cmd_registry_import(Args(clean)), 0)
        before = (self.root / "registry/normalized/requirements.json").read_bytes()
        partial = self.root / "partial.xlsx"
        make_registry(partial, missing_link=True)
        self.assertEqual(flow1c.cmd_registry_import(Args(partial)), 1)
        self.assertEqual((self.root / "registry/normalized/requirements.json").read_bytes(), before)

    def test_fs_without_any_existing_requirement_is_not_usable(self) -> None:
        source = self.root / "missing-all.xlsx"
        make_registry(source, missing_link=True)
        from openpyxl import load_workbook

        workbook = load_workbook(source)
        workbook["Реестр ФС- не удалять"].cell(row=2, column=10, value="SLS-999")
        workbook.save(source)
        self.assertEqual(flow1c.cmd_registry_import(Args(source)), 1)
        self.assertEqual(
            svc_registry.resolve_registry_reference("G-001", product_root=flow1c.ROOT)["state"],
            "no_existing_requirements",
        )

    def test_missing_required_sheet_is_structurally_invalid(self) -> None:
        source = self.root / "missing-sheet.xlsx"
        workbook = Workbook()
        workbook.active.title = "Процессы требования"
        workbook.save(source)
        output = io.StringIO()
        with redirect_stdout(output):
            self.assertEqual(flow1c.cmd_registry_import(Args(source)), 2)
        self.assertEqual(json.loads(output.getvalue())["state"], "BLOCKED")
        self.assertEqual(read_json(self.root / "registry/status.json")["status"], "invalid")

    def test_registry_uses_external_documentation_root(self) -> None:
        documentation_root = self.root / "project-docs"
        documentation_root.mkdir()
        (self.root / ".flow1c.local.json").write_text(
            json.dumps({"documentation_path": str(documentation_root)}), encoding="utf-8"
        )
        source = self.root / "input.xlsx"
        make_registry(source)
        self.assertEqual(flow1c.cmd_registry_import(Args(source)), 0)
        self.assertTrue((documentation_root / "registry/source/requirements.xlsx").is_file())
        self.assertTrue((documentation_root / "registry/normalized/requirements.json").is_file())
        self.assertFalse((self.root / "registry/normalized/requirements.json").exists())

    def test_registry_reconcile_preserves_evidence_and_switches_to_registry(self) -> None:
        reference = ext_scripts_flow1c_policy.resolve_work_reference("USER-TASK-7", None)
        item = svc_work_items.create_provisional_work_item(
            reference,
            title="Order",
            user_brief="Implement SLS-001",
            deviation={"type": "registry_bypass", "reason": "user approved", "actor": "user"},
            requirements=["SLS-001"],
            product_root=flow1c.ROOT,
        )
        old_evidence = item / "evidence" / "old.json"
        write_json(
            old_evidence,
            {"schema_version": 1, "operation": "functional-spec", "code": "USER-TASK-7"},
        )
        source = self.root / "fixed.xlsx"
        make_registry(source)
        exit_code, result = svc_work_items.reconcile_registry(
            "USER-TASK-7", str(source), product_root=flow1c.ROOT
        )
        self.assertEqual(exit_code, 0)
        self.assertEqual(result["state"], "RECONCILED")
        manifest = read_json(item / "manifest.yaml")
        self.assertEqual(manifest["traceability_mode"], "registry")
        self.assertEqual(manifest["requirements"], ["SLS-001"])
        self.assertTrue(old_evidence.is_file())
        self.assertTrue(read_json(old_evidence)["revalidation_required"])

    def test_registry_reconcile_keeps_provisional_on_unmatched_material(self) -> None:
        reference = ext_scripts_flow1c_policy.resolve_work_reference("USER-TASK-8", None)
        item = svc_work_items.create_provisional_work_item(
            reference,
            title="Order",
            user_brief="Implement UNKNOWN-99",
            deviation={"type": "registry_bypass", "reason": "user approved", "actor": "user"},
            requirements=["UNKNOWN-99"],
            product_root=flow1c.ROOT,
        )
        source = self.root / "fixed.xlsx"
        make_registry(source)
        exit_code, result = svc_work_items.reconcile_registry(
            "USER-TASK-8", str(source), product_root=flow1c.ROOT
        )
        self.assertEqual(exit_code, 1)
        self.assertEqual(result["state"], "NEEDS_CONFIRMATION")
        self.assertEqual(read_json(item / "manifest.yaml")["traceability_mode"], "provisional")
        self.assertTrue(Path(result["report_path"]).is_file())


class SetupValidationTests(unittest.TestCase):

    def test_doctor_uses_external_tools_manifest_loader(self) -> None:
        external_manifest = {
            "rlm_tools_bsl": {"requirement": ">=1.34,<2"},
            "bsl_language_server": {"version": "1.0.7"},
        }
        with (
            mock.patch.object(
                ext_scripts_external_tools, "load_manifest", return_value=external_manifest
            ) as loader,
            mock.patch.object(svc_system, "command_path", return_value=""),
            mock.patch.object(ext_scripts_external_tools, "package_version", return_value=None),
            mock.patch.object(svc_context, "CONFIG_FILE", "missing-workflow.json"),
        ):
            output = io.StringIO()
            with redirect_stdout(output):
                flow1c.cmd_doctor(argparse.Namespace(json=True))
        loader.assert_called_once_with(flow1c.ROOT / "config" / "external-tools.json")
        payload = json.loads(output.getvalue())
        manifest_check = next(
            (check for check in payload["checks"] if check["name"] == "External tools manifest")
        )
        self.assertEqual(manifest_check["state"], "OK")

    def test_doctor_json_is_the_machine_readable_readiness_gate(self) -> None:
        output = io.StringIO()
        rows = [("Python", "OK", "3.12.0"), ("Gitea token", "SETUP", "FLOW1C_GITEA_TOKEN")]
        with redirect_stdout(output):
            exit_code = flow1c.emit_doctor_report(rows, json_output=True)
        payload = json.loads(output.getvalue())
        self.assertEqual(exit_code, 0)
        self.assertEqual(payload["state"], "READY")
        self.assertTrue(payload["ready"])
        self.assertEqual(payload["checks"][1]["state"], "SETUP")

    def test_doctor_json_blocks_ready_on_any_error(self) -> None:
        output = io.StringIO()
        with redirect_stdout(output):
            exit_code = flow1c.emit_doctor_report(
                [("RLM configuration index", "ERROR", "stale")], json_output=True
            )
        payload = json.loads(output.getvalue())
        self.assertEqual(exit_code, 1)
        self.assertEqual(payload["state"], "BLOCKED")
        self.assertFalse(payload["ready"])

    def test_git_url_normalization_matches_ssh_and_https(self) -> None:
        ssh_url = "ssh://git@git.example.ru:7822/team/project-docs.git"
        https_url = "https://git.example.ru/team/project-docs.git"
        self.assertEqual(
            svc_system.normalize_git_url(ssh_url), svc_system.normalize_git_url(https_url)
        )

    def test_path_in_another_workflow_clone_is_detected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            other = root / "other-workflow"
            (other / "scripts").mkdir(parents=True)
            (other / ".flow1c.json").write_text("{}", encoding="utf-8")
            (other / "scripts/flow1c.py").write_text("", encoding="utf-8")
            source = other / "external-source"
            source.mkdir()
            detected = svc_system.containing_workflow_root(source)
            self.assertIsNotNone(detected)
            self.assertTrue(os.path.samefile(detected, other))

    def test_local_extension_export_does_not_require_git(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            extension = Path(directory)
            (extension / "CommonModules").mkdir()
            (extension / "CommonModules/Test.bsl").write_text(
                "Процедура Тест() КонецПроцедуры", encoding="utf-8"
            )
            _, state, detail = svc_system.extension_source_state(
                {
                    "extension_mode": "local-export",
                    "extension_path": str(extension),
                    "extension_repository_url": "",
                }
            )
            self.assertEqual((state, detail), ("OK", "local XML/BSL export"))

    def test_nested_1c_source_root_is_resolved_from_checkout(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            checkout = Path(directory)
            source = checkout / "src" / "project_extension"
            source.mkdir(parents=True)
            (source / "Configuration.xml").write_text("<MetaDataObject/>", encoding="utf-8")
            self.assertEqual(svc_system.resolve_1c_source_root(checkout), source.resolve())

    def test_ambiguous_nested_1c_roots_keep_configured_path(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            checkout = Path(directory)
            for name in ("extension_a", "extension_b"):
                source = checkout / "src" / name
                source.mkdir(parents=True)
                (source / "Configuration.xml").write_text("<MetaDataObject/>", encoding="utf-8")
            self.assertEqual(svc_system.resolve_1c_source_root(checkout), checkout.resolve())

    def test_git_extension_requires_matching_origin(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            extension = Path(directory)
            (extension / "Configuration.xml").write_text("<MetaDataObject/>", encoding="utf-8")
            local = {
                "extension_mode": "git",
                "extension_path": str(extension),
                "extension_repository_url": "https://git.example.ru/team/project-extension.git",
            }
            with mock.patch.object(
                svc_system,
                "git_remote_url",
                return_value="ssh://git@git.example.ru:7822/team/project-extension.git",
            ):
                _, state, _ = svc_system.extension_source_state(local)
            self.assertEqual(state, "OK")

    def test_rlm_health_url_replaces_mcp_path(self) -> None:
        self.assertEqual(
            svc_sources.rlm_health_url("http://127.0.0.1:9000/mcp"), "http://127.0.0.1:9000/health"
        )

    def test_redact_url_credentials_removes_embedded_secret(self) -> None:
        self.assertEqual(
            svc_system.redact_url_credentials(
                "https://user:secret@git.example.ru/team/project.git"
            ),
            "https://git.example.ru/team/project.git",
        )

    def test_rlm_endpoint_index_requires_server_side_index(self) -> None:
        with mock.patch.object(
            svc_sources, "rlm_mcp_call", return_value=({"error": "Index not found"}, "")
        ):
            self.assertEqual(
                svc_sources.rlm_endpoint_index_state("http://127.0.0.1:9000/mcp", Path("source")),
                ("ERROR", "Index not found"),
            )

    def test_rlm_endpoint_session_executes_and_closes_smoke_session(self) -> None:
        with mock.patch.object(
            svc_sources,
            "rlm_mcp_call",
            side_effect=[
                ({"session_id": "abc123"}, ""),
                ({"stdout": "FLOW1C_WORKFLOW_RLM_OK\n", "error": None}, ""),
                ({"success": True}, ""),
            ],
        ) as call:
            self.assertEqual(
                svc_sources.rlm_endpoint_session_state("http://127.0.0.1:9000/mcp", Path("source")),
                ("OK", "sandbox execution succeeded and session closed"),
            )
            self.assertEqual(
                [item.args[1] for item in call.call_args_list],
                ["rlm_start", "rlm_execute", "rlm_end"],
            )
            self.assertIn("FLOW1C_WORKFLOW_RLM_OK", call.call_args_list[1].args[2]["code"])

    def test_rlm_session_closes_after_execute_error(self) -> None:
        with mock.patch.object(
            svc_sources,
            "rlm_mcp_call",
            side_effect=[
                ({"session_id": "abc123"}, ""),
                ({"stdout": "", "error": "invalid sandbox code"}, ""),
                ({"success": True}, ""),
            ],
        ) as call:
            result, error = svc_sources.rlm_session_execute(
                "http://127.0.0.1:9000/mcp",
                Path("source"),
                "Find an object",
                "broken()",
                effort="low",
                max_output_chars=1000,
                timeout=10.0,
                domains=["весь каталог"],
            )
        self.assertIsNone(result)
        self.assertEqual(error, "invalid sandbox code")
        self.assertEqual(
            [item.args[1] for item in call.call_args_list], ["rlm_start", "rlm_execute", "rlm_end"]
        )

    def test_rlm_session_retries_with_domains_when_server_requires_selection(self) -> None:
        with mock.patch.object(
            svc_sources,
            "rlm_mcp_call",
            side_effect=[
                ({"error": "Для BSL нужен явный выбор domains"}, ""),
                ({"session_id": "abc123"}, ""),
                ({"stdout": "FLOW1C_WORKFLOW_RLM_OK\n", "error": None}, ""),
                ({"success": True}, ""),
            ],
        ) as call:
            self.assertEqual(
                svc_sources.rlm_endpoint_session_state("http://127.0.0.1:9000/mcp", Path("source")),
                ("OK", "sandbox execution succeeded and session closed"),
            )
        self.assertEqual(
            [item.args[1] for item in call.call_args_list],
            ["rlm_start", "rlm_start", "rlm_execute", "rlm_end"],
        )
        self.assertNotIn("domains", call.call_args_list[0].args[2])
        self.assertEqual(call.call_args_list[1].args[2]["domains"], [])

    def test_rlm_index_requires_fresh_status(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            command = Path(directory) / "rlm-bsl-index.exe"
            command.write_bytes(b"")
            source = Path(directory) / "source"
            source.mkdir()
            completed = subprocess.CompletedProcess(
                args=[], returncode=0, stdout="Status:   fresh\n", stderr=""
            )
            with mock.patch.object(subprocess, "run", return_value=completed):
                self.assertEqual(
                    svc_sources.rlm_index_state(str(command), source, product_root=flow1c.ROOT),
                    ("OK", "fresh"),
                )


class NaturalLanguageGateTests(unittest.TestCase):

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.old_root = flow1c.ROOT
        flow1c.ROOT = self.root
        (self.root / "config").mkdir()
        shutil.copy2(self.old_root / "config" / "stages.json", self.root / "config" / "stages.json")
        from routing_fixture import install_route_assets
        install_route_assets(self.root)
        (self.root / ".agents" / "skills" / "flow1c-functional-spec").mkdir(parents=True, exist_ok=True)
        (self.root / ".agents" / "skills" / "flow1c-functional-spec" / "SKILL.md").write_text(
            "---\nname: flow1c-functional-spec\ndescription: test\n---\nTest skill\n",
            encoding="utf-8",
        )
        (self.root / ".flow1c.json").write_text(
            json.dumps(
                {
                    "version": 1,
                    "project": {"default_branch": "main"},
                    "quality": {"rlm_endpoint": "http://127.0.0.1:9000/mcp"},
                }
            ),
            encoding="utf-8",
        )
        (self.root / ".flow1c.local.json").write_text(
            json.dumps(
                {
                    "documentation_path": str(self.root),
                    "functional_spec_template": str(self.root / "template.docx"),
                }
            ),
            encoding="utf-8",
        )
        (self.root / "template.docx").write_bytes(b"template")

    def tearDown(self) -> None:
        flow1c.ROOT = self.old_root
        self.temp.cleanup()

    def make_work_item(self, *, status: str = "clarification") -> Path:
        item = self.root / "work-items" / "G-001"
        for relative in ("input", "analysis", "specification", "implementation", "evidence"):
            (item / relative).mkdir(parents=True, exist_ok=True)
        (item / "manifest.yaml").write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "code": "G-001",
                    "title": "Test",
                    "status": status,
                    "requirements": ["REQ-001"],
                    "branches": {
                        "documentation": "fs/G-001/specification",
                        "extension": "feature/G-001",
                    },
                    "approvals": {
                        "functional_architect": "pending",
                        "technical_architect": "pending",
                    },
                }
            ),
            encoding="utf-8",
        )
        (item / "input" / "requirements.snapshot.yaml").write_text(
            json.dumps({"REQ-001": {"text": "Requirement"}}), encoding="utf-8"
        )
        return item

    def test_code_is_requested_and_never_invented(self) -> None:
        output = io.StringIO()
        with redirect_stdout(output):
            exit_code = flow1c.cmd_agent_begin(
                type(
                    "A",
                    (),
                    {
                        "operation": "functional-spec",
                        "code": None,
                        "summary": "Напиши ФС",
                        "allow_incomplete_draft": False,
                    },
                )()
            )
        payload = json.loads(output.getvalue())
        self.assertEqual(exit_code, 2)
        self.assertEqual(payload["state"], "BLOCKED")
        self.assertTrue(payload["awaiting_user_input"])
        self.assertIsNone(payload["code"])

    def test_ambiguous_request_gets_one_question_gate(self) -> None:
        output = io.StringIO()
        with redirect_stdout(output):
            exit_code = flow1c.cmd_agent_begin(
                type(
                    "A",
                    (),
                    {
                        "operation": "ambiguous",
                        "code": None,
                        "summary": "Проверь и сделай",
                        "path": [],
                        "allow_incomplete_draft": False,
                    },
                )()
            )
        payload = json.loads(output.getvalue())
        self.assertEqual(exit_code, 1)
        self.assertEqual(payload["state"], "NEEDS_CONFIRMATION")
        self.assertIn("один вопрос", payload["user_message"])

    def test_missing_artifacts_return_absolute_destinations(self) -> None:
        self.make_work_item()
        output = io.StringIO()
        with (
            mock.patch.object(svc_sources, "rlm_readiness", return_value=(True, [])),
            redirect_stdout(output),
        ):
            exit_code = flow1c.cmd_agent_begin(
                type(
                    "A",
                    (),
                    {
                        "operation": "functional-spec",
                        "code": "G-001",
                        "summary": "Напиши ФС",
                        "allow_incomplete_draft": False,
                    },
                )()
            )
        payload = json.loads(output.getvalue())
        self.assertEqual(exit_code, 1)
        self.assertEqual(payload["state"], "NEEDS_INPUT")
        self.assertTrue(
            all((Path(item["path"]).is_absolute() for item in payload["missing_conditional"]))
        )
        self.assertIn("приложить файлы в чат", payload["user_message"])

    def test_formal_functional_spec_still_requests_missing_docx_template(self) -> None:
        self.make_work_item()
        (self.root / ".flow1c.local.json").write_text(
            json.dumps({"documentation_path": str(self.root), "functional_spec_template": ""}),
            encoding="utf-8",
        )
        output = io.StringIO()
        with (
            mock.patch.object(svc_sources, "rlm_readiness", return_value=(True, [])),
            redirect_stdout(output),
        ):
            exit_code = flow1c.cmd_agent_begin(
                type(
                    "A",
                    (),
                    {
                        "operation": "functional-spec",
                        "code": "G-001",
                        "summary": "Напиши ФС",
                        "allow_incomplete_draft": False,
                    },
                )()
            )
        payload = json.loads(output.getvalue())
        self.assertEqual(exit_code, 1)
        self.assertEqual(payload["state"], "NEEDS_INPUT")
        self.assertIn(
            "functional_spec_template", {item["id"] for item in payload["missing_required"]}
        )
        self.assertIn("независимым черновиком", payload["user_message"])

    def test_intake_copies_supported_file_and_records_hash(self) -> None:
        item = self.make_work_item()
        gate = svc_workflow_state.new_gate(
            "functional-spec", "G-001", "NEEDS_INPUT", product_root=flow1c.ROOT
        )
        source = self.root / "meeting.txt"
        source.write_text("Meeting evidence", encoding="utf-8")
        output = io.StringIO()
        args = type(
            "A",
            (),
            {
                "gate_id": gate["gate_id"],
                "category": "meeting_materials",
                "code": None,
                "source": [str(source)],
                "confirm_absence": [],
                "confirm_large": False,
            },
        )()
        with redirect_stdout(output):
            self.assertEqual(flow1c.cmd_artifact_intake(args), 0)
        payload = json.loads(output.getvalue())
        self.assertEqual(payload["state"], "ACCEPTED")
        index = json.loads((item / "input" / "artifacts.json").read_text(encoding="utf-8"))
        self.assertEqual(index["artifacts"][0]["sha256"], sha256(source))
        self.assertFalse(Path(index["artifacts"][0]["relative_path"]).is_absolute())

    def test_intake_rejects_unsupported_file(self) -> None:
        self.make_work_item()
        gate = svc_workflow_state.new_gate(
            "functional-spec", "G-001", "NEEDS_INPUT", product_root=flow1c.ROOT
        )
        source = self.root / "payload.exe"
        source.write_bytes(b"not executable")
        args = type(
            "A",
            (),
            {
                "gate_id": gate["gate_id"],
                "category": "meeting_materials",
                "code": None,
                "source": [str(source)],
                "confirm_absence": [],
                "confirm_large": False,
            },
        )()
        with self.assertRaises(WorkflowError):
            flow1c.cmd_artifact_intake(args)

    def test_confirmed_absence_allows_functional_spec_gate(self) -> None:
        self.make_work_item()
        gate = svc_workflow_state.new_gate(
            "functional-spec", "G-001", "NEEDS_INPUT", product_root=flow1c.ROOT
        )
        categories = ["meeting_materials", "correspondence", "current_process", "document_examples"]
        args = type(
            "A",
            (),
            {
                "gate_id": gate["gate_id"],
                "category": "meeting_materials",
                "code": None,
                "source": [],
                "confirm_absence": categories,
                "confirm_large": False,
            },
        )()
        with redirect_stdout(io.StringIO()):
            self.assertEqual(flow1c.cmd_artifact_intake(args), 0)
        output = io.StringIO()
        with (
            mock.patch.object(svc_sources, "rlm_readiness", return_value=(True, [])),
            redirect_stdout(output),
        ):
            exit_code = flow1c.cmd_agent_begin(
                type(
                    "A",
                    (),
                    {
                        "operation": "functional-spec",
                        "code": "G-001",
                        "summary": "Напиши ФС",
                        "path": [],
                        "allow_incomplete_draft": False,
                    },
                )()
            )
        payload = json.loads(output.getvalue())
        self.assertEqual(exit_code, 0)
        self.assertEqual(payload["state"], "READY")

    def test_registry_intake_without_code_is_reused_from_inbox(self) -> None:
        begin_output = io.StringIO()
        with redirect_stdout(begin_output):
            self.assertEqual(
                flow1c.cmd_agent_begin(
                    type(
                        "A",
                        (),
                        {
                            "operation": "registry",
                            "code": None,
                            "summary": "Импортируй реестр",
                            "path": [],
                            "allow_incomplete_draft": False,
                        },
                    )()
                ),
                1,
            )
        gate = json.loads(begin_output.getvalue())
        source = self.root / "requirements.xlsx"
        source.write_bytes(b"test workbook")
        with redirect_stdout(io.StringIO()):
            self.assertEqual(
                flow1c.cmd_artifact_intake(
                    type(
                        "A",
                        (),
                        {
                            "gate_id": gate["gate_id"],
                            "category": "requirements_workbook",
                            "code": None,
                            "source": [str(source)],
                            "confirm_absence": [],
                            "confirm_large": False,
                        },
                    )()
                ),
                0,
            )
        second_output = io.StringIO()
        with redirect_stdout(second_output):
            self.assertEqual(
                flow1c.cmd_agent_begin(
                    type(
                        "A",
                        (),
                        {
                            "operation": "registry",
                            "code": None,
                            "summary": "Импортируй реестр",
                            "path": [],
                            "allow_incomplete_draft": False,
                        },
                    )()
                ),
                0,
            )
        self.assertEqual(json.loads(second_output.getvalue())["state"], "READY")

    def test_large_intake_requires_explicit_confirmation(self) -> None:
        self.make_work_item()
        gate = svc_workflow_state.new_gate(
            "functional-spec", "G-001", "NEEDS_INPUT", product_root=flow1c.ROOT
        )
        folder = self.root / "batch"
        folder.mkdir()
        (folder / "a.txt").write_text("a", encoding="utf-8")
        (folder / "b.txt").write_text("b", encoding="utf-8")
        args = type(
            "A",
            (),
            {
                "gate_id": gate["gate_id"],
                "category": "meeting_materials",
                "code": None,
                "source": [str(folder)],
                "confirm_absence": [],
                "confirm_large": False,
            },
        )()
        with mock.patch.object(svc_intake, "MAX_INTAKE_FILES", 1), self.assertRaises(WorkflowError):
            flow1c.cmd_artifact_intake(args)

    def test_inbox_package_is_promoted_after_code_assignment(self) -> None:
        self.make_work_item()
        gate = svc_workflow_state.new_gate(
            "functional-spec", None, "NEEDS_CODE", product_root=flow1c.ROOT
        )
        source = self.root / "meeting.txt"
        source.write_text("meeting", encoding="utf-8")
        first_output = io.StringIO()
        with redirect_stdout(first_output):
            self.assertEqual(
                flow1c.cmd_artifact_intake(
                    type(
                        "A",
                        (),
                        {
                            "gate_id": gate["gate_id"],
                            "category": "meeting_materials",
                            "code": None,
                            "source": [str(source)],
                            "confirm_absence": [],
                            "confirm_large": False,
                            "received_via": "file",
                            "promote_intake_id": None,
                        },
                    )()
                ),
                0,
            )
        intake_id = json.loads(first_output.getvalue())["intake_id"]
        promoted_output = io.StringIO()
        with redirect_stdout(promoted_output):
            self.assertEqual(
                flow1c.cmd_artifact_intake(
                    type(
                        "A",
                        (),
                        {
                            "gate_id": gate["gate_id"],
                            "category": "meeting_materials",
                            "code": "G-001",
                            "source": [],
                            "confirm_absence": [],
                            "confirm_large": False,
                            "received_via": "inbox-promotion",
                            "promote_intake_id": intake_id,
                        },
                    )()
                ),
                0,
            )
        payload = json.loads(promoted_output.getvalue())
        self.assertIn("work-items", payload["copied"][0]["relative_path"])
        self.assertFalse((self.root / "inbox" / intake_id / "originals").exists())

    def test_verified_claim_requires_evidence_id_in_output(self) -> None:
        item = self.make_work_item(status="functional_review")
        (item / "reviews").mkdir()
        output_path = item / "reviews" / "functional-review.md"
        output_path.write_text(
            "# Review\n\nТребования проверены, замечаний нет.\n", encoding="utf-8"
        )
        evidence_path = item / "evidence" / "functional-review.json"
        evidence_path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "operation": "functional-review",
                    "code": "G-001",
                    "context_sha256": "a" * 64,
                    "context_evidence_id": "CTX-aaaaaaaaaaaa",
                    "artifacts": [],
                    "rlm_queries": [],
                    "changed_files": [],
                }
            ),
            encoding="utf-8",
        )
        gate = svc_workflow_state.new_gate(
            "functional-review",
            "G-001",
            "READY",
            evidence_path=str(evidence_path),
            output="reviews/functional-review.md",
            product_root=flow1c.ROOT,
        )
        result = io.StringIO()
        with redirect_stdout(result):
            self.assertEqual(
                flow1c.cmd_agent_complete(
                    type("A", (), {"gate_id": gate["gate_id"], "output": None})()
                ),
                2,
            )
        self.assertEqual(json.loads(result.getvalue())["state"], "NON_COMPLIANT")

    def test_invalid_gate_id_is_rejected(self) -> None:
        with self.assertRaises(WorkflowError):
            svc_workflow_state.load_gate("not-a-real-gate", product_root=flow1c.ROOT)

    def test_agent_command_accepts_json_stdin_and_emits_json(self) -> None:
        request = io.StringIO(
            json.dumps(
                {
                    "operation": "functional-spec",
                    "code": None,
                    "summary": "Напиши ФС",
                    "path": [str(self.root / "material.pdf")],
                    "allow_incomplete_draft": False,
                }
            )
        )
        output = io.StringIO()
        with mock.patch.object(sys, "stdin", request), redirect_stdout(output):
            self.assertEqual(flow1c.main(["agent-begin", "--json-stdin"]), 2)
        payload = json.loads(output.getvalue())
        self.assertEqual(payload["state"], "BLOCKED")
        self.assertEqual(payload["presented_paths"], [str(self.root / "material.pdf")])

    def test_stage_rejects_tool_not_in_allowlist(self) -> None:
        gate = svc_workflow_state.new_gate("status", None, "READY", product_root=flow1c.ROOT)
        with self.assertRaises(WorkflowError):
            svc_workflow_state.require_gate_tool(gate, "flow1c_write", product_root=flow1c.ROOT)

    def test_source_query_records_execute_result_not_session_context(self) -> None:
        source = self.root / "extension"
        source.mkdir()
        local = json.loads((self.root / ".flow1c.local.json").read_text(encoding="utf-8"))
        local["extension_path"] = str(source)
        (self.root / ".flow1c.local.json").write_text(json.dumps(local), encoding="utf-8")
        request_id = "a" * 8
        evidence_path = self.root / ".workspace" / "drafts" / request_id / "evidence.json"
        gate = svc_workflow_state.new_gate(
            "consultation",
            None,
            "READY",
            mode="explore",
            request_id=request_id,
            storage_kind="workspace",
            evidence_path=str(evidence_path),
            product_root=flow1c.ROOT,
        )
        write_json(
            evidence_path,
            {
                "schema_version": 1,
                "gate_id": gate["gate_id"],
                "operation": "consultation",
                "rlm_queries": [],
            },
        )
        execute_result = {"stdout": '{"objects": ["Document.Test"]}\n', "error": None}
        args = type(
            "A",
            (),
            {
                "gate_id": gate["gate_id"],
                "source": "extension",
                "query": "Find the test document",
                "code": "print(search_objects('Test'))",
                "reason": "Regression test",
                "effort": "medium",
                "max_chars": 12000,
            },
        )()
        output = io.StringIO()
        with (
            mock.patch.object(svc_sources, "rlm_readiness", return_value=(True, [])),
            mock.patch.object(
                svc_sources, "rlm_session_execute", return_value=(execute_result, "")
            ) as execute,
            redirect_stdout(output),
        ):
            self.assertEqual(flow1c.cmd_source_query(args), 0)
        execute.assert_called_once_with(
            "http://127.0.0.1:9000/mcp",
            source.resolve(),
            "Find the test document",
            "print(search_objects('Test'))",
            effort="medium",
            max_output_chars=12000,
            timeout=120.0,
            domains=["весь каталог"],
        )
        payload = json.loads(output.getvalue())
        evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
        self.assertEqual(payload["result"], execute_result)
        self.assertEqual(evidence["rlm_queries"][0]["result"], execute_result)
        self.assertEqual(evidence["rlm_queries"][0]["code"], "print(search_objects('Test'))")
        self.assertNotIn("session_id", evidence["rlm_queries"][0]["result"])

    def test_source_query_rejects_missing_execution_code(self) -> None:
        source = self.root / "extension"
        source.mkdir()
        local = json.loads((self.root / ".flow1c.local.json").read_text(encoding="utf-8"))
        local["extension_path"] = str(source)
        (self.root / ".flow1c.local.json").write_text(json.dumps(local), encoding="utf-8")
        gate = svc_workflow_state.new_gate(
            "consultation", None, "READY", mode="explore", product_root=flow1c.ROOT
        )
        args = type(
            "A",
            (),
            {
                "gate_id": gate["gate_id"],
                "source": "extension",
                "query": "Find the test document",
                "code": "",
                "reason": "Regression test",
                "effort": "medium",
                "max_chars": 12000,
            },
        )()
        with (
            mock.patch.object(svc_sources, "rlm_readiness", return_value=(True, [])),
            self.assertRaisesRegex(WorkflowError, "execution code is required"),
        ):
            flow1c.cmd_source_query(args)

    def test_git_read_at_ref_cache_distinguishes_start_line_and_limit(self) -> None:
        request_id = "b" * 8
        evidence_path = self.root / ".workspace" / "drafts" / request_id / "evidence.json"
        gate = svc_workflow_state.new_gate(
            "consultation",
            None,
            "READY",
            mode="explore",
            request_id=request_id,
            storage_kind="workspace",
            evidence_path=str(evidence_path),
            product_root=flow1c.ROOT,
        )
        write_json(
            evidence_path,
            {
                "schema_version": 2,
                "gate_id": gate["gate_id"],
                "operation": "consultation",
                "git_analysis": [],
            },
        )
        args = argparse.Namespace(
            gate_id=gate["gate_id"],
            repository="extension",
            action="read-at-ref",
            git_ref="main",
            git_refs=[],
            target_ref="main",
            base_ref="",
            detail="patch",
            paths=[],
            path="Module.bsl",
            cursor="",
            max_files=100,
            start_line=5,
            subject_query="",
            regex=False,
            merges_only=False,
            first_parent=False,
            min_parents=None,
            max_parents=None,
            since="",
            until="",
            include_pr_evidence=True,
            include_patch_evidence=True,
            max_count=20,
            max_chars=1000,
        )

        def read(_repository, _ref, _path, *, start_line, max_chars):
            return {"state": "READ", "content": f"{start_line}:{max_chars}"}

        with (
            mock.patch.object(
                svc_workflow_git_actions, "git_inspection_repository", return_value=self.root
            ),
            mock.patch.object(git_analysis, "read_at_ref", side_effect=read) as reader,
        ):
            results = []
            for start_line, max_chars in ((5, 1000), (10, 1000), (10, 2000), (10, 2000)):
                args.start_line, args.max_chars = (start_line, max_chars)
                output = io.StringIO()
                with redirect_stdout(output):
                    self.assertEqual(flow1c.cmd_agent_git_inspect_v2(args), 0)
                results.append(json.loads(output.getvalue()))
        self.assertEqual(
            [item["content"] for item in results], ["5:1000", "10:1000", "10:2000", "10:2000"]
        )
        self.assertEqual([item["cache_hit"] for item in results], [False, False, False, True])
        self.assertEqual(reader.call_count, 3)

    def test_review_gate_cannot_write_extension(self) -> None:
        self.make_work_item(status="functional_review")
        evidence = self.root / "work-items/G-001/evidence/functional-review.json"
        evidence.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "operation": "functional-review",
                    "code": "G-001",
                    "artifacts": [],
                    "rlm_queries": [],
                    "changed_files": [],
                }
            ),
            encoding="utf-8",
        )
        gate = svc_workflow_state.new_gate(
            "functional-review",
            "G-001",
            "READY",
            evidence_path=str(evidence),
            product_root=flow1c.ROOT,
        )
        args = type(
            "A",
            (),
            {
                "gate_id": gate["gate_id"],
                "target": "extension",
                "path": "x.bsl",
                "content": "",
                "content_stdin": False,
            },
        )()
        with self.assertRaises(WorkflowError):
            flow1c.cmd_agent_write(args)


class FunctionalSectionCommandContractTests(unittest.TestCase):

    def test_section_commands_accept_agent_gate_field(self) -> None:
        parser = flow1c.build_parser()
        for command in (
            "section-catalog",
            "section-save",
            "section-approve",
            "docx-inspect",
            "docx-write-plan",
            "docx-write",
        ):
            args = parser.parse_args([command, "--gate-id", "gate-123"])
            self.assertEqual(args.gate_id, "gate-123", command)

    def test_plan_path_is_limited_to_managed_plan_directory(self) -> None:
        root = MODULE_PATH.parents[1] / ".workspace"
        allowed = root / "docx-plans" / "plan.json"
        self.assertEqual(svc_documents._plan_path(str(allowed), root), allowed.resolve())
        with self.assertRaises(ext_scripts_flow1c_docx.DocxError) as error:
            svc_documents._plan_path(str(root / "unmanaged.json"), root)
        self.assertEqual(error.exception.code, "DOCX_PATH_FORBIDDEN")


if __name__ == "__main__":
    unittest.main()
