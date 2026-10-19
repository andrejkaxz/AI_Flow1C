from __future__ import annotations

import io
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from scripts.flow1c_templates import TemplateService
from scripts.flow1c_templates_policy import TemplateError
from scripts.flow1c_templates_store import atomic_json, new_id, read_json, sha256
from test_document_templates import interpreted_profile, synthetic_fixture

ROOT = Path(__file__).resolve().parents[1]


def replace_document_text(data: bytes, format_name: str, old: str, new: str) -> bytes:
    if format_name == "markdown":
        assert old.encode() in data
        return data.replace(old.encode(), new.encode(), 1)
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(result, "w") as output:
        for info in source.infolist():
            raw = source.read(info.filename)
            if info.filename == "word/document.xml":
                assert old.encode() in raw
                raw = raw.replace(old.encode(), new.encode(), 1)
            output.writestr(info, raw)
    return result.getvalue()


class DocumentResultVerificationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.base = Path(self.temporary.name)
        self.docs = self.base / "docs"
        self.docs.mkdir()
        self.service = TemplateService(ROOT, {"documentation_path": str(self.docs)})

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def prepared(self, format_name: str) -> tuple[Path, dict, bytes]:
        source = synthetic_fixture(self.base, "meeting-minutes", format_name, False)
        received = self.service.intake({"source": str(source), "document_type": "meeting-minutes"})
        candidate = self.service.store.operation_path(received["operation_id"]) / "candidate"
        extraction = read_json(candidate / "extraction.json")
        profile, operations = interpreted_profile(extraction, "meeting-minutes")
        self.service.profile_save({"operation_id": received["operation_id"], "profile": profile})
        active = self.service.activate({"operation_id": received["operation_id"], "expected_revision": None})
        root = self.docs / "drafts" / new_id()
        self.service.resolve({"document_type": "meeting-minutes", "template_id": active["template_id"]}, root)
        plan = self.service.document_plan({"operations": operations}, root)
        if format_name == "markdown":
            from scripts.flow1c_markdown import write_bytes
        else:
            from scripts.flow1c_docx import fill_template as write_bytes
        return root, plan, write_bytes(source, extraction, operations)

    def test_readable_wrong_writer_result_is_rejected_before_publication(self) -> None:
        for format_name in ("markdown", "docx"):
            for old, new in (("Unchanged brand", "Changed brand"),
                             ("Actual user supplied content", "Missing planned value")):
                with self.subTest(format=format_name, mutation=old):
                    root, plan, correct = self.prepared(format_name)
                    wrong = replace_document_text(correct, format_name, old, new)
                    adapter = ("scripts.flow1c_markdown.write_bytes" if format_name == "markdown"
                               else "scripts.flow1c_docx.fill_template")
                    with mock.patch(adapter, return_value=wrong):
                        with self.assertRaises(TemplateError) as error:
                            self.service.document_write({"plan_id": plan["plan_id"]}, root)
                    self.assertEqual(error.exception.code, "TEMPLATE_INTEGRITY_FAILED")
                    self.assertFalse(Path(plan["output"]).exists())
                    self.assertFalse((root / (plan["plan_id"] + ".validation.json")).exists())
                    preserved = Path(error.exception.preserved_state["candidate"])
                    self.assertEqual(preserved.read_bytes(), wrong)
                    result = self.service.document_write({"plan_id": plan["plan_id"]}, root)
                    self.assertTrue(result["ready"])
                    self.assertEqual(result["result_verification"]["status"], "PASSED")

    def test_legacy_result_with_matching_hash_still_checks_actual_content(self) -> None:
        for format_name in ("markdown", "docx"):
            with self.subTest(format=format_name):
                root, plan, _ = self.prepared(format_name)
                result = self.service.document_write({"plan_id": plan["plan_id"]}, root)
                output = Path(result["output"])
                wrong = replace_document_text(output.read_bytes(), format_name,
                                              "Actual user supplied content", "Wrong content")
                output.write_bytes(wrong)
                result.pop("result_verification", None)
                result["output_sha256"] = sha256(output)
                validation = root / (plan["plan_id"] + ".validation.json")
                atomic_json(validation, result)
                record_before = validation.read_bytes()
                with self.assertRaises(TemplateError):
                    self.service.document_validate({"plan_id": plan["plan_id"]}, root)
                self.assertEqual(output.read_bytes(), wrong)
                self.assertEqual(validation.read_bytes(), record_before)

    def test_pending_resume_verifies_before_committing_success_record(self) -> None:
        for format_name in ("markdown", "docx"):
            with self.subTest(format=format_name):
                root, plan, _ = self.prepared(format_name)
                result = self.service.document_write({"plan_id": plan["plan_id"]}, root)
                output = Path(result["output"])
                output.write_bytes(replace_document_text(output.read_bytes(), format_name,
                                                        "Unchanged brand", "Changed brand"))
                result.pop("result_verification", None)
                result["output_sha256"] = sha256(output)
                plan_path = root / "document-plans" / (plan["plan_id"] + ".json")
                saved = read_json(plan_path)
                saved.update(state="WRITE_PENDING", pending_result=result)
                atomic_json(plan_path, saved)
                validation = root / (plan["plan_id"] + ".validation.json")
                validation.unlink()
                with self.assertRaises(TemplateError):
                    self.service.document_write({"plan_id": plan["plan_id"]}, root)
                self.assertFalse(validation.exists())
                self.assertEqual(read_json(plan_path)["state"], "WRITE_PENDING")

    def test_verification_timeout_preserves_candidate_and_allows_retry(self) -> None:
        root, plan, _ = self.prepared("markdown")
        with mock.patch("scripts.flow1c_templates_verify.subprocess.run", side_effect=subprocess.TimeoutExpired("verify", 15)):
            with self.assertRaises(TemplateError) as error:
                self.service.document_write({"plan_id": plan["plan_id"]}, root)
        self.assertTrue(Path(error.exception.preserved_state["candidate"]).is_file())
        self.assertFalse(Path(plan["output"]).exists())
        self.assertTrue(self.service.document_write({"plan_id": plan["plan_id"]}, root)["ready"])

    def test_correct_legacy_record_is_checked_without_rewriting_it(self) -> None:
        root, plan, _ = self.prepared("markdown")
        result = self.service.document_write({"plan_id": plan["plan_id"]}, root)
        validation = root / (plan["plan_id"] + ".validation.json")
        result.pop("result_verification")
        atomic_json(validation, result)
        before = validation.read_bytes()
        checked = self.service.document_validate({"plan_id": plan["plan_id"]}, root)
        self.assertEqual(checked["result_verification"]["status"], "PASSED")
        self.assertEqual(validation.read_bytes(), before)

    def test_verifier_protocol_failure_never_becomes_success(self) -> None:
        from scripts.flow1c_templates_verify import verify
        root, plan, data = self.prepared("markdown")
        output = Path(plan["output"])
        output.write_bytes(data)
        saved = read_json(root / "document-plans" / (plan["plan_id"] + ".json"))
        revision, manifest = self.service.store.revision(saved["pin"])
        source = revision / manifest["source"]
        for payload in (b"[]", b"not json", b'{"status":"PASSED","schema_version":1}', b'{"error":[]}'):
            with self.subTest(payload=payload):
                process = subprocess.CompletedProcess([], 0, stdout=payload, stderr=b"")
                with mock.patch("scripts.flow1c_templates_verify.subprocess.run", return_value=process):
                    with self.assertRaises(TemplateError):
                        verify(source, output, revision / "extraction.json", saved["operations"])


class FormatReadbackTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.base = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def verify_output(self, source: Path, extraction: dict, operations: list, data: bytes) -> dict:
        from scripts.flow1c_templates_verify import verify_local
        accepted = self.base / "extraction.json"
        atomic_json(accepted, extraction)
        output = self.base / ("result" + source.suffix)
        output.write_bytes(data)
        # Verification must not ask either writer to reproduce an expected document.
        with mock.patch("scripts.flow1c_markdown.write_bytes", side_effect=AssertionError("writer called")), \
             mock.patch("scripts.flow1c_docx.fill_template", side_effect=AssertionError("writer called")):
            return verify_local(source, output, accepted, operations)

    def test_markdown_table_expansion_and_positions_preserve_fixed_content(self) -> None:
        from scripts.flow1c_markdown import write_bytes
        from scripts.flow1c_templates_extract import extract
        source = self.base / "table.md"
        source.write_text("# Repeated\n\n| Name | Value |\n| --- | --- |\n| old | old |\n\n"
                          "# Repeated\n\n{{tail}}\n\n```text\n{{fixed-code}}\n```\n", encoding="utf-8")
        extraction = extract(source, "markdown")
        table = next(t for t in extraction["targets"] if t["kind"] == "table")
        field = next(t for t in extraction["targets"] if t["kind"] == "field")
        operations = [{"target_id": table["id"], "kind": "table", "value": [["A|B", "1"], ["C", "2"]]},
                      {"target_id": field["id"], "kind": "text", "value": "Actual [value]"}]
        data = write_bytes(source, extraction, operations)
        self.assertEqual(self.verify_output(source, extraction, operations, data)["operation_count"], 2)
        for wrong in (data.replace(b"Name | Value", b"Renamed | Value"),
                      data.replace(b"| C | 2 |", b"| C | 3 |"),
                      data.replace(b"{{fixed-code}}", b"changed-code")):
            with self.subTest(mutation=wrong[-30:]):
                with self.assertRaises(TemplateError):
                    self.verify_output(source, extraction, operations, wrong)

    def test_docx_split_fields_table_expansion_and_later_field_use_source_addresses(self) -> None:
        from docx import Document
        from scripts.flow1c_docx import fill_template
        from scripts.flow1c_templates_extract import extract
        source = self.base / "table.docx"
        document = Document()
        paragraph = document.add_paragraph("Fixed ")
        paragraph.add_run("{{fi").bold = True
        paragraph.add_run("rst}} and {{second}}")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "old", "old"
        document.add_paragraph("After {{tail}}")
        document.save(source)
        extraction = extract(source, "docx")
        fields = [t for t in extraction["targets"] if t["kind"] == "field"]
        selected = next(t for t in extraction["targets"] if t["kind"] == "table")
        operations = [{"target_id": t["id"], "kind": "text", "value": v}
                      for t, v in zip(fields, ("FIRST", "SECOND", "TAIL"))]
        operations.append({"target_id": selected["id"], "kind": "table", "value": [["A", "1"], ["B", "2"]]})
        data = fill_template(source, extraction, operations)
        self.assertEqual(self.verify_output(source, extraction, operations, data)["status"], "PASSED")
        for old, new in (("Name", "Renamed"), ("SECOND", "FIRST"), ("TAIL", "WRONG")):
            with self.subTest(mutation=old):
                with self.assertRaises(TemplateError):
                    self.verify_output(source, extraction, operations, replace_document_text(data, "docx", old, new))

    def test_docx_section_blocks_are_verified_without_reusing_generation(self) -> None:
        from docx import Document
        from scripts.flow1c_docx import fill_template
        from scripts.flow1c_templates_extract import extract
        for with_exemplar in (False, True):
            with self.subTest(table_exemplar=with_exemplar):
                source = self.base / "section.docx"
                document = Document()
                document.add_heading("Content", 1)
                document.add_paragraph("Old section content")
                if with_exemplar:
                    document.add_table(rows=2, cols=2)
                document.add_heading("Fixed ending", 1)
                document.add_paragraph("Unchanged content")
                document.save(source)
                extraction = extract(source, "docx")
                target = next(t for t in extraction["targets"] if t["kind"] == "section" and t["title"] == "Content")
                value = "First **bold** and `code`\n- Item\n1. Number\n\n| A | B |\n| --- | --- |\n| X |  |\n### Child\nLast"
                operations = [{"target_id": target["id"], "kind": "section", "value": value}]
                data = fill_template(source, extraction, operations)
                self.assertEqual(self.verify_output(source, extraction, operations, data)["status"], "PASSED")
                with self.assertRaises(TemplateError):
                    self.verify_output(source, extraction, operations,
                                       replace_document_text(data, "docx", "Unchanged content", "Changed"))

    def test_docx_immutable_package_parts_cannot_change_with_the_planned_values(self) -> None:
        from docx import Document
        from scripts.flow1c_docx import fill_template
        from scripts.flow1c_templates_extract import extract
        source = self.base / "parts.docx"
        document = Document()
        document.add_paragraph("{{field}}")
        document.sections[0].header.paragraphs[0].text = "Fixed header"
        document.save(source)
        extraction = extract(source, "docx")
        target = next(t for t in extraction["targets"] if t["kind"] == "field")
        operations = [{"target_id": target["id"], "kind": "text", "value": "Actual"}]
        correct = fill_template(source, extraction, operations)
        for part in ("word/header1.xml", "word/styles.xml", "word/_rels/document.xml.rels"):
            with self.subTest(part=part):
                result = io.BytesIO()
                with zipfile.ZipFile(io.BytesIO(correct)) as package, zipfile.ZipFile(result, "w") as output:
                    for info in package.infolist():
                        data = package.read(info.filename)
                        if info.filename == part:
                            data += b"\n"
                        output.writestr(info, data)
                with self.assertRaises(TemplateError):
                    self.verify_output(source, extraction, operations, result.getvalue())
