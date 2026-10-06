from __future__ import annotations

import argparse
import io
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
from scripts.flow1c_templates import TemplateService, readiness
from scripts.flow1c_templates_policy import TemplateError, select_template, validate_profile
from scripts.flow1c_templates_store import LibraryStore, atomic_json, exclusive, new_id, read_json, sha256
from scripts.flow1c_templates_extract import extract
from scripts.flow1c_sections import load_catalog


def synthetic_fixture(folder: Path, kind: str, format_name: str, sample: bool) -> Path:
    """Twenty reproducible synthetic type/format/sample fixtures, with no business data."""
    ids = [s["section_id"] for s in load_catalog(ROOT / "standards/functional-specification-sections.md")["sections"]] if kind == "functional-spec" else ["content"]
    path = folder / f"{kind}-{'sample' if sample else 'clean'}.{'md' if format_name == 'markdown' else 'docx'}"
    if format_name == "markdown":
        text = f"# {kind}\n\nUnchanged brand\n\n"
        for section in ids:
            text += f"## {section}\n\n" + ("SAMPLE value: do not reuse\n\n" if sample else "{{" + section + "}}\n\n")
        path.write_text(text, encoding="utf-8")
    else:
        from docx import Document
        doc = Document()
        doc.add_heading(kind, level=1)
        doc.add_paragraph("Unchanged brand")
        for section in ids:
            doc.add_heading(section, level=2)
            paragraph = doc.add_paragraph()
            if sample:
                paragraph.add_run("SAMPLE value: do not reuse")
            else:
                paragraph.add_run("{{" + section[:2])
                paragraph.add_run(section[2:] + "}}")  # split placeholder runs
        doc.sections[0].header.paragraphs[0].text = "Fixed header"
        doc.save(path)
    return path


def interpreted_profile(extraction: dict, kind: str) -> tuple[dict, list[dict]]:
    variables = [t for t in extraction["targets"] if t["kind"] == "field" or
                 (t["kind"] == "paragraph" and t["text"].startswith("SAMPLE")) or
                 (t["kind"] == "section" and t["part"] == "markdown" and t["text"].strip() == "SAMPLE value: do not reuse")]
    profile = {"schema_version": 1, "source_sha256": extraction["source_sha256"], "rules": [], "questions": [],
        "coverage": [{"target_id": t["id"], "role": "variable" if t in variables else "constant",
                      "purpose": "Ask for actual new content" if t in variables else "Keep corporate structure",
                      "required": t in variables, "required_basis": "user-defined fixture contract" if t in variables else None,
                      "unknown": "Ask user; never reuse example facts"} for t in extraction["targets"]]}
    if kind == "functional-spec":
        canonical = [s["section_id"] for s in load_catalog(ROOT / "standards/functional-specification-sections.md")["sections"]]
        profile["canonical_mapping"] = {section: [t["id"]] for section, t in zip(canonical, variables)}
    operations = [{"target_id": t["id"], "kind": "text" if t["kind"] in {"field", "paragraph"} else t["kind"],
                   "value": "Actual user supplied content", "basis": "Explicit user statement in synthetic test"} for t in variables]
    return profile, operations


class TemplateIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.base = Path(self.temporary.name)
        self.docs = self.base / "docs"
        self.docs.mkdir()
        self.source = self.base / "source"
        self.source.mkdir()
        self.service = TemplateService(ROOT, {"documentation_path": str(self.docs)})

    def tearDown(self):
        self.temporary.cleanup()

    def test_setup_template_import_resumes_in_same_gate_and_preserves_checkpoint(self):
        from argparse import Namespace
        from contextlib import redirect_stdout
        from scripts import flow1c
        from scripts.flow1c_templates_cli import INPUT_FIELDS, command

        checkout = self.base / "workflow"
        checkout.mkdir()
        for folder in ("schemas", "config", "standards"):
            shutil.copytree(ROOT / folder, checkout / folder)
        atomic_json(checkout / flow1c.LOCAL_CONFIG_FILE,
                    {"schema_version": 2, "documentation_path": str(self.docs), "user_settings": {"retain": True}})
        setup_id = new_id()
        source = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        with mock.patch.object(flow1c, "ROOT", checkout):
            flow1c.save_setup_checkpoint({"schema_version": 1, "setup_id": setup_id,
                                         "state": "COMPLETE", "confirmed_values": {"profile": "analysis"}})
            gate = flow1c.new_gate("setup", None, "READY", mode="formal", setup_id=setup_id)
            self.assertIn("flow1c_template", gate["available_actions"])

            def call(action, request):
                args = Namespace(**{key: None for key in INPUT_FIELDS})
                args.action, args.request, args.gate_id = action, request, gate["gate_id"]
                output = io.StringIO()
                with redirect_stdout(output):
                    self.assertEqual(command(flow1c, args), 0)
                return json.loads(output.getvalue())

            received = call("intake", {"source": str(source), "document_type": "meeting-minutes"})
            self.assertEqual(received["state"], "PROFILE_DRAFT")
            operation_id = received["operation_id"]
            resumed = call("resume", {"operation_id": operation_id})
            self.assertEqual(resumed["operation_id"], operation_id)
            service = TemplateService(checkout, {"documentation_path": str(self.docs)})
            extraction = read_json(service.store.operation_path(operation_id) / "candidate/extraction.json")
            profile, _ = interpreted_profile(extraction, "meeting-minutes")
            call("profile-save", {"operation_id": operation_id, "profile": profile})
            active = call("activate", {"operation_id": operation_id, "expected_revision": None, "default": True})
            self.assertEqual(active["state"], "ACTIVE")
            checkpoint = read_json(flow1c.setup_checkpoint_path(setup_id))
            self.assertEqual(checkpoint["state"], "COMPLETE")
            self.assertEqual(checkpoint["confirmed_values"], {"profile": "analysis"})
            self.assertEqual(checkpoint["template_progress"][operation_id]["state"], "ACTIVE")
            self.assertEqual(flow1c.read_json(checkout / flow1c.LOCAL_CONFIG_FILE, {})["user_settings"], {"retain": True})
            self.assertEqual(len(list((checkout / ".workspace/agent-gates").glob("*.json"))), 1)

    def add(self, path: Path, kind: str = "meeting-minutes", **extra):
        request = {"source": str(path), "document_type": kind, **extra}
        received = self.service.intake(request)
        self.assertEqual(received["state"], "PROFILE_DRAFT", received)
        staged = self.service.store.operation_path(received["operation_id"]) / "candidate"
        extraction = read_json(staged / "extraction.json")
        profile, operations = interpreted_profile(extraction, kind)
        saved = self.service.profile_save({"operation_id": received["operation_id"], "profile": profile})
        self.assertEqual(saved["state"], "VALIDATED", saved)
        active = self.service.activate({"operation_id": received["operation_id"], "expected_revision": extra.get("parent_revision"), "default": True})
        return active, profile, operations

    def test_twenty_fixtures_end_to_end_all_types_both_formats_and_examples(self):
        for kind in self.service.types:
            for format_name in ("markdown", "docx"):
                for sample in (False, True):
                    with self.subTest(kind=kind, format=format_name, sample=sample):
                        source = synthetic_fixture(self.source, kind, format_name, sample)
                        before = sha256(source)
                        active, profile, operations = self.add(source, kind)
                        request_root = self.docs / "drafts" / new_id()
                        selected = self.service.resolve({"document_type": kind, "template_id": active["template_id"]}, request_root)
                        plan = self.service.document_plan({"operations": operations}, request_root)
                        check = self.service.document_write({"plan_id": plan["plan_id"]}, request_root)
                        output = Path(check["output"])
                        self.assertEqual(output.suffix, source.suffix)
                        self.assertEqual(sha256(source), before)
                        self.assertEqual(check["integrity"], "PASSED")
                        self.assertEqual(check["layout_state"], "NOT_APPLICABLE" if format_name == "markdown" else "UNVERIFIED")
                        final = extract(output, format_name)
                        visible = " ".join(t["text"] for t in final["targets"])
                        self.assertNotIn("SAMPLE value", visible)
                        self.assertIn("Actual user supplied", visible)
                        source.unlink()
                        self.assertTrue(self.service.document_validate({"plan_id": plan["plan_id"]}, request_root)["ready"])
                        self.assertEqual(self.service.document_write({"plan_id": plan["plan_id"]}, request_root), check)

    def test_pin_survives_update_and_rollback_and_store_restart(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        old, _, _ = self.add(path)
        request_root = self.docs / "drafts" / new_id()
        pin = self.service.resolve({"document_type": "meeting-minutes"}, request_root)["pin"]
        path.write_text(path.read_text(encoding="utf-8") + "\nNew constant\n", encoding="utf-8")
        new, _, _ = self.add(path, template_id=old["template_id"], parent_revision=old["revision_id"])
        self.assertNotEqual(new["revision_id"], old["revision_id"])
        fresh = TemplateService(ROOT, {"documentation_path": str(self.docs)})
        self.assertEqual(fresh.resolve({"document_type": "meeting-minutes"}, request_root)["pin"], pin)
        self.service.activate({"operation_id": old["operation_id"], "expected_revision": new["revision_id"]})
        self.assertEqual(self.service.store.index()["templates"][0]["active_revision"], old["revision_id"])

    def test_failed_candidate_preserves_active_and_profile_answers_resume(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        old, _, _ = self.add(path)
        candidate = self.service.intake({"source": str(path), "document_type": "meeting-minutes",
            "template_id": old["template_id"], "parent_revision": old["revision_id"]})
        extraction = read_json(self.service.store.operation_path(candidate["operation_id"]) / "candidate/extraction.json")
        profile, _ = interpreted_profile(extraction, "meeting-minutes")
        profile["questions"] = [{"target_id": "field", "question": "What is the field purpose?"}]
        failed = self.service.profile_save({"operation_id": candidate["operation_id"], "profile": profile})
        self.assertEqual(failed["state"], "NEEDS_CLARIFICATION")
        self.assertEqual(self.service.resume(candidate["operation_id"])["questions"], profile["questions"])
        self.assertEqual(self.service.store.index()["templates"][0]["active_revision"], old["revision_id"])

    def test_revision_immutability_integrity_and_same_file_rules_dedup(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        old, profile, _ = self.add(path)
        request = {"source": str(path), "document_type": "meeting-minutes", "template_id": old["template_id"], "parent_revision": old["revision_id"], "operation_id": new_id()}
        candidate = self.service.intake(request)
        self.assertEqual(self.service.intake(request), candidate)
        saved = self.service.profile_save({"operation_id": candidate["operation_id"], "profile": profile})
        self.assertEqual(saved["revision_id"], old["revision_id"])
        with self.assertRaises(TemplateError):
            self.service.intake({**request, "variant_name": "changed input"})
        root = self.service.store.revision_path(old["template_id"], old["revision_id"])
        (root / "filling-guide.md").write_text("tampered", encoding="utf-8")
        with self.assertRaises(TemplateError) as error:
            self.service.resolve({"document_type": "meeting-minutes"}, self.docs / "drafts" / new_id())
        self.assertEqual(error.exception.code, "TEMPLATE_INTEGRITY_FAILED")

    def test_rule_correction_same_source_creates_new_revision(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        old, profile, _ = self.add(path)
        candidate = self.service.intake({"source": str(path), "document_type": "meeting-minutes", "template_id": old["template_id"], "parent_revision": old["revision_id"]})
        profile["rules"] = [{"text": "Decisions require an owner and due date", "provenance": "user_defined", "source": "user_statement: decisions need owner and due date"}]
        saved = self.service.profile_save({"operation_id": candidate["operation_id"], "profile": profile})
        self.assertNotEqual(saved["revision_id"], old["revision_id"])
        self.assertEqual(saved["source_sha256"], old["source_sha256"])
        self.assertEqual(self.service.store.index()["templates"][0]["active_revision"], old["revision_id"])

    def test_concurrent_active_switch_conflict_and_lock(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        active, _, _ = self.add(path)
        path.write_text(path.read_text(encoding="utf-8") + "New constant\n", encoding="utf-8")
        updated, _, _ = self.add(path, template_id=active["template_id"], parent_revision=active["revision_id"])
        with self.assertRaises(TemplateError) as error:
            self.service.activate({"operation_id": active["operation_id"], "expected_revision": new_id()})
        self.assertEqual(error.exception.code, "TEMPLATE_REVISION_CONFLICT")
        with exclusive(self.service.store.root):
            with self.assertRaises(TemplateError) as error:
                self.service.activate({"operation_id": active["operation_id"], "expected_revision": updated["revision_id"]})
        self.assertEqual(error.exception.code, "TEMPLATE_REVISION_CONFLICT")

    def test_relocate_preserves_all_hashes_and_rejects_foreign_library(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        active, _, _ = self.add(path)
        destination = self.base / "relocated"
        destination.mkdir()
        transfer = new_id()
        result = self.service.store.relocate(destination, transfer)
        self.assertTrue(result["old_copy_preserved"])
        relocated = LibraryStore(destination)
        self.assertEqual(relocated.index(), self.service.store.index())
        self.assertEqual(self.service.store.relocate(destination, transfer)["library_id"], result["library_id"])
        foreign = self.base / "foreign"
        foreign.mkdir()
        LibraryStore(foreign).index(create=True)
        with self.assertRaises(TemplateError):
            self.service.store.relocate(foreign, new_id())

    def test_markdown_table_escape_code_preservation_and_resources(self):
        path = self.source / "table.md"
        path.write_text("# Data\n\n| Name | Value |\n| --- | --- |\n| {{name}} | {{value}} |\n\n```python\nprint('constant')\n```\n\n![logo](logo.png)\n", encoding="utf-8")
        (self.source / "logo.png").write_bytes(b"synthetic resource bytes")
        received = self.service.intake({"source": str(path), "document_type": "meeting-minutes", "resources_root": str(self.source)})
        extraction = read_json(self.service.store.operation_path(received["operation_id"]) / "candidate/extraction.json")
        profile, operations = interpreted_profile(extraction, "meeting-minutes")
        operations[0]["value"] = "A | B\nC `code`"
        self.service.profile_save({"operation_id": received["operation_id"], "profile": profile})
        self.service.activate({"operation_id": received["operation_id"], "expected_revision": None})
        request_root = self.docs / "drafts" / new_id()
        self.service.resolve({"document_type": "meeting-minutes"}, request_root)
        plan = self.service.document_plan({"operations": operations}, request_root)
        check = self.service.document_write({"plan_id": plan["plan_id"]}, request_root)
        content = Path(check["output"]).read_text(encoding="utf-8")
        self.assertIn("A \\| B<br>C \\`code\\`", content)
        self.assertIn("print('constant')", content)
        self.assertEqual((request_root / "logo.png").read_bytes(), b"synthetic resource bytes")
        (request_root / "logo.png").write_bytes(b"changed")
        with self.assertRaises(TemplateError):
            self.service.document_validate({"plan_id": plan["plan_id"]}, request_root)

    def test_missing_dependency_keeps_received_source_and_resume(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        error = TemplateError("TEMPLATE_DEPENDENCY_UNAVAILABLE", "missing")
        with mock.patch("scripts.flow1c_templates.extract", side_effect=error):
            result = self.service.intake({"source": str(path), "document_type": "meeting-minutes"})
        self.assertEqual(result["state"], "FAILED_RECOVERABLE")
        path.unlink()
        self.assertEqual(self.service.resume(result["operation_id"])["state"], "PROFILE_DRAFT")

    def test_output_no_overwrite_and_tampering_detected(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        _, _, ops = self.add(path)
        root = self.docs / "drafts" / new_id()
        self.service.resolve({"document_type": "meeting-minutes"}, root)
        plan = self.service.document_plan({"operations": ops, "output_name": "result.md"}, root)
        (root / "result.md").write_text("user content", encoding="utf-8")
        with self.assertRaises(TemplateError):
            self.service.document_write({"plan_id": plan["plan_id"]}, root)
        self.assertEqual((root / "result.md").read_text(encoding="utf-8"), "user content")

    def test_interrupted_profile_commit_and_output_publication_resume(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        received = self.service.intake({"source": str(path), "document_type": "meeting-minutes"})
        staged = self.service.store.operation_path(received["operation_id"]) / "candidate"
        profile, operations = interpreted_profile(read_json(staged / "extraction.json"), "meeting-minutes")
        # Crash after directory rename, before registration in the library index.
        original = self.service.store.revision_path(received["template_id"], received["revision_id"])
        from scripts import flow1c_templates
        real_atomic = flow1c_templates.atomic_json
        def fail_index(target, value):
            if target.name == "library.json" and original.exists():
                raise OSError("synthetic interruption")
            return real_atomic(target, value)
        with mock.patch.object(flow1c_templates, "atomic_json", side_effect=fail_index):
            with self.assertRaises(OSError):
                self.service.profile_save({"operation_id": received["operation_id"], "profile": profile})
        self.assertEqual(self.service.resume(received["operation_id"])["state"], "VALIDATED")
        self.service.activate({"operation_id": received["operation_id"], "expected_revision": None})
        root = self.docs / "drafts" / new_id()
        self.service.resolve({"document_type": "meeting-minutes"}, root)
        plan = self.service.document_plan({"operations": operations}, root)
        def fail_validation(target, value):
            if target.name.endswith(".validation.json"):
                raise OSError("synthetic interruption after publication")
            return real_atomic(target, value)
        with mock.patch.object(flow1c_templates, "atomic_json", side_effect=fail_validation):
            with self.assertRaises(OSError):
                self.service.document_write({"plan_id": plan["plan_id"]}, root)
        check = self.service.document_write({"plan_id": plan["plan_id"]}, root)
        self.assertEqual(check["state"], "WRITTEN")
        self.assertTrue(self.service.document_validate({"plan_id": plan["plan_id"]}, root)["ready"])

    def test_docx_fixed_table_header_bookmark_styles_and_package_preservation(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        path = self.source / "complex-supported.docx"
        document = Document()
        document.add_heading("Data", 1)
        paragraph = document.add_paragraph("{{body}}")
        bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), "1")
        bookmark.set(qn("w:name"), "Content")
        paragraph._p.insert(0, bookmark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "1")
        paragraph._p.append(end)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Owner", "Decision"
        table.cell(1, 0).text, table.cell(1, 1).text = "{{owner}}", "{{decision}}"
        document.sections[0].header.paragraphs[0].text = "{{header}}"
        document.add_paragraph("Unchanged list", style="List Bullet")
        document.save(path)
        active, _, operations = self.add(path)
        root = self.docs / "drafts" / new_id()
        self.service.resolve({"document_type": "meeting-minutes"}, root)
        plan = self.service.document_plan({"operations": operations}, root)
        result = self.service.document_write({"plan_id": plan["plan_id"]}, root)
        output = Path(result["output"])
        with zipfile.ZipFile(path) as before, zipfile.ZipFile(output) as after:
            changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
            self.assertEqual(changed, {"word/document.xml", "word/header1.xml"})
            self.assertIn(b'Content', after.read("word/document.xml"))
        self.assertEqual(Document(output).sections[0].header.paragraphs[0].text, "Actual user supplied content")

    def test_future_schema_and_junction_escape_are_blocked(self):
        from scripts.flow1c_templates_store import safe_path
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        self.add(path)
        index_path = self.service.store.root / "library.json"
        index = read_json(index_path)
        index["schema_version"] = 99
        atomic_json(index_path, index)
        before = index_path.read_bytes()
        with self.assertRaises(TemplateError):
            self.service.list()
        self.assertEqual(index_path.read_bytes(), before)
        info = mock.Mock(st_mode=0, st_file_attributes=0x400)
        with mock.patch.object(Path, "lstat", return_value=info):
            with self.assertRaises(TemplateError):
                safe_path(self.docs / "junction" / "escape.md", self.docs)

    def test_formal_fs_rechecks_actual_section_approvals_and_word_policy(self):
        from types import SimpleNamespace
        from scripts.flow1c_templates_cli import check_formal_content
        from scripts.flow1c_sections_policy import ensure_write_authorized, content_sha256, SectionPolicyError
        path = synthetic_fixture(self.source, "functional-spec", "markdown", False)
        active, _, operations = self.add(path, "functional-spec")
        root = self.docs / "work-items" / "synthetic-reference" / "specification" / "template-documents" / new_id()
        self.service.resolve({"document_type": "functional-spec"}, root)
        approved = "Actual user supplied content"
        state = {"state": "APPROVED", "content_sha256": content_sha256(approved), "approved_content_sha256": content_sha256(approved)}
        api = SimpleNamespace(load_manifest=lambda reference: (None, {}), stage_errors=lambda stage, manifest: [],
            load_stages=lambda: {"operations": {"functional-spec": {}}},
            _sections_catalog=lambda: load_catalog(ROOT / "standards/functional-specification-sections.md"),
            _sections_root=lambda reference: self.docs,
            _load_section_current=lambda *args: (None, approved, state), ensure_write_authorized=ensure_write_authorized)
        gate = {"mode": "formal", "operation": "functional-spec", "state": "READY", "work_reference": "synthetic-reference"}
        check_formal_content(api, gate, self.service, {"operations": operations}, root)
        bad_operations = [dict(op) for op in operations]
        bad_operations[0]["value"] = "Changed after approval"
        with self.assertRaises(SectionPolicyError):
            check_formal_content(api, gate, self.service, {"operations": bad_operations}, root)
        self.service.local["template_library"] = {"require_word_for": ["functional-spec"]}
        with self.assertRaises(TemplateError) as error:
            check_formal_content(api, gate, self.service, {"operations": operations}, root)
        self.assertEqual(error.exception.code, "TEMPLATE_POLICY_CONFLICT")


    def test_formal_completion_rechecks_section_approval_after_document_write(self):
        from scripts import flow1c
        from scripts.flow1c_sections_policy import content_sha256

        source = synthetic_fixture(self.source, "functional-spec", "markdown", False)
        _, _, operations = self.add(source, "functional-spec")
        root = self.docs / "work-items" / "synthetic-reference" / "specification" / "template-documents" / new_id()
        self.service.resolve({"document_type": "functional-spec"}, root)
        plan = self.service.document_plan({"operations": operations}, root)
        result = self.service.document_write({"plan_id": plan["plan_id"]}, root)
        approved = "Actual user supplied content"
        section_state = {"state": "APPROVED", "content_sha256": content_sha256(approved),
                         "approved_content_sha256": content_sha256(approved)}
        gate = {"gate_id": new_id(), "mode": "formal", "operation": "functional-spec", "state": "READY",
                "code": "synthetic-reference", "work_reference": "synthetic-reference", "manifest_status": "DRAFT",
                "manifest_approvals": {}, "template_result": result}
        patches = {
            "load_gate": mock.Mock(return_value=gate),
            "load_stages": mock.Mock(return_value={"operations": {"functional-spec": {"allowed_tools": ["flow1c_complete"]}}}),
            "load_manifest": mock.Mock(return_value=(None, {"status": "DRAFT", "approvals": {}})),
            "evidence_for_gate": mock.Mock(return_value=(self.docs / "evidence.json", {})),
            "validate_json_record": mock.Mock(return_value=[]),
            "TemplateService": mock.Mock(return_value=self.service),
            "_sections_root": mock.Mock(return_value=self.docs),
            "_load_section_current": mock.Mock(side_effect=lambda *args: (None, approved, section_state)),
            "publication_snapshot": mock.Mock(return_value={}),
            "save_gate": mock.Mock(),
        }
        with mock.patch.multiple(flow1c, **patches), \
                mock.patch.object(flow1c.templates_cli, "document_root", return_value=root):
            for state, word_required, expected in (("APPROVED", False, 0), ("DRAFT", False, 2), ("APPROVED", True, 2)):
                with self.subTest(section_state=state, word_required=word_required):
                    gate["state"] = "READY"
                    section_state["state"] = state
                    self.service.local["template_library"] = {"require_word_for": ["functional-spec"] if word_required else []}
                    with redirect_stdout(io.StringIO()):
                        code = flow1c.cmd_agent_complete(argparse.Namespace(gate_id=gate["gate_id"], output=None))
                    self.assertEqual(code, expected)
                    if expected:
                        self.assertEqual(gate["state"], "NON_COMPLIANT")
                        self.assertTrue(any("template" in e for e in gate["completion_errors"]))
        self.assertEqual(sha256(Path(result["output"])), result["output_sha256"])

    def test_docx_section_replacement_cannot_remove_section_break(self):
        from docx import Document
        from scripts.flow1c_docx import fill_template

        source = self.source / "section-break.docx"
        document = Document()
        document.add_heading("Data", 1)
        document.add_paragraph("Sample content")
        document.add_section()
        document.add_heading("Following section", 1)
        document.add_paragraph("Preserved text")
        document.save(source)
        extraction = extract(source, "docx")
        target = next(t for t in extraction["targets"] if t["kind"] == "section" and t["title"] == "Data")
        self.assertFalse(target["supported"])
        operations = [{"target_id": target["id"], "kind": "section", "value": "Actual user content"}]
        for old_extraction in (False, True):
            with self.subTest(old_extraction=old_extraction):
                target["supported"] = old_extraction
                with self.assertRaises(TemplateError) as error:
                    fill_template(source, extraction, operations)
                self.assertEqual(error.exception.code, "TEMPLATE_LAYOUT_UNSUPPORTED")
        self.assertEqual(len(Document(source).sections), 2)

    def test_written_document_validation_binds_plan_operations(self):
        from scripts.flow1c_templates_policy import fingerprint

        source = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        _, _, operations = self.add(source)
        root = self.docs / "drafts" / new_id()
        self.service.resolve({"document_type": "meeting-minutes"}, root)
        plan_result = self.service.document_plan({"operations": operations}, root)
        result = self.service.document_write({"plan_id": plan_result["plan_id"]}, root)
        validation_path = root / (plan_result["plan_id"] + ".validation.json")
        legacy_check = read_json(validation_path)
        legacy_check.pop("operations_sha256")
        atomic_json(validation_path, legacy_check)
        self.assertTrue(self.service.document_validate({"plan_id": plan_result["plan_id"]}, root)["ready"])
        atomic_json(validation_path, result)
        plan_path = root / "document-plans" / (plan_result["plan_id"] + ".json")
        plan = read_json(plan_path)
        plan["operations"][0]["value"] = "Different content from the written document"
        plan["content_sha256"] = fingerprint(plan["operations"])
        atomic_json(plan_path, plan)
        for check in (result, legacy_check):
            with self.subTest(legacy="operations_sha256" not in check):
                atomic_json(validation_path, check)
                with self.assertRaises(TemplateError) as error:
                    self.service.document_validate({"plan_id": plan_result["plan_id"]}, root)
                self.assertEqual(error.exception.code, "TEMPLATE_INTEGRITY_FAILED")
        self.assertEqual(sha256(Path(result["output"])), result["output_sha256"])

    def test_large_context_has_lossless_target_guide_and_metadata_pagination(self):
        self.assertEqual(self.service.list()["contracts"]["profile"], json.loads((ROOT / "schemas/document-template-profile.schema.json").read_text(encoding="utf-8")))
        path = self.source / "large.md"
        path.write_text("# Content\n\n" + "constant " * 2100 + "\n\n{{value}}\n", encoding="utf-8")
        received = self.service.intake({"source": str(path), "document_type": "meeting-minutes"})
        staged = self.service.store.operation_path(received["operation_id"]) / "candidate"
        extraction = read_json(staged / "extraction.json")
        extraction["metadata"]["synthetic_large_metadata"] = "m" * 9500
        atomic_json(staged / "extraction.json", extraction)
        profile, _ = interpreted_profile(extraction, "meeting-minutes")
        for item in profile["coverage"]:
            item["purpose"] = "purpose " * 900
        saved = self.service.profile_save({"operation_id": received["operation_id"], "profile": profile})
        self.assertEqual(saved["state"], "VALIDATED")
        pin = {k: saved[k] for k in ("library_id", "document_type", "template_id", "revision_id")}
        offset = text_offset = 0
        fragments = []
        while True:
            page = self.service.inspect({"pin": pin, "offset": offset, "text_offset": text_offset, "limit": 1})
            fragments.append(page["targets"][0]["text"])
            if page["coverage"]["next_text_offset"] is None:
                break
            text_offset = page["coverage"]["next_text_offset"]
        self.assertEqual("".join(fragments), extraction["targets"][0]["text"])
        root, _ = self.service.store.revision(pin)
        for value, continuation, input_offset, expected in (
            ("guide", "guide_next_offset", "guide_offset", (root / "filling-guide.md").read_text(encoding="utf-8")),
            ("metadata_text", "metadata_next_offset", "metadata_offset", json.dumps(extraction["metadata"], ensure_ascii=False)),
        ):
            offset, fragments = 0, []
            while True:
                page = self.service.inspect({"pin": pin, input_offset: offset})
                fragments.append(page[value])
                if page[continuation] is None:
                    break
                offset = page[continuation]
            self.assertEqual("".join(fragments), expected)

    def test_markdown_constant_code_placeholder_and_new_local_resource(self):
        from scripts.flow1c_markdown import write_bytes
        path = self.source / "code.md"
        path.write_text("# Content\n\n{{value}}\n\n```text\n{{constant-code}}\n```\n", encoding="utf-8")
        extraction = extract(path, "markdown")
        field = next(t for t in extraction["targets"] if t["kind"] == "field")
        result = write_bytes(path, extraction, [{"target_id": field["id"], "kind": "text", "value": "Actual"}])
        self.assertIn(b"{{constant-code}}", result)
        path.write_text("# Content\n\nold\n", encoding="utf-8")
        extraction = extract(path, "markdown")
        with self.assertRaises(TemplateError):
            write_bytes(path, extraction, [{"target_id": extraction["targets"][0]["id"], "kind": "section", "value": "![new](missing.png)"}])

    def test_whole_docx_table_replaces_sample_and_preserves_run_style(self):
        from docx import Document
        from scripts.flow1c_docx import fill_template
        path = self.source / "rows.docx"
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
        for cell in table.rows[1].cells:
            cell.paragraphs[0].add_run("{{sample}}").bold = True
        document.save(path)
        extraction = extract(path, "docx")
        target = next(t for t in extraction["targets"] if t["kind"] == "table")
        profile = {"source_sha256": extraction["source_sha256"], "coverage": [
            {"target_id": t["id"], "role": "variable" if t is target else "constant", "purpose": "New rows", "required": False}
            for t in extraction["targets"]], "questions": []}
        validate_profile(profile, extraction)
        output = self.source / "filled.docx"
        output.write_bytes(fill_template(path, extraction, [{"target_id": target["id"], "kind": "table", "value": [["Actual", "1"], ["Other", "2"]]}]))
        result = Document(output)
        self.assertEqual(len(result.tables[0].rows), 3)
        self.assertEqual(result.tables[0].cell(0, 0).text, "Name")
        self.assertEqual(result.tables[0].cell(1, 0).text, "Actual")
        self.assertTrue(result.tables[0].cell(1, 0).paragraphs[0].runs[0].bold)

    def test_interrupted_activation_resumes_and_relocation_cannot_nest(self):
        path = synthetic_fixture(self.source, "meeting-minutes", "markdown", False)
        active, _, _ = self.add(path)
        checkpoint = self.service.store.operation(active["operation_id"])
        checkpoint.update(state="ACTIVATING", ready=False)
        self.service.store.checkpoint(checkpoint)
        self.assertEqual(self.service.resume(active["operation_id"])["state"], "ACTIVE")
        self.assertEqual(self.service.activate({"operation_id": active["operation_id"], "expected_revision": None, "default": True})["state"], "ACTIVE")
        destination = self.service.store.root / "nested-docs"
        destination.mkdir()
        with self.assertRaises(TemplateError):
            self.service.store.relocate(destination, new_id())



    def test_configuration_merge_preserves_concurrent_fields_and_rejects_path_change(self):
        from types import SimpleNamespace
        from scripts.flow1c_templates_cli import save_library_config
        checkout = self.base / "workflow"
        checkout.mkdir()
        path = checkout / ".flow1c.local.json"
        current = {"schema_version": 1, "documentation_path": str(self.docs), "user_updated": {"retain": True}}
        atomic_json(path, current)
        api = SimpleNamespace(ROOT=checkout, LOCAL_CONFIG_FILE=path.name,
                              read_json=lambda file, default: json.loads(file.read_text(encoding="utf-8")) if file.exists() else default,
                              write_json=atomic_json)
        save_library_config(api, {"documentation_path": str(self.docs)}, self.docs, new_id())
        saved = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(saved["user_updated"], current["user_updated"])
        self.assertEqual(saved["schema_version"], 2)
        with self.assertRaises(TemplateError):
            save_library_config(api, {"documentation_path": "stale"}, self.docs, saved["template_library"]["library_id"])
        self.assertEqual(json.loads(path.read_text(encoding="utf-8")), saved)


class TemplatePolicySecurityTests(unittest.TestCase):
    def test_optional_variable_cannot_preserve_sample_and_renderer_never_infers_qa(self):
        from scripts.flow1c_templates_policy import validate_content
        from scripts.flow1c_templates_renderer import render_for_review, RendererCapabilities
        profile = {"coverage": [{"target_id": "sample", "role": "variable", "required": False}]}
        with self.assertRaises(TemplateError):
            validate_content(profile, [])
        self.assertEqual(render_for_review(Path("draft.docx"), Path("pages"))["layout_state"], "UNVERIFIED")
        renderer = mock.Mock()
        renderer.capabilities = RendererCapabilities("fixture", "1", "test-only", ("docx",), True)
        renderer.render.side_effect = TimeoutError("synthetic timeout")
        self.assertEqual(render_for_review(Path("draft.docx"), Path("pages"), renderer)["layout_state"], "FAILED")
    def test_catalog_removal_does_not_remove_stored_pinned_revision(self):
        index = {"library_id": "library", "templates": [{"document_type": "custom", "template_id": "one", "name": "First", "active_revision": "old", "default": False, "revisions": ["old"]}]}
        self.assertEqual(select_template(index, "custom")["revision_id"], "old")
        index["templates"].append({"document_type": "custom", "template_id": "two", "name": "Second", "active_revision": "new", "default": False, "revisions": ["new"]})
        with self.assertRaises(TemplateError) as error:
            select_template(index, "custom")
        self.assertEqual(error.exception.code, "TEMPLATE_VARIANT_AMBIGUOUS")
        self.assertEqual(select_template(index, "custom", template_id="one")["revision_id"], "old")

    def test_profile_overlap_and_policy_conflict(self):
        extraction = {"source_sha256": "hash", "targets": [{"id": "section", "kind": "section", "supported": True, "range": [0, 10]}, {"id": "field", "kind": "field", "supported": True, "range": [2, 5]}]}
        profile = {"source_sha256": "hash", "coverage": [{"target_id": i, "role": "variable", "purpose": "test"} for i in ("section", "field")]}
        with self.assertRaises(TemplateError) as error:
            validate_profile(profile, extraction)
        self.assertEqual(error.exception.code, "TEMPLATE_ANCHOR_AMBIGUOUS")
        profile["coverage"][0]["role"] = "constant"
        with self.assertRaises(TemplateError) as error:
            validate_profile(profile, extraction, policy="functional-spec", canonical_ids=["required"])
        self.assertEqual(error.exception.code, "TEMPLATE_POLICY_CONFLICT")

    def test_markdown_no_office_rlm_git_required_for_readiness(self):
        with tempfile.TemporaryDirectory() as directory:
            result = readiness(ROOT, {"documentation_path": directory}, "markdown", "template-management")
            self.assertTrue(result["ready"])
            self.assertNotIn("docx", [c["component"] for c in result["checks"]])
            self.assertNotIn("rlm", result["required_capabilities"])

    def test_zip_traversal_dtd_and_bounded_markdown_rejected(self):
        from scripts.flow1c_docx import DocxError, validate_docx_package
        from scripts.flow1c_markdown import MAX_MARKDOWN_BYTES, read_source
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name, raw in (("../escape.xml", b"<x/>"), ("safe.xml", b'<!DOCTYPE x [<!ENTITY x "secret">]><x/>')):
                path = root / "bad.docx"
                with zipfile.ZipFile(path, "w") as archive:
                    archive.writestr("[Content_Types].xml", b"<x/>")
                    archive.writestr("word/document.xml", b"<x/>")
                    archive.writestr(name, raw)
                with self.assertRaises(DocxError):
                    validate_docx_package(path)
            large = root / "large.md"
            with large.open("wb") as stream:
                stream.truncate(MAX_MARKDOWN_BYTES + 1)
            with self.assertRaises(TemplateError):
                read_source(large)
            from scripts.flow1c_templates_store import MAX_JSON_BYTES
            state = root / "preserved.json"
            atomic_json(state, {"schema_version": 1, "value": "preserve"})
            before = state.read_bytes()
            with self.assertRaises(TemplateError):
                atomic_json(state, {"schema_version": 1, "value": "x" * MAX_JSON_BYTES})
            self.assertEqual(state.read_bytes(), before)




if __name__ == "__main__":
    unittest.main()
