import importlib.util
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TEST_ROOT = ROOT / ".workspace" / "tests" / "docx-writeback"
POLICY_SPEC = importlib.util.spec_from_file_location("flow1c_sections_policy", ROOT / "scripts" / "flow1c_sections_policy.py")
policy = importlib.util.module_from_spec(POLICY_SPEC)
assert POLICY_SPEC.loader
POLICY_SPEC.loader.exec_module(policy)
DOCX_SPEC = importlib.util.spec_from_file_location("flow1c_docx", ROOT / "scripts" / "flow1c_docx.py")
docx_adapter = importlib.util.module_from_spec(DOCX_SPEC)
assert DOCX_SPEC.loader
DOCX_SPEC.loader.exec_module(docx_adapter)


def section():
    return next(item for item in policy.parse_catalog_markdown((ROOT / "standards" / "functional-specification-sections.md").read_text(encoding="utf-8")).get("sections", []) if item["section_id"] == "technical-implementation")


class DocxWritebackTests(unittest.TestCase):
    def setUp(self):
        TEST_ROOT.mkdir(parents=True, exist_ok=True)

    def tearDown(self):
        for path in TEST_ROOT.glob("flow1c-test-*.docx"):
            try:
                path.unlink()
            except OSError:
                pass

    def make_doc(self, path: Path) -> None:
        document = Document()
        document.add_heading("Общая информация", level=1)
        document.add_paragraph("Сохраняемый текст вне целевого раздела")
        document.add_heading("Техническая реализация", level=1)
        document.add_paragraph("{{PLACEHOLDER: technical-implementation}}")
        document.add_heading("Роли и права", level=1)
        document.add_paragraph("Внешний раздел не изменять")
        document.save(path)

    def make_custom_heading_style_doc(self, path: Path) -> None:
        document = Document()
        style = document.styles.add_style("cont_header3", WD_STYLE_TYPE.PARAGRAPH)
        style_properties = style.element.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "2")
        style_properties.append(outline)
        numbering = OxmlElement("w:numPr")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "12")
        numbering.append(num_id)
        style_properties.append(numbering)

        document.add_paragraph("Техническая реализация", style=style)
        body = document.add_paragraph("Старое содержимое раздела", style=style)
        body_properties = body._p.get_or_add_pPr()
        body_numbering = OxmlElement("w:numPr")
        body_num_id = OxmlElement("w:numId")
        body_num_id.set(qn("w:val"), "0")
        body_numbering.append(body_num_id)
        body_properties.append(body_numbering)
        run_properties = OxmlElement("w:rPr")
        bold = OxmlElement("w:b")
        bold.set(qn("w:val"), "0")
        run_properties.append(bold)
        body_properties.append(run_properties)

        document.add_paragraph("Дополнительный текст раздела")
        document.add_heading("Настройки системы, используемые разработкой", level=3)
        document.add_paragraph("Содержимое следующего раздела")
        document.save(path)

    def test_inspect_write_preserves_source_and_supports_list_and_table(self):
        source = TEST_ROOT / "flow1c-test-spec.docx"
        output = TEST_ROOT / "flow1c-test-spec-filled.docx"
        self.make_doc(source)
        before = docx_adapter.sha256_file(source)
        inspection = docx_adapter.inspect_docx(source, [section()])
        self.assertEqual(inspection["sections"][0]["status"], "FOUND_PLACEHOLDER")
        result = docx_adapter.write_docx(source, output, [{"section": section(), "mode": "replace", "content": "Добавлено\n- Первый пункт\n- Второй пункт\n\n| Поле | Значение |\n| --- | --- |\n| A | B |"}], expected_source_sha256=before)
        self.assertEqual(result["state"], "WRITTEN")
        self.assertEqual(docx_adapter.sha256_file(source), before)
        reopened = Document(output)
        text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
        self.assertIn("Добавлено", text)
        self.assertIn("Первый пункт", text)
        self.assertIn("Внешний раздел не изменять", text)
        self.assertEqual(next(paragraph.style.name for paragraph in reopened.paragraphs if paragraph.text == "Первый пункт"), "List Bullet")
        self.assertEqual(reopened.tables[0].style.name, "Table Grid")

    def test_technical_markdown_renders_as_word_structure(self):
        source = TEST_ROOT / "flow1c-test-technical-markdown.docx"
        output = TEST_ROOT / "flow1c-test-technical-markdown-filled.docx"
        self.make_doc(source)
        content = ("\\### Справочник «Номенклатура»\n"
                   "\\*\\*Модуль объекта:\\*\\*\n"
                   "a. Добавлена процедура \\`add\\_Проверка\\` для реквизита \\`flow1c\\_ProductForm\\`.\n"
                   "· Поле отображается в форме.\n")
        docx_adapter.write_docx(source, output, [{"section": section(), "mode": "replace", "content": content}])
        reopened = Document(output)
        paragraphs = {paragraph.text: paragraph for paragraph in reopened.paragraphs}
        self.assertEqual(paragraphs["1. Справочник «Номенклатура»"].style.name, "Heading 2")
        self.assertTrue(paragraphs["Модуль объекта:"].runs[0].bold)
        self.assertEqual(paragraphs["2. Добавлена процедура «add_Проверка» для реквизита «flow1c_ProductForm»."].style.name, "Heading 2")
        self.assertEqual(paragraphs["Поле отображается в форме."].style.name, "List Bullet")
        self.assertNotIn("###", "\n".join(paragraphs))
        self.assertNotIn("`", "\n".join(paragraphs))

    def test_technical_groups_use_decimal_hierarchy_even_with_letter_numbered_style(self):
        source = TEST_ROOT / "flow1c-test-letter-style.docx"
        output = TEST_ROOT / "flow1c-test-letter-style-filled.docx"
        self.make_doc(source)
        document = Document(source)
        heading_style = document.styles["Heading 2"]
        numbering = OxmlElement("w:numPr")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "7")
        numbering.append(num_id)
        heading_style.element.get_or_add_pPr().append(numbering)
        document.save(source)

        content = ("### Общие модули\n1.1. Новые модули не добавляются.\n"
                   "b. Документ\n2.1. Добавлен реквизит.\n2.2. Добавлена процедура.\n"
                   "c. Регистр\n3.1. Добавлен ресурс.\n"
                   "d. Справочник\n4.1. Изменена форма.\n"
                   "e. Настройки\n5.1. Задано значение.")
        docx_adapter.write_docx(source, output, [{"section": section(), "mode": "replace", "content": content}])
        paragraphs = Document(output).paragraphs
        for index, title in enumerate(("Общие модули", "Документ", "Регистр", "Справочник", "Настройки"), 1):
            group = next(paragraph for paragraph in paragraphs if paragraph.text == f"{index}. {title}")
            self.assertEqual(group.style.name, "Heading 2")
            self.assertEqual(group._p.pPr.numPr.numId.val, 0)
            self.assertTrue(any(paragraph.text.startswith(f"{index}.1. ") for paragraph in paragraphs))
        self.assertFalse(any(paragraph.style.name == "List Number" for paragraph in paragraphs))

    def test_subheadings_follow_numbered_section(self):
        source = TEST_ROOT / "flow1c-test-numbered-section.docx"
        output = TEST_ROOT / "flow1c-test-numbered-section-filled.docx"
        document = Document()
        document.add_heading("6. Техническая реализация", level=1)
        document.add_paragraph("{{заполнить}}")
        document.add_heading("7. Роли и права", level=1)
        document.save(source)
        docx_adapter.write_docx(source, output, [{"section": section(), "mode": "replace",
                                                 "content": "### Номенклатура\n### Марки"}])
        paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
        self.assertIn("6.1. Номенклатура", paragraphs)
        self.assertIn("6.2. Марки", paragraphs)
        self.assertIn("7. Роли и права", paragraphs)

    def test_ambiguous_heading_blocks_write(self):
        source = TEST_ROOT / "flow1c-test-ambiguous.docx"
        document = Document()
        document.add_heading("Техническая реализация", level=1)
        document.add_heading("Техническая реализация", level=1)
        document.save(source)
        inspection = docx_adapter.inspect_docx(source, [section()])
        self.assertEqual(inspection["sections"][0]["status"], "AMBIGUOUS")

    def test_bookmark_heading_preserves_following_sections_and_section_properties(self):
        source = TEST_ROOT / "flow1c-test-bookmark.docx"
        output = TEST_ROOT / "flow1c-test-bookmark-filled.docx"
        document = Document()
        heading = document.add_heading("Техническая реализация", level=1)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), "42")
        start.set(qn("w:name"), "flow1c_section_technical_implementation")
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "42")
        heading._p.insert(0, start)
        heading._p.append(end)
        document.add_paragraph("Старый текст")
        document.add_heading("Роли и права", level=1)
        document.add_paragraph("Содержимое следующего раздела")
        document.sections[0].left_margin = 1234567
        document.save(source)
        expected_margin = Document(source).sections[0].left_margin

        docx_adapter.write_docx(source, output, [{"section": section(), "mode": "replace", "content": "Новый текст"}])

        reopened = Document(output)
        text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
        self.assertIn("Новый текст", text)
        self.assertNotIn("Старый текст", text)
        self.assertIn("Содержимое следующего раздела", text)
        self.assertEqual(reopened.sections[0].left_margin, expected_margin)

    def test_last_heading_replace_keeps_sectpr(self):
        source = TEST_ROOT / "flow1c-test-last.docx"
        output = TEST_ROOT / "flow1c-test-last-filled.docx"
        document = Document()
        document.add_heading("Техническая реализация", level=1)
        document.add_paragraph("{{PLACEHOLDER: technical-implementation}}")
        document.sections[0].right_margin = 1111111
        document.save(source)
        expected_margin = Document(source).sections[0].right_margin

        docx_adapter.write_docx(source, output, [{"section": section(), "mode": "replace", "content": "Новый текст"}])

        reopened = Document(output)
        self.assertIsNotNone(reopened.element.body.sectPr)
        self.assertEqual(reopened.sections[0].right_margin, expected_margin)

    def test_content_control_append_preserves_existing_text(self):
        source = TEST_ROOT / "flow1c-test-control.docx"
        output = TEST_ROOT / "flow1c-test-control-filled.docx"
        document = Document()
        control = OxmlElement("w:sdt")
        properties = OxmlElement("w:sdtPr")
        tag = OxmlElement("w:tag")
        tag.set(qn("w:val"), "flow1c-section:technical-implementation")
        properties.append(tag)
        content = OxmlElement("w:sdtContent")
        paragraph = document.add_paragraph("Существующий текст")
        paragraph._p.getparent().remove(paragraph._p)
        content.append(paragraph._p)
        control.append(properties)
        control.append(content)
        document.element.body.insert(len(document.element.body) - 1, control)
        document.save(source)

        docx_adapter.write_docx(source, output, [{"section": section(), "mode": "append", "content": "Добавленный текст"}])

        inspection = docx_adapter.inspect_docx(output, [section()])
        self.assertIn("Существующий текст", inspection["sections"][0]["existing_text"])
        self.assertIn("Добавленный текст", inspection["sections"][0]["existing_text"])

    def test_custom_heading_style_with_body_override_is_not_reported_empty(self):
        source = TEST_ROOT / "flow1c-test-custom-style.docx"
        self.make_custom_heading_style_doc(source)

        inspection = docx_adapter.inspect_docx(source, [section()])

        self.assertEqual(inspection["sections"][0]["status"], "FOUND_CONTENT")
        self.assertIn("Старое содержимое раздела", inspection["sections"][0]["existing_text"])
        self.assertIn("Дополнительный текст раздела", inspection["sections"][0]["existing_text"])
        self.assertNotIn("Содержимое следующего раздела", inspection["sections"][0]["existing_text"])

    def test_custom_heading_style_replace_and_append_verify_full_written_text(self):
        source = TEST_ROOT / "flow1c-test-custom-source.docx"
        replaced = TEST_ROOT / "flow1c-test-custom-replaced.docx"
        appended = TEST_ROOT / "flow1c-test-custom-appended.docx"
        self.make_custom_heading_style_doc(source)

        docx_adapter.write_docx(source, replaced, [{
            "section": section(), "mode": "replace", "content": "Новый текст\n- Первый пункт\n- Второй пункт",
        }])
        replaced_inspection = docx_adapter.inspect_docx(replaced, [section()])
        self.assertIn("Новый текст", replaced_inspection["sections"][0]["existing_text"])
        self.assertIn("Второй пункт", replaced_inspection["sections"][0]["existing_text"])
        self.assertNotIn("Старое содержимое раздела", replaced_inspection["sections"][0]["existing_text"])
        self.assertIn("Содержимое следующего раздела", "\n".join(p.text for p in Document(replaced).paragraphs))

        docx_adapter.write_docx(source, appended, [{
            "section": section(), "mode": "append", "content": "Добавленный текст",
        }])
        appended_inspection = docx_adapter.inspect_docx(appended, [section()])
        self.assertIn("Старое содержимое раздела", appended_inspection["sections"][0]["existing_text"])
        self.assertIn("Добавленный текст", appended_inspection["sections"][0]["existing_text"])

    def test_lock_and_invalid_package_are_rejected(self):
        lock = TEST_ROOT / "~$flow1c-test-lock.docx"
        lock.write_bytes(b"not a docx")
        with self.assertRaises(docx_adapter.DocxError) as error:
            docx_adapter.validate_docx_package(lock)
        self.assertEqual(error.exception.code, "DOCX_LOCK_FILE")

    def test_write_uses_section_font_paragraph_layout_and_table_geometry(self):
        source = TEST_ROOT / "flow1c-test-layout.docx"
        output = TEST_ROOT / "flow1c-test-layout-filled.docx"
        document = Document()
        document.add_heading("Техническая реализация", level=1)
        body = document.add_paragraph(style="Normal")
        body.paragraph_format.left_indent = Pt(18)
        body.paragraph_format.space_after = Pt(9)
        run = body.add_run("{{заполнить}}")
        run.font.name = "Arial"
        run.font.size = Pt(11)
        table = document.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.columns[0].width = Pt(90)
        table.columns[1].width = Pt(280)
        table.cell(0, 0).text = "Поле"
        table.cell(0, 1).text = "Значение"
        table.cell(1, 0).text = "Образец"
        table.cell(1, 1).text = "Текст"
        document.add_heading("Роли и права", level=1)
        document.add_paragraph("Сохранить границу раздела")
        document.save(source)
        original = docx_adapter.sha256_file(source)

        content = "Новый абзац\n\n| Поле | Значение |\n| --- | --- |\n| A | B |\n| C | D |"
        docx_adapter.write_docx(source, output, [{"section": section(), "mode": "replace", "content": content}])

        reopened = Document(output)
        paragraph = next(item for item in reopened.paragraphs if item.text == "Новый абзац")
        self.assertEqual(paragraph.style.name, "Normal")
        self.assertEqual(paragraph.paragraph_format.left_indent, Pt(18))
        self.assertEqual(paragraph.paragraph_format.space_after, Pt(9))
        self.assertEqual(paragraph.runs[0].font.name, "Arial")
        self.assertEqual(paragraph.runs[0].font.size, Pt(11))
        self.assertEqual(reopened.tables[0].style.name, "Table Grid")
        self.assertEqual(reopened.tables[0].columns[0].width, Pt(90))
        self.assertEqual(reopened.tables[0].columns[1].width, Pt(280))
        self.assertEqual(len(reopened.tables[0].rows), 3)
        self.assertEqual(reopened.tables[0].cell(2, 1).text, "D")
        self.assertIn("Сохранить границу раздела", [item.text for item in reopened.paragraphs])
        self.assertEqual(docx_adapter.sha256_file(source), original)

    def test_table_column_mismatch_fails_before_creating_output(self):
        source = TEST_ROOT / "flow1c-test-table-shape.docx"
        output = TEST_ROOT / "flow1c-test-table-shape-filled.docx"
        document = Document()
        document.add_heading("Техническая реализация", level=1)
        document.add_table(rows=1, cols=2)
        document.save(source)

        with self.assertRaises(docx_adapter.DocxError) as error:
            docx_adapter.write_docx(source, output, [{"section": section(), "mode": "replace", "content": "| A | B | C |\n| --- | --- | --- |\n| 1 | 2 | 3 |"}])

        self.assertEqual(error.exception.code, "DOCX_LAYOUT_UNSUPPORTED")
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
