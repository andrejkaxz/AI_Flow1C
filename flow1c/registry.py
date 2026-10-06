"""Registry import and exact scope resolution."""

from __future__ import annotations

import argparse
import datetime as dt
import re
import shutil
from pathlib import Path
from typing import Any, Iterable

from flow1c import context as runtime
from flow1c import storage as storage
from flow1c.errors import WorkflowError
from flow1c.results import OperationResult


def normalized_header(value: Any) -> str:
    return re.sub("\\s+", " ", str(value or "").strip()).casefold()


def header_map_from_values(values: Iterable[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    for column, value in enumerate(values, start=1):
        key = normalized_header(value)
        if key and key not in result:
            result[key] = column
    return result


def header_map(sheet: Any, row: int) -> dict[str, int]:
    values = next(sheet.iter_rows(min_row=row, max_row=row, values_only=True), ())
    return header_map_from_values(values)


def find_column(headers: dict[str, int], *candidates: str) -> int | None:
    for candidate in candidates:
        found = headers.get(normalized_header(candidate))
        if found:
            return found
    return None


def normalized_cell_value(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return value


def row_cell_value(values: tuple[Any, ...], column: int | None) -> Any:
    if not column:
        return None
    index = column - 1
    return normalized_cell_value(values[index]) if index < len(values) else None


def cell_value(sheet: Any, row: int, column: int | None) -> Any:
    """Compatibility helper for callers that are not streaming a read-only sheet."""
    if not column:
        return None
    return normalized_cell_value(sheet.cell(row, column).value)


def split_ids(value: Any) -> list[str]:
    if value is None:
        return []
    return [item.strip().upper() for item in re.split("[,;\\n]+", str(value)) if item.strip()]


def make_diff(previous: dict[str, Any], current: dict[str, Any]) -> dict[str, list[str]]:
    old_keys = set(previous)
    new_keys = set(current)
    return {
        "added": sorted(new_keys - old_keys),
        "changed": sorted((key for key in old_keys & new_keys if previous[key] != current[key])),
        "missing_from_source": sorted(old_keys - new_keys),
    }


def registry_import(args: argparse.Namespace, *, product_root: Path) -> OperationResult:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import column_index_from_string
    except ImportError as exc:
        raise WorkflowError("openpyxl is required. Run scripts/bootstrap.ps1.") from exc
    config, _ = runtime.load_config(product_root=product_root)
    registry = config.get("registry", {})
    source = Path(args.file).expanduser().resolve()
    if not source.is_file():
        raise WorkflowError(f"Registry file not found: {source}")
    report_path = runtime.project_root(product_root=product_root) / "registry" / "import-report.md"

    def fail_import(kind: str, message: str) -> int:
        storage.write_text(
            report_path,
            "\n".join(
                [
                    "# Registry import report",
                    "",
                    f"Source: `{source.name}`",
                    "",
                    "## Validation",
                    "",
                    f"- ERROR: {message}",
                ]
            ),
        )
        storage.write_json(
            report_path.parent / "status.json",
            {
                "status": "invalid",
                "source_sha256": storage.sha256(source),
                "report_path": str(report_path),
                "updated_at": storage.utc_now(),
            },
        )
        _value = {
            "state": "BLOCKED",
            "report_path": str(report_path),
            "errors": [{"kind": kind, "message": message}],
            "warnings": [],
        }
        return OperationResult(_value, 2)

    try:
        workbook = load_workbook(source, data_only=False, read_only=True)
    except Exception as exc:
        return fail_import("invalid_workbook", f"Cannot open registry workbook: {exc}")
    req_sheet_name = registry.get("requirements_sheet", "Процессы требования")
    fs_sheet_name = registry.get("specifications_sheet", "Реестр ФС- не удалять")
    fatal_errors: list[str] = []
    errors: list[str] = []
    warnings: list[str] = []
    missing_sheets = [
        name for name in (req_sheet_name, fs_sheet_name) if name not in workbook.sheetnames
    ]
    if missing_sheets:
        workbook.close()
        return fail_import(
            "missing_sheet", "Registry sheets not found: " + ", ".join(missing_sheets)
        )
    req_sheet = workbook[req_sheet_name]
    req_header_row = int(registry.get("requirements_header_row", 3))
    req_data_row = int(registry.get("requirements_data_row", req_header_row + 1))
    headers = header_map(req_sheet, req_header_row)
    id_column_letter = str(registry.get("requirement_id_column", "J")).upper()
    id_column = column_index_from_string(id_column_letter)
    req_header_values = next(
        req_sheet.iter_rows(min_row=req_header_row, max_row=req_header_row, values_only=True), ()
    )
    id_header_value = row_cell_value(req_header_values, id_column)
    id_header = normalized_header(id_header_value)
    if id_header not in {"idтребования", "id требования"}:
        fatal_errors.append(
            f"Column {id_column_letter} must contain the requirement ID header; found '{id_header_value}'."
        )
    requirement_text_column = find_column(headers, "Требование")
    process_columns = {
        "code": find_column(headers, "Код БП"),
        "level_1": find_column(headers, "БП Уровень 1"),
        "level_2": find_column(headers, "БП Уровень 2"),
        "level_3": find_column(headers, "БП Уровень 3"),
        "block": find_column(headers, "Block", "Block for status"),
    }
    optional_columns = {
        "source": find_column(headers, "Источник"),
        "registered_at": find_column(headers, "Дата регистрации"),
        "owner": find_column(headers, "Владелец процесса"),
        "analyst": find_column(headers, "Аналитик/ФА (Ах)"),
        "criticality": find_column(headers, "Критичность"),
        "fit_gap": find_column(headers, "Покрывается стандартной функциональностью"),
        "fit_gap_comment": find_column(headers, "Комментарий к fit/gap анализу"),
    }
    requirements: dict[str, Any] = {}
    requirement_occurrences: dict[str, list[dict[str, Any]]] = {}
    duplicate_requirements: list[tuple[str, int]] = []
    req_columns = [
        id_column,
        requirement_text_column,
        *process_columns.values(),
        *optional_columns.values(),
    ]
    req_max_column = max((column for column in req_columns if column))
    req_rows = req_sheet.iter_rows(
        min_row=req_data_row, max_row=req_sheet.max_row, max_col=req_max_column, values_only=True
    )
    for row, values in enumerate(req_rows, start=req_data_row):
        raw_id = row_cell_value(values, id_column)
        if raw_id in (None, ""):
            continue
        requirement_id = str(raw_id).strip().upper()
        if requirement_id.casefold() in {"строку не удалять", "строка не удалять"}:
            continue
        process = {key: row_cell_value(values, column) for key, column in process_columns.items()}
        process = {key: value for key, value in process.items() if value not in (None, "")}
        details = {key: row_cell_value(values, column) for key, column in optional_columns.items()}
        details = {key: value for key, value in details.items() if value not in (None, "")}
        requirement_record = {
            "id": requirement_id,
            "source_row": row,
            "text": row_cell_value(values, requirement_text_column),
            "process": process,
            "details": details,
        }
        requirement_occurrences.setdefault(requirement_id, []).append(requirement_record)
        if requirement_id in requirements:
            duplicate_requirements.append((requirement_id, row))
            continue
        if not runtime.ID_PATTERN.match(requirement_id):
            warnings.append(f"Unusual requirement ID '{requirement_id}' at row {row}.")
        requirements[requirement_id] = requirement_record
    if duplicate_requirements:
        errors.extend(
            (
                f"Duplicate requirement ID {requirement_id} at row {row}."
                for requirement_id, row in duplicate_requirements
            )
        )
    if not requirements:
        fatal_errors.append(
            f"No requirement IDs found in column {id_column_letter}, starting at row {req_data_row}."
        )
    fs_sheet = workbook[fs_sheet_name]
    fs_header_row = int(registry.get("specifications_header_row", 1))
    fs_data_row = int(registry.get("specifications_data_row", fs_header_row + 1))
    fs_headers = header_map(fs_sheet, fs_header_row)
    fs_code_column = find_column(
        fs_headers, "Код разработки (RICEF)", "Код разработки  (RICEF)", "Код ФС"
    )
    fs_title_column = find_column(fs_headers, "Название разработки")
    fs_description_column = find_column(fs_headers, "Описание разработки")
    fs_status_column = find_column(fs_headers, "Статус ФС")
    fs_release_column = find_column(fs_headers, "Плановый релиз")
    fs_requirements_column = find_column(fs_headers, "Код требования", "Коды требований")
    if not fs_code_column:
        fatal_errors.append("The specifications register has no FS/development code column.")
    if not fs_requirements_column:
        fatal_errors.append("The specifications register has no requirement codes column.")
    specifications: dict[str, Any] = {}
    membership: dict[str, str] = {}
    membership_candidates: dict[str, list[str]] = {}
    specification_occurrences: dict[str, list[dict[str, Any]]] = {}
    if fs_code_column and fs_requirements_column:
        fs_columns = [
            fs_code_column,
            fs_title_column,
            fs_description_column,
            fs_status_column,
            fs_release_column,
            fs_requirements_column,
        ]
        fs_max_column = max((column for column in fs_columns if column))
        fs_rows = fs_sheet.iter_rows(
            min_row=fs_data_row, max_row=fs_sheet.max_row, max_col=fs_max_column, values_only=True
        )
        for row, values in enumerate(fs_rows, start=fs_data_row):
            raw_code = row_cell_value(values, fs_code_column)
            if raw_code in (None, ""):
                continue
            code = str(raw_code).strip()
            requirement_ids = split_ids(row_cell_value(values, fs_requirements_column))
            specification_record = {
                "code": code,
                "source_row": row,
                "title": row_cell_value(values, fs_title_column),
                "description": row_cell_value(values, fs_description_column),
                "status_from_registry": row_cell_value(values, fs_status_column),
                "planned_release": row_cell_value(values, fs_release_column),
                "requirements": requirement_ids,
            }
            specification_occurrences.setdefault(code, []).append(specification_record)
            if code in specifications:
                errors.append(f"Duplicate FS code: {code}.")
            for requirement_id in requirement_ids:
                candidates = membership_candidates.setdefault(requirement_id, [])
                if code not in candidates:
                    candidates.append(code)
                if requirement_id not in requirements:
                    errors.append(f"{code} references missing requirement {requirement_id}.")
                if requirement_id in membership:
                    errors.append(
                        f"MVP relation violation: {requirement_id} belongs to both {membership[requirement_id]} and {code}."
                    )
                else:
                    membership[requirement_id] = code
            if code not in specifications:
                specifications[code] = specification_record
    data_root = runtime.project_root(product_root=product_root)
    normalized_dir = data_root / "registry" / "normalized"
    previous_requirements = storage.read_json(normalized_dir / "requirements.json", {})
    previous_specifications = storage.read_json(normalized_dir / "specifications.json", {})
    requirement_diff = make_diff(previous_requirements, requirements)
    specification_diff = make_diff(previous_specifications, specifications)
    impact: dict[str, list[str]] = {}
    changed_ids = set(requirement_diff["changed"]) | set(requirement_diff["missing_from_source"])
    for manifest_path in sorted((data_root / "work-items").glob("*/manifest.yaml")):
        manifest = storage.read_json(manifest_path, {})
        affected = sorted(changed_ids & set(manifest.get("requirements", [])))
        if affected:
            impact[manifest.get("code", manifest_path.parent.name)] = affected
    report_lines = [
        "# Registry import report",
        "",
        f"Source: `{source.name}`",
        f"SHA-256: `{storage.sha256(source)}`",
        f"Imported at: `{dt.datetime.now().astimezone().isoformat(timespec='seconds')}`",
        "",
        f"Requirements: {len(requirements)}",
        f"Specifications: {len(specifications)}",
        "",
        "## Changes",
        "",
        f"- Added requirements: {', '.join(requirement_diff['added']) or 'none'}",
        f"- Changed requirements: {', '.join(requirement_diff['changed']) or 'none'}",
        f"- Missing from source: {', '.join(requirement_diff['missing_from_source']) or 'none'}",
        f"- Added specifications: {', '.join(specification_diff['added']) or 'none'}",
        f"- Changed specifications: {', '.join(specification_diff['changed']) or 'none'}",
        "",
        "## Impact",
        "",
    ]
    if impact:
        report_lines.extend((f"- {code}: {', '.join(ids)}" for code, ids in sorted(impact.items())))
        report_lines.append("")
        report_lines.append(
            "No PR was created. A user decision is required for each affected specification."
        )
    else:
        report_lines.append("No existing work item is affected.")
    report_lines.extend(["", "## Validation", ""])
    report_lines.extend((f"- ERROR: {message}" for message in [*fatal_errors, *errors]))
    report_lines.extend((f"- WARNING: {message}" for message in warnings))
    if not fatal_errors and (not errors) and (not warnings):
        report_lines.append("No validation issues found.")
    report_path = data_root / "registry" / "import-report.md"
    storage.write_text(report_path, "\n".join(report_lines))
    workbook.close()

    def issue_record(message: str, *, warning: bool = False) -> dict[str, Any]:
        row_match = re.search("\\brow (\\d+)\\b", message, flags=re.IGNORECASE)
        kind = "registry_warning" if warning else "registry_validation"
        lowered = message.casefold()
        if "duplicate requirement" in lowered:
            kind = "duplicate_requirement"
        elif "duplicate fs" in lowered:
            kind = "duplicate_specification"
        elif "missing requirement" in lowered:
            kind = "missing_requirement"
        record: dict[str, Any] = {"kind": kind, "message": message}
        if row_match:
            record["source_row"] = int(row_match.group(1))
        return record

    result = {
        "state": "BLOCKED" if fatal_errors else "PARTIAL" if errors else "READY",
        "report_path": str(report_path),
        "errors": [issue_record(message) for message in [*fatal_errors, *errors]],
        "warnings": [issue_record(message, warning=True) for message in warnings],
    }
    if fatal_errors:
        storage.write_json(
            report_path.parent / "status.json",
            {
                "status": "invalid",
                "source_sha256": storage.sha256(source),
                "report_path": str(report_path),
                "updated_at": storage.utc_now(),
            },
        )
        _value = result
        return OperationResult(_value, 2)
    source_target = data_root / "registry" / "source" / "requirements.xlsx"
    source_target.parent.mkdir(parents=True, exist_ok=True)
    if source != source_target.resolve():
        shutil.copy2(source, source_target)
    target_dir = data_root / "registry" / ("partial" if errors else "normalized")
    storage.write_json(target_dir / "requirements.json", requirements)
    storage.write_json(target_dir / "specifications.json", specifications)
    storage.write_json(target_dir / "links.json", {"requirement_to_specification": membership})
    if errors:
        storage.write_json(
            target_dir / "scope-index.json",
            {
                "schema_version": 1,
                "source_sha256": storage.sha256(source),
                "requirements": requirement_occurrences,
                "specifications": specification_occurrences,
                "requirement_to_specifications": membership_candidates,
                "errors": [issue_record(message) for message in errors],
            },
        )
    storage.write_json(
        data_root / "registry" / "status.json",
        {
            "status": "verified" if not errors else "partial",
            "source_sha256": storage.sha256(source),
            "report_path": str(report_path),
            "errors_count": len(errors),
            "warnings_count": len(warnings),
            "updated_at": storage.utc_now(),
        },
    )
    result.update(requirements=len(requirements), specifications=len(specifications))
    _value = result
    return OperationResult(_value, 1 if errors else 0)


def registry_scope_index(data_root: Path | None = None, *, product_root: Path) -> dict[str, Any]:
    """Load the latest usable registry view without replacing the last verified indexes."""
    root = data_root or runtime.project_root(product_root=product_root)
    status = storage.read_json(root / "registry" / "status.json", {})
    source_status = str(status.get("status") or "absent")
    if source_status == "invalid":
        return {"state": "invalid", "source_status": source_status, "status": status}
    if source_status == "partial":
        index = storage.read_json(root / "registry" / "partial" / "scope-index.json", {})
        if not isinstance(index, dict) or index.get("schema_version") != 1:
            return {"state": "invalid", "source_status": "invalid", "status": status}
        return {**index, "state": "usable", "source_status": source_status, "status": status}
    requirements = storage.read_json(root / "registry" / "normalized" / "requirements.json", {})
    specifications = storage.read_json(root / "registry" / "normalized" / "specifications.json", {})
    links = storage.read_json(root / "registry" / "normalized" / "links.json", {}).get(
        "requirement_to_specification", {}
    )
    legacy_verified = source_status == "absent" and bool(requirements) and bool(specifications)
    if (
        source_status != "verified"
        and (not legacy_verified)
        or not isinstance(requirements, dict)
        or (not isinstance(specifications, dict))
    ):
        return {"state": "absent", "source_status": source_status, "status": status}
    return {
        "schema_version": 1,
        "state": "usable",
        "source_status": "verified" if legacy_verified else source_status,
        "source_sha256": status.get("source_sha256"),
        "status": status,
        "requirements": {key: [value] for key, value in requirements.items()},
        "specifications": {key: [value] for key, value in specifications.items()},
        "requirement_to_specifications": {key: [value] for key, value in links.items()},
        "errors": [],
    }


def resolve_registry_reference(
    value: str, reference_kind: str = "auto", data_root: Path | None = None, *, product_root: Path
) -> dict[str, Any]:
    """Resolve an FS number or requirement ID against the latest scoped registry view."""
    raw = str(value or "").strip()
    kind = str(reference_kind or "auto").strip().casefold()
    if kind not in runtime.REFERENCE_KINDS:
        raise WorkflowError(f"Unknown reference_kind: {reference_kind}")
    index = registry_scope_index(data_root, product_root=product_root)
    if index["state"] != "usable":
        return {
            "state": index["state"],
            "requested_reference": raw,
            "reference_kind": kind,
            "source_status": index.get("source_status"),
        }
    requirements = index.get("requirements", {})
    specifications = index.get("specifications", {})
    requirement_key = raw.upper()
    requirement_match = requirement_key in requirements
    specification_match = raw in specifications
    if kind == "auto" and requirement_match and specification_match:
        return {
            "state": "ambiguous",
            "reason": "reference_kind_collision",
            "requested_reference": raw,
            "reference_kind": kind,
            "candidates": ["requirement", "specification"],
            "source_status": index["source_status"],
        }
    resolved_kind = kind
    if kind == "auto":
        resolved_kind = (
            "requirement"
            if requirement_match
            else "specification" if specification_match else "auto"
        )
    if resolved_kind == "requirement":
        if not requirement_match:
            return {
                "state": "not_found",
                "requested_reference": raw,
                "reference_kind": resolved_kind,
                "source_status": index["source_status"],
            }
        candidates = list(
            dict.fromkeys(index.get("requirement_to_specifications", {}).get(requirement_key, []))
        )
        if not candidates:
            return {
                "state": "not_assigned",
                "requested_reference": raw,
                "reference_kind": resolved_kind,
                "source_status": index["source_status"],
            }
        if len(candidates) > 1:
            return {
                "state": "ambiguous",
                "reason": "multiple_specifications",
                "requested_reference": raw,
                "reference_kind": resolved_kind,
                "candidates": candidates,
                "source_status": index["source_status"],
            }
        code = candidates[0]
    elif resolved_kind == "specification":
        if not specification_match:
            return {
                "state": "not_found",
                "requested_reference": raw,
                "reference_kind": resolved_kind,
                "source_status": index["source_status"],
            }
        code = raw
    else:
        return {
            "state": "not_found",
            "requested_reference": raw,
            "reference_kind": kind,
            "source_status": index["source_status"],
        }
    occurrences = specifications.get(code, [])
    linked_ids = list(
        dict.fromkeys(
            (
                requirement_id
                for occurrence in occurrences
                if isinstance(occurrence, dict)
                for requirement_id in occurrence.get("requirements", [])
            )
        )
    )
    existing_ids = [
        requirement_id for requirement_id in linked_ids if requirement_id in requirements
    ]
    if not existing_ids:
        return {
            "state": "no_existing_requirements",
            "requested_reference": raw,
            "reference_kind": resolved_kind,
            "code": code,
            "source_status": index["source_status"],
        }
    first = next((item for item in occurrences if isinstance(item, dict)), {})
    selected_requirements = {key: requirements[key][0] for key in existing_ids}
    return {
        "state": "resolved",
        "requested_reference": raw,
        "reference_kind": resolved_kind,
        "code": code,
        "requirements": existing_ids,
        "requirement_records": selected_requirements,
        "specification": {**first, "code": code, "requirements": existing_ids},
        "source_status": index["source_status"],
        "source_sha256": index.get("source_sha256") or index.get("status", {}).get("source_sha256"),
    }
